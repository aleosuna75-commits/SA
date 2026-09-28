# -*- coding: utf-8 -*-
"""
DASHBOARD EN HTML (un solo archivo, sin conexion a internet): indices (HParametros) y reservas (RRC / SONR / RFV),
real y proyectado. Se abre con doble clic en cualquier navegador actual (Chrome, Edge, Firefox o Safari de 2021 en
adelante) y se ve bien en pantalla, celular e impresion.

Genera salidas/Dashboard_Indices_Reservas.html a partir de las mismas salidas que dashboard.py (la version Excel):
  * pestaña "Indices"   : filtros (ramo, indice o LAG, periodo); indicadores; evolucion mensual real vs proyeccion
                          con banda al 80%; comparativo por ramo; patron de desarrollo (LAG 1-10); resumen del ramo.
  * pestaña "Reservas"  : filtros (reserva, concepto, ramo, moneda); indicadores; mensual 2025-2027; por ramo;
                          historico 2022-2027; por concepto.
  * pestaña "Analisis"  : tablas con mapa de calor (indices por ramo, totales por concepto) y modelo por tipo.
Cada grafica tiene tooltip (en las de lineas tambien con las flechas del teclado; en barras, con Tab) y una vista
de tabla equivalente.

Uso: lo llama proyeccion_reservas.py al final (GENERAR_DASHBOARD) o se ejecuta directo (F5 en VSCode).
"""
from __future__ import annotations

import json
import math
import sys

import openpyxl
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
try:
    import dashboard as dx  # noqa: E402  (lectores de las salidas, paleta y utilerias)
except ModuleNotFoundError as _e:
    if _e.name == "dashboard":
        raise SystemExit("Falta dashboard.py en la misma carpeta que dashboard_html.py. Los archivos del zip "
                         "(proyeccion_reservas.py, dashboard.py, dashboard_html.py, excel_fiel.py, tipo_cambio.py) "
                         "deben estar juntos en una sola carpeta.") from _e
    raise

SALIDA_HTML = dx.SALIDAS / "Dashboard_Indices_Reservas.html"
PRIMER_PERIODO_MENSUAL = dx.PRIMER_PERIODO_MENSUAL


def _redondear(v, dec):
    if v is None or (isinstance(v, float) and not math.isfinite(v)):
        return None
    return round(float(v), dec)


def leer_resumen() -> dict:
    """Pares (etiqueta, valor) de la hoja Resumen del diagnostico."""
    if not dx.ARCHIVO_DIAGNOSTICO.exists():
        return {}
    wb = openpyxl.load_workbook(dx.ARCHIVO_DIAGNOSTICO, read_only=True, data_only=True)
    out = {}
    if "Resumen" in wb.sheetnames:
        for a, b, *_ in wb["Resumen"].iter_rows(values_only=True):
            if a:
                out[str(a).strip()] = b
    wb.close()
    return out


MESES_CORTOS = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


def _clave_serie(r: dict) -> str:
    return f"{r['Libro']}|{r['Serie']}|{r['Ramo']}" if r["Libro"] == "HPARAM" else f"{r['Grupo']}|{r['Serie']}|{r['Ramo']}"


def leer_tendencias() -> dict:
    """Modelo por serie proyectada con 'Tendencia historica' (hojas Series_Modelos y Estacionalidad del diagnostico):
    {clave: [pendiente en la escala del modelo, ventana, log (1/0), valor de la recta en el ultimo mes,
             factores estacionales ene..dic en la escala del modelo o None]}."""
    if not dx.ARCHIVO_DIAGNOSTICO.exists():
        return {}
    wb = openpyxl.load_workbook(dx.ARCHIVO_DIAGNOSTICO, read_only=True, data_only=True)
    out, estacional = {}, {}
    if "Estacionalidad" in wb.sheetnames:
        filas = wb["Estacionalidad"].iter_rows(values_only=True)
        enc = [str(h) for h in next(filas)]
        for f in filas:
            r = dict(zip(enc, f))
            if r.get("Libro") is None or r.get("Credibilidad aplicada") in (None, 0):
                continue
            log = r.get("Transformacion") == "log"
            factores = [r.get(m) for m in MESES_CORTOS]
            if any(v is None for v in factores):
                continue
            estacional[_clave_serie(r)] = [_redondear(math.log1p(v) if log else v, 8) for v in factores]
    if "Series_Modelos" in wb.sheetnames:
        filas = wb["Series_Modelos"].iter_rows(values_only=True)
        enc = [str(h) for h in next(filas)]
        for f in filas:
            r = dict(zip(enc, f))
            if r.get("Modelo") != "Tendencia historica" or r.get("Tendencia mensual") is None:
                continue
            log = r.get("Transformacion") == "log"
            pend = math.log1p(r["Tendencia mensual"]) if log else r["Tendencia mensual"]
            clave = _clave_serie(r)
            out[clave] = [pend, r.get("Ventana (meses)"), 1 if log else 0, r.get("Ajuste tendencia ultimo mes"),
                          estacional.get(clave)]
    wb.close()
    return out


def preparar_datos() -> dict:
    """Lee las salidas de la proyeccion y arma el diccionario que se incrusta en el HTML."""
    registros_i, ramos_i, ultimo = dx.leer_indices()
    intervalos = dx.leer_intervalos()
    registros_m, tc, ramos_m = dx.leer_montos(ultimo)
    metodo = dx.leer_metodo()
    fin = max(r[3] for r in registros_i)
    fin_m = max(r[4] for r in registros_m)
    periodos_i = dx.rango(dx.PRIMER_PERIODO_INDICES, fin)
    periodos_m = dx.rango(dx.PRIMER_PERIODO_MONTOS, fin_m)
    pos_i = {p: k for k, p in enumerate(periodos_i)}
    pos_m = {p: k for k, p in enumerate(periodos_m)}

    ind, li, ls = {}, {}, {}
    for tipo, ramo, serie, per, v in registros_i:
        if per not in pos_i:
            continue
        ind.setdefault(ramo, {}).setdefault(serie, [None] * len(periodos_i))[pos_i[per]] = _redondear(v, 8)
        if tipo == "Proyección":
            a, b = intervalos.get((serie, ramo, per), (None, None))
            li.setdefault(ramo, {}).setdefault(serie, [None] * len(periodos_i))[pos_i[per]] = _redondear(a, 8)
            ls.setdefault(ramo, {}).setdefault(serie, [None] * len(periodos_i))[pos_i[per]] = _redondear(b, 8)

    mon = {}
    for libro, reserva, conc, ramo, per, tipo, v in registros_m:
        if per in pos_m:
            mon.setdefault(reserva, {}).setdefault(conc, {}).setdefault(ramo, [None] * len(periodos_m))[pos_m[per]] = \
                _redondear(v, 0)

    conceptos = {"RRC": ["NETO", "BRUTO", "BEL", "GTO", "IRR", "MR"], "SONR": ["NETO", "BRUTO", "BEL", "IRR", "MR"],
                 "RFV": ["NETO", "BRUTO", "IRR", "RCONT"]}
    conceptos = {k: [c for c in v if c in mon.get(k, {})] for k, v in conceptos.items() if k in mon}
    filas_metodo = []
    for d in metodo:
        filas_metodo.append({
            "tipo": d.get("Tipo"), "libro": d.get("Libro"), "modelo": d.get("Modelo"),
            "series": d.get("Series con backtest"), "cortes": d.get("Cortes por serie (mediana)"),
            "err_modelo": _redondear(d.get("Error % modelo (mediana)"), 1),
            "err_ultimo": _redondear(d.get("Error % ultimo valor (mediana)"), 1),
            "mejora": _redondear(d.get("% series en que el modelo mejora al ultimo valor"), 6),
        })
    modelo_por_tipo = {}
    for d in filas_metodo:
        modelo_por_tipo.setdefault(d["tipo"], d["modelo"])
    resumen = leer_resumen()
    ventana = resumen.get("Ventana de tendencia (meses)")
    try:
        ventana = int(float(ventana))
    except (TypeError, ValueError):
        ventana = None                                   # toda la historia
    persistencia = None
    try:                                             # "indices: ... con persistencia 0.8; montos y LAGs: ..."
        texto = str(resumen.get("Desviacion del ultimo mes", ""))
        persistencia = float(texto.split("persistencia")[1].split(";")[0].strip()) if "persistencia" in texto else None
    except (IndexError, ValueError):
        persistencia = None
    sin_tc = [p for p in periodos_m if not tc.get(p)]
    if sin_tc:
        print(f"   Aviso: sin tipo de cambio en la BD para {sin_tc}; en MXN esos meses se muestran como s/d")
    return {
        "ultimo": ultimo, "fin": fin, "fin_m": fin_m,
        "p_dic": (ultimo // 100) * 100 + 12, "p_12": dx.mover(ultimo, -12), "primer_mensual": PRIMER_PERIODO_MENSUAL,
        "periodos_i": periodos_i, "periodos_m": periodos_m,
        "ramos_i": ramos_i, "indices": dx.INDICES, "lags": dx.LAGS,
        "ind": ind, "li": li, "ls": ls,
        "ramos_m": {k: ramos_m.get(k, []) for k in conceptos}, "conceptos": conceptos, "mon": mon,
        "tc": {str(p): _redondear(t, 6) for p, t in tc.items()},
        "metodo": filas_metodo, "modelo_por_tipo": modelo_por_tipo, "ventana_tendencia": ventana,
        "persistencia_indices": persistencia,
        "tend": leer_tendencias(),
        "generado": datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


def generar(ruta_salida: Path = SALIDA_HTML) -> Path:
    print("Generando dashboard HTML ...", flush=True)
    datos = preparar_datos()
    json_datos = json.dumps(datos, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = (PLANTILLA.replace("__DATOS__", json_datos)
            .replace("__PIE__", dx.PIE_PAGINA)
            .replace("__GENERADO__", datos["generado"]))
    tmp = ruta_salida.with_suffix(".tmp.html")
    tmp.write_text(html, encoding="utf-8")
    tmp.replace(ruta_salida)
    print(f"   Dashboard HTML: {ruta_salida.name} ({ruta_salida.stat().st_size / 1024:,.0f} KB)", flush=True)
    return ruta_salida


# =============================================================================
# PLANTILLA (HTML + CSS + JS, todo en el archivo)
# =============================================================================
PLANTILLA = r"""<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Índices y reservas</title>
<meta name="description" content="Dashboard de índices de reservas y montos (RRC, SONR, RFV): real y proyección">
<style>
:root {
  color-scheme: light;
  --page: #f9f9f7; --surface: #fcfcfb; --surface-2: #f3f3f0;
  --ink: #0b0b0b; --ink-2: #52514e; --muted: #6f6e69;
  --grid: #e1e0d9; --axis: #c3c2b7; --border: rgba(11,11,11,0.10);
  --real: #2a78d6; --proy: #eb6834; --deemph: #c3c8d4; --deemph-ink: #898781;
  --ord-1: #86b6ef; --ord-2: #3987e5; --ord-3: #184f95;
  --div-neg: #2a78d6; --div-mid: #f0efec; --div-pos: #e34948;
  --focus: #2a78d6;
  --shadow: 0 1px 2px rgba(11,11,11,0.06);
}
@media (prefers-color-scheme: dark) {
  :root:where(:not([data-theme="light"])) {
    color-scheme: dark;
    --page: #0d0d0d; --surface: #1a1a19; --surface-2: #232322;
    --ink: #ffffff; --ink-2: #c3c2b7; --muted: #898781;
    --grid: #2c2c2a; --axis: #383835; --border: rgba(255,255,255,0.10);
    --real: #3987e5; --proy: #d95926; --deemph: #4a4a47; --deemph-ink: #a9a89f;
    --ord-1: #9ec5f4; --ord-2: #5598e7; --ord-3: #256abf;
    --div-neg: #3987e5; --div-mid: #383835; --div-pos: #e66767;
    --focus: #3987e5; --shadow: none;
  }
}
:root[data-theme="dark"] {
  color-scheme: dark;
  --page: #0d0d0d; --surface: #1a1a19; --surface-2: #232322;
  --ink: #ffffff; --ink-2: #c3c2b7; --muted: #898781;
  --grid: #2c2c2a; --axis: #383835; --border: rgba(255,255,255,0.10);
  --real: #3987e5; --proy: #d95926; --deemph: #4a4a47; --deemph-ink: #a9a89f;
  --ord-1: #9ec5f4; --ord-2: #5598e7; --ord-3: #256abf;
  --div-neg: #3987e5; --div-mid: #383835; --div-pos: #e66767;
  --focus: #3987e5; --shadow: none;
}
* { box-sizing: border-box; }
html, body { margin: 0; background: var(--page); color: var(--ink);
  font: 14px/1.45 system-ui, -apple-system, "Segoe UI", sans-serif; }
body { padding: 0 16px 32px; }
a { color: inherit; }
button, select { font: inherit; color: inherit; }
.encabezado { max-width: 1320px; margin: 0 auto; padding: 20px 0 8px; display: flex; flex-wrap: wrap;
  align-items: flex-end; justify-content: space-between; gap: 12px 24px; }
.encabezado h1 { font-size: 22px; margin: 0 0 4px; font-weight: 650; letter-spacing: -0.01em; }
.encabezado .sub { color: var(--ink-2); margin: 0; }
.acciones { display: flex; gap: 8px; align-items: center; flex-wrap: wrap; }
.tabs { display: flex; gap: 4px; background: var(--surface-2); border-radius: 10px; padding: 4px; }
.tabs button { border: 0; background: transparent; padding: 7px 14px; border-radius: 7px; cursor: pointer;
  color: var(--ink-2); font-weight: 500; }
.tabs button[aria-selected="true"] { background: var(--surface); color: var(--ink); box-shadow: var(--shadow);
  border: 1px solid var(--border); }
.tabs button:focus-visible, .btn:focus-visible, select:focus-visible, .plot:focus-visible {
  outline: 2px solid var(--focus); outline-offset: 2px; }
.btn { border: 1px solid var(--border); background: var(--surface); border-radius: 8px; padding: 6px 10px;
  cursor: pointer; color: var(--ink-2); }
.btn:hover { color: var(--ink); }
main { max-width: 1320px; margin: 0 auto; }
.vista[hidden] { display: none; }
.filtros { display: flex; flex-wrap: wrap; gap: 10px 16px; align-items: flex-end; padding: 12px 0 4px; }
.filtro label { display: block; font-size: 12px; color: var(--ink-2); margin-bottom: 4px; }
.filtro select { min-width: 150px; max-width: 100%; padding: 7px 30px 7px 10px; border-radius: 8px;
  border: 1px solid var(--border); background: var(--surface); color: var(--ink); appearance: none;
  background-image: linear-gradient(45deg, transparent 50%, var(--muted) 50%),
    linear-gradient(135deg, var(--muted) 50%, transparent 50%);
  background-position: calc(100% - 15px) 50%, calc(100% - 10px) 50%; background-size: 5px 5px; background-repeat: no-repeat; }
.filtros .nota { color: var(--ink-2); font-size: 12px; align-self: center; margin-left: auto; }
.aviso { color: #a33a0e; font-size: 13px; margin: 4px 0 0; }
:root[data-theme="dark"] .aviso { color: #f0a070; }
@media (prefers-color-scheme: dark) { :root:where(:not([data-theme="light"])) .aviso { color: #f0a070; } }
.kpis { display: grid; grid-template-columns: repeat(auto-fit, minmax(190px, 1fr)); gap: 12px; margin: 12px 0; }
.kpi { background: var(--surface); border: 1px solid var(--border); border-radius: 12px; padding: 14px 16px 12px;
  box-shadow: var(--shadow); min-width: 0; }
.kpi .lbl { color: var(--ink-2); font-size: 12.5px; margin-bottom: 6px; }
.kpi .val { font-size: 28px; font-weight: 600; letter-spacing: -0.01em; line-height: 1.1; }
.kpi .dlt { color: var(--ink-2); font-size: 12.5px; margin-top: 6px; }
.rejilla { display: grid; grid-template-columns: repeat(auto-fit, minmax(min(100%, 440px), 1fr)); gap: 14px; }
.tarjeta { background: var(--surface); border: 1px solid var(--border); border-radius: 12px; padding: 14px 16px 12px;
  box-shadow: var(--shadow); min-width: 0; display: flex; flex-direction: column; }
.tarjeta.ancha { grid-column: 1 / -1; }
.tarjeta header { display: flex; justify-content: space-between; align-items: flex-start; gap: 10px; margin-bottom: 6px; }
.tarjeta h2 { font-size: 15px; margin: 0; font-weight: 600; }
.tarjeta .sub { color: var(--ink-2); font-size: 12.5px; margin: 2px 0 0; }
.tarjeta .btn { font-size: 12px; padding: 4px 9px; flex: none; }
.cuerpo { position: relative; flex: 1; min-height: 0; }
.plot { display: block; width: 100%; height: auto; overflow: visible; }
.plot text { font-family: inherit; }
.tick { fill: var(--muted); font-size: 11px; font-variant-numeric: tabular-nums; }
.cat { fill: var(--ink-2); font-size: 11px; font-variant-numeric: tabular-nums; }
.grid { stroke: var(--grid); stroke-width: 1; shape-rendering: crispEdges; }
.base { stroke: var(--axis); stroke-width: 1; shape-rendering: crispEdges; }
.linea { fill: none; stroke-width: 2; stroke-linejoin: round; stroke-linecap: round; }
.marcador { stroke: var(--surface); stroke-width: 2; }
.etq { fill: var(--ink-2); font-size: 11px; }
.etq.fuerte { fill: var(--ink); font-weight: 600; }
.cruz { stroke: var(--muted); stroke-width: 1; shape-rendering: crispEdges; pointer-events: none; }
.barra { transition: filter .1s; }
.barra:hover, .barra.activa, .barra:focus-visible { filter: brightness(1.12); }
.barra:focus-visible { outline: 2px solid var(--focus); outline-offset: 1px; }
.hit { fill: transparent; }
.leyenda { display: flex; flex-wrap: wrap; gap: 6px 16px; margin-top: 8px; color: var(--ink-2); font-size: 12.5px; }
.leyenda .item { display: inline-flex; align-items: center; gap: 7px; }
.leyenda svg { width: 22px; height: 12px; flex: none; }
.tooltip { position: absolute; z-index: 5; pointer-events: none; background: var(--surface); color: var(--ink);
  border: 1px solid var(--border); border-radius: 8px; padding: 8px 10px; box-shadow: 0 4px 14px rgba(0,0,0,0.12);
  font-size: 12.5px; min-width: 150px; max-width: 260px; display: none; }
.tooltip .tt-titulo { color: var(--ink-2); margin-bottom: 4px; }
.tooltip .tt-fila { display: flex; align-items: center; gap: 8px; margin: 2px 0; }
.tooltip .tt-fila .llave { width: 16px; height: 0; border-top: 2px solid; flex: none; }
.tooltip .tt-fila .llave.dash { border-top-style: dashed; }
.tooltip .tt-fila .llave.caja { height: 10px; border: 0; border-radius: 2px; }
.tooltip .tt-fila b { font-weight: 600; font-variant-numeric: tabular-nums; }
.tooltip .tt-fila span { color: var(--ink-2); }
table.datos { width: 100%; border-collapse: collapse; font-size: 12.5px; font-variant-numeric: tabular-nums; }
table.datos th, table.datos td { padding: 5px 7px; border-bottom: 1px solid var(--grid); text-align: right;
  white-space: nowrap; }
table.datos th:first-child, table.datos td:first-child, table.datos .txt { text-align: left; }
table.datos thead th { color: var(--ink-2); font-weight: 600; position: sticky; top: 0; background: var(--surface);
  white-space: normal; vertical-align: bottom; z-index: 1; }
table.datos th:first-child, table.datos td:first-child { position: sticky; left: 0; background: var(--surface); }
table.datos thead th:first-child { z-index: 2; }
table.datos tbody tr:hover { background: var(--surface-2); }
.tabla-envoltura { max-height: 360px; overflow: auto; border: 1px solid var(--grid); border-radius: 8px; }
.analisis h2 { font-size: 16px; margin: 20px 0 8px; }
.analisis p { color: var(--ink-2); max-width: 90ch; }
.celda-calor { color: var(--ink); }
.pie { max-width: 1320px; margin: 28px auto 0; color: var(--ink-2); font-size: 12.5px; border-top: 1px solid var(--grid);
  padding-top: 12px; display: flex; flex-wrap: wrap; justify-content: space-between; gap: 6px 16px; }
.sr-only { position: absolute; width: 1px; height: 1px; overflow: hidden; clip: rect(0 0 0 0); }
@media (max-width: 600px) {
  .kpi .val { font-size: 24px; }
  table.datos { font-size: 12px; }
  table.datos th, table.datos td { padding: 4px 6px; }
  .encabezado h1 { font-size: 19px; }
  .filtro select { min-width: 130px; }
}
@media print {
  body { background: #fff; padding: 0; }
  .acciones, .tarjeta .btn, .filtros .nota { display: none; }
  .tarjeta, .kpi { break-inside: avoid; box-shadow: none; }
}
/* al imprimir siempre la paleta clara (va al final para ganar a las reglas del tema oscuro) */
@media print {
  :root, :root[data-theme="dark"], :root:where(:not([data-theme="light"])) {
    color-scheme: light;
    --page: #ffffff; --surface: #ffffff; --surface-2: #f3f3f0;
    --ink: #0b0b0b; --ink-2: #52514e; --muted: #6f6e69;
    --grid: #e1e0d9; --axis: #c3c2b7; --border: rgba(11,11,11,0.10);
    --real: #2a78d6; --proy: #eb6834; --deemph: #c3c8d4; --deemph-ink: #898781;
    --ord-1: #86b6ef; --ord-2: #3987e5; --ord-3: #184f95;
    --div-neg: #2a78d6; --div-mid: #f0efec; --div-pos: #e34948; --shadow: none;
  }
}
</style>
</head>
<body>
<script id="datos" type="application/json">__DATOS__</script>
<header class="encabezado">
  <div>
    <h1>Índices y reservas: real y proyección</h1>
    <p class="sub" id="subtitulo"></p>
  </div>
  <div class="acciones">
    <div class="tabs" role="tablist" aria-label="Vistas">
      <button role="tab" id="tab-indices" aria-selected="true" aria-controls="vista-indices">Índices</button>
      <button role="tab" id="tab-reservas" aria-selected="false" aria-controls="vista-reservas">Reservas</button>
      <button role="tab" id="tab-analisis" aria-selected="false" aria-controls="vista-analisis">Análisis</button>
    </div>
    <button class="btn" id="tema" type="button" title="Cambiar tema">Tema: sistema</button>
  </div>
</header>
<main>
  <section class="vista" id="vista-indices" role="tabpanel" aria-labelledby="tab-indices">
    <div class="filtros">
      <div class="filtro"><label for="f-ramo">Ramo</label><select id="f-ramo"></select></div>
      <div class="filtro"><label for="f-serie">Índice o LAG</label><select id="f-serie"></select></div>
      <div class="filtro"><label for="f-periodo">Periodo del comparativo por ramo</label><select id="f-periodo"></select></div>
      <div class="nota" id="nota-ind"></div>
    </div>
    <div class="kpis" id="kpis-ind"></div>
    <div class="rejilla" id="rejilla-ind"></div>
  </section>
  <section class="vista" id="vista-reservas" role="tabpanel" aria-labelledby="tab-reservas" hidden>
    <div class="filtros">
      <div class="filtro"><label for="f-reserva">Reserva</label><select id="f-reserva"></select></div>
      <div class="filtro"><label for="f-concepto">Concepto</label><select id="f-concepto"></select></div>
      <div class="filtro"><label for="f-ramo-res">Ramo</label><select id="f-ramo-res"></select></div>
      <div class="filtro"><label for="f-moneda">Moneda</label><select id="f-moneda"><option>USD</option><option>MXN</option></select></div>
      <div class="nota" id="nota-res"></div>
    </div>
    <div class="kpis" id="kpis-res"></div>
    <p class="sub" id="cedido" style="margin:0 0 10px;color:var(--ink-2)"></p>
    <div class="rejilla" id="rejilla-res"></div>
  </section>
  <section class="vista analisis" id="vista-analisis" role="tabpanel" aria-labelledby="tab-analisis" hidden></section>
</main>
<footer class="pie">
  <span>__PIE__</span>
  <span>Generado el __GENERADO__ por proyeccion_reservas.py · azul = real, naranja punteado = proyección, gris fino = modelo ajustado sobre la historia: recta de tendencia más el patrón por mes del año cuando la serie lo tiene; la proyección arranca del último real y, en los índices, su desviación respecto al modelo se desvanece hacia el nivel del último año (solo en las series que se proyectan directo: índices, LAGs y BEL/BRUTO por ramo en USD) · banda = intervalo al 80%</span>
</footer>
<script>
'use strict';
const D = JSON.parse(document.getElementById('datos').textContent);
const MESES = ['ene','feb','mar','abr','may','jun','jul','ago','sep','oct','nov','dic'];
const eti = p => `${MESES[p % 100 - 1]}-${String(p).slice(2, 4)}`;
const etiLarga = p => `${MESES[p % 100 - 1]} ${Math.floor(p / 100)}`;
const NF = {};
const nf = d => (NF[d] ||= new Intl.NumberFormat('es-MX', { minimumFractionDigits: d, maximumFractionDigits: d }));
const fmt = (v, d) => { if (v == null || !isFinite(v)) return 's/d'; if (Math.abs(v) < Math.pow(10, -d) / 2) v = 0; return nf(d).format(v); };
const varPct = (base, v) => (base > 0 && v != null && isFinite(v)) ? v / base - 1 : null;   // variacion solo con base positiva
const decimalesPaso = (ts, minimo) => { const paso = ts.length > 1 ? Math.abs(ts[1] - ts[0]) : 1; return Math.max(minimo || 0, Math.min(6, Math.ceil(-Math.log10(paso) - 1e-9))); };
const fmtPct = v => (v == null || !isFinite(v)) ? 's/d' : (v > 0 ? '+' : v < 0 ? '−' : '') + nf(1).format(Math.abs(v) * 100) + ' %';
const cssVar = n => `var(${n})`;                    // las marcas usan la variable: siguen al tema y a la impresion
const cssHex = n => getComputedStyle(document.documentElement).getPropertyValue(n).trim();
const el = (tag, attrs = {}, ...hijos) => {
  const e = document.createElement(tag);
  for (const [k, v] of Object.entries(attrs)) {
    if (k === 'class') e.className = v; else if (k === 'text') e.textContent = v; else if (v != null) e.setAttribute(k, v);
  }
  for (const h of hijos) if (h != null) e.append(h);
  return e;
};
const svg = (tag, attrs = {}) => {
  const e = document.createElementNS('http://www.w3.org/2000/svg', tag);
  for (const [k, v] of Object.entries(attrs)) if (v != null) e.setAttribute(k, v);
  return e;
};
const idx = (arr, v) => arr.indexOf(v);
// Curva del modelo ajustado sobre la historia: recta de tendencia del diagnostico (valor en el ultimo mes y
// pendiente) mas el patron por mes del ano, si la serie lo tiene, dibujada sobre la ventana de la regresion y hasta
// el ultimo dato real de la serie. La proyeccion es esa misma curva trasladada al ultimo real (anclaje), por eso no
// se prolonga al horizonte. escala = factor para pasar a las unidades de la grafica; periodos = AAAAMM de cada
// posicion (para saber el mes del ano).
function rectaModelo(clave, vals, periodos, escala) {
  const p = D.tend[clave]; if (!p) return null;
  const [pend, ventana, log, ajuste, est] = p;
  const iAncla = ultimoFinito(vals); if (iAncla < 0 || ajuste == null) return null;
  const out = new Array(periodos.length).fill(null);
  for (let i = Math.max(0, iAncla - ventana + 1); i <= iAncla; i++) {
    const v = ajuste + pend * (i - iAncla) + (est ? est[(periodos[i] % 100) - 1] : 0);
    out[i] = (log ? Math.exp(v) : v) * escala;
  }
  return out;
}
const tieneEstacionalidad = clave => !!(D.tend[clave] && D.tend[clave][4]);
const nombreModelo = clave => {
  const v = D.tend[clave] ? D.tend[clave][1] : D.ventana_tendencia;       // ventana propia de la serie (20, 32, 36...)
  return tieneEstacionalidad(clave) ? `Modelo ajustado (tendencia y patrón del año, ${v ? v + ' meses' : 'toda la historia'})`
    : (v ? `Tendencia (últimos ${v} meses)` : 'Tendencia (toda la historia)');
};
const ultimoFinito = arr => { for (let i = arr.length - 1; i >= 0; i--) if (arr[i] != null) return i; return -1; };

// ---------------------------------------------------------------- tema
const btnTema = document.getElementById('tema');
let tema = '';
try { tema = localStorage.getItem('tema') || ''; } catch (e) { tema = ''; }
function aplicarTema() {
  if (tema) document.documentElement.dataset.theme = tema; else delete document.documentElement.dataset.theme;
  btnTema.textContent = 'Tema: ' + (tema === 'light' ? 'claro' : tema === 'dark' ? 'oscuro' : 'sistema');
  redibujarTodo();
}
btnTema.addEventListener('click', () => {
  tema = tema === '' ? 'light' : tema === 'light' ? 'dark' : '';
  try { localStorage.setItem('tema', tema); } catch (e) { /* sin almacenamiento */ }
  aplicarTema();
});
window.matchMedia('(prefers-color-scheme: dark)').addEventListener('change', () => { if (!tema) redibujarTodo(); });

// ---------------------------------------------------------------- escalas
function ticks(min, max, n = 5) {
  if (!(isFinite(min) && isFinite(max))) return [0];
  if (min === max) { min -= 1; max += 1; }
  const span = max - min, paso0 = span / n, mag = Math.pow(10, Math.floor(Math.log10(paso0)));
  const cand = [1, 2, 2.5, 5, 10].map(m => m * mag);
  const paso = cand.find(c => c >= paso0) || cand[cand.length - 1];
  const out = [];
  for (let v = Math.ceil(min / paso) * paso; v <= max + paso * 1e-9; v += paso) out.push(+v.toFixed(10));
  return out;
}
function dominio(vals, incluirCero) {
  const f = vals.filter(v => v != null && isFinite(v));
  if (!f.length) return [0, 1];
  let lo = Math.min(...f), hi = Math.max(...f);
  if (incluirCero) { lo = Math.min(0, lo); hi = Math.max(0, hi); }
  const pad = (hi - lo || Math.abs(hi) || 1) * 0.06;
  lo = incluirCero && lo >= 0 ? 0 : lo - pad;
  hi = hi + pad;
  const t = ticks(lo, hi, 5);
  return [Math.min(lo, t[0]), Math.max(hi, t[t.length - 1])];
}
function pathBarra(x, y, w, h, lado) {
  // barra con el extremo del dato redondeado (4px) y la base (el cero) recta
  const r = Math.min(4, Math.abs(w) / 2, Math.abs(h) / 2);
  if (lado === 'arriba') return `M${x},${y + h} V${y + r} Q${x},${y} ${x + r},${y} H${x + w - r} Q${x + w},${y} ${x + w},${y + r} V${y + h} Z`;
  if (lado === 'abajo') return `M${x},${y} V${y + h - r} Q${x},${y + h} ${x + r},${y + h} H${x + w - r} Q${x + w},${y + h} ${x + w},${y + h - r} V${y} Z`;
  if (lado === 'izquierda') return `M${x + w},${y} H${x + r} Q${x},${y} ${x},${y + r} V${y + h - r} Q${x},${y + h} ${x + r},${y + h} H${x + w} Z`;
  return `M${x},${y} H${x + w - r} Q${x + w},${y} ${x + w},${y + r} V${y + h - r} Q${x + w},${y + h} ${x + w - r},${y + h} H${x} Z`;
}

// ---------------------------------------------------------------- tooltip
function crearTooltip(cuerpo) {
  const tt = el('div', { class: 'tooltip', role: 'status', 'aria-live': 'polite' });
  cuerpo.append(tt);
  return {
    mostrar(px, py, titulo, filas) {
      tt.replaceChildren(el('div', { class: 'tt-titulo', text: titulo }));
      for (const f of filas) {
        const llave = el('span', { class: 'llave' + (f.dash ? ' dash' : '') + (f.caja ? ' caja' : '') });
        llave.style.borderColor = f.color; if (f.caja) llave.style.background = f.color;
        tt.append(el('div', { class: 'tt-fila' }, llave, el('b', { text: f.valor }), el('span', { text: f.nombre })));
      }
      tt.style.display = 'block';
      const W = cuerpo.clientWidth, w = tt.offsetWidth, h = tt.offsetHeight;
      let x = px + 14; if (x + w > W - 4) x = px - w - 14; if (x < 0) x = 4;
      let y = py - h / 2; if (y < 0) y = 4; if (y + h > cuerpo.clientHeight) y = Math.max(4, cuerpo.clientHeight - h - 4);
      tt.style.left = x + 'px'; tt.style.top = y + 'px';
    },
    ocultar() { tt.style.display = 'none'; },
  };
}

// ---------------------------------------------------------------- tarjeta con grafica y vista de tabla
function tarjeta(titulo, sub, dibujar, tabla, ancha) {
  const cuerpo = el('div', { class: 'cuerpo' });
  const btn = el('button', { class: 'btn', type: 'button', 'aria-pressed': 'false', text: 'Tabla' });
  const h2 = el('h2', { text: titulo });
  const card = el('article', { class: 'tarjeta' + (ancha ? ' ancha' : '') },
    el('header', {}, el('div', {}, h2, el('p', { class: 'sub', text: sub })), btn), cuerpo);
  let modoTabla = false;
  const pintar = () => {
    cuerpo.replaceChildren();
    if (modoTabla) cuerpo.append(el('div', { class: 'tabla-envoltura' }, tabla()));
    else dibujar(cuerpo, Math.max(240, cuerpo.clientWidth || card.clientWidth - 32));
  };
  btn.addEventListener('click', () => {
    modoTabla = !modoTabla; btn.textContent = modoTabla ? 'Gráfica' : 'Tabla'; btn.setAttribute('aria-pressed', String(modoTabla)); pintar();
  });
  card._pintar = pintar;
  return card;
}
function tablaDatos(encabezados, filas) {
  const t = el('table', { class: 'datos' });
  t.append(el('thead', {}, el('tr', {}, ...encabezados.map(h => el('th', { scope: 'col', text: h })))));
  t.append(el('tbody', {}, ...filas.map(f => el('tr', {}, ...f.map((c, i) => el(i === 0 ? 'th' : 'td', i === 0 ? { scope: 'row', text: c } : { text: c }))))));
  return t;
}
function leyenda(items) {
  const l = el('div', { class: 'leyenda' });
  for (const it of items) {
    const s = svg('svg', { viewBox: '0 0 22 12', 'aria-hidden': 'true' });
    if (it.caja) s.append(svg('rect', { x: 0, y: 1, width: 22, height: 10, rx: 2, fill: it.color, 'fill-opacity': it.opacidad ?? 1 }));
    else s.append(svg('line', { x1: 0, y1: 6, x2: 22, y2: 6, stroke: it.color, 'stroke-width': it.fino ? 1.25 : 2, 'stroke-dasharray': it.dash ? '5 3' : null, 'stroke-linecap': 'round' }));
    if (it.marcador) s.append(svg('circle', { cx: 11, cy: 6, r: 3.5, fill: it.color }));
    l.append(el('span', { class: 'item' }, s, el('span', { text: it.nombre })));
  }
  return l;
}

// ---------------------------------------------------------------- grafica de lineas (con banda, cruz y tooltip)
function graficaLineas(cuerpo, W, o) {
  // o = {labels, series:[{nombre, valores, color, dash, marcadores, banda:{lo,hi}}], fmt, cadaX, etiquetasFin, incluirCero}
  const anchoTexto = t => 8 + 6.6 * String(t).length;
  let margenFin = 18;
  if (o.etiquetasFin) for (const se of o.series) if (se.etiquetaFin) { const i = ultimoFinito(se.valores); if (i >= 0) margenFin = Math.max(margenFin, anchoTexto(o.fmt(se.valores[i])) + 4); }
  const m = { l: 52, r: margenFin, t: 14, b: 30 }, H = o.alto || 250;
  const s = svg('svg', { class: 'plot', viewBox: `0 0 ${W} ${H}`, width: W, height: H, tabindex: 0, role: 'group', 'aria-roledescription': 'gráfica', 'aria-label': o.aria || '' });
  const n = o.labels.length, x0 = m.l, x1 = W - m.r, y0 = H - m.b, y1 = m.t;
  const X = i => n > 1 ? x0 + (x1 - x0) * i / (n - 1) : (x0 + x1) / 2;
  const todos = [];
  for (const se of o.series) { todos.push(...se.valores); if (se.banda) { todos.push(...se.banda.lo, ...se.banda.hi); } }
  const [lo, hi] = dominio(todos, o.incluirCero);
  const Y = v => y0 - (y0 - y1) * (v - lo) / (hi - lo);
  const ts = ticks(lo, hi, 5), decT = decimalesPaso(ts, o.decTick || 0);
  for (const t of ts) {
    if (t < lo || t > hi) continue;
    s.append(svg('line', { class: 'grid', x1: x0, x2: x1, y1: Y(t), y2: Y(t) }));
    const tx = svg('text', { class: 'tick', x: x0 - 8, y: Y(t) + 4, 'text-anchor': 'end' }); tx.textContent = fmt(t, decT); s.append(tx);
  }
  s.append(svg('line', { class: 'base', x1: x0, x2: x1, y1: y0, y2: y0 }));
  const auto = Math.max(1, Math.ceil(n / Math.max(2, Math.floor((x1 - x0) / 62))));
  const cada = Math.max(o.cadaX || 1, Math.ceil(auto / (o.cadaX || 1)) * (o.cadaX || 1));
  for (let i = 0; i < n; i += cada) {
    const tx = svg('text', { class: 'tick', x: X(i), y: y0 + 18, 'text-anchor': 'middle' }); tx.textContent = o.labels[i]; s.append(tx);
  }
  for (const se of o.series) {
    if (se.banda) {
      let d = '', back = [];
      for (let i = 0; i < n; i++) {
        const a = se.banda.lo[i], b = se.banda.hi[i];
        if (a == null || b == null) continue;
        d += (d ? 'L' : 'M') + X(i) + ',' + Y(b) + ' '; back.push(X(i) + ',' + Y(a));
      }
      if (d) s.append(svg('path', { d: d + 'L' + back.reverse().join(' L') + ' Z', fill: se.color, 'fill-opacity': 0.12, stroke: 'none' }));
    }
  }
  for (const se of o.series) {
    let d = '', abierto = false;
    for (let i = 0; i < n; i++) {
      const v = se.valores[i];
      if (v == null) { abierto = false; continue; }
      d += (abierto ? 'L' : 'M') + X(i) + ',' + Y(v) + ' '; abierto = true;
    }
    if (d) s.append(svg('path', { class: 'linea', d, stroke: se.color, 'stroke-width': se.fino ? 1.25 : null, 'stroke-dasharray': se.dash ? '6 4' : null }));
    if (se.marcadores) for (let i = 0; i < n; i++) if (se.valores[i] != null)
      s.append(svg('circle', { class: 'marcador', cx: X(i), cy: Y(se.valores[i]), r: 4, fill: se.color }));
  }
  if (o.etiquetasFin) {
    const puestas = [];
    for (const se of o.series) {
      if (!se.etiquetaFin) continue;
      const i = ultimoFinito(se.valores); if (i < 0) continue;
      const texto = o.fmt(se.valores[i]), ancho = anchoTexto(texto);
      let x = X(i) + 6, y = Y(se.valores[i]), ancla = 'start';
      if (x + ancho > W) { x = X(i); y -= 9; ancla = 'middle'; }         // sin lugar a la derecha: encima del punto
      for (const p of puestas) if (Math.abs(p.x - x) < ancho && Math.abs(p.y - y) < 12) y = p.y + (y >= p.y ? 12 : -12);
      puestas.push({ x, y });
      const tx = svg('text', { class: 'etq fuerte', x, y: y + 4, 'text-anchor': ancla }); tx.textContent = texto; s.append(tx);
    }
  }
  // capa de interaccion: cruz + tooltip con todas las series
  const cruz = svg('line', { class: 'cruz', y1: y1, y2: y0, x1: 0, x2: 0, visibility: 'hidden' });
  const puntos = svg('g'); s.append(cruz, puntos);
  const hit = svg('rect', { class: 'hit', x: x0, y: y1, width: x1 - x0, height: y0 - y1 }); s.append(hit);
  const tt = crearTooltip(cuerpo);
  let iSel = -1;
  const tieneDato = i => i >= 0 && i < n && o.series.some(se => se.valores[i] != null);
  const primerIndiceConDato = () => { for (const se of o.series) { const i = ultimoFinito(se.valores); if (i >= 0) return i; } return -1; };
  const mostrar = i0 => {
    if (i0 < 0 || i0 >= n) return ocultar();
    let i = i0;                                   // si el mes no tiene dato, el mas cercano que si (hasta 6 meses)
    for (let d = 1; !tieneDato(i) && d <= 6; d++) { if (tieneDato(i0 - d)) i = i0 - d; else if (tieneDato(i0 + d)) i = i0 + d; }
    if (!tieneDato(i)) return ocultar();
    iSel = i; cruz.setAttribute('x1', X(i)); cruz.setAttribute('x2', X(i)); cruz.setAttribute('visibility', 'visible');
    puntos.replaceChildren();
    const filas = [];
    for (const se of o.series) {
      const v = se.valores[i]; if (v == null) continue;
      puntos.append(svg('circle', { class: 'marcador', cx: X(i), cy: Y(v), r: 4.5, fill: se.color }));
      let valor = o.fmt(v);
      if (se.banda && se.banda.lo[i] != null && se.banda.hi[i] != null && !(se.banda.lo[i] === v && se.banda.hi[i] === v))
        valor += ` (${o.fmt(se.banda.lo[i])} – ${o.fmt(se.banda.hi[i])})`;
      filas.push({ valor, nombre: se.nombre, color: se.color, dash: se.dash });
    }
    const rect = s.getBoundingClientRect(), esc = rect.width / W;
    const py = filas.length ? Y(o.series.find(se => se.valores[i] != null).valores[i]) * esc : (y0 + y1) / 2 * esc;
    if (filas.length) tt.mostrar(X(i) * esc, py, o.tituloX ? o.tituloX(i) : o.labels[i], filas); else tt.ocultar();
  };
  const ocultar = () => { iSel = -1; cruz.setAttribute('visibility', 'hidden'); puntos.replaceChildren(); tt.ocultar(); };
  const desdePuntero = ev => {
    const rect = s.getBoundingClientRect(); const px = (ev.clientX - rect.left) * W / rect.width;
    const i = Math.max(0, Math.min(n - 1, Math.round((px - x0) / ((x1 - x0) || 1) * (n - 1))));
    mostrar(i);
  };
  hit.addEventListener('pointermove', desdePuntero); hit.addEventListener('pointerdown', desdePuntero);
  hit.addEventListener('pointerleave', ocultar);
  s.addEventListener('keydown', ev => {
    if (ev.key === 'ArrowRight' || ev.key === 'ArrowLeft') { ev.preventDefault(); const base = iSel < 0 ? primerIndiceConDato() : iSel; if (base >= 0) mostrar(Math.max(0, Math.min(n - 1, base + (ev.key === 'ArrowRight' ? 1 : -1)))); }
    else if (ev.key === 'Escape') ocultar();
  });
  s.addEventListener('focus', () => { if (iSel < 0) mostrar(primerIndiceConDato()); });
  s.addEventListener('blur', ocultar);
  cuerpo.append(s);
  if (o.series.length > 1) {
    const items = o.series.map(se => ({ nombre: se.nombre, color: se.color, dash: se.dash, marcador: se.marcadores, fino: se.fino }));
    if (o.series.some(se => se.banda)) items.push({ nombre: 'Intervalo al 80 %', color: cssVar('--proy'), caja: true, opacidad: 0.3 });
    cuerpo.append(leyenda(items));
  }
}

// ---------------------------------------------------------------- barras horizontales (agrupadas; con resaltado)
function graficaBarrasH(cuerpo, W, o) {
  // o = {categorias, series:[{nombre, valores, color | colores[]}], fmt, etiquetar: i => bool, leyendaItems}
  const k = o.series.length, grosor = Math.min(24, k > 1 ? 11 : 16), gap = 2, filaH = k * grosor + (k - 1) * gap + 10;
  const m = { l: 56, r: 56, t: 20, b: 8 }, n = o.categorias.length, H = m.t + m.b + n * filaH;
  const s = svg('svg', { class: 'plot', viewBox: `0 0 ${W} ${H}`, width: W, height: H, role: 'group', 'aria-roledescription': 'gráfica', 'aria-label': o.aria || '' });
  const x0 = m.l, x1 = W - m.r;
  const todos = o.series.flatMap(se => se.valores);
  const [lo, hi] = dominio(todos, true);
  const X = v => x0 + (x1 - x0) * (v - lo) / (hi - lo);
  const ts = ticks(lo, hi, 4), decT = decimalesPaso(ts, 0);
  for (const t of ts) {
    if (t < lo || t > hi) continue;
    s.append(svg('line', { class: 'grid', y1: m.t - 6, y2: H - m.b, x1: X(t), x2: X(t) }));
    const tx = svg('text', { class: 'tick', x: X(t), y: m.t - 9, 'text-anchor': 'middle' }); tx.textContent = fmt(t, decT); s.append(tx);
  }
  s.append(svg('line', { class: 'base', x1: X(0), x2: X(0), y1: m.t - 6, y2: H - m.b }));
  const tt = crearTooltip(cuerpo), etiquetas = svg('g');
  o.categorias.forEach((cat, i) => {
    const yFila = m.t + i * filaH + 5;
    const tx = svg('text', { class: 'cat', x: x0 - 8, y: yFila + (filaH - 10) / 2 + 4, 'text-anchor': 'end' }); tx.textContent = cat; s.append(tx);
    o.series.forEach((se, j) => {
      const v = se.valores[i]; if (v == null) return;
      const y = yFila + j * (grosor + gap), color = se.colores ? se.colores[i] : se.color, neg = v < 0;
      const xa = X(Math.min(0, v)), xb = X(Math.max(0, v)), xDato = neg ? xa : xb;
      const p = svg('path', { class: 'barra', d: pathBarra(xa, y, Math.max(xb - xa, 0.5), grosor, neg ? 'izquierda' : 'derecha'), fill: color });
      p.setAttribute('tabindex', 0); p.setAttribute('role', 'img'); p.setAttribute('aria-label', `${cat}, ${se.nombre}: ${o.fmt(v)}`);
      const mostrar = () => {
        const rect = s.getBoundingClientRect(), esc = rect.width / W;
        tt.mostrar(xDato * esc, (y + grosor / 2) * esc, cat, [{ valor: o.fmt(v), nombre: se.nombre, color, caja: true }]);
      };
      p.addEventListener('pointermove', mostrar); p.addEventListener('focus', mostrar);
      p.addEventListener('pointerleave', tt.ocultar); p.addEventListener('blur', tt.ocultar);
      s.append(p);
      if (o.etiquetar && o.etiquetar(i, j)) {
        const e = svg('text', { class: 'etq fuerte', x: xDato + (neg ? -6 : 6), y: y + grosor / 2 + 4, 'text-anchor': neg ? 'end' : 'start' }); e.textContent = o.fmt(v); etiquetas.append(e);
      }
    });
  });
  s.append(etiquetas);                     // las etiquetas van encima de todas las barras
  cuerpo.append(s);
  if (o.leyendaItems) cuerpo.append(leyenda(o.leyendaItems));
}

// ---------------------------------------------------------------- columnas (agrupadas o en un solo carril)
function graficaColumnas(cuerpo, W, o) {
  // o = {categorias, series:[{nombre, valores, color}], fmt, compartido (una sola columna por categoria), cadaX}
  const m = { l: 52, r: 12, t: 14, b: 30 }, H = o.alto || 250, n = o.categorias.length;
  const s = svg('svg', { class: 'plot', viewBox: `0 0 ${W} ${H}`, width: W, height: H, role: 'group', 'aria-roledescription': 'gráfica', 'aria-label': o.aria || '' });
  const x0 = m.l, x1 = W - m.r, y0 = H - m.b, y1 = m.t;
  const carril = (x1 - x0) / n, k = o.compartido ? 1 : o.series.length, gap = 2;
  const grosor = Math.min(24, Math.max(3, (carril * 0.72 - (k - 1) * gap) / k));
  const anchoGrupo = k * grosor + (k - 1) * gap;
  const todos = o.series.flatMap(se => se.valores);
  const [lo, hi] = dominio(todos, true);
  const Y = v => y0 - (y0 - y1) * (v - lo) / (hi - lo);
  const ts = ticks(lo, hi, 5), decT = decimalesPaso(ts, 0);
  for (const t of ts) {
    if (t < lo || t > hi) continue;
    s.append(svg('line', { class: 'grid', x1: x0, x2: x1, y1: Y(t), y2: Y(t) }));
    const tx = svg('text', { class: 'tick', x: x0 - 8, y: Y(t) + 4, 'text-anchor': 'end' }); tx.textContent = fmt(t, decT); s.append(tx);
  }
  s.append(svg('line', { class: 'base', x1: x0, x2: x1, y1: Y(0), y2: Y(0) }));
  const auto = Math.max(1, Math.ceil(n / Math.max(2, Math.floor((x1 - x0) / 60))));
  const cada = Math.max(o.cadaX || 1, Math.ceil(auto / (o.cadaX || 1)) * (o.cadaX || 1));
  for (let i = 0; i < n; i += cada) {
    const tx = svg('text', { class: 'tick', x: x0 + carril * (i + 0.5), y: y0 + 18, 'text-anchor': 'middle' }); tx.textContent = o.categorias[i]; s.append(tx);
  }
  const tt = crearTooltip(cuerpo), etiquetas = svg('g');
  o.categorias.forEach((cat, i) => {
    const xg = x0 + carril * (i + 0.5) - anchoGrupo / 2;
    const etiquetables = o.series.map((se, j) => se.valores[i] != null && o.etiquetar && o.etiquetar(i, j) ? j : -1).filter(j => j >= 0);
    // con columnas angostas y dos etiquetas en el mismo grupo, solo se etiqueta la ultima serie (la tabla trae el resto)
    const jEtiqueta = grosor < 30 && etiquetables.length > 1 ? etiquetables[etiquetables.length - 1] : null;
    o.series.forEach((se, j) => {
      const v = se.valores[i]; if (v == null) return;
      const x = xg + (o.compartido ? 0 : j * (grosor + gap)), neg = v < 0;
      const ya = Y(Math.max(0, v)), yb = Y(Math.min(0, v)), yDato = neg ? yb : ya;
      const p = svg('path', { class: 'barra', d: pathBarra(x, ya, grosor, Math.max(yb - ya, 0.5), neg ? 'abajo' : 'arriba'), fill: se.color });
      p.setAttribute('tabindex', 0); p.setAttribute('role', 'img'); p.setAttribute('aria-label', `${o.tituloX ? o.tituloX(i) : cat}, ${se.nombre}: ${o.fmt(v)}`);
      const mostrar = () => {
        const rect = s.getBoundingClientRect(), esc = rect.width / W;
        tt.mostrar((x + grosor / 2) * esc, yDato * esc, o.tituloX ? o.tituloX(i) : cat, [{ valor: o.fmt(v), nombre: se.nombre, color: se.color, caja: true }]);
      };
      p.addEventListener('pointermove', mostrar); p.addEventListener('focus', mostrar);
      p.addEventListener('pointerleave', tt.ocultar); p.addEventListener('blur', tt.ocultar);
      s.append(p);
      if (etiquetables.includes(j) && (jEtiqueta == null || j === jEtiqueta)) {
        const e = svg('text', { class: 'etq fuerte', x: x + grosor / 2, y: neg ? yDato + 13 : yDato - 6, 'text-anchor': 'middle' }); e.textContent = o.fmt(v); etiquetas.append(e);
      }
    });
  });
  s.append(etiquetas);
  cuerpo.append(s);
  if (o.series.length > 1) cuerpo.append(leyenda(o.series.map(se => ({ nombre: se.nombre, color: se.color, caja: true }))));
}

// ---------------------------------------------------------------- mapa de calor para tablas
function colorDivergente(v, lim) {
  // v en [-lim, lim] -> azul (negativo) · gris (0) · rojo (positivo), mezclando con el neutro
  if (v == null || !isFinite(v)) return null;
  const t = Math.max(-1, Math.min(1, v / lim));
  const mezcla = (a, b, f) => {
    const pa = a.match(/\w\w/g).map(h => parseInt(h, 16)), pb = b.match(/\w\w/g).map(h => parseInt(h, 16));
    return 'rgb(' + pa.map((c, i) => Math.round(c + (pb[i] - c) * f)).join(',') + ')';
  };
  const mid = cssHex('--div-mid'), pole = t < 0 ? cssHex('--div-neg') : cssHex('--div-pos');
  return mezcla(mid, pole, Math.abs(t) * 0.75);
}

// ---------------------------------------------------------------- estado
const E = {
  ramo: D.ramos_i.includes('10') ? '10' : D.ramos_i[0], serie: D.indices[0], periodo: D.fin,
  reserva: 'RRC', concepto: 'NETO', ramoRes: 'Todos', moneda: 'USD', vista: 'indices',
};
const tarjetasVivas = [];
function redibujarTodo() {
  for (const c of tarjetasVivas) if (c.isConnected && c._pintar) c._pintar();
  if (E.vista === 'analisis') pintarAnalisis();          // los colores del mapa de calor dependen del tema
}
let redim; window.addEventListener('resize', () => { clearTimeout(redim); redim = setTimeout(redibujarTodo, 120); });

function llenarSelect(sel, opciones, valor) {
  sel.replaceChildren();
  for (const op of opciones) {
    if (op.grupo) { const g = el('optgroup', { label: op.grupo }); for (const o of op.opciones) g.append(el('option', { value: o.v, text: o.t })); sel.append(g); }
    else sel.append(el('option', { value: op.v, text: op.t }));
  }
  sel.value = String(valor);
}
function kpi(lbl, val, dlt) { return el('div', { class: 'kpi' }, el('div', { class: 'lbl', text: lbl }), el('div', { class: 'val', text: val }), dlt ? el('div', { class: 'dlt', text: dlt }) : null); }

// ---------------------------------------------------------------- vista: indices
const P_I = D.periodos_i, iUlt = idx(P_I, D.ultimo), iDic = idx(P_I, D.p_dic), iFin = idx(P_I, D.fin), i12 = idx(P_I, D.p_12);
const valorInd = (ramo, serie, i) => ((D.ind[ramo] || {})[serie] || [])[i] ?? null;
const serieInd = (ramo, serie) => (D.ind[ramo] || {})[serie] || new Array(P_I.length).fill(null);
const promedio12 = arr => { const v = arr.slice(i12 + 1, iUlt + 1).filter(x => x != null); return v.length ? v.reduce((a, b) => a + b, 0) / v.length : null; };
const decInd = serie => 4;

function pintarIndices() {
  const { ramo, serie } = E, iPer = idx(P_I, E.periodo);
  const vals = serieInd(ramo, serie), esLag = D.lags.includes(serie);
  document.getElementById('nota-ind').textContent = `Real hasta ${eti(D.ultimo)} · proyección ${eti(P_I[iUlt + 1])} a ${eti(D.fin)} · ${esLag ? 'LAGs' : 'Índices'}: ${D.modelo_por_tipo[esLag ? 'lag' : 'indice'] || 'n/d'}`;
  const ult = vals[iUlt], pDic = vals[iDic], pFin = vals[iFin], prom = promedio12(vals);
  const f4 = v => fmt(v, 4), conPatron = tieneEstacionalidad(`HPARAM|${serie}|${ramo}`);
  document.getElementById('kpis-ind').replaceChildren(
    kpi(`Último real (${eti(D.ultimo)})`, f4(ult)),
    kpi('Promedio real 12 meses', f4(prom), `${eti(P_I[i12 + 1])} a ${eti(D.ultimo)}`),
    kpi(`Proyección ${eti(D.p_dic)}`, f4(pDic), varPct(ult, pDic) == null ? null : fmtPct(varPct(ult, pDic)) + ' vs último real' + (conPatron ? ' (incluye el mes del año)' : '')),
    kpi(`Proyección ${eti(D.fin)}`, f4(pFin), varPct(ult, pFin) == null ? null : fmtPct(varPct(ult, pFin)) + ' vs último real' + (conPatron ? ' (incluye el mes del año)' : '')),
    kpi(`Variación a 12 meses (${eti(P_I[iUlt + 12])} vs ${eti(D.ultimo)})`, fmtPct(varPct(ult, vals[iUlt + 12])), conPatron ? 'mismo mes del año: sin el efecto estacional' : 'mismo mes del año siguiente'),
  );
  const real = vals.map((v, i) => i <= iUlt ? v : null);
  const proy = vals.map((v, i) => i >= iUlt ? v : null);
  const lo = (D.li[ramo] || {})[serie] || [], hi = (D.ls[ramo] || {})[serie] || [];
  const banda = { lo: P_I.map((_, i) => i === iUlt ? vals[i] : (lo[i] ?? null)), hi: P_I.map((_, i) => i === iUlt ? vals[i] : (hi[i] ?? null)) };
  const rejilla = document.getElementById('rejilla-ind'); rejilla.replaceChildren(); tarjetasVivas.length = 0;

  const nombreTend = nombreModelo(`HPARAM|${serie}|${ramo}`);
  const tendI = rectaModelo(`HPARAM|${serie}|${ramo}`, real, P_I, 1);
  const c1 = tarjeta(`Evolución mensual · ${serie} · ramo ${ramo}`, (!tendI ? 'Real y proyección con banda al 80 %' : conPatron
    ? 'Real, modelo ajustado (tendencia + patrón del año) y proyección con banda al 80 %' : 'Real, línea de tendencia y proyección con banda al 80 %')
    + (!esLag && D.persistencia_indices != null && D.persistencia_indices < 1 ? ` · la desviación del último mes se desvanece hacia el nivel del último año (persistencia ${D.persistencia_indices})` : ''),
    (cuerpo, W) => graficaLineas(cuerpo, W, {
      labels: P_I.map(eti), tituloX: i => etiLarga(P_I[i]), fmt: v => fmt(v, decInd(serie)), cadaX: 6, etiquetasFin: true,
      aria: `Evolución mensual de ${serie} del ramo ${ramo}`,
      series: [
        ...(tendI ? [{ nombre: nombreTend, valores: tendI, color: cssVar('--deemph-ink'), fino: true }] : []),
        { nombre: 'Real', valores: real, color: cssVar('--real'), etiquetaFin: true },
        { nombre: 'Proyección', valores: proy, color: cssVar('--proy'), dash: true, banda, etiquetaFin: true },
      ],
    }),
    () => tablaDatos(['Mes', 'Real', 'Proyección', 'Banda inferior', 'Banda superior', ...(tendI ? [nombreTend] : [])],
      P_I.map((p, i) => [etiLarga(p), fmt(real[i], 4), fmt(i > iUlt ? proy[i] : null, 4), fmt(i > iUlt ? lo[i] : null, 4), fmt(i > iUlt ? hi[i] : null, 4), ...(tendI ? [fmt(tendI[i], 4)] : [])])));

  const valoresRamo = D.ramos_i.map(r => valorInd(r, serie, iPer));
  const c2 = tarjeta(`${serie} por ramo · ${etiLarga(E.periodo)}${iPer > iUlt ? ' (proyección)' : ' (real)'}`, 'El ramo seleccionado se resalta en azul',
    (cuerpo, W) => graficaBarrasH(cuerpo, W, {
      categorias: D.ramos_i, fmt: v => fmt(v, 2), aria: `${serie} por ramo en ${etiLarga(E.periodo)}`,
      series: [{ nombre: serie, valores: valoresRamo, colores: D.ramos_i.map(r => r === ramo ? cssVar('--real') : cssVar('--deemph')) }],
      etiquetar: i => D.ramos_i[i] === ramo,
      leyendaItems: [{ nombre: 'Ramo seleccionado', color: cssVar('--real'), caja: true }, { nombre: 'Otros ramos', color: cssVar('--deemph'), caja: true }],
    }),
    () => tablaDatos(['Ramo', `${serie} ${etiLarga(E.periodo)}`], D.ramos_i.map((r, i) => [r, fmt(valoresRamo[i], 4)])));
  c2.querySelector('.cuerpo').style.minHeight = '0';

  const cortes = [[D.p_12 - 100, 'real'], [D.p_12, 'real'], [D.ultimo, 'real'], [D.fin, 'proyección']].filter(([p]) => idx(P_I, p) >= 0);
  const ordinales = [cssVar('--ord-1'), cssVar('--ord-2'), cssVar('--ord-3')];
  const seriesLag = cortes.map(([p, tipo], k) => ({
    nombre: `${eti(p)} (${tipo})`, color: tipo === 'proyección' ? cssVar('--proy') : ordinales[k] || ordinales[2], dash: tipo === 'proyección', marcadores: true,
    valores: D.lags.map(l => valorInd(ramo, l, idx(P_I, p))),
  }));
  const c3 = tarjeta(`Patrón de desarrollo acumulado (LAG 1–10) · ramo ${ramo}`, 'Tres cortes reales y la proyección a dic-27',
    (cuerpo, W) => graficaLineas(cuerpo, W, { labels: D.lags.map(l => l.replace('LAG ', 'L')), tituloX: i => D.lags[i], fmt: v => fmt(v, 2), cadaX: 1, incluirCero: true, series: seriesLag, aria: `Patrón LAG 1 a 10 del ramo ${ramo}` }),
    () => tablaDatos(['LAG', ...seriesLag.map(s => s.nombre)], D.lags.map((l, i) => [l, ...seriesLag.map(s => fmt(s.valores[i], 4))])));

  const filasRes = [...D.indices, ...D.lags.slice(0, 3)].map(se => {
    const v = serieInd(ramo, se); const u = v[iUlt], pf = v[iFin];
    return { se, u, prom: promedio12(v), pd: v[iDic], pf, var: varPct(u, pf) };
  });
  const c4 = tarjeta(`Resumen de índices · ramo ${ramo}`, `Últimos reales y proyección, con la variación a ${eti(D.fin)} (en las series con patrón incluye el mes del año)`,
    (cuerpo) => {
      const t = el('table', { class: 'datos' });
      t.append(el('thead', {}, el('tr', {}, ...['Índice', `Real ${eti(D.ultimo)}`, 'Prom. 12 m', `Proy. ${eti(D.p_dic)}`, `Proy. ${eti(D.fin)}`, 'Var. vs real'].map(h => el('th', { scope: 'col', text: h })))));
      const tb = el('tbody');
      for (const f of filasRes) {
        const cv = el('td', { text: fmtPct(f.var) }); const bg = colorDivergente(f.var, 0.3); if (bg) { cv.style.background = bg; cv.className = 'celda-calor'; }
        tb.append(el('tr', {}, el('th', { scope: 'row', text: f.se }), el('td', { text: fmt(f.u, 4) }), el('td', { text: fmt(f.prom, 4) }), el('td', { text: fmt(f.pd, 4) }), el('td', { text: fmt(f.pf, 4) }), cv));
      }
      t.append(tb); cuerpo.append(el('div', { class: 'tabla-envoltura', style: 'max-height:none;border:0' }, t));
    },
    () => tablaDatos(['Índice', `Real ${eti(D.ultimo)}`, 'Prom. 12 m', `Proy. ${eti(D.p_dic)}`, `Proy. ${eti(D.fin)}`, 'Var. vs real'],
      filasRes.map(f => [f.se, fmt(f.u, 4), fmt(f.prom, 4), fmt(f.pd, 4), fmt(f.pf, 4), fmtPct(f.var)])));
  c4.querySelector('.btn').remove();
  for (const c of [c1, c2, c3, c4]) { rejilla.append(c); tarjetasVivas.push(c); }
  redibujarTodo();
}

// ---------------------------------------------------------------- vista: reservas
const P_M = D.periodos_m, jUlt = idx(P_M, D.ultimo), jDic = idx(P_M, D.p_dic), jFin = idx(P_M, D.fin_m), j12 = idx(P_M, D.p_12), jMen = Math.max(0, idx(P_M, D.primer_mensual));
const tcDe = p => D.tc[String(p)] || null;
function montoSerie(reserva, concepto, ramo) {
  // suma por ramo (o "Todos") en la moneda elegida, en millones
  const porRamo = (D.mon[reserva] || {})[concepto] || {};
  const ramos = ramo === 'Todos' ? Object.keys(porRamo) : [ramo];
  return P_M.map((p, i) => {
    let s = null;
    for (const r of ramos) { const v = (porRamo[r] || [])[i]; if (v != null) s = (s || 0) + v; }
    if (s == null) return null;
    const f = E.moneda === 'MXN' ? tcDe(p) : 1;
    if (f == null) return null;                 // sin TC en la BD: mejor s/d que una cifra falsa
    return s * f / 1e6;
  });
}
function pintarReservas() {
  const { reserva, concepto, moneda } = E; const ramoRes = E.ramoRes;
  const conceptos = D.conceptos[reserva] || [], ramos = D.ramos_m[reserva] || [];
  const aviso = document.getElementById('nota-res');
  aviso.textContent = `Cifras en millones de ${moneda}${moneda === 'MXN' ? ' (USD × TC del mes)' : ''} · real hasta ${eti(D.ultimo)} · montos: ${D.modelo_por_tipo.nivel || 'n/d'}; razones: ${D.modelo_por_tipo.razon || 'n/d'}`;
  const vals = montoSerie(reserva, concepto, ramoRes);
  const u = vals[jUlt], pd = vals[jDic], pf = vals[jFin], v12 = vals[j12], p12 = vals[jUlt + 12];
  const f1 = v => fmt(v, 1);
  document.getElementById('kpis-res').replaceChildren(
    kpi(`Real ${eti(D.ultimo)} (M ${moneda})`, f1(u)),
    kpi(`Proyección ${eti(D.p_dic)} (M ${moneda})`, f1(pd), varPct(u, pd) == null ? null : fmtPct(varPct(u, pd)) + ' vs último real'),
    kpi(`Proyección ${eti(D.fin_m)} (M ${moneda})`, f1(pf), varPct(u, pf) == null ? null : fmtPct(varPct(u, pf)) + ' vs último real'),
    kpi('Crecimiento proyectado 12 meses', fmtPct(varPct(u, p12)), `${eti(D.ultimo)} a ${eti(P_M[jUlt + 12])}, mismo mes`),
    kpi('Crecimiento real 12 meses', fmtPct(varPct(v12, u)), `${eti(D.p_12)} a ${eti(D.ultimo)}`),
  );
  const irr = montoSerie(reserva, 'IRR', ramoRes), bruto = montoSerie(reserva, 'BRUTO', ramoRes);
  const ced = i => bruto[i] > 0 && irr[i] != null ? irr[i] / bruto[i] : null;
  document.getElementById('cedido').textContent = conceptos.includes('IRR') && conceptos.includes('BRUTO')
    ? `% cedido (IRR / BRUTO) · ${eti(D.ultimo)}: ${ced(jUlt) == null ? 's/d' : nf(1).format(ced(jUlt) * 100) + ' %'} · ${eti(D.fin_m)}: ${ced(jFin) == null ? 's/d' : nf(1).format(ced(jFin) * 100) + ' %'}` : '';
  const rejilla = document.getElementById('rejilla-res'); rejilla.replaceChildren(); tarjetasVivas.length = 0;
  const etiqRamo = ramoRes === 'Todos' ? 'todos los ramos' : `ramo ${ramoRes}`;
  const real = vals.map((v, i) => i <= jUlt ? v : null), proy = vals.map((v, i) => i > jUlt ? v : null);
  const proyLinea = vals.map((v, i) => i >= jUlt ? v : null);
  // decimales de etiquetas y tooltip segun la magnitud de la serie (ramos chicos: menos de 1 M USD)
  const decGraf = arr => { const m = Math.max(0, ...arr.filter(v => v != null).map(Math.abs)); return m < 1 ? 3 : m < 10 ? 2 : m < 100 ? 1 : 0; };
  const decV = decGraf(vals);

  const c1 = tarjeta(`${reserva} ${concepto} mensual · ${etiqRamo}`, `${eti(P_M[jMen])} a ${eti(D.fin_m)}, millones de ${moneda}`,
    (cuerpo, W) => graficaColumnas(cuerpo, W, {
      categorias: P_M.slice(jMen).map(eti), tituloX: i => etiLarga(P_M[jMen + i]), fmt: v => fmt(v, decV), compartido: true, cadaX: 3,
      aria: `${reserva} ${concepto} mensual`,
      series: [{ nombre: 'Real', valores: real.slice(jMen), color: cssVar('--real') }, { nombre: 'Proyección', valores: proy.slice(jMen), color: cssVar('--proy') }],
      etiquetar: (i) => jMen + i === jUlt || jMen + i === jFin,
    }),
    () => tablaDatos(['Mes', 'Tipo', `M ${moneda}`], P_M.slice(jMen).map((p, k) => [etiLarga(p), jMen + k <= jUlt ? 'Real' : 'Proyección', fmt(vals[jMen + k], 2)])));

  const porRamoU = ramos.map(r => montoSerie(reserva, concepto, r)[jUlt]), porRamoF = ramos.map(r => montoSerie(reserva, concepto, r)[jFin]);
  const c2 = tarjeta(`${reserva} ${concepto} por ramo`, `${eti(D.ultimo)} real y ${eti(D.fin_m)} proyectado, millones de ${moneda}`,
    (cuerpo, W) => graficaBarrasH(cuerpo, W, {
      categorias: ramos, fmt: v => fmt(v, 1), aria: `${reserva} ${concepto} por ramo`,
      series: [{ nombre: `${eti(D.ultimo)} (real)`, valores: porRamoU, color: cssVar('--real') }, { nombre: `${eti(D.fin_m)} (proyección)`, valores: porRamoF, color: cssVar('--proy') }],
      etiquetar: (i, j) => ramoRes !== 'Todos' && ramos[i] === ramoRes,
      leyendaItems: [{ nombre: `${eti(D.ultimo)} (real)`, color: cssVar('--real'), caja: true }, { nombre: `${eti(D.fin_m)} (proyección)`, color: cssVar('--proy'), caja: true }],
    }),
    () => tablaDatos(['Ramo', `${eti(D.ultimo)} real`, `${eti(D.fin_m)} proyección`, 'Variación'], ramos.map((r, i) => [r, fmt(porRamoU[i], 2), fmt(porRamoF[i], 2), fmtPct(varPct(porRamoU[i], porRamoF[i]))])));

  const nombreTendR = nombreModelo(`${reserva}|${concepto}|${ramoRes}`);
  // la curva del modelo solo existe para las series que se modelan directo: BEL (RRC, SONR) o BRUTO (RFV) por ramo, en USD
  const tendR = (moneda === 'USD' && ramoRes !== 'Todos') ? rectaModelo(`${reserva}|${concepto}|${ramoRes}`, real, P_M, 1e-6) : null;
  const c3 = tarjeta(`${reserva} ${concepto} histórico y proyección · ${etiqRamo}`, `${eti(P_M[0])} a ${eti(D.fin_m)}, millones de ${moneda}`,
    (cuerpo, W) => graficaLineas(cuerpo, W, {
      labels: P_M.map(eti), tituloX: i => etiLarga(P_M[i]), fmt: v => fmt(v, decV), cadaX: 6, etiquetasFin: true, incluirCero: true,
      aria: `${reserva} ${concepto} histórico y proyección`,
      series: [...(tendR ? [{ nombre: nombreTendR, valores: tendR, color: cssVar('--deemph-ink'), fino: true }] : []),
        { nombre: 'Real', valores: real, color: cssVar('--real'), etiquetaFin: true }, { nombre: 'Proyección', valores: proyLinea, color: cssVar('--proy'), dash: true, etiquetaFin: true }],
    }),
    () => { const dec = Math.max(...vals.filter(v => v != null)) < 10 ? 4 : 2; return tablaDatos(['Mes', 'Tipo', `M ${moneda}`, ...(tendR ? [nombreTendR] : [])], P_M.map((p, i) => [etiLarga(p), i <= jUlt ? 'Real' : 'Proyección', fmt(vals[i], dec), ...(tendR ? [fmt(tendR[i], dec)] : [])])); });

  const porConcU = conceptos.map(c => montoSerie(reserva, c, ramoRes)[jUlt]), porConcF = conceptos.map(c => montoSerie(reserva, c, ramoRes)[jFin]);
  const c4 = tarjeta(`${reserva} por concepto · ${etiqRamo}`, `${eti(D.ultimo)} real y ${eti(D.fin_m)} proyectado, millones de ${moneda}`,
    (cuerpo, W) => graficaColumnas(cuerpo, W, {
      categorias: conceptos, fmt: v => fmt(v, decGraf([...porConcU, ...porConcF])), cadaX: 1, aria: `${reserva} por concepto`,
      series: [{ nombre: `${eti(D.ultimo)} (real)`, valores: porConcU, color: cssVar('--real') }, { nombre: `${eti(D.fin_m)} (proyección)`, valores: porConcF, color: cssVar('--proy') }],
      etiquetar: (i) => conceptos[i] === concepto,
    }),
    () => tablaDatos(['Concepto', `${eti(D.ultimo)} real`, `${eti(D.fin_m)} proyección`, 'Variación'], conceptos.map((c, i) => [c, fmt(porConcU[i], 2), fmt(porConcF[i], 2), fmtPct(varPct(porConcU[i], porConcF[i]))])));
  for (const c of [c1, c2, c3, c4]) { rejilla.append(c); tarjetasVivas.push(c); }
  redibujarTodo();
}

// ---------------------------------------------------------------- vista: analisis
function pintarAnalisis() {
  const v = document.getElementById('vista-analisis'); v.replaceChildren();
  v.append(el('h2', { text: `1. Índices por ramo: real ${eti(D.ultimo)} contra proyección ${eti(D.fin)} (en las series con patrón la variación incluye el mes del año)` }));
  const enc1 = ['Ramo']; for (const s of D.indices) enc1.push(`${s} real`, `${s} proy.`, 'Var.');
  const t1 = el('table', { class: 'datos' }); t1.append(el('thead', {}, el('tr', {}, ...enc1.map(h => el('th', { scope: 'col', text: h })))));
  const tb1 = el('tbody');
  for (const r of D.ramos_i) {
    const tr = el('tr', {}, el('th', { scope: 'row', text: r }));
    for (const s of D.indices) {
      const a = valorInd(r, s, iUlt), b = valorInd(r, s, iFin), va = varPct(a, b);
      const cv = el('td', { text: fmtPct(va) }); const bg = colorDivergente(va, 0.3); if (bg) cv.style.background = bg;
      tr.append(el('td', { text: fmt(a, 4) }), el('td', { text: fmt(b, 4) }), cv);
    }
    tb1.append(tr);
  }
  t1.append(tb1); v.append(el('div', { class: 'tabla-envoltura', style: 'max-height:none' }, t1));

  v.append(el('h2', { text: '2. Reservas: totales por concepto (millones de USD)' }));
  const cortes = [...new Set([Math.floor(D.p_12 / 100) * 100 - 100 + 12, Math.floor(D.p_12 / 100) * 100 + 12, D.ultimo, D.p_dic, D.fin_m])].filter(p => idx(P_M, p) >= 0);
  const enc2 = ['Reserva', 'Concepto', ...cortes.map(p => `${eti(p)} ${p <= D.ultimo ? '(real)' : '(proy.)'}`), 'Crec. real 12 m', 'Crec. anual proy.'];
  const t2 = el('table', { class: 'datos' }); t2.append(el('thead', {}, el('tr', {}, ...enc2.map(h => el('th', { scope: 'col', text: h })))));
  const tb2 = el('tbody'); const monedaGuardada = E.moneda; E.moneda = 'USD';
  for (const res of Object.keys(D.conceptos)) for (const c of D.conceptos[res]) {
    const s = montoSerie(res, c, 'Todos'); const u = s[jUlt], f = s[jFin], v12 = s[j12];
    const tr = el('tr', {}, el('th', { scope: 'row', text: res }), el('td', { class: 'txt', text: c }));
    for (const p of cortes) tr.append(el('td', { text: fmt(s[idx(P_M, p)], 1) }));
    const g12 = varPct(v12, u), gp = u > 0 && f > 0 ? Math.pow(f / u, 12 / (jFin - jUlt)) - 1 : null;
    for (const g of [g12, gp]) { const cv = el('td', { text: fmtPct(g) }); const bg = colorDivergente(g, 0.5); if (bg) cv.style.background = bg; tr.append(cv); }
    tb2.append(tr);
  }
  E.moneda = monedaGuardada;
  t2.append(tb2); v.append(el('div', { class: 'tabla-envoltura', style: 'max-height:none' }, t2));

  v.append(el('h2', { text: '3. Modelo por tipo de serie y error del backtest' }));
  v.append(el('p', { text: 'Suavizamiento exponencial (familia Holt-Winters). Error % = suma de errores absolutos / suma de valores reales, re-proyectando desde 16, 12 y 8 meses antes del final; mediana entre las series de cada tipo. Fianzas tiene 20 meses de historia y solo alcanza un corte.' }));
  const nombres = { nivel: 'Montos', indice: 'Índices', razon: 'Razones', lag: 'LAGs' }, libros = { DANOS: 'Daños', FIANZAS: 'Fianzas', HPARAM: 'HParametros' };
  const t3 = tablaDatos(['Tipo de serie', 'Modelo', 'Series con backtest', 'Cortes', 'Error % modelo', 'Error % último valor', 'Series en que el modelo mejora al último valor'],
    D.metodo.map(d => [`${nombres[d.tipo] || d.tipo} · ${libros[d.libro] || d.libro}`, d.modelo || '', d.series ?? '', d.cortes ?? '', fmt(d.err_modelo, 1), fmt(d.err_ultimo, 1), d.mejora == null ? 's/d' : nf(0).format(Math.round(d.mejora * 100)) + ' %']));
  for (const tr of t3.querySelectorAll('tbody tr')) tr.children[1].classList.add('txt');
  v.append(el('div', { class: 'tabla-envoltura', style: 'max-height:none' }, t3));
  v.append(el('p', { text: 'Nota: un LAG en 0 después de que el patrón acumulado ya superó 50 % es un marcador de dato faltante en la BD y se muestra como s/d, igual que en la proyección.' }));
}

// ---------------------------------------------------------------- filtros y pestañas
function iniciar() {
  document.getElementById('subtitulo').textContent = `Real ${eti(P_I[0])} a ${eti(D.ultimo)} y proyección ${eti(P_I[iUlt + 1])} a ${eti(D.fin)} · HParametros_2026 y BD_Montos (RRC, SONR, RFV)`;
  const fRamo = document.getElementById('f-ramo'), fSerie = document.getElementById('f-serie'), fPer = document.getElementById('f-periodo');
  llenarSelect(fRamo, D.ramos_i.map(r => ({ v: r, t: r })), E.ramo);
  llenarSelect(fSerie, [{ grupo: 'Índices', opciones: D.indices.map(s => ({ v: s, t: s })) }, { grupo: 'LAGs', opciones: D.lags.map(s => ({ v: s, t: s })) }], E.serie);
  llenarSelect(fPer, [...P_I].reverse().map(p => ({ v: p, t: `${etiLarga(p)}${p > D.ultimo ? ' (proy.)' : ''}` })), E.periodo);
  fRamo.addEventListener('change', () => { E.ramo = fRamo.value; pintarIndices(); });
  fSerie.addEventListener('change', () => { E.serie = fSerie.value; pintarIndices(); });
  fPer.addEventListener('change', () => { E.periodo = +fPer.value; pintarIndices(); });

  const fRes = document.getElementById('f-reserva'), fConc = document.getElementById('f-concepto'), fRR = document.getElementById('f-ramo-res'), fMon = document.getElementById('f-moneda');
  const llenarDependientes = () => {
    const cs = D.conceptos[E.reserva] || []; if (!cs.includes(E.concepto)) E.concepto = cs[0];
    llenarSelect(fConc, cs.map(c => ({ v: c, t: c })), E.concepto);
    const rs = ['Todos', ...(D.ramos_m[E.reserva] || [])]; if (!rs.includes(E.ramoRes)) E.ramoRes = 'Todos';
    llenarSelect(fRR, rs.map(r => ({ v: r, t: r })), E.ramoRes);
  };
  llenarSelect(fRes, Object.keys(D.conceptos).map(r => ({ v: r, t: r })), E.reserva); llenarDependientes(); fMon.value = E.moneda;
  fRes.addEventListener('change', () => { E.reserva = fRes.value; llenarDependientes(); pintarReservas(); });
  fConc.addEventListener('change', () => { E.concepto = fConc.value; pintarReservas(); });
  fRR.addEventListener('change', () => { E.ramoRes = fRR.value; pintarReservas(); });
  fMon.addEventListener('change', () => { E.moneda = fMon.value; pintarReservas(); });

  const tabs = document.querySelectorAll('.tabs [role=tab]');
  const mostrarVista = nombre => {
    E.vista = nombre;
    for (const t of tabs) { const on = t.id === 'tab-' + nombre; t.setAttribute('aria-selected', String(on)); document.getElementById('vista-' + nombre.replace('tab-', '')).hidden = false; }
    for (const v of document.querySelectorAll('.vista')) v.hidden = v.id !== 'vista-' + nombre;
    if (nombre === 'indices') pintarIndices(); else if (nombre === 'reservas') pintarReservas(); else pintarAnalisis();
    try { localStorage.setItem('vista', nombre); } catch (e) { /* sin almacenamiento */ }
  };
  for (const t of tabs) t.addEventListener('click', () => mostrarVista(t.id.replace('tab-', '')));
  const lista = [...tabs];
  for (const t of lista) t.addEventListener('keydown', ev => {
    if (ev.key !== 'ArrowRight' && ev.key !== 'ArrowLeft') return;
    ev.preventDefault();
    const k = (lista.indexOf(t) + (ev.key === 'ArrowRight' ? 1 : lista.length - 1)) % lista.length;
    lista[k].focus(); mostrarVista(lista[k].id.replace('tab-', ''));
  });
  let vistaInicial = 'indices';
  try { vistaInicial = localStorage.getItem('vista') || 'indices'; } catch (e) { /* sin almacenamiento */ }
  if (!['indices', 'reservas', 'analisis'].includes(vistaInicial)) vistaInicial = 'indices';
  aplicarTema();
  mostrarVista(vistaInicial);
}
iniciar();
</script>
</body>
</html>
"""


if __name__ == "__main__":
    generar()
