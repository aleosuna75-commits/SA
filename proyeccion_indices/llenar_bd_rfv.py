# -*- coding: utf-8 -*-
"""
Llenado de BD_ RFV (ramos de Fianzas) a partir de Res_Rvas_2025 y Res_Rvas_2026.

Replica exactamente el criterio con el que se lleno manualmente la hoja
BD_Montos_RRC_SONR (ramos de Danos) del archivo "BD_ BEL - IRR - MR":

  * Fuente: hoja BacktestingFIANZAS de cada Res_Rvas, columnas del escenario
    "SAP" (fila 1 = "SAP", REAL / "12REALES - BEL REAL"), mismo periodo.
  * Se copia de forma transpuesta: en el Backtesting los ramos estan en filas y
    los periodos en columnas; en la BD los periodos van en filas y los ramos en
    columnas (RAM_130 ... RAM_170).
  * Mapeo de conceptos (bloques del Backtesting -> CONCEPTO de la BD):
        NETO            -> RFV NETO
        BRUTO           -> RFV BRUTO
        IRR             -> RFV IRR
        CONTIGENCIA SDO -> RCONT
  * Moneda: la BD esta en USD (igual que la BD de Danos). Si el Backtesting
    dice "Montos en MXN" (Res_Rvas_2025) cada monto se divide entre el TC del
    periodo tomado de la columna TC de la propia BD (mismo tratamiento que se
    le dio a SONR 2025 en la hoja de referencia). Si dice "Montos en USD"
    (Res_Rvas_2026) se copia tal cual.
  * Periodos sin cifra real en SAP (p. ej. 202609-202612) se dejan en 0, igual
    que en la hoja de referencia.
  * RCONT: la contingencia SAP de 2026 (BASE!L) ya viene en USD y se copia tal
    cual. En Res_Rvas_2025 el bloque SAP de contingencia de BacktestingFIANZAS
    esta vacio; el saldo total real de la reserva de contingencia (USD) esta en
    la hoja BacktestingCATAS_ ("RESERVA DE CONTINGENCIA" / "SALDO TOTAL", que
    viene del archivo FIA; el renglon siguiente liga a BASE!L14 con diferencias
    menores a 0.1%) y se reparte por ramo con la mezcla observada en 2026. Esas
    celdas llevan un comentario con su procedencia (ver RCONT_COMPLETAR_CON_TOTAL_SAP).
  * 202512 se convierte con el TC de la BD (18.008). En Res_Rvas_2025 la fila 7
    de BacktestingFIANZAS trae 18.08 capturado a mano (error de dedo: RRC, SONR
    y la columna TC de la BD dicen 18.008); el script lo reporta.

Uso (VSCode): abrir este archivo y ejecutar (F5 / "Run Python File").
Los paquetes que falten se instalan solos en el interprete activo.
"""
from __future__ import annotations

import importlib
import os
import subprocess
import sys


def _asegurar_paquetes(paquetes: dict[str, str]) -> None:
    """Instala (con el mismo interprete que ejecuta el script) lo que falte."""
    faltantes = []
    for modulo, nombre_pip in paquetes.items():
        try:
            importlib.import_module(modulo)
        except ImportError:
            faltantes.append(nombre_pip)
    if faltantes:
        print(f"Instalando paquetes faltantes: {', '.join(faltantes)} ...", flush=True)
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", *faltantes])
        except subprocess.CalledProcessError as e:
            raise SystemExit(f"No se pudieron instalar {faltantes} (sin internet, proxy o permisos). Instalalos "
                             f"manualmente con: {sys.executable} -m pip install {' '.join(faltantes)}") from e
        # si pip instalo en la carpeta de usuario (Python 'para todos los usuarios' en Windows), agregarla
        import site
        usuario = site.getusersitepackages()
        if os.path.isdir(usuario) and usuario not in sys.path:
            site.addsitedir(usuario)
        importlib.invalidate_caches()


_asegurar_paquetes({"openpyxl": "openpyxl>=3.1"})

import re  # noqa: E402
from pathlib import Path  # noqa: E402

import openpyxl  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from excel_fiel import guardar_libro, verificar_escritura  # noqa: E402
from tipo_cambio import TC_FCST  # noqa: E402

# =============================================================================
# CONFIGURACION
# =============================================================================
CARPETA = Path(__file__).resolve().parent
ENTRADAS = CARPETA / "entradas"
SALIDAS = CARPETA / "salidas"

ARCHIVOS_RES_RVAS = [ENTRADAS / "Res_Rvas_2025.xlsx", ENTRADAS / "Res_Rvas_2026.xlsx"]
ARCHIVO_BD_RFV = ENTRADAS / "BD_ RFV.xlsx"          # plantilla (no se modifica)
ARCHIVO_SALIDA = SALIDAS / "BD_ RFV.xlsx"            # BD llena

HOJA_BACKTEST = "BacktestingFIANZAS"
HOJA_BD = "BD_Montos_RRC_SONR"
ESCENARIO_FUENTE = "SAP"                             # fila 1 del Backtesting

# CONCEPTO en la BD  ->  etiqueta del bloque en la columna B del Backtesting
MAPEO_CONCEPTOS = {
    "RFV NETO": "NETO",
    "RFV BRUTO": "BRUTO",
    "RFV IRR": "IRR",
    "RCONT": "CONTIGENCIA SDO",
}

# En Res_Rvas_2026 la formula SAP de CONTIGENCIA SDO (=[n]BASE!$L$17) NO se
# divide entre el TC, a diferencia de BRUTO / IRR (=[n]RRC!$I$41/AR$7). Se
# verifico que BASE!L ya viene en USD (los archivos fuente de ene-26 y feb-26
# tienen el mismo MXN y BASE!L cambia exactamente en la razon de TC), por lo
# que se copia tal cual.  True = dividir entre el TC de la fila 7.
RCONT_DIVIDIR_ENTRE_TC_EN_ARCHIVOS_USD = False

# En Res_Rvas_2025 el bloque SAP de CONTIGENCIA SDO de BacktestingFIANZAS esta
# vacio. El saldo total real de la reserva de contingencia (en USD) si existe en
# la hoja BacktestingCATAS_ ("RESERVA DE CONTINGENCIA" / "SALDO TOTAL"), sin
# desglose por ramo.
#   True  = se usa ese total y se reparte por ramo con la mezcla observada en los
#           meses que si traen desglose (2026); cada celda lleva un comentario.
#   False = se deja en 0 (copia estricta de BacktestingFIANZAS).
RCONT_COMPLETAR_CON_TOTAL_SAP = True

# Columna TC: se actualiza con el supuesto de Inversiones (tipo_cambio.py: FCST 2026 / FCST 2027) en los
# meses que trae. 202601-202608 = TC real de SAP (la plantilla traia una interpolacion en 202606-202608).
# No afecta los montos de 2026 (ya vienen en USD); 2025 se convierte con el TC de la BD (no cambia).
ACTUALIZAR_TC_CON_FCST = True
PREFIJO_HOJA_TOTAL_CONTINGENCIA = "BacktestingCATAS"

FORMATO_NUMERO = "#,##0"                             # igual que la BD de Danos
TOLERANCIA = 0.01                                    # USD, para validaciones


# =============================================================================
# LECTURA DEL BACKTESTING
# =============================================================================
def _norm(texto) -> str:
    return re.sub(r"\s+", " ", str(texto or "")).strip().upper()


def _es_periodo(valor) -> bool:
    return isinstance(valor, (int, float)) and not isinstance(valor, bool) and 190001 <= valor <= 209912 \
        and float(valor).is_integer()


def leer_backtesting(ruta: Path) -> dict:
    """Devuelve {'moneda', 'tc': {periodo: tc}, 'datos': {(bloque, periodo, ramo): valor},
    'cierre': {(bloque, ramo): valor REAL CIERRE col C}, 'periodo_cierre'}."""
    wb = openpyxl.load_workbook(ruta, data_only=True)
    if HOJA_BACKTEST not in wb.sheetnames:
        raise ValueError(f"{ruta.name}: no existe la hoja {HOJA_BACKTEST}")
    ws = wb[HOJA_BACKTEST]

    etiqueta_moneda = _norm(ws["B3"].value)
    if "MXN" in etiqueta_moneda:
        moneda = "MXN"
    elif "USD" in etiqueta_moneda:
        moneda = "USD"
    else:
        raise ValueError(f"{ruta.name}: no se pudo identificar la moneda en B3 ({ws['B3'].value!r})")

    # Columnas del escenario fuente (fila 1)
    cols_fuente = [c for c in range(1, ws.max_column + 1) if _norm(ws.cell(1, c).value) == ESCENARIO_FUENTE]
    if not cols_fuente:
        raise ValueError(f"{ruta.name}: no hay columnas '{ESCENARIO_FUENTE}' en la fila 1")

    # Bloques: fila cuyo texto en col B es una etiqueta conocida
    etiquetas = {_norm(v) for v in MAPEO_CONCEPTOS.values()}
    bloques = {}
    for r in range(1, ws.max_row + 1):
        etiqueta = _norm(ws.cell(r, 2).value)
        if etiqueta in etiquetas and etiqueta not in bloques:
            bloques[etiqueta] = r
    faltan = etiquetas - set(bloques)
    if faltan:
        raise ValueError(f"{ruta.name}: no se encontraron los bloques {sorted(faltan)}")

    tc = {}
    datos = {}
    cierre = {}
    periodo_cierre = ws.cell(bloques[_norm("NETO")], 3).value
    for etiqueta, r0 in bloques.items():
        # filas de ramos: debajo del encabezado mientras la col B sea un codigo numerico
        filas_ramo = []
        r = r0 + 1
        while isinstance(ws.cell(r, 2).value, (int, float)):
            filas_ramo.append((r, int(ws.cell(r, 2).value)))
            r += 1
        if not filas_ramo:
            raise ValueError(f"{ruta.name}: el bloque {etiqueta} no tiene ramos")
        for c in cols_fuente:
            per = ws.cell(r0, c).value
            if not _es_periodo(per):
                continue
            per = int(per)
            tc[per] = ws.cell(7, c).value
            for rr, ramo in filas_ramo:
                v = ws.cell(rr, c).value
                datos[(etiqueta, per, ramo)] = float(v) if isinstance(v, (int, float)) else 0.0
        for rr, ramo in filas_ramo:
            v = ws.cell(rr, 3).value
            cierre[(etiqueta, ramo)] = float(v) if isinstance(v, (int, float)) else 0.0
    periodos = sorted({k[1] for k in datos})
    return {"moneda": moneda, "tc": tc, "datos": datos, "cierre": cierre,
            "periodo_cierre": periodo_cierre, "periodos": periodos, "archivo": ruta.name,
            "contingencia_total": leer_total_contingencia(wb)}


def leer_total_contingencia(wb) -> dict:
    """Saldo total SAP de la reserva de contingencia ({periodo: USD}) de la hoja BacktestingCATAS*:
    renglon 'SALDO TOTAL' bajo el titulo 'RESERVA DE CONTINGENCIA'. Vacio si no existe."""
    for nombre in wb.sheetnames:
        if not nombre.upper().startswith(PREFIJO_HOJA_TOTAL_CONTINGENCIA.upper()):
            continue
        ws = wb[nombre]
        etiqueta_moneda = " ".join(_norm(ws.cell(r, 2).value) for r in range(1, 5))
        if "USD" not in etiqueta_moneda:
            continue
        titulo = None
        for r in range(1, ws.max_row + 1):
            texto = _norm(ws.cell(r, 2).value)
            if texto.startswith("RESERVA DE CONTINGENCIA"):
                titulo = r
            elif titulo and texto == "SALDO TOTAL" and r - titulo <= 3:
                fila_per = r - 1
                total = {}
                for c in range(3, ws.max_column + 1):
                    per, v = ws.cell(fila_per, c).value, ws.cell(r, c).value
                    # solo la primera tabla: mas a la derecha hay otras con los mismos periodos (variaciones)
                    if _es_periodo(per) and int(per) not in total and isinstance(v, (int, float)):
                        total[int(per)] = float(v)
                if total:
                    return {"hoja": nombre, "fila": r, "valores": total}
    return {}


# =============================================================================
# LLENADO DE LA BD
# =============================================================================
def llenar():
    if not ARCHIVO_BD_RFV.exists():
        raise FileNotFoundError(f"No existe la plantilla {ARCHIVO_BD_RFV}")
    fuentes = [leer_backtesting(p) for p in ARCHIVOS_RES_RVAS]
    for f in fuentes:
        print(f"[fuente] {f['archivo']}: moneda {f['moneda']}, periodos SAP {f['periodos'][0]}-{f['periodos'][-1]}")

    SALIDAS.mkdir(parents=True, exist_ok=True)
    verificar_escritura([ARCHIVO_SALIDA])
    wb = openpyxl.load_workbook(ARCHIVO_BD_RFV)          # la plantilla no se modifica
    ws = wb[HOJA_BD]

    encabezados = {_norm(ws.cell(3, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value}
    col_concepto, col_periodo, col_tc = encabezados["CONCEPTO"], encabezados["PERIODO"], encabezados["TC"]
    cols_ramo = {int(h.split("_")[1]): c for h, c in encabezados.items() if h.startswith("RAM_")}

    resumen = []
    advertencias = []
    if ACTUALIZAR_TC_CON_FCST:
        cambios = 0
        for r in range(4, ws.max_row + 1):
            periodo = ws.cell(r, col_periodo).value
            if _es_periodo(periodo) and int(periodo) in TC_FCST and ws.cell(r, col_tc).value != TC_FCST[int(periodo)]:
                ws.cell(r, col_tc).value = TC_FCST[int(periodo)]
                cambios += 1
        if cambios:
            advertencias.append(f"TC actualizado con el supuesto de Inversiones (tipo_cambio.py) en {cambios} renglones")
    for r in range(4, ws.max_row + 1):
        concepto = _norm(ws.cell(r, col_concepto).value)
        periodo = ws.cell(r, col_periodo).value
        if concepto not in MAPEO_CONCEPTOS or not _es_periodo(periodo):
            continue
        periodo = int(periodo)
        bloque = _norm(MAPEO_CONCEPTOS[concepto])
        # archivo del mismo anio que el periodo
        fuente = next((f for f in fuentes if periodo in f["periodos"] and str(periodo)[:4] in f["archivo"]), None) \
            or next((f for f in fuentes if periodo in f["periodos"]), None)
        if fuente is None:
            raise ValueError(f"Periodo {periodo} no existe en las columnas {ESCENARIO_FUENTE} de ningun Res_Rvas")
        tc_bd = ws.cell(r, col_tc).value
        tc_fuente = fuente["tc"].get(periodo)

        for ramo, c in cols_ramo.items():
            v = fuente["datos"].get((bloque, periodo, ramo), 0.0)
            if fuente["moneda"] == "MXN":
                if not tc_bd:
                    raise ValueError(f"Fila {r}: falta TC en la BD para convertir MXN a USD")
                v = v / tc_bd
            elif bloque == "CONTIGENCIA SDO" and RCONT_DIVIDIR_ENTRE_TC_EN_ARCHIVOS_USD:
                v = v / tc_fuente
            celda = ws.cell(r, c)
            celda.value = v
            celda.number_format = FORMATO_NUMERO
        if fuente["moneda"] == "MXN" and tc_fuente and abs(tc_fuente - tc_bd) > 1e-9:
            advertencias.append(
                f"{concepto} {periodo}: TC de la BD = {tc_bd} (usado) vs TC en {fuente['archivo']} "
                f"fila 7 = {tc_fuente}")
        resumen.append((concepto, periodo, r, fuente["archivo"], fuente["moneda"]))

    # ---------------------------------- RCONT sin desglose SAP: total SAP x mezcla
    if RCONT_COMPLETAR_CON_TOTAL_SAP:
        filas_rcont = {int(ws.cell(r, col_periodo).value): r for r in range(4, ws.max_row + 1)
                       if _norm(ws.cell(r, col_concepto).value) == "RCONT"
                       and _es_periodo(ws.cell(r, col_periodo).value)}
        con_desglose = {p: {ramo: ws.cell(r, c).value or 0.0 for ramo, c in cols_ramo.items()}
                        for p, r in filas_rcont.items() if any(ws.cell(r, c).value for c in cols_ramo.values())}
        suma_total = sum(sum(v.values()) for v in con_desglose.values())
        mezcla = {ramo: sum(v[ramo] for v in con_desglose.values()) / suma_total for ramo in cols_ramo} \
            if suma_total else {}
        for p, r in sorted(filas_rcont.items()):
            if p in con_desglose or not mezcla:
                continue
            # mismo archivo que surte el resto del periodo (anio del periodo); la hoja trae ademas
            # columnas de presupuesto (p.ej. 202612 en el libro 2025) que no se deben usar
            fuente = next((f for f in fuentes if p in f["periodos"] and str(p)[:4] in f["archivo"]
                           and f["contingencia_total"] and p in f["contingencia_total"]["valores"]), None)
            if fuente is None:
                continue
            info = fuente["contingencia_total"]
            total = info["valores"][p]
            for ramo, c in cols_ramo.items():
                celda = ws.cell(r, c)
                celda.value = total * mezcla[ramo]
                celda.number_format = FORMATO_NUMERO
                celda.comment = Comment(
                    f"Imputado: saldo total SAP de contingencia {p} = {total:,.2f} USD "
                    f"({fuente['archivo']} > {info['hoja']} fila {info['fila']}) x mezcla por ramo de "
                    f"{min(con_desglose)}-{max(con_desglose)} ({mezcla[ramo]:.2%}). BacktestingFIANZAS "
                    f"no trae desglose SAP para este mes.", "llenar_bd_rfv.py")
                celda.comment.width, celda.comment.height = 320, 110
            advertencias.append(f"RCONT {p}: total SAP {total:,.0f} USD de {info['hoja']} repartido por ramo "
                                f"con la mezcla {', '.join(f'{k}:{v:.1%}' for k, v in mezcla.items())}")

    # ------------------------------------------------------------ validaciones
    valores = {}
    for r in range(4, ws.max_row + 1):
        concepto = _norm(ws.cell(r, col_concepto).value)
        periodo = ws.cell(r, col_periodo).value
        if concepto in MAPEO_CONCEPTOS and _es_periodo(periodo):
            for ramo, c in cols_ramo.items():
                valores[(concepto, int(periodo), ramo)] = ws.cell(r, c).value or 0.0
    errores_identidad = [
        k for k in valores if k[0] == "RFV NETO"
        and abs(valores[k] - (valores[("RFV BRUTO",) + k[1:]] - valores[("RFV IRR",) + k[1:]])) > TOLERANCIA
    ]
    if errores_identidad:
        raise AssertionError(f"No se cumple RFV NETO = RFV BRUTO - RFV IRR en {errores_identidad[:5]}")

    # Cruce del cierre (col C 'REAL CIERRE' del archivo del anio siguiente)
    inv = {_norm(v): k for k, v in MAPEO_CONCEPTOS.items()}
    for f in fuentes:
        per = f["periodo_cierre"]
        if not _es_periodo(per):
            continue
        for (bloque, ramo), v_cierre in f["cierre"].items():
            if bloque == "CONTIGENCIA SDO":
                continue      # la columna 'REAL CIERRE' de contingencia viene de un archivo de presupuesto
            clave = (inv[bloque], int(per), ramo)
            if clave in valores and abs(valores[clave] - v_cierre) > 1.0:
                advertencias.append(
                    f"Cierre {per} {inv[bloque]} RAM_{ramo}: BD = {valores[clave]:,.0f} vs 'REAL CIERRE' "
                    f"de {f['archivo']} = {v_cierre:,.0f} (dif {valores[clave] - v_cierre:,.0f})")

    guardar_libro(wb, ARCHIVO_SALIDA, original=ARCHIVO_BD_RFV)   # atomico: si algo fallo antes, no hay salida

    print("\nResumen (suma de ramos, USD; valores finales escritos):")
    for concepto, periodo, r, archivo, moneda in resumen:
        total = sum(ws.cell(r, c).value or 0.0 for c in cols_ramo.values())
        origen = f"{archivo} ({moneda})"
        if ws.cell(r, next(iter(cols_ramo.values()))).comment is not None:
            origen = "BacktestingCATAS_ (USD, total repartido por ramo)"
        print(f"  {concepto:<10} {periodo}  {origen:<46} {total:>18,.0f}")
    if advertencias:
        print("\nADVERTENCIAS / CRUCES A REVISAR:")
        for a in advertencias:
            print("  -", a)
    print(f"\nListo: {ARCHIVO_SALIDA}")
    return ARCHIVO_SALIDA


if __name__ == "__main__":
    llenar()
