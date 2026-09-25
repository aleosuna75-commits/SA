# -*- coding: utf-8 -*-
"""
DASHBOARDS EN EXCEL: indices (HParametros) y reservas (RRC / SONR / RFV), real y proyectado.

Genera salidas/Dashboard_Indices_Reservas.xlsx a partir de las salidas de proyeccion_reservas.py:
  * "Dashboard Índices"  : selectores de ramo, indice/LAG y periodo; indicadores; evolucion mensual real vs
                           proyeccion con banda al 80%; comparativo por ramo; patron de desarrollo (LAGs);
                           resumen de los 4 indices del ramo.
  * "Dashboard Reservas" : selectores de reserva, concepto, ramo y moneda (USD/MXN con el TC de la BD);
                           indicadores; mensual 2025-2027; historico 2022-2027; por ramo; por concepto.
  * "Análisis"           : tablas con mapas de calor (indices por ramo y totales de reservas) y metodo.
  * "BD_Indices", "BD_Reservas": bases de datos (tablas de Excel con filtros).
Los dashboards son interactivos sin macros: las listas desplegables alimentan formulas (SUMIFS) y las
graficas se recalculan al cambiar la seleccion.

Uso: se ejecuta solo al final de proyeccion_reservas.py (GENERAR_DASHBOARD) o directamente (F5 en VSCode).
"""
from __future__ import annotations

import importlib
import math
import os
import re
import subprocess
import sys


def _asegurar_paquetes(paquetes: dict[str, str]) -> None:
    faltantes = []
    for modulo, nombre_pip in paquetes.items():
        try:
            importlib.import_module(modulo)
        except ImportError:
            faltantes.append(nombre_pip)
    if faltantes:
        print(f"Instalando paquetes faltantes: {', '.join(faltantes)} ...", flush=True)
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", *faltantes])
        import site
        usuario = site.getusersitepackages()
        if os.path.isdir(usuario) and usuario not in sys.path:
            site.addsitedir(usuario)
        importlib.invalidate_caches()


_asegurar_paquetes({"openpyxl": "openpyxl>=3.1"})

from pathlib import Path  # noqa: E402

import openpyxl  # noqa: E402
from openpyxl.chart import BarChart, LineChart, Reference, Series  # noqa: E402
from openpyxl.chart.series import SeriesLabel  # noqa: E402
from openpyxl.chart.shapes import GraphicalProperties  # noqa: E402
from openpyxl.chart.text import RichText  # noqa: E402
from openpyxl.drawing.line import LineProperties  # noqa: E402
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor  # noqa: E402
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties  # noqa: E402
from openpyxl.formatting.rule import ColorScaleRule  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: E402
from openpyxl.utils.cell import coordinate_from_string  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.hyperlink import Hyperlink  # noqa: E402
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from excel_fiel import guardar_libro  # noqa: E402

# =============================================================================
# CONFIGURACION
# =============================================================================
CARPETA = Path(__file__).resolve().parent
SALIDAS = CARPETA / "salidas"
ARCHIVO_DANOS = SALIDAS / "BD_ BEL - IRR - MR_Proyeccion.xlsx"
ARCHIVO_RFV = SALIDAS / "BD_ RFV_Proyeccion.xlsx"
ARCHIVO_DIAGNOSTICO = SALIDAS / "Diagnostico_Proyeccion.xlsx"
SALIDA_DASHBOARD = SALIDAS / "Dashboard_Indices_Reservas.xlsx"

HOJA_PARAMETROS = "HParametros_2026"
HOJA_MONTOS = "BD_Montos_RRC_SONR"
INDICES = ["Ind Sin RRC", "Ind sin RRC 99.5%", "Ind Sin SONR Media", "Ind Sin SONR 99.5%"]
LAGS = [f"LAG {i}" for i in range(1, 11)]
PRIMER_PERIODO_INDICES = 202101
PRIMER_PERIODO_MONTOS = 202201
PRIMER_PERIODO_MENSUAL = 202501          # grafica de columnas mensual de reservas

# Colores (paleta validada: azul = real, naranja = proyeccion; gris = contexto)
AZUL, NARANJA, GRIS_BANDA, GRIS_BARRA = "2A78D6", "EB6834", "B5B3AC", "C3C8D4"
AZULES_ORDINALES = ["86B6EF", "3987E5", "184F95"]
FONDO, PANEL, BANDA_KPI, NAV_ACTIVO = "EEF1F6", "FFFFFF", "C9D3EE", "DDE3F3"
TEXTO, TEXTO_2, BORDE = "1F2937", "52514E", "D5DAE5"
FUENTE = "Aptos Narrow"
ANCHO_REJILLA = 4.3                      # ancho de las 30 columnas de la rejilla de los dashboards (D:AG)
PIE_PAGINA = "Información de uso interno de Grupo Peña Verde"
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


# =============================================================================
# UTILERIAS
# =============================================================================
def etiqueta(p: int) -> str:
    return f"{MESES[p % 100 - 1]}-{str(p // 100)[2:]}"


def mover(p: int, meses: int) -> int:
    i = (p // 100) * 12 + p % 100 - 1 + meses
    return (i // 12) * 100 + i % 12 + 1


def rango(p1: int, p2: int) -> list[int]:
    out, p = [], p1
    while p <= p2:
        out.append(p)
        p = mover(p, 1)
    return out


def numero(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    if isinstance(v, str):
        t = v.replace("\xa0", " ").strip().replace(",", "")
        try:
            return float(t)
        except ValueError:
            return None
    return None


def norm(t) -> str:
    return re.sub(r"\s+", " ", str(t or "").replace("\xa0", " ")).strip().upper()


# =============================================================================
# LECTURA DE LAS SALIDAS DE LA PROYECCION
# =============================================================================
def leer_indices():
    wb = openpyxl.load_workbook(ARCHIVO_DANOS, read_only=True, data_only=True)
    ws = wb[HOJA_PARAMETROS]
    filas = ws.iter_rows(min_row=3, values_only=True)
    enc = next(filas)
    col = {}
    for i, h in enumerate(enc):
        if h is not None and str(h).strip() not in col:
            col[str(h).strip()] = i
    registros, activos = [], []
    for f in filas:
        if not f or len(f) <= col["Ramo"]:
            continue
        tipo, fecha, ramo = f[col["Tipo de Indice"]], f[col["Fecha"]], f[col["Ramo"]]
        if tipo not in ("Real", "Proyección") or not isinstance(fecha, (int, float)) or ramo is None:
            continue
        ramo = str(ramo).strip()
        fecha = int(fecha)
        if tipo == "Proyección" and ramo not in activos:
            activos.append(ramo)
        max_previo = 0.0
        for serie in INDICES + LAGS:
            v = numero(f[col[serie]]) if col[serie] < len(f) else None
            if serie in LAGS:          # 0 despues de que el patron acumulado supero 50% = marcador de faltante
                if v == 0 and max_previo > 0.5:
                    v = None
                if v is not None:
                    max_previo = max(max_previo, v)
            if v is not None:
                registros.append((tipo, ramo, serie, fecha, v))
    wb.close()
    ultimo = max(r[3] for r in registros if r[0] == "Real")
    registros = [r for r in registros if r[1] in activos and r[3] >= PRIMER_PERIODO_INDICES]
    return registros, activos, ultimo


def leer_intervalos() -> dict:
    """{(serie, ramo, periodo): (LI, LS)} de los indices (hoja Pronosticos_Drivers del diagnostico)."""
    if not ARCHIVO_DIAGNOSTICO.exists():
        return {}
    wb = openpyxl.load_workbook(ARCHIVO_DIAGNOSTICO, read_only=True, data_only=True)
    ws = wb["Pronosticos_Drivers"]
    filas = ws.iter_rows(values_only=True)
    enc = [str(h) for h in next(filas)]
    i = {h: k for k, h in enumerate(enc)}
    li_col = next(k for h, k in i.items() if h.startswith("LI"))
    ls_col = next(k for h, k in i.items() if h.startswith("LS"))
    out = {}
    for f in filas:
        if f[i["Libro"]] == "HPARAM":
            out[(f[i["Serie"]], str(f[i["Ramo"]]), int(f[i["Periodo"]]))] = (numero(f[li_col]), numero(f[ls_col]))
    wb.close()
    return out


def leer_montos(ultimo: int):
    registros, tc = [], {}
    ramos = {}
    for libro, ruta in (("DANOS", ARCHIVO_DANOS), ("FIANZAS", ARCHIVO_RFV)):
        wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
        ws = wb[HOJA_MONTOS]
        filas = ws.iter_rows(min_row=3, values_only=True)
        enc = [str(h).strip() if h is not None else None for h in next(filas)]
        c_conc, c_per, c_tc = enc.index("CONCEPTO"), enc.index("PERIODO"), enc.index("TC")
        c_ramos = {h.split("_", 1)[1]: k for k, h in enumerate(enc) if h and h.startswith("RAM_")}
        for f in filas:
            if not f or f[c_conc] is None or not isinstance(f[c_per], (int, float)):
                continue
            per = int(f[c_per])
            if per < PRIMER_PERIODO_MONTOS:
                continue
            concepto = norm(f[c_conc])
            reserva, conc = ("RFV", "RCONT") if concepto == "RCONT" else concepto.split(" ", 1)
            ramos.setdefault(reserva, list(c_ramos))
            t = numero(f[c_tc])
            if t:
                tc.setdefault(per, t)
            tipo = "Real" if per <= ultimo else "Proyección"
            for ramo, k in c_ramos.items():
                v = numero(f[k])
                registros.append((libro, reserva, conc, ramo, per, tipo, v if v is not None else 0.0))
        wb.close()
    return registros, tc, ramos


def leer_metodo() -> list:
    if not ARCHIVO_DIAGNOSTICO.exists():
        return []
    wb = openpyxl.load_workbook(ARCHIVO_DIAGNOSTICO, read_only=True, data_only=True)
    if "Validacion_Metodo" not in wb.sheetnames:
        return []
    filas = wb["Validacion_Metodo"].iter_rows(values_only=True)
    enc = [str(h) for h in next(filas)]
    out = []
    for f in filas:
        d = dict(zip(enc, f))
        if d.get("Seleccionado"):
            out.append(d)
    wb.close()
    return out


# =============================================================================
# ESTILO
# =============================================================================
def relleno(color):
    return PatternFill("solid", fgColor=color)


def fuente(tam=10, negrita=False, color=TEXTO, cursiva=False):
    return Font(name=FUENTE, size=tam, bold=negrita, color=color, italic=cursiva)


LADO = Side(style="thin", color=BORDE)
BORDE_CAJA = Border(left=LADO, right=LADO, top=LADO, bottom=LADO)


def pintar(ws, rango_celdas, color):
    for fila in ws[rango_celdas]:
        for c in fila:
            c.fill = relleno(color)


def caja(ws, rango_celdas, color=PANEL):
    """Panel con relleno y contorno fino."""
    filas = ws[rango_celdas]
    for i, fila in enumerate(filas):
        for j, c in enumerate(fila):
            c.fill = relleno(color)
            c.border = Border(left=LADO if j == 0 else None, right=LADO if j == len(fila) - 1 else None,
                              top=LADO if i == 0 else None, bottom=LADO if i == len(filas) - 1 else None)


def preparar_hoja(ws, ancho_nav=27, columnas=30, ancho=ANCHO_REJILLA, filas=48):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.column_dimensions["A"].width = 1.5
    ws.column_dimensions["B"].width = ancho_nav
    ws.column_dimensions["C"].width = 1.5
    for k in range(4, 4 + columnas):
        ws.column_dimensions[get_column_letter(k)].width = ancho
    ws.column_dimensions[get_column_letter(4 + columnas)].width = 1.5
    pintar(ws, f"A1:{get_column_letter(4 + columnas)}{filas}", FONDO)


def texto_ejes(tam=800, color=TEXTO_2):
    cp = CharacterProperties(sz=tam, solidFill=color, latin=None)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def colocar(ws, ch, desde: str, hasta: str, margen_px: int = 5):
    """Ancla la grafica a las celdas (se estira con ellas): de la esquina superior izquierda de `desde` a la
    inferior derecha de `hasta`, con un margen para no tapar el contorno del panel."""
    emu = 9525
    c1, f1 = coordinate_from_string(desde)
    c2, f2 = coordinate_from_string(hasta)
    k2 = column_index_from_string(c2)
    ancho = ws.column_dimensions[c2].width if c2 in ws.column_dimensions else 8.43
    alto = ws.row_dimensions[f2].height if f2 in ws.row_dimensions and ws.row_dimensions[f2].height else 15
    ancla = TwoCellAnchor()
    ancla._from = AnchorMarker(col=column_index_from_string(c1) - 1, colOff=margen_px * emu,
                               row=f1 - 1, rowOff=margen_px * emu)
    ancla.to = AnchorMarker(col=k2 - 1, colOff=max(int(ancho * 7 + 5) - margen_px, 0) * emu,
                            row=f2 - 1, rowOff=max(int(alto * 4 / 3) - margen_px, 0) * emu)
    ws.add_chart(ch, ancla)


def estilo_grafica(ch, formato_y="General", leyenda=True):
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.plot_area.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.number_format = formato_y
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E6E8EE", w=6350))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="C9CCD4", w=6350))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.txPr = texto_ejes()
    ch.y_axis.txPr = texto_ejes()
    ch.display_blanks = "gap"
    if leyenda:
        ch.legend.position = "b"
        ch.legend.txPr = texto_ejes(850, TEXTO)
    else:
        ch.legend = None
    return ch


def linea(serie, color, ancho_pt=2.0, punteada=False):
    serie.graphicalProperties.line.solidFill = color
    serie.graphicalProperties.line.width = int(ancho_pt * 12700)
    if punteada:
        serie.graphicalProperties.line.dashStyle = "dash"
    serie.marker.symbol = "none"
    serie.smooth = False


def barra(serie, color):
    serie.graphicalProperties.solidFill = color
    serie.graphicalProperties.line.solidFill = PANEL


# =============================================================================
# CONSTRUCCION
# =============================================================================
def generar(ruta_salida: Path = SALIDA_DASHBOARD) -> Path:
    print("Generando dashboards ...", flush=True)
    registros_i, ramos_i, ultimo = leer_indices()
    intervalos = leer_intervalos()
    registros_m, tc, ramos_m = leer_montos(ultimo)
    metodo = leer_metodo()
    proc = {d.get("Tipo"): d.get("Procedimiento") or "n/d" for d in metodo}
    fin = max(r[3] for r in registros_i)
    fin_m = max(r[4] for r in registros_m)
    p_dic_actual = (ultimo // 100) * 100 + 12
    p_12 = mover(ultimo, -12)

    wb = openpyxl.Workbook()
    wd_i = wb.active
    wd_i.title = "Dashboard Índices"
    wd_r = wb.create_sheet("Dashboard Reservas")
    wa = wb.create_sheet("Análisis")
    wbi = wb.create_sheet("BD_Indices")
    wbr = wb.create_sheet("BD_Reservas")
    wl = wb.create_sheet("Listas")
    wc = wb.create_sheet("Calc")

    def nombre(n, ref):
        wb.defined_names[n] = DefinedName(n, attr_text=ref)

    # ------------------------------------------------------------ bases de datos
    wbi.append(["Tipo", "Ramo", "Serie", "Periodo", "Valor", "LI 80%", "LS 80%"])
    for tipo, ramo, serie, per, v in registros_i:
        li, ls = intervalos.get((serie, ramo, per), (None, None)) if tipo == "Proyección" else (None, None)
        wbi.append([tipo, ramo, serie, per, v, li, ls])
    n_i = wbi.max_row
    wbr.append(["Libro", "Reserva", "Concepto", "Ramo", "Periodo", "Tipo", "Valor USD", "TC"])
    for libro, reserva, conc, ramo, per, tipo, v in registros_m:
        wbr.append([libro, reserva, conc, ramo, per, tipo, v, tc.get(per)])
    n_r = wbr.max_row
    for ws, n, ref, nombre_tabla in ((wbi, n_i, f"A1:G{n_i}", "TablaIndices"), (wbr, n_r, f"A1:H{n_r}", "TablaReservas")):
        tabla = Table(displayName=nombre_tabla, ref=ref)
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(tabla)
        ws.freeze_panes = "A2"
        for k in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(k)].width = 14
        for fila in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            for c in fila:
                c.number_format = "@"
    for c in wbi["E"][1:] + wbi["F"][1:] + wbi["G"][1:]:
        c.number_format = "0.0000"
    for c in wbr["G"][1:]:
        c.number_format = "#,##0"
    for c in wbr["D"][1:]:
        c.number_format = "@"
    for n_, col in (("I_Tipo", "A"), ("I_Ramo", "B"), ("I_Serie", "C"), ("I_Per", "D"), ("I_Val", "E"),
                    ("I_LI", "F"), ("I_LS", "G")):
        nombre(n_, f"BD_Indices!${col}$2:${col}${n_i}")
    for n_, col in (("R_Res", "B"), ("R_Conc", "C"), ("R_Ramo", "D"), ("R_Per", "E"), ("R_Tipo", "F"), ("R_Val", "G")):
        nombre(n_, f"BD_Reservas!${col}$2:${col}${n_r}")

    # ------------------------------------------------------------ listas (selectores)
    conceptos = {"RRC": ["NETO", "BRUTO", "BEL", "GTO", "IRR", "MR"], "SONR": ["NETO", "BRUTO", "BEL", "IRR", "MR"],
                 "RFV": ["NETO", "BRUTO", "IRR", "RCONT"]}
    periodos_sel = rango(PRIMER_PERIODO_INDICES, fin)
    listas = {
        "L_RamosInd": ramos_i, "L_SeriesInd": INDICES + LAGS, "L_PeriodosInd": periodos_sel,
        "L_Reservas": list(conceptos), "L_Unidades": ["USD", "MXN"],
        **{f"Conceptos_{k}": v for k, v in conceptos.items()},
        **{f"Ramos_{k}": ["Todos"] + ramos_m.get(k, []) for k in conceptos},
        **{f"RamosSin_{k}": ramos_m.get(k, []) for k in conceptos},
    }
    for j, (n_, valores) in enumerate(listas.items(), start=1):
        letra = get_column_letter(j)
        wl.cell(1, j, n_).font = Font(bold=True)
        for i, v in enumerate(valores, start=2):
            c = wl.cell(i, j, v)
            if isinstance(v, str):
                c.number_format = "@"
        nombre(n_, f"Listas!${letra}$2:${letra}${len(valores) + 1}")
    j_tc = len(listas) + 2
    wl.cell(1, j_tc, "Periodo").font = Font(bold=True)
    wl.cell(1, j_tc + 1, "TC").font = Font(bold=True)
    periodos_m = rango(PRIMER_PERIODO_MONTOS, fin_m)
    for i, p in enumerate(periodos_m, start=2):
        wl.cell(i, j_tc, p)
        wl.cell(i, j_tc + 1, tc.get(p))
    lp, lt = get_column_letter(j_tc), get_column_letter(j_tc + 1)
    nombre("TC_Per", f"Listas!${lp}$2:${lp}${len(periodos_m) + 1}")
    nombre("TC_Val", f"Listas!${lt}$2:${lt}${len(periodos_m) + 1}")
    wl.sheet_state = "hidden"

    # ------------------------------------------------------------ dashboard indices
    ws = wd_i
    preparar_hoja(ws)
    ultima_col = get_column_letter(33)
    caja(ws, "B2:B46")
    ws["B3"] = "◆ ÍNDICES Y RESERVAS"
    ws["B3"].font = fuente(13, True)
    menu = [("Dashboard Índices", "'Dashboard Índices'!A1"), ("Dashboard Reservas", "'Dashboard Reservas'!A1"),
            ("Análisis", "'Análisis'!A1"), ("Base de datos: índices", "'BD_Indices'!A1"),
            ("Base de datos: reservas", "'BD_Reservas'!A1")]

    def navegacion(hoja, activo):
        for k, (texto, destino) in enumerate(menu):
            c = hoja.cell(5 + k, 2, ("▸ " if texto == activo else "   ") + texto)
            c.hyperlink = Hyperlink(ref=c.coordinate, location=destino)
            c.font = fuente(10.5, texto == activo, TEXTO if texto == activo else TEXTO_2)
            c.fill = relleno(NAV_ACTIVO if texto == activo else PANEL)
            if k == 0:
                c.border = Border(left=LADO, right=LADO, top=LADO)
            else:
                c.border = Border(left=LADO, right=LADO)

    navegacion(ws, "Dashboard Índices")

    def selector(hoja, fila, titulo, valor, formula_lista, nombre_celda, formato="@"):
        hoja.cell(fila, 2, titulo).font = fuente(9, True, TEXTO_2)
        c = hoja.cell(fila + 1, 2, valor)
        c.number_format = formato
        c.font = fuente(11, True, AZUL)
        c.fill = relleno("F4F6FB")
        c.border = Border(left=Side(style="thin", color=AZUL), right=Side(style="thin", color=AZUL),
                          top=Side(style="thin", color=AZUL), bottom=Side(style="thin", color=AZUL))
        c.alignment = Alignment(horizontal="center")
        # en el XML la formula de la validacion va sin "=" (con "=" Excel la considera dañada)
        dv = DataValidation(type="list", formula1=formula_lista.lstrip("="), allow_blank=False,
                            showDropDown=False, showErrorMessage=True)
        dv.error, dv.errorTitle = "Elige un valor de la lista", "Valor no válido"
        hoja.add_data_validation(dv)
        dv.add(c.coordinate)
        nombre(nombre_celda, f"'{hoja.title}'!${get_column_letter(c.column)}${c.row}")
        return c

    ramo_ini = "10" if "10" in ramos_i else ramos_i[0]
    selector(ws, 11, "RAMO", ramo_ini, "=L_RamosInd", "SelRamoInd")
    selector(ws, 14, "ÍNDICE / LAG", INDICES[0], "=L_SeriesInd", "SelIndice")
    selector(ws, 17, "PERIODO (COMPARATIVO POR RAMO)", fin, "=L_PeriodosInd", "SelPeriodoInd", "0")
    notas = [f"Real hasta {etiqueta(ultimo)}", f"Proyección {etiqueta(mover(ultimo, 1))} a {etiqueta(fin)}",
             "", "Cambia los selectores (listas", "desplegables) y las gráficas", "se actualizan solas.", "",
             f"Índices: {proc.get('indice', 'n/d')}", f"LAGs: {proc.get('lag', 'n/d')}",
             "(métodos validados fuera de muestra)"]
    for k, t in enumerate(notas):
        ws.cell(21 + k, 2, t).font = fuente(9, k < 2, TEXTO_2, k >= 3)

    ws["D2"] = "Dashboard de Índices de Reservas"
    ws["D2"].font = fuente(18, True)
    ws["D3"] = (f"Seguimiento mensual {etiqueta(PRIMER_PERIODO_INDICES)} a {etiqueta(ultimo)} (real) y proyección "
                f"{etiqueta(mover(ultimo, 1))} a {etiqueta(fin)} · HParametros_2026")
    ws["D3"].font = fuente(10, False, TEXTO_2)

    # KPI (formulas)
    crit = "I_Ramo,SelRamoInd,I_Serie,SelIndice"

    def valor_indice(per_expr):
        return f'IFERROR(SUMIFS(I_Val,{crit},I_Per,{per_expr})/(COUNTIFS({crit},I_Per,{per_expr})>0),NA())'

    kpis = [
        (f"Último real ({etiqueta(ultimo)})", "=" + valor_indice(ultimo), "0.0000"),
        ("Promedio real 12 meses",
         f'=IFERROR(AVERAGEIFS(I_Val,I_Tipo,"Real",{crit},I_Per,">"&{p_12}),NA())', "0.0000"),
        (f"Proyección {etiqueta(p_dic_actual)}", "=" + valor_indice(p_dic_actual), "0.0000"),
        (f"Proyección {etiqueta(fin)}", "=" + valor_indice(fin), "0.0000"),
        (f"Var. {etiqueta(fin)} vs {etiqueta(ultimo)}", "", "+0.0%;-0.0%;0.0%"),
    ]
    pintar(ws, f"D5:{ultima_col}8", BANDA_KPI)
    for k, (tit, form, fmt) in enumerate(kpis):
        col = 4 + k * 6
        ws.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col + 5)
        ws.merge_cells(start_row=7, start_column=col, end_row=8, end_column=col + 5)
        ws.cell(5, col, tit).font = fuente(9.5, True, TEXTO_2)
        ws.cell(5, col).alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        c = ws.cell(7, col, f'=IFERROR({form[1:]},"s/d")' if form else None)   # s/d = sin dato
        c.number_format = fmt
        c.font = fuente(20, True)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.cell(7, 28).value = '=IFERROR(V7/D7-1,"s/d")'      # proyeccion al cierre (V7) contra ultimo real (D7)

    # paneles
    caja(ws, "D10:S27")
    caja(ws, "U10:AG27")
    caja(ws, "D29:S46")
    caja(ws, "U29:AG46")
    ws["D10"] = '="Evolución mensual · "&SelIndice&" · Ramo "&SelRamoInd'
    ws["U10"] = (f'=SelIndice&" por ramo · "&IFERROR(INDEX(Calc!$B:$B,MATCH(SelPeriodoInd,Calc!$A:$A,0)),'
                 f'SelPeriodoInd)&IF(SelPeriodoInd>{ultimo}," (proyección)"," (real)")')
    ws["D29"] = '="Patrón de desarrollo acumulado (LAG 1-10) · Ramo "&SelRamoInd'
    ws["U29"] = '="Resumen de índices · Ramo "&SelRamoInd'
    for ref in ("D10", "U10", "D29", "U29"):
        ws[ref].font = fuente(11.5, True)
        ws[ref].alignment = Alignment(indent=1, vertical="center")
    ws.row_dimensions[10].height = 22
    ws.row_dimensions[29].height = 22

    # Calc: serie mensual del indice seleccionado
    wc["A1"], wc["B1"], wc["C1"], wc["D1"], wc["E1"], wc["F1"] = ("Periodo", "Mes", "Real", "Proyección",
                                                                  "Banda 80% inferior", "Banda 80% superior")
    per_i = rango(PRIMER_PERIODO_INDICES, fin)
    for k, p in enumerate(per_i, start=2):
        wc.cell(k, 1, p)
        wc.cell(k, 2, etiqueta(p))
        base = f"{crit},I_Per,$A{k}"
        wc.cell(k, 3, f'=IF(COUNTIFS(I_Tipo,"Real",{base})=0,NA(),SUMIFS(I_Val,I_Tipo,"Real",{base}))')
        proy = f'COUNTIFS(I_Tipo,"Proyección",{base})'
        wc.cell(k, 4, f'=IF($A{k}={ultimo},C{k},IF({proy}=0,NA(),SUMIFS(I_Val,I_Tipo,"Proyección",{base})))')
        wc.cell(k, 5, f'=IF($A{k}={ultimo},C{k},IF({proy}=0,NA(),SUMIFS(I_LI,I_Tipo,"Proyección",{base})))')
        wc.cell(k, 6, f'=IF($A{k}={ultimo},C{k},IF({proy}=0,NA(),SUMIFS(I_LS,I_Tipo,"Proyección",{base})))')
    n_ci = len(per_i) + 1
    ch = LineChart()
    ch.add_data(Reference(wc, min_col=3, max_col=6, min_row=1, max_row=n_ci), titles_from_data=True)
    ch.set_categories(Reference(wc, min_col=2, min_row=2, max_row=n_ci))
    linea(ch.series[0], AZUL, 2.0)
    linea(ch.series[1], NARANJA, 2.0, punteada=True)
    linea(ch.series[2], GRIS_BANDA, 1.0)
    linea(ch.series[3], GRIS_BANDA, 1.0)
    estilo_grafica(ch, "0.00")
    ch.x_axis.tickLblSkip = 6
    ch.x_axis.tickMarkSkip = 6
    colocar(ws, ch, "D11", "S27")

    # Calc: comparativo por ramo (seleccionado resaltado)
    wc["H1"], wc["I1"], wc["J1"], wc["K1"] = "Ramo", "Valor", "Ramo seleccionado", "Otros ramos"
    for k, ramo in enumerate(ramos_i, start=2):
        wc.cell(k, 8, ramo).number_format = "@"
        base = f'I_Ramo,$H{k},I_Serie,SelIndice,I_Per,SelPeriodoInd'
        wc.cell(k, 9, f"=IF(COUNTIFS({base})=0,NA(),SUMIFS(I_Val,{base}))")
        wc.cell(k, 10, f"=IF($H{k}=SelRamoInd,I{k},NA())")
        wc.cell(k, 11, f"=IF($H{k}<>SelRamoInd,I{k},NA())")
    n_cr = len(ramos_i) + 1
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap, ch.gapWidth = "bar", "clustered", 100, 55
    ch.add_data(Reference(wc, min_col=10, max_col=11, min_row=1, max_row=n_cr), titles_from_data=True)
    ch.set_categories(Reference(wc, min_col=8, min_row=2, max_row=n_cr))
    barra(ch.series[0], AZUL)
    barra(ch.series[1], GRIS_BARRA)
    estilo_grafica(ch, "0.00")
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.crosses = "max"
    colocar(ws, ch, "U11", "AG27")

    # Calc: patron de LAGs (instantaneas)
    instantaneas = [mover(ultimo, -24), mover(ultimo, -12), ultimo, fin]
    titulos = [f"{etiqueta(p)} (real)" for p in instantaneas[:3]] + [f"{etiqueta(fin)} (proyección)"]
    wc["M1"] = "LAG"
    for j, t in enumerate(titulos):
        wc.cell(1, 14 + j, t)
    for k, lag in enumerate(LAGS, start=2):
        wc.cell(k, 13, lag)
        for j, p in enumerate(instantaneas):
            base = f'I_Ramo,SelRamoInd,I_Serie,$M{k},I_Per,{p}'
            wc.cell(k, 14 + j, f"=IF(COUNTIFS({base})=0,NA(),SUMIFS(I_Val,{base}))")
    ch = LineChart()
    ch.add_data(Reference(wc, min_col=14, max_col=17, min_row=1, max_row=11), titles_from_data=True)
    ch.set_categories(Reference(wc, min_col=13, min_row=2, max_row=11))
    for s, color in zip(ch.series[:3], AZULES_ORDINALES):
        linea(s, color, 2.0)
        s.marker.symbol = "circle"
        s.marker.size = 5
        s.marker.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill=color))
    linea(ch.series[3], NARANJA, 2.0, punteada=True)
    estilo_grafica(ch, "0.00")
    colocar(ws, ch, "D30", "S46")

    # Tabla resumen del ramo (formulas); cada dato ocupa 2 columnas x 2 filas de la rejilla
    cab = [("Índice", 21, 24), (f"Real\n{etiqueta(ultimo)}", 25, 26), (f"Proy.\n{etiqueta(p_dic_actual)}", 27, 28),
           (f"Proy.\n{etiqueta(fin)}", 29, 30), ("Var. vs\nreal", 31, 32)]
    linea_fina = Border(bottom=Side(style="thin", color=BORDE))
    for t, c1_, c2_ in cab:
        ws.merge_cells(start_row=31, start_column=c1_, end_row=31, end_column=c2_)
        cel = ws.cell(31, c1_, t)
        cel.font = fuente(9, True, TEXTO_2)
        cel.alignment = Alignment(horizontal="left" if c1_ == 21 else "right", vertical="bottom", wrap_text=True,
                                  indent=1 if c1_ == 21 else 0)
        for k in range(c1_, c2_ + 1):
            ws.cell(31, k).border = linea_fina
    ws.row_dimensions[31].height = 30
    for k, serie in enumerate(INDICES + ["LAG 1", "LAG 2", "LAG 3"]):
        f_ = 32 + k * 2
        base = f'I_Ramo,SelRamoInd,I_Serie,"{serie}"'
        forms = [serie] + [f"=IFERROR(SUMIFS(I_Val,{base},I_Per,{p})/(COUNTIFS({base},I_Per,{p})>0),NA())"
                           for p in (ultimo, p_dic_actual, fin)]
        forms.append(f"=IFERROR({get_column_letter(29)}{f_}/{get_column_letter(25)}{f_}-1,NA())")
        for (_, c1_, c2_), form in zip(cab, forms):
            ws.merge_cells(start_row=f_, start_column=c1_, end_row=f_ + 1, end_column=c2_)
            cel = ws.cell(f_, c1_, form)
            cel.font = fuente(10, c1_ == 31)
            cel.number_format = "+0.0%;-0.0%;0.0%" if c1_ == 31 else "0.0000"
            cel.alignment = Alignment(horizontal="left" if c1_ == 21 else "right", vertical="center",
                                      indent=1 if c1_ == 21 else 0)
            for k2 in range(c1_, c2_ + 1):
                ws.cell(f_ + 1, k2).border = linea_fina
    ws.conditional_formatting.add(
        "AE32:AF45", ColorScaleRule(start_type="num", start_value=-0.3, start_color="6DA7EC", mid_type="num",
                                    mid_value=0, mid_color="F0EFEC", end_type="num", end_value=0.3,
                                    end_color="E66767"))
    ws.cell(46, 21, "#N/D = sin dato para ese ramo/índice").font = fuente(8, False, TEXTO_2, True)

    # ------------------------------------------------------------ dashboard reservas
    ws = wd_r
    preparar_hoja(ws)
    caja(ws, "B2:B46")
    ws["B3"] = "◆ ÍNDICES Y RESERVAS"
    ws["B3"].font = fuente(13, True)
    navegacion(ws, "Dashboard Reservas")
    selector(ws, 11, "RESERVA", "RRC", "=L_Reservas", "SelReserva")
    selector(ws, 14, "CONCEPTO", "NETO", '=INDIRECT("Conceptos_"&SelReserva)', "SelConcepto")
    selector(ws, 17, "RAMO", "Todos", '=INDIRECT("Ramos_"&SelReserva)', "SelRamoRes")
    selector(ws, 20, "MONEDA", "USD", "=L_Unidades", "SelUnidad")
    ws["B22"] = ('=IF(AND(ISNUMBER(MATCH(SelConcepto,INDIRECT("Conceptos_"&SelReserva),0)),'
                 'ISNUMBER(MATCH(SelRamoRes,INDIRECT("Ramos_"&SelReserva),0))),"",'
                 '"⚠ Concepto o ramo no existe en "&SelReserva&": elígelo de nuevo")')
    ws["B22"].font = fuente(9, True, "B4541F")
    ws["B22"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[22].height = 26
    notas = [f"Real hasta {etiqueta(ultimo)}", f"Proyección {etiqueta(mover(ultimo, 1))} a {etiqueta(fin_m)}",
             "Cifras en millones.", "MXN = USD × TC del mes", "(TC FCST de Inversiones).", "",
             f"Montos: {proc.get('nivel', 'n/d')}", f"Razones: {proc.get('razon', 'n/d')}",
             "NETO = BRUTO - IRR"]
    for k, t in enumerate(notas):
        ws.cell(24 + k, 2, t).font = fuente(9, k < 2, TEXTO_2, k >= 3)
    ws["D2"] = "Dashboard de Reservas: Real y Proyección"
    ws["D2"].font = fuente(18, True)
    ws["D3"] = (f"RRC y SONR (Daños) y RFV (Fianzas) por concepto y ramo · real {etiqueta(PRIMER_PERIODO_MONTOS)} "
                f"a {etiqueta(ultimo)} y proyección {etiqueta(mover(ultimo, 1))} a {etiqueta(fin_m)}")
    ws["D3"].font = fuente(10, False, TEXTO_2)

    wc["X1"] = '=IF(SelRamoRes="Todos","*",SelRamoRes)'
    nombre("CritRamoRes", "Calc!$X$1")
    unidad = 'IF(SelUnidad="MXN",INDEX(TC_Val,MATCH({p},TC_Per,0)),1)'

    def monto(conc_expr, per_expr, ramo_expr="CritRamoRes"):
        base = f"R_Res,SelReserva,R_Conc,{conc_expr},R_Ramo,{ramo_expr},R_Per,{per_expr}"
        return (f"IF(COUNTIFS({base})=0,NA(),SUMIFS(R_Val,{base})*{unidad.format(p=per_expr)}/1000000)")

    pintar(ws, f"D5:{ultima_col}8", BANDA_KPI)
    kpis_r = [
        ('="Real "&"' + etiqueta(ultimo) + '"&" (M "&SelUnidad&")"', "=" + monto("SelConcepto", ultimo), "#,##0.0"),
        (f'="Proy. {etiqueta(p_dic_actual)} (M "&SelUnidad&")"', "=" + monto("SelConcepto", p_dic_actual), "#,##0.0"),
        (f'="Proy. {etiqueta(fin_m)} (M "&SelUnidad&")"', "=" + monto("SelConcepto", fin_m), "#,##0.0"),
        ("Crec. anual proyectado", f"=IFERROR((P7/D7)^(12/{len(rango(ultimo, fin_m)) - 1})-1,NA())",
         "+0.0%;-0.0%;0.0%"),
        ("Crec. real 12 meses",
         f"=IFERROR(D7/({monto('SelConcepto', p_12)})-1,NA())", "+0.0%;-0.0%;0.0%"),
    ]
    for k, (tit, form, fmt) in enumerate(kpis_r):
        col = 4 + k * 6
        ws.merge_cells(start_row=5, start_column=col, end_row=6, end_column=col + 5)
        ws.merge_cells(start_row=7, start_column=col, end_row=8, end_column=col + 5)
        ws.cell(5, col, tit).font = fuente(9.5, True, TEXTO_2)
        ws.cell(5, col).alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
        c = ws.cell(7, col, f'=IFERROR({form[1:]},"s/d")' if form else None)   # s/d = sin dato
        c.number_format = fmt
        c.font = fuente(20, True)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    # En USD y MXN el % cedido es el mismo
    ws["D9"] = (f'="% cedido (IRR/BRUTO) {etiqueta(fin_m)}: "&IFERROR(TEXT({monto(chr(34) + "IRR" + chr(34), fin_m)}/'
                f'{monto(chr(34) + "BRUTO" + chr(34), fin_m)},"0.0%"),"n/d")&"   ·   '
                f'% cedido {etiqueta(ultimo)}: "&IFERROR(TEXT({monto(chr(34) + "IRR" + chr(34), ultimo)}/'
                f'{monto(chr(34) + "BRUTO" + chr(34), ultimo)},"0.0%"),"n/d")')
    ws["D9"].font = fuente(9.5, True, TEXTO_2)

    caja(ws, "D11:S27")
    caja(ws, "U11:AG27")
    caja(ws, "D29:S46")
    caja(ws, "U29:AG46")
    ws["D11"] = f'=SelReserva&" "&SelConcepto&" mensual {etiqueta(PRIMER_PERIODO_MENSUAL)}-{etiqueta(fin_m)} · Ramo "&SelRamoRes&" (M "&SelUnidad&")"'
    ws["U11"] = f'=SelReserva&" "&SelConcepto&" por ramo · {etiqueta(ultimo)} vs {etiqueta(fin_m)} (M "&SelUnidad&")"'
    ws["D29"] = f'=SelReserva&" "&SelConcepto&" histórico y proyección · Ramo "&SelRamoRes&" (M "&SelUnidad&")"'
    ws["U29"] = f'=SelReserva&" por concepto · {etiqueta(ultimo)} vs {etiqueta(fin_m)} · Ramo "&SelRamoRes&" (M "&SelUnidad&")"'
    for ref in ("D11", "U11", "D29", "U29"):
        ws[ref].font = fuente(11.5, True)
        ws[ref].alignment = Alignment(indent=1, vertical="center")
    ws.row_dimensions[11].height = 22
    ws.row_dimensions[29].height = 22

    # Calc: serie mensual de reservas
    c0 = 25   # columna Y
    for j, t in enumerate(["Periodo", "Mes", "Monto", "Real", "Proyección", "Proyección (columnas)"]):
        wc.cell(1, c0 + j, t)
    per_m = rango(PRIMER_PERIODO_MONTOS, fin_m)
    for k, p in enumerate(per_m, start=2):
        L = get_column_letter
        wc.cell(k, c0, p)
        wc.cell(k, c0 + 1, etiqueta(p))
        wc.cell(k, c0 + 2, "=" + monto("SelConcepto", f"${L(c0)}{k}"))
        m = f"{L(c0 + 2)}{k}"
        wc.cell(k, c0 + 3, f"=IF(${L(c0)}{k}<={ultimo},{m},NA())")
        wc.cell(k, c0 + 4, f"=IF(${L(c0)}{k}>={ultimo},{m},NA())")
        wc.cell(k, c0 + 5, f"=IF(${L(c0)}{k}>{ultimo},{m},NA())")
    n_m = len(per_m) + 1
    k_ini = per_m.index(PRIMER_PERIODO_MENSUAL) + 2 if PRIMER_PERIODO_MENSUAL in per_m else 2
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap, ch.gapWidth = "col", "clustered", 100, 45
    real = Reference(wc, min_col=c0 + 3, min_row=k_ini, max_row=n_m)
    proy = Reference(wc, min_col=c0 + 5, min_row=k_ini, max_row=n_m)
    s_real, s_proy = Series(real), Series(proy)
    s_real.tx = SeriesLabel(v="Real")
    s_proy.tx = SeriesLabel(v="Proyección")
    ch.series.extend([s_real, s_proy])
    ch.set_categories(Reference(wc, min_col=c0 + 1, min_row=k_ini, max_row=n_m))
    barra(ch.series[0], AZUL)
    barra(ch.series[1], NARANJA)
    estilo_grafica(ch, "#,##0")
    ch.x_axis.tickLblSkip = 3
    colocar(ws, ch, "D12", "S27")

    ch = LineChart()
    ch.add_data(Reference(wc, min_col=c0 + 3, max_col=c0 + 4, min_row=1, max_row=n_m), titles_from_data=True)
    ch.set_categories(Reference(wc, min_col=c0 + 1, min_row=2, max_row=n_m))
    linea(ch.series[0], AZUL, 2.0)
    linea(ch.series[1], NARANJA, 2.0, punteada=True)
    estilo_grafica(ch, "#,##0")
    ch.x_axis.tickLblSkip = 6
    ch.x_axis.tickMarkSkip = 6
    colocar(ws, ch, "D30", "S46")

    # Calc: por ramo (lista dependiente de la reserva)
    c1 = 32   # AF
    for j, t in enumerate(["Ramo", f"{etiqueta(ultimo)} (real)", f"{etiqueta(fin_m)} (proyección)"]):
        wc.cell(1, c1 + j, t)
    max_ramos = max(len(v) for v in ramos_m.values())
    for k in range(2, max_ramos + 2):
        L = get_column_letter
        wc.cell(k, c1, f'=IFERROR(INDEX(INDIRECT("RamosSin_"&SelReserva),{k - 1}),"")')
        r_ = f"{L(c1)}{k}"
        wc.cell(k, c1 + 1, f'=IF({r_}="",NA(),{monto("SelConcepto", ultimo, r_)})')
        wc.cell(k, c1 + 2, f'=IF({r_}="",NA(),{monto("SelConcepto", fin_m, r_)})')
    ch = BarChart()
    ch.type, ch.grouping, ch.gapWidth = "bar", "clustered", 50
    ch.add_data(Reference(wc, min_col=c1 + 1, max_col=c1 + 2, min_row=1, max_row=max_ramos + 1), titles_from_data=True)
    ch.set_categories(Reference(wc, min_col=c1, min_row=2, max_row=max_ramos + 1))
    barra(ch.series[0], AZUL)
    barra(ch.series[1], NARANJA)
    estilo_grafica(ch, "#,##0")
    ch.x_axis.scaling.orientation = "maxMin"
    ch.y_axis.crosses = "max"
    colocar(ws, ch, "U12", "AG27")

    # Calc: por concepto
    c2 = 36   # AJ
    for j, t in enumerate(["Concepto", f"{etiqueta(ultimo)} (real)", f"{etiqueta(fin_m)} (proyección)"]):
        wc.cell(1, c2 + j, t)
    for k in range(2, 8):
        L = get_column_letter
        wc.cell(k, c2, f'=IFERROR(INDEX(INDIRECT("Conceptos_"&SelReserva),{k - 1}),"")')
        cc = f"{L(c2)}{k}"
        wc.cell(k, c2 + 1, f'=IF({cc}="",NA(),{monto(cc, ultimo)})')
        wc.cell(k, c2 + 2, f'=IF({cc}="",NA(),{monto(cc, fin_m)})')
    ch = BarChart()
    ch.type, ch.grouping, ch.gapWidth = "col", "clustered", 60
    ch.add_data(Reference(wc, min_col=c2 + 1, max_col=c2 + 2, min_row=1, max_row=7), titles_from_data=True)
    ch.set_categories(Reference(wc, min_col=c2, min_row=2, max_row=7))
    barra(ch.series[0], AZUL)
    barra(ch.series[1], NARANJA)
    estilo_grafica(ch, "#,##0")
    colocar(ws, ch, "U30", "AG46")
    wc.sheet_state = "hidden"

    # ------------------------------------------------------------ analisis (valores)
    construir_analisis(wa, registros_i, ramos_i, registros_m, ultimo, p_dic_actual, fin, fin_m, p_12, metodo)
    for hoja, area in ((wd_i, "A1:AH47"), (wd_r, "A1:AH47"), (wa, None)):
        hoja.sheet_properties.tabColor = "8EA9DB"
        hoja.page_setup.orientation = "landscape"
        hoja.page_setup.paperSize = hoja.PAPERSIZE_LETTER
        hoja.sheet_properties.pageSetUpPr.fitToPage = True
        hoja.page_setup.fitToWidth = 1
        hoja.page_setup.fitToHeight = 1 if area else 0
        hoja.print_options.horizontalCentered = True
        hoja.page_margins.left = hoja.page_margins.right = 0.3
        hoja.page_margins.top = hoja.page_margins.bottom = 0.4
        if area:
            hoja.print_area = area
    wd_i.sheet_view.tabSelected = True
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    clasificar(wb)
    guardar_libro(wb, ruta_salida)
    print(f"   Dashboard: {ruta_salida.name}", flush=True)
    return ruta_salida


def clasificar(wb) -> None:
    """El dashboard trae los mismos datos que la BD: hereda su etiqueta de sensibilidad (propiedades MSIP) y
    lleva el mismo pie de pagina de uso interno."""
    try:
        origen = openpyxl.load_workbook(ARCHIVO_DANOS, read_only=True)
        for prop in origen.custom_doc_props:
            if prop.name.startswith("MSIP_Label_"):
                wb.custom_doc_props.append(prop)
        origen.close()
    except Exception as e:  # noqa: BLE001
        print(f"   Aviso: no se pudo copiar la etiqueta de sensibilidad de {ARCHIVO_DANOS.name}: {e!r}")
    for ws in wb.worksheets:
        pie = ws.oddFooter.center
        pie.text, pie.font, pie.size, pie.color = PIE_PAGINA, "Aptos", 10, "000000"


def construir_analisis(ws, registros_i, ramos_i, registros_m, ultimo, p_dic, fin, fin_m, p_12, metodo):
    preparar_hoja(ws, ancho_nav=16, columnas=16, ancho=11.5, filas=80)
    ws.column_dimensions["C"].width = 11.5
    ws["B2"] = "Análisis: real vs proyección"
    ws["B2"].font = fuente(18, True)
    ws["B3"] = "Valores fijos calculados al generar el archivo (se actualizan en cada corrida de la proyección)."
    ws["B3"].font = fuente(10, False, TEXTO_2)
    for k, (t, d) in enumerate([("‹ Dashboard Índices", "'Dashboard Índices'!A1"),
                                ("‹ Dashboard Reservas", "'Dashboard Reservas'!A1")]):
        c = ws.cell(4, 2 + k * 4, t)
        c.hyperlink = Hyperlink(ref=c.coordinate, location=d)
        c.font = fuente(10, True, AZUL)

    val_i = {(r[1], r[2], r[3]): r[4] for r in registros_i}
    fila = 6
    ws.cell(fila, 2, f"1. Índices por ramo: real {etiqueta(ultimo)} vs proyección {etiqueta(fin)}").font = fuente(12, True)
    fila += 1
    cab = ["Ramo"]
    for serie in INDICES:
        cab += [f"{serie}\nreal {etiqueta(ultimo)}", f"{serie}\nproy. {etiqueta(fin)}", "Var. %"]
    for j, t in enumerate(cab):
        c = ws.cell(fila, 2 + j, t)
        c.font = fuente(9, True, TEXTO_2)
        c.fill = relleno(BANDA_KPI)
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[fila].height = 42
    ini = fila + 1
    for ramo in ramos_i:
        fila += 1
        ws.cell(fila, 2, ramo).font = fuente(10, True)
        ws.cell(fila, 2).fill = relleno(PANEL)
        for k, serie in enumerate(INDICES):
            a, b = val_i.get((ramo, serie, ultimo)), val_i.get((ramo, serie, fin))
            for j, v in enumerate([a, b, (b / a - 1) if a and b is not None else None]):
                c = ws.cell(fila, 3 + k * 3 + j, v)
                c.number_format = "+0.0%;-0.0%;0.0%" if j == 2 else "0.0000"
                c.fill = relleno(PANEL)
                c.font = fuente(10)
    for k in range(len(INDICES)):
        col = get_column_letter(5 + k * 3)
        ws.conditional_formatting.add(
            f"{col}{ini}:{col}{fila}",
            ColorScaleRule(start_type="num", start_value=-0.3, start_color="6DA7EC", mid_type="num", mid_value=0,
                           mid_color="F0EFEC", end_type="num", end_value=0.3, end_color="E66767"))
        for col_v in (get_column_letter(3 + k * 3), get_column_letter(4 + k * 3)):
            ws.conditional_formatting.add(
                f"{col_v}{ini}:{col_v}{fila}",
                ColorScaleRule(start_type="min", start_color="EAF2FD", end_type="max", end_color="6DA7EC"))

    fila += 3
    ws.cell(fila, 2, "2. Reservas: totales por concepto (millones USD)").font = fuente(12, True)
    fila += 1
    cortes = [mover(ultimo, -20) // 100 * 100 + 12, mover(ultimo, -8) // 100 * 100 + 12, ultimo, p_dic, fin_m]
    cortes = list(dict.fromkeys(cortes))
    cab = ["Reserva", "Concepto"] + [f"{etiqueta(p)}{' (real)' if p <= ultimo else ' (proy.)'}" for p in cortes] + \
          ["Crec. real 12m", "Crec. anualizado proy."]
    for j, t in enumerate(cab):
        c = ws.cell(fila, 2 + j, t)
        c.font = fuente(9, True, TEXTO_2)
        c.fill = relleno(BANDA_KPI)
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[fila].height = 32
    tot = {}
    for _, reserva, conc, _, per, _, v in registros_m:
        tot[(reserva, conc, per)] = tot.get((reserva, conc, per), 0.0) + v
    ini = fila + 1
    orden = [("RRC", c) for c in ["NETO", "BRUTO", "BEL", "GTO", "IRR", "MR"]] + \
            [("SONR", c) for c in ["NETO", "BRUTO", "BEL", "IRR", "MR"]] + [("RFV", c) for c in ["NETO", "BRUTO", "IRR", "RCONT"]]
    meses_proy = len(rango(ultimo, fin_m)) - 1
    for reserva, conc in orden:
        fila += 1
        ws.cell(fila, 2, reserva).font = fuente(10, True)
        ws.cell(fila, 3, conc).font = fuente(10)
        valores = [tot.get((reserva, conc, p)) for p in cortes]
        for j, v in enumerate(valores):
            c = ws.cell(fila, 4 + j, None if v is None else v / 1e6)
            c.number_format = "#,##0.0"
            c.font = fuente(10, j >= 3)
        v12, vu, vf = tot.get((reserva, conc, p_12)), tot.get((reserva, conc, ultimo)), tot.get((reserva, conc, fin_m))
        c = ws.cell(fila, 4 + len(cortes), (vu / v12 - 1) if v12 and vu else None)
        c.number_format = "+0.0%;-0.0%;0.0%"
        c = ws.cell(fila, 5 + len(cortes), ((vf / vu) ** (12 / meses_proy) - 1) if vu and vf and vu > 0 and vf > 0 else None)
        c.number_format = "+0.0%;-0.0%;0.0%"
        for j in range(2, 6 + len(cortes)):
            ws.cell(fila, j).fill = relleno(PANEL)
    for col in (get_column_letter(4 + len(cortes)), get_column_letter(5 + len(cortes))):
        ws.conditional_formatting.add(
            f"{col}{ini}:{col}{fila}",
            ColorScaleRule(start_type="num", start_value=-0.5, start_color="6DA7EC", mid_type="num", mid_value=0,
                           mid_color="F0EFEC", end_type="num", end_value=0.5, end_color="E66767"))

    fila += 3
    ws.cell(fila, 2, "3. Método de proyección (validación fuera de muestra)").font = fuente(12, True)
    fila += 1
    # columnas: B tipo | C:D procedimiento | E WAPE | F AvgRelMAE | G sesgo | H:M criterio
    posiciones = [(2, 2), (3, 4), (5, 5), (6, 6), (7, 7), (8, 13)]
    cab = ["Tipo de serie", "Procedimiento", "WAPE %", "AvgRelMAE", "Sesgo agregado % 13-16", "Criterio"]
    for (c1_, c2_), t in zip(posiciones, cab):
        ws.merge_cells(start_row=fila, start_column=c1_, end_row=fila, end_column=c2_)
        c = ws.cell(fila, c1_, t)
        c.font = fuente(9, True, TEXTO_2)
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        for k in range(c1_, c2_ + 1):
            ws.cell(fila, k).fill = relleno(BANDA_KPI)
    ws.row_dimensions[fila].height = 32
    nombres = {"nivel": "Montos", "indice": "Índices", "razon": "Razones", "lag": "LAGs"}
    for d in metodo:
        fila += 1
        vals = [nombres.get(d.get("Tipo"), d.get("Tipo")), d.get("Procedimiento"), d.get("WAPE %"), d.get("AvgRelMAE"),
                d.get("Sesgo agregado % 13-16"), str(d.get("Seleccionado", "")).replace("SI: ", "")]
        for j, ((c1_, c2_), v) in enumerate(zip(posiciones, vals)):
            ws.merge_cells(start_row=fila, start_column=c1_, end_row=fila, end_column=c2_)
            c = ws.cell(fila, c1_, v)
            c.font = fuente(10, j == 1)
            c.number_format = "0.000" if j == 3 else ("0.0" if j in (2, 4) else "General")
            c.alignment = Alignment(wrap_text=j in (1, 5), vertical="center")
            for k in range(c1_, c2_ + 1):
                ws.cell(fila, k).fill = relleno(PANEL)
        ws.row_dimensions[fila].height = 30
    ws.freeze_panes = "A5"


if __name__ == "__main__":
    generar()
