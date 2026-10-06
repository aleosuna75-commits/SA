# -*- coding: utf-8 -*-
"""Insumos para los scripts de valuacion del area (RRC y SONR) tomados de la BD proyectada.

Lee ``BD_ BEL - IRR - MR_Proyeccion.xlsx`` (la salida de ``proyeccion_reservas.py``) y arma, con la forma que
esperan los scripts del area, las tablas que antes venian de CSV o de constantes:

================================  ==========================================================================
Insumo del script                 De donde sale en la BD
================================  ==========================================================================
ParametrosMens (RRC)              HParametros: Ind Sin RRC / Ind sin RRC 99.5% (``Real`` hasta el ultimo mes real
                                  y ``Proyección`` despues; en los ramos de RAMOS_IS_FA, el IS (FA) de la hoja
                                  IS_FA en los meses proyectados) y el bloque FACTOR GTO de la hoja de montos como
                                  Ind. Gasto, los dos por mes (``-1`` a ``-12``). Duracion y retencion para el MR
                                  (Pesos_dur, Resto Monedas_dur, Pesos_ret, Resto Monedas_ret) no estan en la BD:
                                  se toman del CSV del area si se indica.
IS_Cat (RRC, ramos 71 y 73)        Ind Sin RRC de TEV e Hidro por mes (Real / Proyección).
ParamSONR (SONR)                   HParametros: Ind Sin SONR Media / 99.5% y LAG 1 a LAG 10 por mes y ramo (con el
                                  IS (FA) en los ramos de RAMOS_IS_FA en los meses proyectados) y Factor_Ret =
                                  1 - IRR / BEL de SONR de la hoja de montos (en los meses proyectados, con la
                                  CESION y el FACTOR MR proyectados).
FACTOR MR (RRC y SONR)             Bloque FACTOR MR de la hoja de montos (MR / PND o PD) por mes y ramo: el MR de los
                                  scripts es PND del contrato x FACTOR MR (config MR_DESDE = "BD").
Tipo de cambio USD                 Columna TC de la hoja de montos (real hasta el ultimo mes y pronostico despues).
Escenario base (0)                 Montos reales de diciembre del ano anterior de la hoja de montos.
Saldos proyectados                 Montos proyectados de la hoja de montos (BEL por FND y sus derivados).
================================  ==========================================================================

Los montos y factores proyectados de la hoja de montos son formulas: se leen de los valores que Excel guardo
(abrir y guardar la BD en Excel) y, si no los hay, del ``Diagnostico_Proyeccion.xlsx`` de la misma corrida.
Las claves de ramo son las de los scripts (10, 31, 35, 39, 40, 50, 60, 71, 73, 80, 90, 100, 110).
"""
from __future__ import annotations

import math
import re
import unicodedata
import warnings
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

VERSION_AOD = "2026-10-06h"      # version que piden los scripts _aod_BD.py (parchar_aod.py)
HOJA_MONTOS = "BD_Montos_RRC_SONR"
TC_BD_DESDE = 202601             # desde este mes el TC MXN / USD de los scripts _aod sale de la columna TC de la BD
                                 # (TC_Real_Esti.xlsx: FCST, columna J, en 2026; FCST 2027, columna M, en 2027)
MONEDA_USD, MONEDA_MXN = 31, 1   # cMON_Id / MonedaOri de la base de valuacion
HOJA_IS_FA = "IS_FA"
HOJA_PE = "PE_RAMO"                 # prima tomada del mes por ramo (USD): real (PExRamo) y FCST
HOJA_DIAG_INDICADORES = "Indicadores_Ramo"
HOJA_DIAG_MONTOS = "Montos_Proyectados"
RAMOS_SCRIPT = [10, 31, 35, 39, 40, 50, 60, 71, 73, 80, 90, 100, 110]
RAMOS_SONR = [10, 31, 35, 39, 40, 50, 60, 80, 90, 100, 110]          # 71 y 73 no tienen parametros de SONR
ETIQUETA_HP = {71: ("TEV", "Tev"), 73: ("Hidro",)}                     # clave del script -> etiqueta en HParametros
RAMO_BD = {31: "30", 35: "34", 39: "37"}                               # clave del script -> columna RAM_ de la BD
RAMOS_IS_FA = {"RRC": (40, 50, 80, 90), "SONR": (40, 50, 80, 90)}      # mismos ramos que el BEL por FND de la BD
PARAMETROS_HP = {"IS_RRC": "Ind Sin RRC", "IS_RRC_99": "Ind sin RRC 99.5%",
                 "IS_SONR": "Ind Sin SONR Media", "IS_SONR_99": "Ind Sin SONR 99.5%"}
LAGS = [f"LAG {k}" for k in range(1, 11)]
SIGNO = {"RRC": {"BEL": -1, "BELG": -1, "IRR": -1, "MR": -1, "BRUTO": 1, "NETO": 1},   # convencion de la salida del
         "SONR": {"BEL": 1, "IRR": 1, "MR": 1, "BRUTO": 1, "NETO": 1}}                 # area (RRC_esc / SONR_esc)
CONCEPTO_BD = {"BELG": "GTO"}                                          # tipo de monto del area -> concepto de la BD
COLUMNAS_ESC = ["Reserva", "Escenario", "Tipo de Monto", "Ramo", "Periodo", "Monto_MXN", "Monto_USD", "TC"]
UMBRAL_CERO = 1.0                                                      # USD: debajo, el monto cuenta como cero


def norm(t) -> str:
    return re.sub(r"\s+", " ", str(t or "").replace("\xa0", " ")).strip().upper()


def mes_mas(periodo: int, k: int) -> int:
    """Periodo AAAAMM mas k meses (k puede ser negativo)."""
    i = (periodo // 100) * 12 + (periodo % 100 - 1) + k
    return (i // 12) * 100 + i % 12 + 1


def rango_meses(p0: int, p1: int) -> list[int]:
    out, p = [], p0
    while p <= p1:
        out.append(p)
        p = mes_mas(p, 1)
    return out


def _num(v):
    if isinstance(v, bool) or v is None:
        return math.nan
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(float(v)) else math.nan
    t = str(v).strip().replace("\xa0", "").replace(" ", "").replace(",", "")
    try:
        return float(t)
    except ValueError:
        return math.nan


def _interpolar(valores: dict, periodos: list[int], huecos_reg: list, etiqueta: str) -> dict:
    """Serie mensual sobre periodos: huecos interiores interpolados linealmente, orillas con el vecino mas cercano.
    Los huecos se anotan en huecos_reg como (etiqueta, primer mes, ultimo mes, n)."""
    xs = [valores.get(p, math.nan) for p in periodos]
    n = len(xs)
    idx = [i for i, x in enumerate(xs) if not math.isnan(x)]
    if not idx:
        return {p: math.nan for p in periodos}
    out = list(xs)
    huecos = []
    for i in range(n):
        if math.isnan(out[i]):
            huecos.append(periodos[i])
            ant = max((j for j in idx if j < i), default=None)
            sig = min((j for j in idx if j > i), default=None)
            if ant is not None and sig is not None:
                out[i] = xs[ant] + (xs[sig] - xs[ant]) * (i - ant) / (sig - ant)
            else:
                out[i] = xs[ant if ant is not None else sig]
    if huecos:
        huecos_reg.append((etiqueta, huecos[0], huecos[-1], len(huecos)))
    return dict(zip(periodos, out))


def _sin_acentos(t: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", str(t)) if not unicodedata.combining(c)).upper()


NOMBRE_BD = "BD_ BEL - IRR - MR_Proyeccion.xlsx"


def es_nombre_bd(nombre: str) -> bool:
    """True si el archivo es la BD proyectada: BD_ BEL - IRR - MR_Proyeccion.xlsx, sin importar acentos, mayusculas ni
    una marca de copia como " (1)" o " - copia". No cuenta las variantes de escenario (_PE12, _PE18, ProyeccionP), el
    BD_ RFV ni el archivo temporal ~$ que Excel deja mientras esta abierto."""
    n = re.sub(r"\s+", " ", _sin_acentos(nombre).strip())
    if n.startswith("~$") or not n.endswith(".XLSX"):
        return False
    base = re.sub(r"( ?\(\d+\)| - COPIA( ?\(\d+\))?)$", "", n[:-5].strip())
    return base == re.sub(r"\s+", " ", _sin_acentos(NOMBRE_BD[:-5]))


def buscar_bd(ruta=None, carpetas=()) -> Path:
    """Ruta de la BD proyectada. La indicada (argumento o config_local.RUTA_BD) si existe; si no, la primera que
    encuentra en este orden: para cada carpeta de ``carpetas``, su ``salidas`` (donde la deja proyeccion_reservas.py), la
    carpeta misma, la ``salidas`` de la carpeta de arriba y la de arriba. Dentro de ese orden gana el nombre exacto sobre
    las copias (" (1)", " - copia"). Si hay otras candidatas lo avisa. FileNotFoundError con las carpetas revisadas si
    no encuentra ninguna."""
    if ruta and Path(ruta).is_file():
        return Path(ruta).resolve()
    revisar = []
    if ruta:
        revisar.append(Path(ruta).resolve().parent)
    for c in carpetas:
        c = Path(c).resolve()
        for d in (c / "salidas", c, c.parent / "salidas", c.parent):
            if d not in revisar:
                revisar.append(d)
    hallados = []
    for d in revisar:
        try:
            archivos = sorted(f for f in d.iterdir() if f.is_file() and es_nombre_bd(f.name)) if d.is_dir() else []
        except OSError:
            archivos = []
        hallados += archivos
    if hallados:
        exactos = [f for f in hallados if _sin_acentos(f.name) == _sin_acentos(NOMBRE_BD)]
        bd = (exactos or hallados)[0].resolve()
        if ruta:
            print(f"   AVISO: no esta la BD en {ruta}; se usa {bd}")
        otras = [str(f) for f in hallados if f.resolve() != bd]
        if otras:
            print(f"   AVISO: se usa la BD {bd}; tambien hay: " + "; ".join(otras[:4]) + (" ..." if len(otras) > 4 else ""))
        return bd
    donde = "\n      ".join(str(d) for d in revisar) or "(ninguna carpeta indicada)"
    raise FileNotFoundError(("No esta la BD proyectada" + (f" en {ruta}" if ruta else "") +
                             ". Busque 'BD_ BEL - IRR - MR_Proyeccion.xlsx' en:\n      " + donde +
                             "\n   Copiala a la carpeta de los scripts o pon su ruta en config_local.py (RUTA_BD)."))


def importar_o_instalar(modulo: str, nombre_pip: str | None = None):
    """El modulo; si falta, lo instala con pip en el mismo Python que corre el script, como proyeccion_reservas.py.
    None si no se pudo instalar (sin internet, proxy o permisos): verificar_arranque lo reporta."""
    import importlib
    import os
    import site
    import subprocess
    import sys
    try:
        return importlib.import_module(modulo)
    except ImportError:
        pass
    print(f"Instalando {nombre_pip or modulo} ...", flush=True)
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", nombre_pip or modulo])
    except (subprocess.CalledProcessError, OSError):
        return None
    usuario = site.getusersitepackages()        # (pip pudo instalar en la carpeta del usuario)
    if os.path.isdir(usuario) and usuario not in sys.path:
        site.addsitedir(usuario)
    importlib.invalidate_caches()
    try:
        return importlib.import_module(modulo)
    except ImportError:
        return None


def _ramos(x) -> tuple:
    """Lista de ramos de la configuracion como enteros: acepta (31,), [31], 31, (31) o '31'."""
    if x is None:
        return ()
    if not isinstance(x, (list, tuple, set, frozenset)):
        x = (x,)
    out = []
    for v in x:
        try:
            out.append(int(float(v)))
        except (TypeError, ValueError):
            pass
    return tuple(out)


def _tabla_area(x, nombre: str):
    """(DataFrame, nombre) de una tabla del area dada ya leida (DataFrame) o como ruta de su CSV; (None, nombre) si no
    hay o la ruta no existe."""
    if x is None:
        return None, nombre
    if isinstance(x, pd.DataFrame):
        return x, nombre
    ruta = Path(str(x))
    return (pd.read_csv(ruta), ruta.name) if ruta.is_file() else (None, ruta.name)


DRIVER_ACCESS = "Microsoft Access Driver (*.mdb, *.accdb)"
ARCHIVOS_AREA = {           # reserva -> (variable de config_local, archivo dentro de esa carpeta o None si es el archivo)
    "RRC": [("CATALOGOS", None)] + [("CSV_AUXILIARES_RRC", f) for f in (
        "LlavesPol.csv", "AjManuales.csv", "Subramo.csv", "CesionPI.csv", "AFUN.csv", "zFrecuencias.csv",
        "TablaCesion_Esc1.csv", "Cesion ID Esp.csv")] + [("PPTO_TECNICO_RRC", None)],
    "SONR": [("CSV_AUXILIARES_SONR", "Subramo.csv"), ("PPTO_TECNICO_SONR", None)],
}
SALIDAS_AREA = {"RRC": ("RRC_esc.xlsx", "Parametros_usados_RRC.xlsx"),
                "SONR": ("SONR_esc.xlsx", "Parametros_usados_SONR.xlsx")}


def verificar_arranque(cfg, reserva: str, pyodbc_mod, carpeta) -> Path:
    """Antes de empezar, revisa lo que el script necesita y la BD no trae (pyodbc, el controlador de Access, la base
    Access, los archivos del area y que las salidas no esten abiertas en Excel) y junta todo lo que falte en un solo
    mensaje, para no enterarse a la mitad de la corrida. Regresa la ruta de la BD proyectada."""
    import sys
    problemas, bd = [], None
    try:
        bd = buscar_bd(getattr(cfg, "RUTA_BD", None), [carpeta])
    except FileNotFoundError as e:
        problemas.append(str(e))
    bits = 64 if sys.maxsize > 2 ** 32 else 32
    if pyodbc_mod is None:
        problemas.append("Falta pyodbc (conecta con la base Access) y no se pudo instalar solo. En la terminal de VS "
                         f"Code corre:  \"{sys.executable}\" -m pip install pyodbc")
    else:
        try:
            drivers = list(pyodbc_mod.drivers())
        except Exception:                       # (un pyodbc sin drivers(): se deja pasar y lo dira la conexion)
            drivers = None
        if drivers is not None and DRIVER_ACCESS not in drivers:
            problemas.append(f"Windows no tiene el controlador de ODBC '{DRIVER_ACCESS}' de {bits} bits (tu Python es "
                             f"de {bits} bits y deben coincidir). Instala el 'Microsoft Access Database Engine 2016 "
                             f"Redistributable' de {bits} bits. Controladores que hay: {', '.join(drivers) or 'ninguno'}")
    acc = getattr(cfg, "ACCESS_DBQ", None)
    if not acc or not Path(str(acc)).is_file():
        problemas.append(f"No se encuentra la base Access: {acc} (ACCESS_DBQ). Revisa la ruta y la conexion a la red "
                         "o a la VPN")
    pendientes = list(ARCHIVOS_AREA.get(reserva, []))
    if reserva == "SONR" and getattr(cfg, "MES", None) == 12:      # en diciembre el SONR no usa el presupuesto
        pendientes = [x for x in pendientes if x[0] != "PPTO_TECNICO_SONR"]
    if reserva == "RRC" and str(getattr(cfg, "MR_DESDE", "BD")).upper() != "BD":
        pendientes.append(("CSV_DURACION_RRC", None))
    carpetas_mal = set()
    for var, archivo in pendientes:
        base = getattr(cfg, var, None)
        if not base:
            problemas.append(f"Falta {var} en config_local.py")
            continue
        if archivo is None:
            if not Path(str(base)).is_file():
                problemas.append(f"No esta el archivo {base} ({var})")
        elif not Path(str(base)).is_dir():
            if var not in carpetas_mal:
                carpetas_mal.add(var)
                problemas.append(f"No existe la carpeta {base} ({var})")
        elif not (Path(str(base)) / archivo).is_file():
            problemas.append(f"No esta {archivo} en {base} ({var})")
    salida = Path(str(getattr(cfg, "CARPETA_SALIDA", None) or carpeta))
    for nombre in SALIDAS_AREA.get(reserva, ()):
        r = salida / nombre
        if r.is_file():
            try:
                with open(r, "r+b"):
                    pass
            except PermissionError:
                problemas.append(f"{r} esta abierto (en Excel?): cierralo, el script lo va a reescribir")
    if problemas:
        raise SystemExit(f"\nNo se puede correr el {reserva}; antes hay que resolver esto:\n" +
                         "\n".join(f"  {i}. {t}" for i, t in enumerate(problemas, 1)) +
                         f"\nLas rutas se cambian en {Path(carpeta) / 'config_local.py'}")
    return bd


class InsumosBD:
    """Lee la BD proyectada y entrega los insumos de los scripts del area."""

    def __init__(self, ruta_bd, ruta_diagnostico=None, usar_is_fa: bool = True, ramos_is_fa: dict | None = None,
                 tipo_real: str = "Real", tipo_proyeccion: str = "Proyección", verbose: bool = True):
        self.ruta_bd = Path(ruta_bd)
        if not self.ruta_bd.exists():
            raise FileNotFoundError(f"No esta la BD proyectada: {self.ruta_bd}")
        diag = [Path(ruta_diagnostico)] if ruta_diagnostico else []      # el indicado y, si no existe, el que esta
        diag.append(self.ruta_bd.with_name("Diagnostico_Proyeccion.xlsx"))  # junto a la BD
        self.ruta_diagnostico = next((d for d in diag if d.is_file()), diag[0])
        self.usar_is_fa = usar_is_fa
        self.ramos_is_fa = {k: _ramos(v) for k, v in (ramos_is_fa if ramos_is_fa is not None else RAMOS_IS_FA).items()}
        self.tipo_real, self.tipo_proyeccion = norm(tipo_real), norm(tipo_proyeccion)
        self.avisos: list[str] = []
        self._huecos: list[tuple] = []          # (etiqueta, desde, hasta, n) de los meses llenados con los vecinos
        self._series: dict = {}
        self._diag: dict | None = None
        self.fuentes: list[dict] = []          # (tablas del area) de donde salio cada celda: BD o area
        self._ret_fuera: dict = {}             # Factor_Ret fuera de [0, 1] (IRR mayor que el BEL) por ramo
        self._mr_negativo: dict = {}           # (reserva, ramo) -> meses con MR negativo en la BD (se usa 0)
        self._fnd: dict = {}                   # (ramo, mes) -> FND de la BD o nan (fnd_contrato)
        self._fnd_uso: dict = {}               # (reserva, ramo, mes) -> [con FND de la BD, del area por antiguedad,
                                               #                          del area porque la BD no trae FND]
        self._fnd_fuera: dict = {}             # ramo -> meses con FND de la BD fuera de [0, 1] (se acota)
        self._fnd_detalle: dict = {}           # (reserva, ramo, mes) -> FND del area, resultante, factor y prima / PRIMA N AÑOS
        self._fnd_incompleto: dict = {}        # reserva -> {mes: ultimo CALMONTH de la base} (sin prima del mes)
        self._pe_faltan: set = set()           # meses pedidos a la hoja PE_RAMO que no trae (contratos de 2027)
        self._sonr_nivel: dict = {}            # (ramo, mes) -> BEL del metodo sobre los contratos de la BD / BEL de la BD
        self._sonr_sin_nivel: dict = {}        # ramo -> meses en que el SONR queda con el metodo tal cual (y por que)
        wb = openpyxl.load_workbook(self.ruta_bd, read_only=True, data_only=True)
        try:
            self.hoja_hp = next((n for n in wb.sheetnames if norm(n).startswith("HPARAMETROS")), None)
            if self.hoja_hp is None or HOJA_MONTOS not in wb.sheetnames:
                raise ValueError(f"La BD debe traer las hojas HParametros_* y {HOJA_MONTOS}; trae {wb.sheetnames}")
            self._leer_hp(wb[self.hoja_hp])
            self._leer_montos(wb[HOJA_MONTOS])
            self._leer_is_fa(wb[HOJA_IS_FA] if HOJA_IS_FA in wb.sheetnames else None)
            self._leer_pe(wb[HOJA_PE] if HOJA_PE in wb.sheetnames else None)
        finally:
            wb.close()
        self.anio, self.mes = self.ultimo_real // 100, self.ultimo_real % 100
        self._validar_valores()
        if verbose:
            for linea in self.resumen():
                print("   " + linea, flush=True)

    def _validar_valores(self):
        """Comprueba al inicio que los montos proyectados (formulas) tengan valor guardado o diagnostico, para no fallar a
        media corrida: prueba BEL, IRR y MR de RRC y SONR del primer y ultimo mes proyectado en todos los ramos."""
        if self.primer_proyectado is None:
            return
        for p in (self.primer_proyectado, self.ultimo_proyectado):
            if p not in self.periodos_con_montos:
                continue
            for reserva in ("RRC", "SONR"):
                for conc in ("BEL", "IRR", "MR"):
                    for clave in RAMOS_SCRIPT:
                        if (norm(f"{reserva} {conc}"), p, self.ramo_bd(clave)) in self.montos:
                            self.monto(f"{reserva} {conc}", p, clave)       # (levanta ValueError con el mensaje claro)
        if self._diag and (self._diag["montos"] or self._diag["ind"]):
            self.avisos.append(f"los montos y factores proyectados se leyeron de {self.ruta_diagnostico.name} (la BD no "
                               "trae los valores guardados de las formulas)")

    # ------------------------------------------------------------------ lectura
    @staticmethod
    def _encabezado(ws, marcas: tuple, max_filas: int = 12):
        for i, fila in enumerate(ws.iter_rows(min_row=1, max_row=max_filas, values_only=True), 1):
            nombres = [norm(x) for x in fila]
            if all(norm(m) in nombres for m in marcas):
                return i, {norm(x): j for j, x in enumerate(fila) if x is not None}
        raise ValueError(f"No se encontro el encabezado con {marcas} en la hoja {ws.title}")

    @staticmethod
    def _clave_ramo(v) -> int | None:
        t = str(v).strip() if v is not None else ""
        if t.isdigit():
            return int(t)
        for clave, etiquetas in ETIQUETA_HP.items():
            if norm(t) in {norm(e) for e in etiquetas}:
                return clave
        return None

    def _leer_hp(self, ws):
        fila, col = self._encabezado(ws, ("Tipo de Indice", "Fecha", "Ramo", "LAG 1"))
        c_tipo, c_fecha, c_ramo = col[norm("Tipo de Indice")], col[norm("Fecha")], col[norm("Ramo")]
        c_par = {k: col.get(norm(n)) for k, n in PARAMETROS_HP.items()}
        c_lag = {k: col.get(norm(k)) for k in LAGS}
        faltan = [n for k, n in PARAMETROS_HP.items() if c_par[k] is None] + [k for k in LAGS if c_lag[k] is None]
        if faltan:
            raise ValueError(f"En {ws.title} faltan las columnas {faltan}")
        self.hp: dict = {}                       # (tipo, periodo, clave) -> {param: valor}
        self.tipos: dict = {}
        for r in ws.iter_rows(min_row=fila + 1, values_only=True):
            tipo = norm(r[c_tipo]) if c_tipo < len(r) else ""
            if not tipo:
                continue
            periodo = _num(r[c_fecha]) if c_fecha < len(r) else math.nan
            clave = self._clave_ramo(r[c_ramo]) if c_ramo < len(r) else None
            if math.isnan(periodo) or clave is None:
                continue
            periodo = int(periodo)
            d = {k: _num(r[c]) if c < len(r) else math.nan for k, c in c_par.items()}
            lags = [_num(r[c]) if c < len(r) else math.nan for c in c_lag.values()]
            max_previo = 0.0                       # (como el modelo principal) un LAG en 0 despues de que el patron
            for k, v in enumerate(lags):           # acumulado supero 50 % es marcador de faltante, en toda la corrida
                if not math.isnan(v) and v == 0 and (k == 0 or max_previo > 0.5):
                    lags[k] = math.nan
                elif not math.isnan(v):
                    max_previo = max(max_previo, v)
            d.update(dict(zip(LAGS, lags)))
            self.hp[(tipo, periodo, clave)] = d
            self.tipos.setdefault(tipo, set()).add(periodo)
        if self.tipo_real not in self.tipos:
            raise ValueError(f"{ws.title} no trae renglones '{self.tipo_real}'; tipos: {sorted(self.tipos)}")
        reales = self.tipos[self.tipo_real]
        proy = self.tipos.get(self.tipo_proyeccion, set())
        self.ultimo_real = mes_mas(min(proy), -1) if proy else max(reales)
        self.primer_proyectado = min(proy) if proy else None
        self.ultimo_proyectado = max(proy) if proy else None
        self.primer_periodo_hp = min(reales)

    def _leer_montos(self, ws):
        fila, col = self._encabezado(ws, ("CONCEPTO", "PERIODO", "TC"))
        self.col_ramo_bd = {h[4:]: c for h, c in col.items() if h.startswith("RAM_")}
        self.col_bloque = {}                     # (nombre, ramo_bd) -> columna
        for h, c in col.items():
            m = re.match(r"^(.*\S)\s+(\d+)$", h)
            if m and not h.startswith("RAM_") and m.group(2) in self.col_ramo_bd:
                self.col_bloque[(m.group(1), m.group(2))] = c
        c_con, c_per, c_tc = col["CONCEPTO"], col["PERIODO"], col["TC"]
        self.montos: dict = {}                   # (concepto, periodo, ramo_bd) -> valor o None
        self.bloques: dict = {}                  # (nombre, reserva, periodo, ramo_bd) -> valor o None
        self.tc: dict = {}
        self.periodos_montos: set = set()
        self.periodos_con_montos: set = set()    # meses con algun monto (los de prima sin montos quedan fuera)
        for r in ws.iter_rows(min_row=fila + 1, values_only=True):
            concepto = norm(r[c_con]) if c_con < len(r) else ""
            periodo = _num(r[c_per]) if c_per < len(r) else math.nan
            if not concepto or math.isnan(periodo):
                continue
            periodo = int(periodo)
            self.periodos_montos.add(periodo)
            tc = _num(r[c_tc]) if c_tc < len(r) else math.nan
            if not math.isnan(tc) and tc > 0:
                self.tc.setdefault(periodo, tc)
            for ramo_bd, c in self.col_ramo_bd.items():
                v = r[c] if c < len(r) else None
                self.montos[(concepto, periodo, ramo_bd)] = v
                if v is not None and v != "":
                    self.periodos_con_montos.add(periodo)
            if concepto in ("RRC BEL", "SONR BEL"):
                reserva = concepto.split()[0]
                for (nombre, ramo_bd), c in self.col_bloque.items():
                    self.bloques[(nombre, reserva, periodo, ramo_bd)] = r[c] if c < len(r) else None
        if not self.tc:
            raise ValueError(f"La columna TC de {ws.title} esta vacia")
        if self.periodos_con_montos:               # la historia y la proyeccion son contiguas: desde el primer mes con monto
            primero = min(self.periodos_con_montos)
            self.periodos_con_montos = {p for p in self.periodos_montos if p >= primero}

    def _leer_is_fa(self, ws):
        self.is_fa: dict = {}                    # (reserva, clave, periodo) -> {"media": v, "99.5": v}
        if ws is None:
            self.avisos.append(f"la BD no trae la hoja {HOJA_IS_FA}: el IS (FA) no se usa")
            return
        try:
            fila, col = self._encabezado(ws, ("Ramo", "MesProc", "IS RRC", "IS SONR"))
        except ValueError:
            self.avisos.append(f"la hoja {HOJA_IS_FA} no trae Ramo, MesProc, IS RRC e IS SONR: el IS (FA) no se usa")
            return
        c = {k: col.get(norm(k)) for k in ("Ramo", "MesProc", "IS RRC", "IS SONR", "IS RRC 99.5%", "IS SONR 99.5%")}
        for r in ws.iter_rows(min_row=fila + 1, values_only=True):
            clave = self._clave_ramo(r[c["Ramo"]]) if c["Ramo"] < len(r) else None
            periodo = _num(r[c["MesProc"]]) if c["MesProc"] < len(r) else math.nan
            if clave is None or math.isnan(periodo):
                continue
            periodo = int(periodo)
            for reserva, k_m, k_99 in (("RRC", "IS RRC", "IS RRC 99.5%"), ("SONR", "IS SONR", "IS SONR 99.5%")):
                media = _num(r[c[k_m]]) if c[k_m] is not None and c[k_m] < len(r) else math.nan
                p99 = _num(r[c[k_99]]) if c[k_99] is not None and c[k_99] < len(r) else math.nan
                if not math.isnan(media):
                    self.is_fa.setdefault((reserva, clave, periodo), {"media": media, "99.5": p99})   # primer renglon
        if not self.is_fa:
            self.avisos.append(f"la hoja {HOJA_IS_FA} no trae indices: el IS (FA) no se usa")

    def _leer_pe(self, ws):
        """Hoja PE_RAMO: prima tomada (PE) del mes por ramo en USD; real (PExRamo) hasta el ultimo mes real y FCST
        despues. self.pe = {(clave del script, mes): PE}."""
        self.pe: dict = {}
        if ws is None:
            return
        try:
            fila, col = self._encabezado(ws, ("PERIODO", "FUENTE"))
        except ValueError:
            return
        c_per = col[norm("PERIODO")]
        c_ram = {clave: col.get(norm(f"PE {self.ramo_bd(clave)}")) for clave in RAMOS_SCRIPT}
        for r in ws.iter_rows(min_row=fila + 1, values_only=True):
            per = _num(r[c_per]) if c_per < len(r) else math.nan
            if math.isnan(per):
                continue
            for clave, c in c_ram.items():
                if c is not None and c < len(r):
                    v = _num(r[c])
                    if not math.isnan(v):
                        self.pe[(clave, int(per))] = v

    def _cargar_diag(self):
        if self._diag is not None:
            return self._diag
        self._diag = {"ind": {}, "montos": {}}
        if not self.ruta_diagnostico.exists():
            return self._diag
        wb = openpyxl.load_workbook(self.ruta_diagnostico, read_only=True, data_only=True)
        try:
            if HOJA_DIAG_INDICADORES in wb.sheetnames:
                filas = wb[HOJA_DIAG_INDICADORES].iter_rows(values_only=True)
                enc = [str(h) if h is not None else "" for h in next(filas)]
                for f in filas:
                    d = dict(zip(enc, f))
                    if d.get("Reserva") in ("RRC", "SONR") and d.get("Periodo"):
                        self._diag["ind"][(d["Reserva"], str(d["Ramo"]), int(d["Periodo"]))] = d
            if HOJA_DIAG_MONTOS in wb.sheetnames:
                filas = wb[HOJA_DIAG_MONTOS].iter_rows(values_only=True)
                next(filas)
                for libro, concepto, periodo, ramo, valor in filas:
                    if libro == "DANOS" and periodo:
                        self._diag["montos"][(norm(concepto), int(periodo), str(ramo))] = valor
        finally:
            wb.close()
        return self._diag

    # ------------------------------------------------------------------ parametros de HParametros
    @staticmethod
    def ramo_bd(clave: int) -> str:
        return RAMO_BD.get(clave, str(clave))

    def _tipo_de(self, periodo: int) -> str:
        return self.tipo_real if periodo <= self.ultimo_real else self.tipo_proyeccion

    def serie_cruda(self, param: str, clave: int) -> dict:
        """Serie mensual del parametro sin llenar huecos: {periodo: valor o nan} (Real hasta el ultimo mes real,
        Proyección despues)."""
        fin = self.ultimo_proyectado or self.ultimo_real
        periodos = rango_meses(self.primer_periodo_hp, fin)
        return {p: self.hp.get((self._tipo_de(p), p, clave), {}).get(param, math.nan) for p in periodos}

    def serie(self, param: str, clave: int) -> dict:
        """Serie mensual del parametro (IS_RRC, IS_RRC_99, IS_SONR, IS_SONR_99 o 'LAG k') para el ramo, Real hasta el
        ultimo mes real y Proyección despues, con los huecos llenados; {periodo: valor} (nan si el ramo no lo trae)."""
        k = (param, clave)
        if k not in self._series:
            fin = self.ultimo_proyectado or self.ultimo_real
            periodos = rango_meses(self.primer_periodo_hp, fin)
            crudo = {p: self.hp.get((self._tipo_de(p), p, clave), {}).get(param, math.nan) for p in periodos}
            self._series[k] = _interpolar(crudo, periodos, self._huecos, f"{param} ramo {clave}")
        return self._series[k]

    def parametro(self, param: str, clave: int, periodo: int) -> float:
        s = self.serie(param, clave)
        if periodo in s:
            return s[periodo]
        if not s:
            return math.nan
        return s[min(s)] if periodo < min(s) else s[max(s)]      # fuera del rango: el extremo mas cercano

    def indice(self, reserva: str, clave: int, periodo: int, cual: str = "media") -> float:
        """IS media o 99.5 de la reserva para el ramo y mes: el IS (FA) en los ramos de RAMOS_IS_FA en los meses
        proyectados (si lo trae), si no el de HParametros."""
        if self.usar_is_fa and periodo > self.ultimo_real and clave in self.ramos_is_fa.get(reserva, ()):
            fa = self.is_fa.get((reserva, clave, periodo))
            if fa and not math.isnan(fa.get(cual, math.nan)):
                return fa[cual]
        base = "IS_RRC" if reserva == "RRC" else "IS_SONR"
        return self.parametro(base if cual == "media" else base + "_99", clave, periodo)

    def fuente_indice(self, reserva: str, clave: int, periodo: int) -> str:
        if self.usar_is_fa and periodo > self.ultimo_real and clave in self.ramos_is_fa.get(reserva, ()) \
                and (reserva, clave, periodo) in self.is_fa:
            return "IS (FA)"
        return self._tipo_de(periodo)

    # ------------------------------------------------------------------ valores de la hoja de montos
    def monto(self, concepto: str, periodo: int, clave: int) -> float:
        """Monto USD del concepto ('RRC BEL', 'SONR IRR', ...) del ramo y mes: el valor que guardo Excel; si la celda
        es formula sin valor, el del diagnostico; si tampoco, error."""
        ramo_bd = self.ramo_bd(clave)
        if periodo not in self.periodos_con_montos:        # (meses de prima sin montos: no hay dato)
            return math.nan
        v = _num(self.montos.get((norm(concepto), periodo, ramo_bd)))
        if math.isnan(v) and (norm(concepto), periodo, ramo_bd) in self.montos:
            v = _num(self._cargar_diag()["montos"].get((norm(concepto), periodo, ramo_bd)))
            if math.isnan(v):
                raise ValueError(f"{concepto} {periodo} ramo {clave}: la celda es formula sin valor guardado. Abre la BD "
                                 f"en Excel y guardala, o pon {self.ruta_diagnostico.name} junto a la BD")
        return v

    def bloque(self, nombre: str, reserva: str, clave: int, periodo: int) -> float:
        ramo_bd = self.ramo_bd(clave)
        v = _num(self.bloques.get((nombre, reserva, periodo, ramo_bd)))
        if math.isnan(v) and (nombre, reserva, periodo, ramo_bd) in self.bloques:
            d = self._cargar_diag()["ind"].get((reserva, ramo_bd, periodo))
            v = _num(d.get(nombre)) if d else math.nan
        return v

    def factor_gasto(self, clave: int, periodo: int) -> float:
        """Ind. Gasto del RRC = FACTOR GTO (GTO / PND) del ramo y mes."""
        v = self.bloque("FACTOR GTO", "RRC", clave, periodo)
        if math.isnan(v) and periodo <= self.ultimo_real and periodo in self.periodos_con_montos:
            bel, gto = self.monto("RRC BEL", periodo, clave), self.monto("RRC GTO", periodo, clave)
            is_ = self.parametro("IS_RRC", clave, periodo)
            v = gto * is_ / bel if bel and not math.isnan(bel) and abs(bel) >= UMBRAL_CERO and not math.isnan(is_) else math.nan
        return v

    def _factor_mr_crudo(self, reserva: str, clave: int, periodo: int) -> float:
        v = self.bloque("FACTOR MR", reserva, clave, periodo)
        if math.isnan(v) and periodo <= self.ultimo_real and periodo in self.periodos_con_montos:
            bel, mr = self.monto(f"{reserva} BEL", periodo, clave), self.monto(f"{reserva} MR", periodo, clave)
            is_ = self.parametro("IS_RRC" if reserva == "RRC" else "IS_SONR", clave, periodo)
            v = mr * is_ / bel if bel and not math.isnan(bel) and abs(bel) >= UMBRAL_CERO and not math.isnan(is_) else math.nan
        return v

    def factor_mr(self, reserva: str, clave: int, periodo: int) -> float:
        """FACTOR MR = MR / PND (RRC) o MR / PD (SONR) del ramo y mes: el bloque de la hoja de montos (en un mes real sin
        valor guardado, MR / (BEL / IS) de los montos); en un mes sin BEL, el del mes mas cercano con dato."""
        k = ("FACTOR_MR", reserva, clave)
        if k not in self._series:
            periodos = sorted(self.periodos_con_montos)
            crudo = {p: self._factor_mr_crudo(reserva, clave, p) for p in periodos}
            self._series[k] = _interpolar(crudo, periodos, self._huecos, f"FACTOR MR {reserva} ramo {clave}")
        s = self._series[k]
        if periodo in s:
            return s[periodo]
        return s[min(s)] if periodo < min(s) else s[max(s)]

    def _factor_ret_crudo(self, clave: int, periodo: int) -> float:
        bel = _num(self.montos.get(("SONR BEL", periodo, self.ramo_bd(clave))))
        irr = _num(self.montos.get(("SONR IRR", periodo, self.ramo_bd(clave))))
        if not math.isnan(bel) and not math.isnan(irr) and bel >= UMBRAL_CERO:
            return 1.0 - irr / bel
        if periodo > self.ultimo_real and periodo in self.periodos_con_montos:
            try:
                bel_p = self.monto("SONR BEL", periodo, clave)
            except ValueError:
                bel_p = math.nan
            if math.isnan(bel_p) or abs(bel_p) < UMBRAL_CERO:  # la BD no proyecta IBNR en ese mes: no hay retencion que leer
                return math.nan
            ces, fm = self.bloque("CESION", "SONR", clave, periodo), self.bloque("FACTOR MR", "SONR", clave, periodo)
            is_b = self.indice("SONR", clave, periodo)
            if not math.isnan(ces) and not math.isnan(fm) and not math.isnan(is_b) and is_b > 0:
                return 1.0 - ces * (1.0 + fm / is_b)           # IRR / BEL = CESION x BRUTO / BEL = CESION x (1 + FM / IS)
        return math.nan

    def factor_ret(self, clave: int, periodo: int) -> float:
        """Factor de retencion del SONR (IRR = BEL x (1 - Factor_Ret)) = 1 - IRR / BEL del ramo y mes; en un mes sin
        BEL (p. ej. el ramo 80 en la historia) el del mes mas cercano con dato; acotado a [0, 1]."""
        k = ("FACTOR_RET", clave)
        if k not in self._series:
            periodos = sorted(self.periodos_con_montos)
            crudo = {p: self._factor_ret_crudo(clave, p) for p in periodos}
            s = _interpolar(crudo, periodos, self._huecos, f"Factor_Ret SONR ramo {clave}")
            fuera = [p for p, v in s.items() if not math.isnan(v) and not 0 <= v <= 1]
            if fuera:
                self.avisos.append(f"Factor_Ret SONR ramo {clave}: {len(fuera)} mes(es) con IRR mayor que el BEL en la BD "
                                   f"({fuera[0]} a {fuera[-1]}): retencion 0 (todo cedido). Revisalo; con "
                                   f"RAMOS_FACTOR_RET_CSV = ({clave},) se toma el Factor_Ret del area")
            self._series[k] = {p: (min(max(v, 0.0), 1.0) if not math.isnan(v) else v) for p, v in s.items()}
        s = self._series[k]
        if periodo in s:
            return s[periodo]
        return s[min(s)] if periodo < min(s) else s[max(s)]

    # ------------------------------------------------------------------ tablas para los scripts
    def parametros_rrc(self, anio: int | None = None, csv_duracion=None) -> pd.DataFrame:
        """ParametrosMens del RRC: un renglon por ramo con 'IS Bel Media-m', 'IS Bel 99.5%-m', 'Ind. Gasto-m' y
        'FACTOR MR-m' (MR / PND de la BD) para m = 1..12 del ano (mes de valuacion), mas Pesos_dur, Resto Monedas_dur,
        Pesos_ret y Resto Monedas_ret del CSV del area (csv_duracion) si se indica."""
        anio = anio or self.anio
        filas = []
        for clave in RAMOS_SCRIPT:
            f = {"Ramo": clave}
            for m in range(1, 13):
                p = anio * 100 + m
                f[f"IS Bel Media-{m}"] = self.indice("RRC", clave, p, "media")
                f[f"IS Bel 99.5%-{m}"] = self.indice("RRC", clave, p, "99.5")
                f[f"Ind. Gasto-{m}"] = self.factor_gasto(clave, p)
                f[f"FACTOR MR-{m}"] = self.factor_mr("RRC", clave, p)
            filas.append(f)
        df = pd.DataFrame(filas)
        dur = ["Pesos_dur", "Resto Monedas_dur", "Pesos_ret", "Resto Monedas_ret"]
        if csv_duracion and Path(csv_duracion).exists():
            csv = pd.read_csv(csv_duracion)
            if all(c in csv.columns for c in dur + ["Ramo"]):
                csv = csv[["Ramo"] + dur].drop_duplicates("Ramo")
                csv["Ramo"] = pd.to_numeric(csv["Ramo"], errors="coerce")
                df = df.merge(csv, how="left", on="Ramo")
                sin = df.loc[df[dur].isna().any(axis=1), "Ramo"].tolist()
                if sin:
                    self.avisos.append(f"duracion / retencion del MR: sin dato en {Path(csv_duracion).name} para los "
                                       f"ramos {sin}")
            else:
                self.avisos.append(f"{Path(csv_duracion).name} no trae las columnas {dur}: el MR del RRC no se puede "
                                   "calcular")
        else:
            self.avisos.append("sin CSV de duracion / retencion (ParametrosMens del area): solo hace falta para el MR "
                               "del RRC con MR_DESDE = 'AREA' (con 'BD' no se usa)")
        for c in dur:
            if c not in df.columns:
                df[c] = math.nan
        return df

    def is_cat(self, csv_is_cat=None) -> pd.DataFrame:
        """IS_Cat del RRC: 'IS Bel Media' = mes (CALMONTH), '71' y '73' = Ind Sin RRC de TEV e Hidro de ese mes. Los
        meses que HParametros no trae (anteriores a la hoja o vacios en ella) salen del CSV del area si se indica; si no,
        los anteriores llevan el primer mes de HParametros y los vacios se llenan con los meses vecinos."""
        fin = self.ultimo_proyectado or self.ultimo_real
        periodos = rango_meses(self.primer_periodo_hp, fin)
        csv_val: dict = {}
        if csv_is_cat and Path(csv_is_cat).exists():
            csv = pd.read_csv(csv_is_cat)
            if all(c in csv.columns for c in ("IS Bel Media", "71", "73")):
                for m, a, b in zip(pd.to_numeric(csv["IS Bel Media"], errors="coerce"),
                                   pd.to_numeric(csv["71"], errors="coerce"), pd.to_numeric(csv["73"], errors="coerce")):
                    if not math.isnan(m):
                        csv_val[int(m)] = {71: a, 73: b}
        cols = {}
        for clave in (71, 73):
            crudo = self.serie_cruda("IS_RRC", clave)
            del_csv = [p for p in periodos if math.isnan(crudo[p]) and not math.isnan(csv_val.get(p, {}).get(clave, math.nan))]
            for p in del_csv:
                crudo[p] = csv_val[p][clave]
            if del_csv:
                self.avisos.append(f"IS_Cat ramo {clave}: {len(del_csv)} mes(es) sin indice en HParametros ({del_csv[0]} a "
                                   f"{del_csv[-1]}) tomados de {Path(csv_is_cat).name}")
            lleno = _interpolar(crudo, periodos, self._huecos, f"IS_Cat ramo {clave}")
            cols[clave] = [lleno[p] for p in periodos]
        df = pd.DataFrame({"IS Bel Media": periodos, "71": cols[71], "73": cols[73]})
        previos = pd.DataFrame({"IS Bel Media": [m for m in sorted(csv_val) if m < self.primer_periodo_hp]})
        if len(previos):
            previos["71"] = [csv_val[m][71] for m in previos["IS Bel Media"]]
            previos["73"] = [csv_val[m][73] for m in previos["IS Bel Media"]]
            df = pd.concat([previos, df], ignore_index=True)
        else:
            inicio = mes_mas(self.primer_periodo_hp, -12 * 12)
            extra = rango_meses(inicio, mes_mas(self.primer_periodo_hp, -1))
            df = pd.concat([pd.DataFrame({"IS Bel Media": extra, "71": df["71"].iloc[0], "73": df["73"].iloc[0]}), df],
                           ignore_index=True)
            self.avisos.append(f"IS_Cat (71 y 73): los meses anteriores a {self.primer_periodo_hp} llevan el indice de "
                               "ese primer mes (no hay CSV del area con los anteriores)")
        return df.reset_index(drop=True)

    def param_sonr(self, desde: int | None = None, hasta: int | None = None, csv_param_sonr=None,
                   ramos_factor_ret_csv: tuple = ()) -> pd.DataFrame:
        """ParamSONR: un renglon por mes y ramo con Llave 'AAAAMM-ramo', Factor_Ret, Factor_MR (MR / PD de la BD), Ind
        Sin SONR Media, Ind Sin SONR 99.5% y LAG 1 a LAG 10 (los parametros del mes de valuacion). En los ramos de ramos_factor_ret_csv el Factor_Ret
        sale del ParamSONR del area (csv_param_sonr, por Llave; en un mes que el CSV no trae, el ultimo que trae)."""
        desde = desde or self.primer_periodo_hp
        hasta = hasta or self.ultimo_proyectado or self.ultimo_real
        fr_csv: dict = {}
        if ramos_factor_ret_csv:
            if csv_param_sonr and Path(csv_param_sonr).exists():
                csv = pd.read_csv(csv_param_sonr)
                if "Llave" in csv.columns and "Factor_Ret" in csv.columns:
                    for llave, fr in zip(csv["Llave"].astype(str), pd.to_numeric(csv["Factor_Ret"], errors="coerce")):
                        partes = llave.split("-")
                        if len(partes) == 2 and partes[0].isdigit() and partes[1].isdigit() and not math.isnan(fr):
                            fr_csv[(int(partes[1]), int(partes[0]))] = float(fr)
                if not fr_csv:
                    self.avisos.append(f"{Path(csv_param_sonr).name} no trae Llave y Factor_Ret: el Factor_Ret de los ramos "
                                       f"{list(ramos_factor_ret_csv)} sale de la BD")
                else:
                    self.avisos.append(f"Factor_Ret de los ramos {list(ramos_factor_ret_csv)} tomado de "
                                       f"{Path(csv_param_sonr).name} (RAMOS_FACTOR_RET_CSV), no de la BD")
            else:
                self.avisos.append(f"RAMOS_FACTOR_RET_CSV = {list(ramos_factor_ret_csv)} pero no esta el CSV del ParamSONR "
                                   "del area: el Factor_Ret sale de la BD")

        def factor_ret_de(clave, p):
            if clave in ramos_factor_ret_csv and fr_csv:
                meses = sorted(m for (c, m) in fr_csv if c == clave)
                if meses:
                    m = p if p in meses else max([m for m in meses if m <= p] or [meses[0]])
                    return fr_csv[(clave, m)]
            return self.factor_ret(clave, p)
        filas = []
        primer_dato = {}
        for clave in RAMOS_SONR:
            cr = self.serie_cruda("IS_SONR", clave)
            con = [p for p, v in cr.items() if not math.isnan(v)]
            primer_dato[clave] = min(con) if con else None
        for p in rango_meses(desde, hasta):
            for clave in RAMOS_SONR:
                if primer_dato[clave] is None or p < primer_dato[clave]:   # (antes del primer dato del ramo no hay parametro)
                    continue
                lags = {k: self.parametro(k, clave, p) for k in LAGS}
                media = self.indice("SONR", clave, p, "media")
                if math.isnan(media) and all(math.isnan(v) for v in lags.values()):
                    continue
                filas.append({"Llave": f"{p}-{clave}", "AñoMes": p, "Ramo": clave,
                              "Factor_Ret": factor_ret_de(clave, p), "Factor_MR": self.factor_mr("SONR", clave, p),
                              "Ind Sin SONR Media": media, "Ind Sin SONR 99.5%": self.indice("SONR", clave, p, "99.5"),
                              **lags, "Fuente IS": self.fuente_indice("SONR", clave, p)})
        return pd.DataFrame(filas)

    # ------------------------------------------------------------------ tablas del area con nuestros indices
    # Toman la tabla del area tal cual viene (sus renglones y columnas) y cambian cada celda por la nuestra solo donde
    # la BD trae el dato; donde no lo trae se queda el valor del area. No se llenan huecos con meses vecinos.
    def _hp_crudo(self, param: str, clave: int, periodo: int) -> float:
        return _num(self.hp.get((self._tipo_de(periodo), periodo, clave), {}).get(param, math.nan))

    def _is_crudo(self, reserva: str, clave: int, periodo: int, cual: str = "media"):
        """(valor, fuente) del IS sin llenar huecos: el IS (FA) en RAMOS_IS_FA en los meses proyectados, si no el de
        HParametros (Real hasta el ultimo mes real, Proyección despues); nan si la BD no lo trae."""
        if self.usar_is_fa and periodo > self.ultimo_real and clave in self.ramos_is_fa.get(reserva, ()):
            fa = self.is_fa.get((reserva, clave, periodo))
            if fa and not math.isnan(fa.get(cual, math.nan)):
                return fa[cual], "IS (FA)"
        base = ("IS_RRC" if reserva == "RRC" else "IS_SONR") + ("" if cual == "media" else "_99")
        return self._hp_crudo(base, clave, periodo), "BD " + self._tipo_de(periodo).capitalize()

    def _bd_modela(self, reserva: str, clave: int, periodo: int) -> bool:
        """True si la hoja de montos trae BEL de la reserva en el ramo y mes; si no, sus factores no son dato."""
        try:
            v = self.monto(f"{reserva} BEL", periodo, clave)
        except ValueError:
            return False
        return not math.isnan(v) and abs(v) >= UMBRAL_CERO

    def _factor_crudo(self, tipo: str, reserva: str, clave: int, periodo: int) -> float:
        """GTO (GTO / PND), MR (MR / PND o PD; un MR negativo cuenta como 0) o RET (1 - IRR / BEL del SONR, acotado a
        [0, 1]) del ramo y mes, sin llenar huecos; nan solo si la BD no modela la reserva en ese ramo y mes (sin BEL).
        Un 0 de la BD es dato: es lo que modela la BD."""
        if not self._bd_modela(reserva, clave, periodo):
            return math.nan
        if tipo == "GTO":
            v = self.factor_gasto(clave, periodo)
        elif tipo == "MR":
            v = self._factor_mr_crudo(reserva, clave, periodo)
            if not math.isnan(v) and v < 0:
                self._mr_negativo.setdefault((reserva, clave), []).append(periodo)
                v = 0.0
        else:
            v = self._factor_ret_crudo(clave, periodo)
            if not math.isnan(v) and not 0 <= v <= 1:
                self._ret_fuera.setdefault(clave, []).append(periodo)
                v = min(max(v, 0.0), 1.0)
        return v

    def _anotar(self, tabla: str, clave, periodo, columna: str, valor, fuente: str, cuenta: dict, sin: dict):
        self.fuentes.append({"Tabla": tabla, "Ramo": clave, "Mes": periodo, "Columna": columna,
                             "Valor": valor, "Fuente": fuente})
        if fuente.startswith("Area") and (valor is None or (isinstance(valor, float) and math.isnan(valor))) \
                and "formula" not in fuente:
            fuente = "Sin dato (ni la BD ni el area)"
            self.fuentes[-1]["Fuente"] = fuente
        cuenta[fuente] = cuenta.get(fuente, 0) + 1
        if fuente.startswith(("Area", "Sin dato")):
            motivo = "por RAMOS_FACTOR_RET_AREA" if "RAMOS_FACTOR_RET_AREA" in fuente else "la BD no lo trae"
            sin.setdefault((clave, columna.split("-")[0].strip(), motivo), set()).add(periodo)

    def _aviso_tabla(self, tabla: str, cuenta: dict, sin: dict, nota_mr: str = ""):
        """Un aviso con el conteo de celdas por fuente y uno por ramo y meses con lo que se quedo del area."""
        total = sum(cuenta.values())
        partes = ", ".join(f"{n} {f}" for f, n in sorted(cuenta.items(), key=lambda x: -x[1]))
        self.avisos.append(f"{tabla}: {total} celdas; {partes}")
        grupos: dict = {}
        for (clave, col, motivo), meses in sin.items():
            grupos.setdefault((clave, tuple(sorted(meses)), motivo), []).append(col)
        sin_dato = {(f["Ramo"], f["Columna"].split("-")[0].strip()) for f in self.fuentes
                    if f["Tabla"] == tabla and f["Fuente"].startswith("Sin dato")}
        for (clave, meses, motivo), cols in sorted(grupos.items(), key=lambda x: (x[0][0] or 0, x[0][1], x[0][2])):
            rango = f"{meses[0]}" if len(meses) == 1 else f"{meses[0]} a {meses[-1]}"
            nota = nota_mr if any(c.upper().startswith(("FACTOR MR", "FACTOR_MR")) for c in cols) else ""
            vacias = [c for c in cols if (clave, c) in sin_dato]
            if vacias:
                self.avisos.append(f"{tabla}: ramo {clave}, {len(meses)} mes(es) ({rango}): {', '.join(cols)} sin dato "
                                   "en la BD ni en la tabla del area (vacio, como en el script original)")
            else:
                self.avisos.append(f"{tabla}: ramo {clave}, {len(meses)} mes(es) ({rango}): {', '.join(cols)} del area "
                                   f"({motivo}){nota}")

    def _aviso_mr_negativo(self, reserva: str):
        for (res, clave), ps in sorted(self._mr_negativo.items()):
            if res == reserva:
                ps = sorted(set(ps))
                self.avisos.append(f"FACTOR MR {reserva} ramo {clave}: {len(ps)} mes(es) con MR negativo en la BD "
                                   f"({ps[0]} a {ps[-1]}): se usa 0")
        self._mr_negativo = {k: v for k, v in self._mr_negativo.items() if k[0] != reserva}

    @staticmethod
    def _ramo_area(v) -> int | None:
        try:
            return int(float(v))
        except (TypeError, ValueError):
            return None

    def parametros_rrc_area(self, anio: int, area) -> pd.DataFrame:
        """ParametrosMens del area (DataFrame o ruta) con nuestros indices donde la BD los trae, m = 1..12 del ano:
        IS Bel Media-m e IS Bel 99.5%-m (IS de FA en RAMOS_IS_FA en los meses proyectados), Ind. Gasto-m (FACTOR GTO;
        si no, el Ind. Gasto del area) y FACTOR MR-m (MR / PND; si no, vacio y el script usa su formula de capital).
        Duracion, retencion y demas columnas quedan como vienen."""
        df, nombre = _tabla_area(area, "ParametrosMens del area")
        if df is None:
            raise FileNotFoundError(f"No se pudo leer {nombre}")
        df = df.copy()
        for m in range(1, 13):
            for c in (f"IS Bel Media-{m}", f"IS Bel 99.5%-{m}", f"Ind. Gasto-{m}", f"FACTOR MR-{m}"):
                if c not in df.columns:
                    df[c] = math.nan
                df[c] = pd.to_numeric(df[c], errors="coerce").astype(float)
        cuenta, sin, tabla = {}, {}, f"ParametrosMens RRC {anio}"
        for i, v in df["Ramo"].items():
            clave = self._ramo_area(v)
            for m in range(1, 13):
                p = anio * 100 + m
                for col, cual in ((f"IS Bel Media-{m}", "media"), (f"IS Bel 99.5%-{m}", "99.5")):
                    x, fuente = self._is_crudo("RRC", clave, p, cual) if clave else (math.nan, "")
                    if math.isnan(x):
                        x, fuente = df.at[i, col], "Area (la BD no lo trae)"
                    df.at[i, col] = x
                    self._anotar(tabla, clave, p, col, x, fuente, cuenta, sin)
                g = self._factor_crudo("GTO", "RRC", clave, p) if clave else math.nan
                fuente = "BD FACTOR GTO"
                if math.isnan(g):
                    g = _num(df.at[i, "Ind. Gasto"]) if "Ind. Gasto" in df.columns else math.nan
                    fuente = "Area (la BD no lo trae)"
                df.at[i, f"Ind. Gasto-{m}"] = g
                self._anotar(tabla, clave, p, f"Ind. Gasto-{m}", g, fuente, cuenta, sin)
                fm = self._factor_crudo("MR", "RRC", clave, p) if clave else math.nan
                df.at[i, f"FACTOR MR-{m}"] = fm
                self._anotar(tabla, clave, p, f"FACTOR MR-{m}", fm,
                             "BD FACTOR MR" if not math.isnan(fm) else "Area (formula de capital)", cuenta, sin)
        self._aviso_tabla(tabla, cuenta, sin, ": su MR sale de la formula de capital del script")
        self._aviso_mr_negativo("RRC")
        return df

    def is_cat_area(self, area) -> pd.DataFrame:
        """IS_Cat del area (DataFrame o ruta; 'IS Bel Media' = mes del contrato, '71' y '73') con el Ind Sin RRC de TEV
        e Hidro de la BD en los meses que la BD trae; los demas meses quedan como vienen. Los meses que la BD trae y la
        tabla del area no, se agregan."""
        df, nombre = _tabla_area(area, "IS_Cat del area")
        if df is None:
            df = pd.DataFrame(columns=["IS Bel Media", "71", "73"])
        df = df.copy()
        df.columns = [str(c) for c in df.columns]
        for c in ("71", "73"):
            if c not in df.columns:
                df[c] = math.nan
            df[c] = pd.to_numeric(df[c], errors="coerce").astype(float)
        meses = pd.to_numeric(df["IS Bel Media"], errors="coerce")
        fin = self.ultimo_proyectado or self.ultimo_real
        ya = set(meses.dropna().astype(int))
        nuevos = [p for p in rango_meses(self.primer_periodo_hp, fin) if p not in ya
                  and any(not math.isnan(self._is_crudo("RRC", c, p)[0]) for c in (71, 73))]
        if nuevos:
            df = pd.concat([df, pd.DataFrame({"IS Bel Media": nuevos})], ignore_index=True)
            meses = pd.to_numeric(df["IS Bel Media"], errors="coerce")
        cuenta, sin, tabla = {}, {}, "IS_Cat RRC (71 y 73)"
        for i, m in meses.items():
            if math.isnan(m):
                continue
            p = int(m)
            for clave in (71, 73):
                x, fuente = self._is_crudo("RRC", clave, p)
                if math.isnan(x):
                    x, fuente = df.at[i, str(clave)], "Area (la BD no lo trae)"
                df.at[i, str(clave)] = x
                self._anotar(tabla, clave, p, str(clave), x, fuente, cuenta, sin)
        self._aviso_tabla(tabla, cuenta, sin)
        if nuevos:
            self.avisos.append(f"{tabla}: {len(nuevos)} mes(es) que la tabla del area no traia ({nuevos[0]} a "
                               f"{nuevos[-1]}) se agregaron con el indice de la BD; con la tabla del area salian vacios")
        return df.sort_values("IS Bel Media", kind="stable").reset_index(drop=True)

    def param_sonr_area(self, area, ramos_factor_ret_area: tuple = (), anio: int | None = None) -> pd.DataFrame:
        """ParamSONR del area (DataFrame o ruta; Llave 'AAAAMM-ramo') con nuestros parametros donde la BD los trae: Ind
        Sin SONR Media y 99.5% (IS de FA en RAMOS_IS_FA en los meses proyectados), LAG 1 a LAG 10, Factor_Ret (1 - IRR /
        BEL del SONR; en los ramos de ramos_factor_ret_area, el del area) y Factor_MR (MR / PD; si no, vacio y el script
        usa su formula). Las llaves de los ramos de SONR que la tabla del area no trae en sus meses se agregan."""
        df, nombre = _tabla_area(area, "ParamSONR del area")
        if df is None:
            df = pd.DataFrame(columns=["Llave"])
        df = df.copy()
        cols = ["Factor_Ret", "Factor_MR", "Ind Sin SONR Media", "Ind Sin SONR 99.5%"] + LAGS
        for c in cols:
            if c not in df.columns:
                df[c] = math.nan
            df[c] = pd.to_numeric(df[c], errors="coerce").astype(float)

        ramos_factor_ret_area = _ramos(ramos_factor_ret_area)

        def partes(llave):
            t = str(llave).split("-")
            if len(t) != 2 or not t[0].strip().isdigit():
                return None, None
            r = _num(t[1])
            return (int(t[0]), int(r)) if math.isfinite(r) and r == int(r) else (None, None)
        llaves = [partes(x) for x in df["Llave"]]
        exactas = {(p, c) for (p, c), x in zip(llaves, df["Llave"]) if p and str(x) == f"{p}-{c}"}
        raras = [str(x) for (p, c), x in zip(llaves, df["Llave"]) if not p or str(x) != f"{p}-{c}"]
        if raras:
            self.avisos.append(f"ParamSONR del area: {len(raras)} llave(s) con otro formato que AAAAMM-ramo (p. ej. "
                               f"'{raras[0]}'); el script solo cruza las de ese formato")
        meses = sorted({p for p, _ in llaves if p} | (set(rango_meses(anio * 100 + 1, anio * 100 + 12)) if anio else set())) \
            or rango_meses(self.anio * 100 + 1, self.anio * 100 + 12)
        parecidas = {pc: i for i, pc in zip(df.index, llaves) if pc[0]}         # (fila del area con otra llave)
        por_ramo: dict = {}                                                    # ramo -> [(mes, fila)] de la tabla del area
        for i, (pp, cc) in zip(df.index, llaves):
            if pp:
                por_ramo.setdefault(cc, []).append((pp, i))

        respaldo_anterior: dict = {}                                         # (mes, ramo) -> mes del area que se uso

        def fila_area(p, c):
            """Fila del area para la llave: la misma (con otro formato); si la tabla no trae el ano, la del mismo mes de un
            ano anterior (el archivo de 2026 para 2027, como previo); si tampoco, la del ultimo mes del ramo hasta p (si no,
            la primera)."""
            if (p, c) in parecidas:
                return parecidas[(p, c)]
            meses_ramo = dict(por_ramo.get(c, []))
            for k in range(1, 6) if p // 100 not in {m // 100 for m in meses_ramo} else ():
                if mes_mas(p, -12 * k) in meses_ramo:
                    respaldo_anterior[(p, c)] = mes_mas(p, -12 * k)
                    return meses_ramo[mes_mas(p, -12 * k)]
            filas = sorted(por_ramo.get(c, []))
            antes = [i for m, i in filas if m <= p]
            return antes[-1] if antes else (filas[0][1] if filas else None)

        def ret_disponible(p, c):
            fa = fila_area(p, c)
            del_area = fa is not None and not math.isnan(_num(df.at[fa, "Factor_Ret"]))
            if c in ramos_factor_ret_area:
                return del_area
            return del_area or not math.isnan(self._factor_crudo("RET", "SONR", c, p))
        extra = [(p, c) for p in meses for c in RAMOS_SONR if (p, c) not in exactas
                 and (not math.isnan(self._is_crudo("SONR", c, p)[0]) or fila_area(p, c) is not None) and ret_disponible(p, c)]
        self._ret_fuera = {}
        if extra:
            nuevas = pd.DataFrame({"Llave": [f"{p}-{c}" for p, c in extra]})
            respaldo = [fila_area(p, c) for p, c in extra]
            for c in cols:
                nuevas[c] = [df.at[i, c] if i is not None else math.nan for i in respaldo]
            df = pd.concat([df, nuevas], ignore_index=True)
            llaves += extra
            for c in cols:
                df[c] = pd.to_numeric(df[c], errors="coerce").astype(float)
        cuenta, sin, tabla = {}, {}, "ParamSONR"
        for i, (p, clave) in zip(df.index, llaves):
            if p is None:
                continue
            valores = {"Ind Sin SONR Media": self._is_crudo("SONR", clave, p, "media"),
                       "Ind Sin SONR 99.5%": self._is_crudo("SONR", clave, p, "99.5")}
            for k in LAGS:
                valores[k] = (self._hp_crudo(k, clave, p), "BD " + self._tipo_de(p).capitalize())
            valores["Factor_Ret"] = ((math.nan, "") if clave in ramos_factor_ret_area
                                     else (self._factor_crudo("RET", "SONR", clave, p), "BD 1 - IRR / BEL"))
            valores["Factor_MR"] = (self._factor_crudo("MR", "SONR", clave, p), "BD FACTOR MR")
            for col, (x, fuente) in valores.items():
                if math.isnan(x):
                    x = df.at[i, col]
                    fuente = ("Area (formula de capital)" if col == "Factor_MR" else
                              "Area (RAMOS_FACTOR_RET_AREA)" if col == "Factor_Ret" and clave in ramos_factor_ret_area
                              else "Area (la BD no lo trae)")
                df.at[i, col] = x
                self._anotar(tabla, clave, p, col, x, fuente, cuenta, sin)
        self._aviso_tabla(tabla, cuenta, sin, ": su MR sale de la formula Desviacion / -BC x BC2 del script")
        self._aviso_mr_negativo("SONR")
        for clave, ps in sorted(self._ret_fuera.items()):
            self.avisos.append(f"Factor_Ret SONR ramo {clave}: {len(ps)} mes(es) con IRR mayor que el BEL en la BD "
                               f"({min(ps)} a {max(ps)}): retencion 0 (todo cedido). Si no es correcto, pon {clave} en "
                               "RAMOS_FACTOR_RET_AREA para tomar el del area")
        if extra:
            ramos_extra = sorted({c for _, c in extra})
            previas = sorted({m for pc, m in respaldo_anterior.items() if pc in set(extra)})
            self.avisos.append(f"ParamSONR: {len(extra)} llave(s) que la tabla del area no traia con el formato AAAAMM-ramo "
                               f"(ramos {ramos_extra}) se agregaron con los parametros de la BD; lo que la BD no trae, de la "
                               "tabla del area: la misma llave, el mismo mes de un ano que la tabla si trae"
                               + (f" (sus meses {previas[0]} a {previas[-1]}, solo como previo)" if previas else "")
                               + " o el ultimo mes del ramo")
        return df

    def fnd_bd(self, clave: int, periodo: int) -> float:
        """FND de la BD del ramo y mes: el bloque FD/FND del RRC (PND / PRIMA de los ultimos 12 meses), acotado a [0, 1];
        nan si la BD no modela el RRC del ramo en ese mes (sin BEL) o no trae el FND."""
        k = (clave, periodo)
        if k not in self._fnd:
            v = self.bloque("FD/FND", "RRC", clave, periodo) if self._bd_modela("RRC", clave, periodo) else math.nan
            if not math.isnan(v) and not 0 <= v <= 1:
                self._fnd_fuera.setdefault(clave, []).append(periodo)
                v = min(max(v, 0.0), 1.0)
            self._fnd[k] = v
        return self._fnd[k]

    def fnd_contratos(self, reserva: str, df: pd.DataFrame, periodo, col_ramo: str, col_calmonth: str, col_fnd: str,
                      col_prima: str, col_tc: str, modo="ESCALAR") -> pd.Series:
        """FND de cada contrato (mismo indice que df) con nuestro FND por ramo y mes de valuacion (fnd_bd), sobre los
        contratos de los ultimos 12 meses (CALMONTH en (periodo - 12 meses, periodo]), que es la prima sobre la que la BD
        mide su FND (PND / PRIMA N AÑOS):
          modo "ESCALAR" (o True): se conserva el perfil por contrato del area (un contrato del mes queda mas por
            devengar que uno de hace 11 meses) y se escala para que el FND del ramo, ponderado por la prima en pesos del
            contrato (col_prima x col_tc), sea el de la BD; los que llegarian a mas de 1 se quedan en 1 y el resto se
            reescala. Si el ramo no tiene FND del area con que escalar, todos llevan el FND de la BD.
          modo "PLANO": todos los contratos del ramo llevan el FND de la BD.
        Los contratos de mas de 12 meses, los ramos sin FND en la BD y los meses que la base todavia no trae (el CALMONTH
        mas reciente es anterior al mes de valuacion: falta la prima nueva) quedan con el FND del area (col_fnd).
        RRC: es el PORC_ND; SONR: el FND con que sale la prima devengada (1 - FND)."""
        modo = "PLANO" if str(modo).upper() == "PLANO" else "ESCALAR"
        periodo = int(periodo)
        fnd_area = pd.to_numeric(df[col_fnd], errors="coerce").astype(float)
        out = fnd_area.copy()
        cm = pd.to_numeric(df[col_calmonth], errors="coerce")
        ramos = df[col_ramo].map(self._ramo_area)
        en_ventana = (cm > mes_mas(periodo, -12)) & (cm <= periodo)
        if cm.notna().any() and cm.max() < periodo:                     # la base no trae prima de ese mes
            self._fnd_incompleto.setdefault(reserva, {})[periodo] = int(cm.max())
            en_ventana = en_ventana & False
        prima = pd.to_numeric(df[col_prima], errors="coerce") * pd.to_numeric(df[col_tc], errors="coerce")
        for clave in sorted({int(c) for c in ramos.dropna().unique()}):
            del_ramo = ramos == clave
            uso = self._fnd_uso.setdefault((reserva, clave, periodo), [0, 0, 0])
            uso[1] += int((del_ramo & ~en_ventana).sum())
            sel = del_ramo & en_ventana
            if not sel.any():
                continue
            f_bd = self.fnd_bd(clave, periodo)
            if math.isnan(f_bd):
                uso[2] += int(sel.sum())
                continue
            uso[0] += int(sel.sum())
            p, fa = prima[sel], fnd_area[sel].fillna(0.0).clip(0.0, 1.0)
            ok = p.notna()
            nuevo = pd.Series(f_bd, index=p.index)
            tot = float(p[ok].sum())
            base = float((p[ok] * fa[ok]).sum())
            k = math.nan
            if modo == "ESCALAR" and ok.any() and tot != 0 and base != 0 and f_bd * tot / base > 0:
                objetivo, topados = f_bd * tot, pd.Series(False, index=p.index)
                for _ in range(30):
                    libres = ok & ~topados
                    base_l = float((p[libres] * fa[libres]).sum())
                    resto = objetivo - float(p[ok & topados].sum())
                    if base_l == 0 or resto / base_l <= 0:
                        break
                    k = resto / base_l
                    nuevo[libres] = fa[libres] * k
                    nuevos_topes = libres & (nuevo > 1.0)
                    if not nuevos_topes.any():
                        break
                    nuevo[nuevos_topes] = 1.0
                    topados = topados | nuevos_topes
            out[sel] = nuevo
            tc_usd = self.tc_de(periodo)
            p12 = self.bloque("PRIMA N AÑOS", "RRC", clave, periodo)
            self._fnd_detalle[(reserva, clave, periodo)] = {
                "FND del area (ponderado)": base / tot if tot else math.nan,
                "FND resultante (ponderado)": float((p[ok] * nuevo[ok]).sum()) / tot if tot else math.nan,
                "Factor de escala": k if modo == "ESCALAR" else math.nan,
                "Prima del script / PRIMA N AÑOS de la BD": abs(tot / tc_usd) / p12 if tot and tc_usd and p12 else math.nan}
        return out

    def tabla_fnd(self) -> pd.DataFrame:
        """Uso del FND por reserva, ramo y mes de valuacion: el FND de la BD y cuantos contratos lo tomaron."""
        filas = [{"Reserva": r, "Ramo": c, "Mes": p, "FND de la BD": self._fnd.get((c, p), math.nan),
                  **self._fnd_detalle.get((r, c, p), {}),
                  "Contratos con FND de la BD": n_bd, "Contratos de mas de 12 meses o mes sin prima (FND del area)": n_ant,
                  "Contratos sin FND en la BD (FND del area)": n_sin}
                 for (r, c, p), (n_bd, n_ant, n_sin) in sorted(self._fnd_uso.items(),
                                                              key=lambda x: (x[0][0], x[0][1] or 0, x[0][2]))]
        return pd.DataFrame(filas)

    def _avisos_fnd(self) -> list[str]:
        out = []
        for reserva in sorted({r for r, _, _ in self._fnd_uso}):
            usos = {k: v for k, v in self._fnd_uso.items() if k[0] == reserva}
            n_bd, n_ant, n_sin = (sum(v[i] for v in usos.values()) for i in range(3))
            sin = sorted({c for (r, c, p), v in usos.items() if v[2] and c is not None})
            out.append(f"FND {reserva}: {n_bd} contrato(s) con el FND de la BD de su ramo y mes; con el del area, {n_ant} "
                       f"de mas de 12 meses o de meses sin prima en la base" +
                       (f" y {n_sin} de ramos sin FND en la BD ({sin})" if n_sin else ""))
            inc = self._fnd_incompleto.get(reserva, {})
            if inc:
                out.append(f"FND {reserva}: en {len(inc)} mes(es) de valuacion ({min(inc)} a {max(inc)}) la base Access todavia no "
                           f"trae la prima del mes (la ultima es {max(inc.values())}): se valua el portafolio que ya esta, con el "
                           "FND del area y sin la prima nueva; no es la reserva completa de esos meses")
            dif = sorted({c for (r, c, p), d in self._fnd_detalle.items() if r == reserva and p <= self.ultimo_real
                          and abs(d.get("Prima del script / PRIMA N AÑOS de la BD", 1.0) - 1) > 0.10})
            if dif:
                out.append(f"FND {reserva}: en los ramos {dif} la prima de 12 meses de la base Access difiere en mas de 10 % de "
                           "la PRIMA N AÑOS de la BD (sobre la que se midio el FND); ver la hoja FND")
        for clave, ps in sorted(self._fnd_fuera.items()):
            ps = sorted(set(ps))
            out.append(f"FND ramo {clave}: {len(ps)} mes(es) con FND de la BD fuera de [0, 1] ({ps[0]} a {ps[-1]}): se acota")
        return out

    # ------------------------------------------------------------------ tipo de cambio para los scripts _aod
    def tc_bd_meses(self, desde: int = TC_BD_DESDE) -> dict:
        """{mes: TC MXN / USD de la BD} desde el mes indicado."""
        return {p: float(v) for p, v in sorted(self.tc.items()) if p >= desde and not math.isnan(_num(v))}

    def tc_usd_area(self, tabla: pd.DataFrame, desde: int = TC_BD_DESDE) -> pd.DataFrame:
        """TC_USD de los scripts (cTCAD_FecAMD, cTCAD_Mnt = MXN por USD de la base Access) con el TC de la BD desde
        `desde`: reemplaza los meses que la base trae y agrega los que no; los anteriores quedan como vienen. Ordenada por
        mes."""
        df = tabla.copy()
        nuestros = self.tc_bd_meses(desde)
        f = pd.to_numeric(df["cTCAD_FecAMD"], errors="coerce")
        df["cTCAD_Mnt"] = pd.to_numeric(df["cTCAD_Mnt"], errors="coerce").astype(float)
        m = f.isin(list(nuestros))
        antes = df.loc[m, "cTCAD_Mnt"].to_numpy()
        df.loc[m, "cTCAD_Mnt"] = [nuestros[int(x)] for x in f[m]]
        cambiados = int((abs(antes - df.loc[m, "cTCAD_Mnt"].to_numpy()) > 1e-12).sum())
        ya = set(f[m].astype(int))
        nuevos = [p for p in nuestros if p not in ya]
        if nuevos:
            df = pd.concat([df, pd.DataFrame({"cTCAD_FecAMD": nuevos, "cTCAD_Mnt": [nuestros[p] for p in nuevos]})],
                           ignore_index=True)
        orden = pd.to_numeric(df["cTCAD_FecAMD"], errors="coerce").to_numpy()
        df = df.iloc[np.argsort(orden, kind="stable")].reset_index(drop=True)
        if nuestros and (cambiados or nuevos):
            self.avisos.append(f"TC USD: {len(nuestros)} mes(es) de la BD ({min(nuestros)} a {max(nuestros)}); "
                               f"{cambiados} cambian el de la base Access y {len(nuevos)} se agregan. Los meses "
                               f"anteriores a {desde}, de la base")
        return df

    def tc_monedas_area(self, tabla: pd.DataFrame, desde: int = TC_BD_DESDE, hasta: int | None = None) -> pd.DataFrame:
        """Tabla de tipos de cambio de todas las monedas de la base Access (cTCAD_FecAMD, cMON_Id, cTCAD_Mnt, Llave) con
        el TC de la BD en el dolar desde `desde` (reemplaza o agrega el renglon del mes) y el peso en 1 en los meses que se
        agregan. Las demas monedas quedan como vienen: en los meses que la base no trae no tienen tipo de cambio."""
        df = tabla.copy()
        nuestros = self.tc_bd_meses(desde)
        f = pd.to_numeric(df["cTCAD_FecAMD"], errors="coerce")
        mon = pd.to_numeric(df["cMON_Id"], errors="coerce")
        df["cTCAD_Mnt"] = pd.to_numeric(df["cTCAD_Mnt"], errors="coerce").astype(float)
        m = (mon == MONEDA_USD) & f.isin(list(nuestros))
        df.loc[m, "cTCAD_Mnt"] = [nuestros[int(x)] for x in f[m]]
        hay = {(int(a), int(b)) for a, b in zip(f, mon) if not (math.isnan(a) or math.isnan(b))}
        pesos = df.loc[(mon == MONEDA_MXN) & f.notna(), ["cTCAD_FecAMD", "cTCAD_Mnt"]].dropna()
        if len(pesos):
            ult_mxn = float(pesos.loc[pd.to_numeric(pesos["cTCAD_FecAMD"], errors="coerce").idxmax(), "cTCAD_Mnt"])
            if abs(ult_mxn - 1.0) > 1e-9:
                self.avisos.append(f"TC: el peso (cMON_Id {MONEDA_MXN}) del ultimo mes de la base vale {ult_mxn}, no 1; en los "
                                   "meses que se agregan el peso va en 1")
        nuevos = [(p, MONEDA_USD, v) for p, v in nuestros.items() if (p, MONEDA_USD) not in hay]
        nuevos += [(p, MONEDA_MXN, 1.0) for p in nuestros if (p, MONEDA_MXN) not in hay]      # (un peso es un peso)
        if nuevos:
            df = pd.concat([df, pd.DataFrame({"cTCAD_FecAMD": [p for p, _, _ in nuevos], "cMON_Id": [c for _, c, _ in nuevos],
                                              "cTCAD_Mnt": [v for _, _, v in nuevos],
                                              "Llave": [f"{p}-{c}" for p, c, _ in nuevos]})], ignore_index=True)
        ultimo = int(f.max()) if f.notna().any() else None
        otras = sorted({int(b) for a, b in zip(f, mon) if ultimo is not None and a == ultimo and not math.isnan(b)}
                       - {MONEDA_USD, MONEDA_MXN})
        sin = [p for p in nuestros if ultimo is not None and p > ultimo and (hasta is None or p <= hasta)]
        if otras and sin:
            self.avisos.append(f"TC de otras monedas: la base Access no trae {len(otras)} moneda(s) ({otras}) de {sin[0]} a "
                               f"{sin[-1]}; la BD solo trae el dolar, asi que esos contratos salen vacios en esos meses")
        return df

    # ------------------------------------------------------------------ scripts _aod_2027: sin la base Access
    def tc_monedas_bd(self) -> pd.DataFrame:
        """Tabla de tipos de cambio con la forma de la de la base Access (cTCAD_FecAMD, cMON_Id, cTCAD_Mnt, Llave) armada
        con la columna TC de la BD: el dolar con el TC de la BD y el peso en 1, en todos los meses de la BD."""
        filas = []
        for p, v in sorted(self.tc.items()):
            if not math.isnan(_num(v)):
                filas += [(p, MONEDA_USD, float(v)), (p, MONEDA_MXN, 1.0)]
        return pd.DataFrame({"cTCAD_FecAMD": [a for a, _, _ in filas], "cMON_Id": [b for _, b, _ in filas],
                             "cTCAD_Mnt": [c for _, _, c in filas], "Llave": [f"{a}-{b}" for a, b, _ in filas]})

    def _contratos_bd(self, meses: list[int], ramos) -> list[dict]:
        """Un "contrato" por ramo y mes con la prima tomada de la BD (hoja PE_RAMO, USD): lo que los scripts de 2027 valuan en
        lugar de los contratos de la base. Se avisa si faltan meses de prima."""
        if not self.pe:
            raise SystemExit(f"La BD {self.ruta_bd.name} no trae la hoja {HOJA_PE} (prima por ramo y mes): sin ella no hay "
                             "contratos para 2027. Vuelve a generar la BD con proyeccion_reservas.py")
        out, sin = [], []
        for p in meses:
            hay = False
            for clave in ramos:
                pe = self.pe.get((clave, p))
                if pe is None or math.isnan(pe):
                    continue
                hay = True
                if pe == 0:
                    continue
                out.append({"clave": clave, "p": p, "pe": pe})
            if not hay:
                sin.append(p)
        self._pe_faltan.update(sin)
        return out

    def contratos_rrc_bd(self, desde_excl, hasta) -> pd.DataFrame:
        """Contratos del RRC armados con la prima de la BD para CALMONTH en (desde_excl, hasta], con las columnas que deja la
        consulta de la base y el ramo ya como clave del script: prima tomada negativa (convencion de la base), en dolares
        (MonedaOri 31), proporcional (TipoRea 1, frecuencia 1), vigencia de 12 meses desde el mes y la cesion de
        cesion_rrc_bd del mes que se valua (PrimaCedidaOri = cesion x prima, asi el script calcula esa cesion)."""
        meses = rango_meses(mes_mas(int(desde_excl), 1), int(hasta))
        filas = []
        ces_ramo = {c: self.cesion_rrc_bd(c, int(hasta)) for c in RAMOS_SCRIPT}
        for d in self._contratos_bd(meses, RAMOS_SCRIPT):
            clave, p, pe = d["clave"], d["p"], d["pe"]
            tc = self.tc_de(p)
            ces = ces_ramo[clave]
            ini = pd.Timestamp(year=p // 100, month=p % 100, day=1)
            filas.append({"SRamo": clave, "Pais": math.nan, "TipoRea": 1, "OfiRepPt": 1, "MonedaOri": MONEDA_USD,
                          "CorrTom": 0, "CiaTom": 0, "CtoTom": 0, "Susc": p // 100, "Período": 1, "CALMONTH": p,
                          "IniVig": ini, "FinVig": ini + pd.DateOffset(months=12),
                          "PrimaTomadaOri": -pe, "PmaTom_sEROri": -pe, "PrimaCedidaOri": ces * pe,
                          "PrimaTomadaNal": -pe * tc, "PmaTom_sERNal": -pe * tc, "PrimaCedidaNal": ces * pe * tc,
                          "Ramo": clave})
        cols = ["SRamo", "Pais", "TipoRea", "OfiRepPt", "MonedaOri", "CorrTom", "CiaTom", "CtoTom", "Susc", "Período",
                "CALMONTH", "IniVig", "FinVig", "PrimaTomadaOri", "PmaTom_sEROri", "PrimaCedidaOri", "PrimaTomadaNal",
                "PmaTom_sERNal", "PrimaCedidaNal", "Ramo"]
        df = pd.DataFrame(filas, columns=cols)
        df["Pais"] = df["Pais"].astype(float)
        return df

    def cesion_rrc_bd(self, clave: int, periodo: int) -> float:
        """Cesion de los contratos de la BD en el RRC: IRR / BEL de la BD del ramo en el mes que se valua. El script calcula
        IRR = BEL x cesion y la BD IRR = BRUTO x CESION (IRR / BRUTO), asi el IRR del script es el de la BD; si la BD no
        trae BEL del RRC en el mes, el bloque CESION."""
        if self._bd_modela("RRC", clave, periodo):
            irr = self.monto("RRC IRR", periodo, clave)
            if not math.isnan(irr):
                return irr / self.monto("RRC BEL", periodo, clave)
        v = self.bloque("CESION", "RRC", clave, periodo)
        return 0.0 if math.isnan(v) else v

    def nivel_sonr_bd(self, tabla: pd.DataFrame, periodo: int, tc_usd: pd.DataFrame) -> pd.Series:
        """Prima Dev del metodo propio del SONR escalada por ramo para que el BEL del mes sea el de la BD (FD SONR x
        PEACUMULADA x IS). Con un contrato por ramo y mes toda la prima cae en la ventana del ano en que se registra (Susc =
        ano del mes) y no en la de su ano de suscripcion, y el metodo da un BEL mas alto que el de la BD; el IRR (Factor_Ret)
        y el MR (Factor_MR o formula) salen de esa Prima Dev, igual que en el script. tabla: Tbase del mes con 'Ramo',
        'Prima Dev', 'LAG' (1 - LAG k) e 'Ind Sin SONR Media'; tc_usd: la tabla TC_USD del script (el BEL en pesos se pasa
        a dolares con ella). Los ramos que la BD no modela en el mes (sin BEL) quedan con el metodo tal cual, con aviso."""
        dev = pd.to_numeric(tabla["Prima Dev"], errors="coerce").astype(float)
        bel = dev * pd.to_numeric(tabla["LAG"], errors="coerce") * pd.to_numeric(tabla["Ind Sin SONR Media"], errors="coerce")
        fila_tc = tc_usd.loc[pd.to_numeric(tc_usd["cTCAD_FecAMD"], errors="coerce") == int(periodo), "cTCAD_Mnt"]
        tc = _num(fila_tc.iloc[-1]) if len(fila_tc) else math.nan
        out = dev.copy()
        for ramo, idx in tabla.groupby("Ramo").groups.items():
            clave = int(_num(ramo))
            actual = float(bel.loc[idx].sum(skipna=True)) / tc if math.isfinite(tc) and tc else math.nan
            if not self._bd_modela("SONR", clave, periodo):
                self._sonr_sin_nivel.setdefault((clave, "la BD no trae BEL de SONR"), []).append(periodo)
                continue
            if not math.isfinite(tc) or not tc:
                self._sonr_sin_nivel.setdefault((clave, "el script no trae tipo de cambio del mes"), []).append(periodo)
                continue
            if not math.isfinite(actual) or abs(actual) < UMBRAL_CERO:
                self._sonr_sin_nivel.setdefault((clave, "el metodo no tiene prima ese mes"), []).append(periodo)
                continue
            meta = self.monto("SONR BEL", periodo, clave)
            out.loc[idx] = dev.loc[idx] * (meta / actual)
            self._sonr_nivel[(clave, periodo)] = actual / meta
        return out

    def contratos_sonr_bd(self, desde, hasta) -> pd.DataFrame:
        """Contratos del SONR armados con la prima de la BD para CALMONTH en [desde, hasta], con las columnas de la consulta
        de la base: prima tomada positiva (la consulta la cambia de signo), en dolares, proporcional, frecuencia 1."""
        meses = rango_meses(int(desde), int(hasta))
        filas = []
        for d in self._contratos_bd(meses, RAMOS_SCRIPT):
            clave, p, pe = d["clave"], d["p"], d["pe"]
            ini = pd.Timestamp(year=p // 100, month=p % 100, day=1)
            filas.append({"CALMONTH": p, "Ramo_filt": clave, "Pais": math.nan, "TipoRea": 1, "CorrTom": 0, "CiaTom": 0,
                          "CtoTom": 0, "Susc": p // 100, "MonedaOri": MONEDA_USD, "IniVig": ini,
                          "FinVig": ini + pd.DateOffset(months=12), "Periodo": 1, "PmaTomOri": pe,
                          "PmaTomNal": pe * self.tc_de(p)})
        cols = ["CALMONTH", "Ramo_filt", "Pais", "TipoRea", "CorrTom", "CiaTom", "CtoTom", "Susc", "MonedaOri", "IniVig",
                "FinVig", "Periodo", "PmaTomOri", "PmaTomNal"]
        df = pd.DataFrame(filas, columns=cols)
        df["Pais"] = df["Pais"].astype(float)
        return df

    def escenario_base_area(self, reserva: str, tabla, anio: int, previo_ano_anterior: bool = False) -> pd.DataFrame:
        """Escenarios 0 y 1 para un ano sin la base: del Escenario_base del area se quedan el escenario 0 de diciembre del
        ano anterior y el escenario 1 (presupuesto) de los meses del ano; si no trae el escenario 0, sale de los montos de la
        BD de ese diciembre (proyectados si es despues del ultimo mes real). Los renglones de otros anos se quitan."""
        df = tabla.copy() if isinstance(tabla, pd.DataFrame) else pd.DataFrame(columns=COLUMNAS_ESC)
        for c in COLUMNAS_ESC:
            if c not in df.columns:
                df[c] = math.nan
        per = pd.to_numeric(df["Periodo"], errors="coerce")
        esc = pd.to_numeric(df["Escenario"], errors="coerce")
        base = (anio - 1) * 100 + 12
        e0 = df[(esc == 0) & (per == base)]
        e1 = df[(esc == 1) & (per // 100 == anio)]
        quitados = len(df) - len(e0) - len(e1)
        if e0.empty:
            e0 = self.escenario_base(reserva, anio)
            self.avisos.append(f"Escenario 0 {reserva}: {base} de los montos de la BD"
                               + (" (proyectados)" if base > self.ultimo_real else "") + "; el Escenario_base del area no lo trae")
        if e1.empty and previo_ano_anterior:
            e1 = df[(esc == 1) & (per // 100 == anio - 1)].copy()
            if len(e1):
                e1["Periodo"] = pd.to_numeric(e1["Periodo"], errors="coerce").astype(int) + 100
                quitados -= len(e1)
                self.avisos.append(f"Escenario 1 {reserva}: el Escenario_base del area no trae presupuesto de {anio}; va el "
                                   f"presupuesto de {anio - 1} con los meses de {anio} (solo como previo, no es el presupuesto "
                                   f"{anio}): {len(e1)} renglon(es)")
        if e1.empty:
            self.avisos.append(f"Escenario 1 {reserva}: el Escenario_base del area no trae presupuesto de {anio}; no se incluye")
        if quitados:
            self.avisos.append(f"Escenario_base {reserva}: {quitados} renglon(es) de otros anos se quitaron")
        return pd.concat([e0[COLUMNAS_ESC], e1[COLUMNAS_ESC]], ignore_index=True)

    def exportar_tablas(self, ruta, tablas: dict):
        """Escribe en un xlsx las tablas que uso un script, la hoja Fuentes (de donde salio cada celda) y los avisos."""
        avisos = self.avisos_texto()
        with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
            pd.DataFrame({"Dato": ["BD", "Ultimo mes real", "Proyeccion", "IS (FA)"] + [f"Aviso {i + 1}" for i in range(len(avisos))],
                          "Valor": [str(self.ruta_bd), self.ultimo_real,
                                    f"{self.primer_proyectado} a {self.ultimo_proyectado}",
                                    f"RRC {self.ramos_is_fa.get('RRC')}, SONR {self.ramos_is_fa.get('SONR')}"
                                    if self.is_fa and self.usar_is_fa else "no"] + avisos}).to_excel(
                xw, sheet_name="Resumen", index=False)
            for nombre, df in tablas.items():
                df.to_excel(xw, sheet_name=str(nombre)[:31], index=False)
            if self.fuentes:
                pd.DataFrame(self.fuentes).to_excel(xw, sheet_name="Fuentes", index=False)
            if self._fnd_uso:
                self.tabla_fnd().to_excel(xw, sheet_name="FND", index=False)

    def tc_tabla(self) -> pd.DataFrame:
        """TC USD por mes como la tabla de tipo de cambio del area: cTCAD_FecAMD (AAAAMM), cTCAD_Mnt."""
        return pd.DataFrame({"cTCAD_FecAMD": sorted(self.tc), "cTCAD_Mnt": [self.tc[p] for p in sorted(self.tc)]})

    def tc_dict(self) -> dict:
        return dict(sorted(self.tc.items()))

    def tc_de(self, periodo: int) -> float:
        if periodo in self.tc:
            return self.tc[periodo]
        anteriores = [p for p in self.tc if p <= periodo]
        return self.tc[max(anteriores)] if anteriores else self.tc[min(self.tc)]

    def saldos(self, reserva: str, periodos: list[int], escenario: int, tipos: tuple | None = None) -> pd.DataFrame:
        """Saldos de la BD (reales o proyectados) por ramo, mes y tipo de monto en la forma del escenario base de los
        scripts (COLUMNAS_ESC), con la convencion de signos del area (SIGNO)."""
        tipos = tipos or tuple(SIGNO[reserva])
        filas = []
        for p in periodos:
            tc = self.tc_de(p)
            for clave in RAMOS_SCRIPT:
                for tipo in tipos:
                    concepto = f"{reserva} {CONCEPTO_BD.get(tipo, tipo)}"
                    if (norm(concepto), p, self.ramo_bd(clave)) not in self.montos:
                        continue
                    usd = self.monto(concepto, p, clave)
                    usd = 0.0 if math.isnan(usd) or abs(usd) < UMBRAL_CERO else usd * SIGNO[reserva][tipo]
                    filas.append({"Reserva": reserva, "Escenario": escenario, "Tipo de Monto": tipo, "Ramo": clave,
                                  "Periodo": p, "Monto_MXN": usd * tc, "Monto_USD": usd, "TC": tc})
        return pd.DataFrame(filas, columns=COLUMNAS_ESC)

    def escenario_base(self, reserva: str, anio: int | None = None) -> pd.DataFrame:
        """Escenario 0 (ano base): saldos reales de diciembre del ano anterior."""
        anio = anio or self.anio
        return self.saldos(reserva, [(anio - 1) * 100 + 12], 0)

    def saldos_proyectados(self, reserva: str, desde: int | None = None, hasta: int | None = None,
                           escenario: int = 1) -> pd.DataFrame:
        """Saldos proyectados de la BD (BEL por FND y derivados) de desde a hasta, para los meses que el script no
        valua (los 'meses faltantes' del escenario 3 o el escenario 1)."""
        desde = desde or self.primer_proyectado
        hasta = hasta or self.ultimo_proyectado
        if desde is None or hasta is None:
            return pd.DataFrame(columns=COLUMNAS_ESC)
        return self.saldos(reserva, rango_meses(desde, hasta), escenario)

    # ------------------------------------------------------------------ informes
    def avisos_texto(self) -> list[str]:
        """Avisos, con los meses llenados agrupados por rango (un aviso por rango, con los parametros que abarca)."""
        grupos: dict = {}
        for etiqueta, d, h, n in self._huecos:
            grupos.setdefault((d, h, n), []).append(etiqueta)
        out = list(self.avisos) + self._avisos_fnd()
        if self._sonr_nivel:
            r = sorted(self._sonr_nivel.values())
            out.append(f"SONR: BEL del metodo propio llevado al de la BD en {len(r)} ramo(s) x mes (SONR_NIVEL_BD): con un "
                       f"contrato por ramo y mes el metodo daba entre {r[0]:.2f} y {r[-1]:.2f} veces el BEL de la BD (mediana "
                       f"{r[len(r) // 2]:.2f}); el IRR y el MR siguen al BEL")
        for (clave, motivo), ps in sorted(self._sonr_sin_nivel.items()):
            out.append(f"SONR ramo {clave}: {len(ps)} mes(es) ({min(ps)} a {max(ps)}) con el metodo propio tal cual sobre la "
                       f"prima de la BD, sin llevarlo al BEL de la BD: {motivo}")
        if self._pe_faltan:
            f = sorted(self._pe_faltan)
            out.append(f"Prima de la BD: {len(f)} mes(es) sin PE en la hoja {HOJA_PE} ({f[0]} a {f[-1]}): esos meses no tienen "
                       "contratos (en el IBNR, los anos de LAG mas viejos van sin prima)")
        for (d, h, n), etiquetas in sorted(grupos.items()):
            rango = f"{d}" if d == h else f"{d} a {h}"
            out.append(f"{n} mes(es) sin dato ({rango}) llenados con los meses vecinos en {len(etiquetas)} serie(s): "
                       + ", ".join(etiquetas[:6]) + (" ..." if len(etiquetas) > 6 else ""))
        return out

    def resumen(self) -> list[str]:
        n_fa = len({(r, c) for (r, c, _) in self.is_fa})
        lineas = [f"Insumos de {self.ruta_bd}: HParametros '{self.hoja_hp}' con {len(self.hp)} renglones; tipos "
                  f"{', '.join(sorted(self.tipos))}; ultimo mes real {self.ultimo_real}; proyeccion "
                  f"{self.primer_proyectado} a {self.ultimo_proyectado}",
                  f"Hoja de montos: {len(self.periodos_montos)} meses ({min(self.periodos_montos)} a "
                  f"{max(self.periodos_montos)}), ramos {', '.join(sorted(self.col_ramo_bd, key=int))}, TC de "
                  f"{min(self.tc)} a {max(self.tc)}",
                  (f"IS (FA): {n_fa} series (reserva x ramo), {min(p for (_, _, p) in self.is_fa)} a "
                   f"{max(p for (_, _, p) in self.is_fa)}; se usa en los meses proyectados de RRC "
                   f"{self.ramos_is_fa.get('RRC')} y SONR {self.ramos_is_fa.get('SONR')}" if self.is_fa and self.usar_is_fa
                   else "IS (FA): no se usa")]
        return lineas

    def exportar(self, ruta, anio: int | None = None, csv_duracion=None, csv_is_cat=None, csv_param_sonr=None,
                 ramos_factor_ret_csv: tuple = ()):
        """Escribe en un xlsx las tablas que reciben los scripts (para auditar que insumo se uso)."""
        anio = anio or self.anio
        tablas = [("ParametrosMens_RRC", self.parametros_rrc(anio, csv_duracion)),
                  ("IS_Cat", self.is_cat(csv_is_cat)),
                  ("ParamSONR", self.param_sonr(csv_param_sonr=csv_param_sonr, ramos_factor_ret_csv=ramos_factor_ret_csv)),
                  ("TC_USD", self.tc_tabla()),
                  ("Escenario_base", pd.concat([self.escenario_base("RRC", anio), self.escenario_base("SONR", anio)],
                                               ignore_index=True)),
                  ("Saldos_proyectados", pd.concat([self.saldos_proyectados("RRC"), self.saldos_proyectados("SONR")],
                                                   ignore_index=True))]
        avisos = self.avisos_texto()               # (ya con los que levantaron las tablas)
        with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
            pd.DataFrame({"Dato": ["BD", "Ultimo mes real", "Proyeccion", "IS (FA)"] + [f"Aviso {i + 1}" for i in range(len(avisos))],
                          "Valor": [str(self.ruta_bd), self.ultimo_real,
                                    f"{self.primer_proyectado} a {self.ultimo_proyectado}",
                                    "si" if self.is_fa and self.usar_is_fa else "no"] + avisos}).to_excel(
                xw, sheet_name="Resumen", index=False)
            for nombre, df in tablas:
                df.to_excel(xw, sheet_name=nombre, index=False)


if __name__ == "__main__":
    # Uso: python insumos_bd.py [ruta de la BD] [xlsx de salida]. Sin argumentos (p. ej. con el boton Run de VS Code)
    # toma la BD y los CSV de config_local.py, o busca la BD junto a este archivo y en salidas/, y escribe
    # Parametros_usados.xlsx en CARPETA_SALIDA. Solo arma y revisa los insumos: las reservas se corren con
    # reforecastRRC_v11_insumosBD.py y ReforecastSONR_v4_insumosBD.py.
    import sys
    aqui = Path(__file__).resolve().parent
    sys.path.insert(0, str(aqui))
    try:
        import config_local as cfg
    except ImportError:
        cfg = None
    try:
        bd = buscar_bd(sys.argv[1] if len(sys.argv) > 1 else getattr(cfg, "RUTA_BD", None), [aqui])
    except FileNotFoundError as e:
        raise SystemExit(str(e))
    salida = (Path(sys.argv[2]) if len(sys.argv) > 2 else
              Path(getattr(cfg, "CARPETA_SALIDA", None) or aqui / "salidas_area") / "Parametros_usados.xlsx")
    ins = InsumosBD(bd, getattr(cfg, "RUTA_DIAGNOSTICO", None), usar_is_fa=getattr(cfg, "USAR_IS_FA", True),
                    ramos_is_fa=getattr(cfg, "RAMOS_IS_FA", None))
    salida.parent.mkdir(parents=True, exist_ok=True)
    try:
        ins.exportar(salida, getattr(cfg, "ANIO", None), csv_duracion=getattr(cfg, "CSV_DURACION_RRC", None),
                     csv_is_cat=getattr(cfg, "CSV_IS_CAT", None), csv_param_sonr=getattr(cfg, "CSV_PARAM_SONR", None),
                     ramos_factor_ret_csv=tuple(getattr(cfg, "RAMOS_FACTOR_RET_CSV", ()) or ()))
    except PermissionError:
        raise SystemExit(f"No se pudo escribir {salida}: cierralo en Excel y vuelve a correr")
    print(f"   Tablas en {salida}")
    for a in ins.avisos_texto():
        print("   AVISO:", a)
    print("   Listo. Este archivo solo arma los insumos; las reservas se corren con reforecastRRC_v11_insumosBD.py y "
          "ReforecastSONR_v4_insumosBD.py")
