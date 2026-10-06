# -*- coding: utf-8 -*-
"""Corre ReforecastRRC_aod_BD.py y ReforecastSONR_aod_BD.py (salida de parchar_aod.py) sin la base Access ni los
archivos del area, y comprueba contrato por contrato y mes por mes que los indices que usan son los de la BD proyectada
(o los de la tabla del area donde la BD no trae el dato).

Uso: python simulacion_aod.py <carpeta con los *_aod_BD.py> ["<BD proyectada>"] [carpeta de trabajo]

Las rutas de Windows de los scripts se redirigen por nombre de archivo a una carpeta
de trabajo con tablas del area inventadas, con valores faciles de reconocer (p. ej. IS 0.111), para distinguir que
celda salio del area. Simula pyodbc (tipo de cambio y contratos) y mec_devengamiento (FND calibrado sin desplazamiento).
Corre dos veces: con tipo de cambio en la base para todo el ano y solo hasta el ultimo mes real (prueba el aviso). Los
montos que salen NO son reservas.
"""
from __future__ import annotations

import builtins
import math
import re
import runpy
import shutil
import sys
import tempfile
import traceback
import types
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

AQUI = Path(__file__).resolve().parent
SCRIPTS = AQUI.parent
sys.path.insert(0, str(SCRIPTS))
sys.path.insert(0, str(AQUI))
from insumos_bd import InsumosBD, LAGS, NOMBRE_BD, RAMOS_SCRIPT, RAMOS_SONR, buscar_bd, mes_mas, rango_meses  # noqa: E402
from simulacion_sin_access import RAMO_NOMBRE, SUBRAMOS, contratos  # noqa: E402

AREA = {"IS": 0.111, "IS99": 0.222, "GTO": 0.0777, "CAT71": 0.333, "CAT73": 0.444, "IS_SONR": 0.666,
        "IS_SONR99": 0.777, "LAG": 0.0888, "RET": 0.555}            # valores del area inventados (reconocibles)
ORIG = {"read_csv": pd.read_csv, "read_excel": pd.read_excel, "to_excel": pd.DataFrame.to_excel,
        "ExcelWriter": pd.ExcelWriter, "exists": Path.exists, "is_file": Path.is_file, "print": builtins.print}


def _es_windows(p) -> bool:
    s = str(p)
    return "\\" in s or bool(re.match(r"^[A-Za-z]:", s))


def _base(p) -> str:
    return str(p).replace("\\", "/").split("/")[-1]


def tablas_area(carpeta: Path, anio: int):
    """Las tablas que leen los scripts, con los nombres de archivo del area."""
    carpeta.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Valores"
    ws["J2"], ws["K2"], ws["L2"], ws["M2"] = "SR", "Ramo", "Desc", "Nota"
    ws["O2"], ws["P2"], ws["Q2"], ws["R2"], ws["S2"], ws["T2"] = "País", "TerrSAP", "Region", "x", "y", "z"
    fila = 3
    for clave, subs in SUBRAMOS.items():
        for sr in subs:
            ws.cell(fila, 10, sr)
            ws.cell(fila, 11, RAMO_NOMBRE[clave])
            fila += 1
    for i, (pais, terr) in enumerate([(1, "R01"), (2, "R03"), (3, "R05"), (4, "R04")], start=3):
        ws.cell(i, 15, pais)
        ws.cell(i, 16, terr)
    wb.save(carpeta / "CentralizadoCatálogos_SIRECySAP.xlsx")
    pm = pd.DataFrame({"Ramo": RAMOS_SCRIPT, "Pesos_dur": 1.1, "Resto Monedas_dur": 1.2, "Pesos_ret": 0.8,
                       "Resto Monedas_ret": 0.9, "Ind. Gasto": AREA["GTO"],
                       **{f"IS Bel Media-{m}": AREA["IS"] for m in range(1, 13)},
                       **{f"IS Bel 99.5%-{m}": AREA["IS99"] for m in range(1, 13)}})
    pm.to_csv(carpeta / "ParametrosMensPPTO_3+9.csv", index=False)
    meses_cat = rango_meses(201501, anio * 100 + 12)
    pd.DataFrame({"IS Bel Media": meses_cat, "71": AREA["CAT71"], "73": AREA["CAT73"]}).to_csv(carpeta / "IS_Cat.csv", index=False)
    pd.DataFrame({"IS Bel Media": meses_cat, "71": 0.5, "73": 0.5}).to_csv(carpeta / "IS_Cat_PPTO.csv", index=False)
    pd.DataFrame({"Llave": ["000000-0-0-0"]}).to_csv(carpeta / "LlavesPol.csv", index=False)
    pd.DataFrame(columns=["SRamo", "Ramo"]).to_csv(carpeta / "AjManuales.csv", index=False)
    pd.DataFrame({"CeBe": [f"A{r:03d}" for r in RAMOS_SCRIPT], "Ramo": RAMOS_SCRIPT}).to_csv(carpeta / "Subramo.csv", index=False)
    pd.DataFrame({"Ramo": RAMOS_SCRIPT, str(anio): 0.05}).to_csv(carpeta / "CesionPI.csv", index=False)
    pd.DataFrame({"AFUN": ["LN04001"], "Linea de Negocio": ["Norte"]}).to_csv(carpeta / "AFUN.csv", index=False)
    pd.DataFrame({"Llave": ["1-1-1-1"], "Periodo": [1]}).to_csv(carpeta / "zFrecuencias.csv", index=False)
    pd.DataFrame({"Llave": ["10-R01-1"], str(anio): 0.3}).to_csv(carpeta / "TablaCesion_Esc1.csv", index=False)
    pd.DataFrame({"Llave": ["1-1-1-2024-1"], "Porcentaje Cedido": [0.2]}).to_csv(carpeta / "Cesion ID Esp.csv", index=False)
    for reserva, ramos in (("RRC", RAMOS_SCRIPT), ("SONR", RAMOS_SONR)):
        filas = [{"Reserva": reserva, "Escenario": e, "Tipo de Monto": t, "Ramo": r, "Periodo": p, "Monto_MXN": 0.0,
                  "Monto_USD": 1000.0, "TC": 0.0}
                 for e, ps in ((0, [(anio - 1) * 100 + 12]), (1, rango_meses(anio * 100 + 1, anio * 100 + 12)))
                 for p in ps for r in ramos for t in ("BEL", "IRR", "MR", "BRUTO", "NETO")]
        pd.DataFrame(filas).to_csv(carpeta / f"Escenario_base_{reserva}.csv", index=False)
    pd.DataFrame([{"Ramo": r, "NoLAG": k} for r in RAMOS_SONR for k in range(1, 11)]).to_csv(
        carpeta / "TablaBase_MetodoPropio.csv", index=False)
    pd.DataFrame({"Ramo": [10], "NoLAG": [1]}).to_csv(carpeta / "TablaBase_MetodoPropio_ext.csv", index=False)
    pd.DataFrame([{"Llave": f"{anio * 100 + m}-{r}", "Factor_Ret": AREA["RET"], "Ind Sin SONR Media": AREA["IS_SONR"],
                   "Ind Sin SONR 99.5%": AREA["IS_SONR99"], **{k: AREA["LAG"] for k in LAGS}}
                  for m in range(1, 13) for r in RAMOS_SONR]).to_csv(carpeta / "ParamSONR2026_3+9.csv", index=False)
    for n in ("PNDmes.csv", "FrecCol.csv"):
        pd.DataFrame({"x": [1]}).to_csv(carpeta / n, index=False)


def redirigir(area: Path, salida: Path, capturas: dict):
    """Rutas de Windows -> carpeta de trabajo, por nombre de archivo."""
    salida.mkdir(parents=True, exist_ok=True)

    def leer(nombre):
        def f(p, *a, **k):
            if isinstance(p, (str, Path)) and _es_windows(p):
                p = area / _base(p)
                if not ORIG["exists"](p):
                    raise FileNotFoundError(f"simulacion: no hay tabla inventada para {p.name}")
            return ORIG[nombre](p, *a, **k)
        return f

    def to_excel(self, destino, *a, **k):
        if isinstance(destino, (str, Path)) and _es_windows(destino):
            capturas[_base(destino)] = self.copy()
            destino = salida / _base(destino)
        return ORIG["to_excel"](self, destino, *a, **k)

    def writer(p, *a, **k):
        if isinstance(p, (str, Path)) and _es_windows(p):
            p = salida / _base(p)
        return ORIG["ExcelWriter"](p, *a, **k)

    def existe(nombre):
        def f(self, *a, **k):
            if _es_windows(self):
                return ORIG[nombre](area / _base(self))
            return ORIG[nombre](self, *a, **k)
        return f
    pd.read_csv, pd.read_excel = leer("read_csv"), leer("read_excel")
    pd.DataFrame.to_excel, pd.ExcelWriter = to_excel, writer
    Path.exists, Path.is_file = existe("exists"), existe("is_file")


def restaurar():
    pd.read_csv, pd.read_excel = ORIG["read_csv"], ORIG["read_excel"]
    pd.DataFrame.to_excel, pd.ExcelWriter = ORIG["to_excel"], ORIG["ExcelWriter"]
    Path.exists, Path.is_file = ORIG["exists"], ORIG["is_file"]


def simular_access(ins: InsumosBD, tc_hasta: int):
    tc_filas = []
    for p in rango_meses(mes_mas(ins.primer_periodo_hp, -12), tc_hasta):
        usd = ins.tc_de(p)
        for mon, v in ((1, 1.0), (31, usd), (2, usd * 1.08)):
            tc_filas.append({"cTCAD_FecAMD": p, "cMON_Id": mon, "cTCAD_Mnt": v, "Llave": f"{p}-{mon}"})
    tc = pd.DataFrame(tc_filas)
    gonz = contratos(ins.anio, rango_meses(mes_mas(ins.ultimo_real, -11 * 12), ins.ultimo_real))

    class _Conn:
        def cursor(self):
            return self

        def close(self):
            pass
    falso = types.ModuleType("pyodbc")
    falso.connect = lambda conn_str: _Conn()
    sys.modules["pyodbc"] = falso
    mec = types.ModuleType("mec_devengamiento")            # FND calibrado sin desplazamiento (delta 0)
    mec.cargar_delta = lambda carpeta: {}
    mec.antiguedad_registro = lambda mes_val, calmonth: (int(mes_val) // 100 - int(calmonth) // 100) * 12 + int(mes_val) % 100 - int(calmonth) % 100
    mec.fnd_desplazado = lambda ramo, valor, k, delta: valor
    sys.modules["mec_devengamiento"] = mec

    def read_sql(sql, conn, *a, **k):
        if "aMOT_MovTipCambio" in sql:
            if "cMON_Id = 31" in sql:
                return tc[tc["cMON_Id"] == 31][["cTCAD_FecAMD", "cTCAD_Mnt"]].reset_index(drop=True)
            return tc.copy()
        if "aMOG_MovGonzalo" in sql:
            if "Ramo_filt" in sql:
                m = re.search(r"aPog_MesProc\) >= (\d+) And Val\(aPog_MesProc\) <= (\d+)", sql)
                sel = gonz[(gonz["CALMONTH"] >= int(m.group(1))) & (gonz["CALMONTH"] <= int(m.group(2)))]
                cols = ['CALMONTH', 'Ramo_filt', 'Pais', 'TipoRea', 'CorrTom', 'CiaTom', 'CtoTom', 'Susc', 'MonedaOri',
                        'IniVig', 'FinVig', 'Periodo', 'PmaTomOri', 'PmaTomNal']
                return sel[cols].reset_index(drop=True)
            m = re.search(r"aPog_MesProc\) > (\d+) and Val\(aPog_MesProc\) <= (\d+)\)", sql)
            sel = gonz[(gonz["CALMONTH"] > int(m.group(1))) & (gonz["CALMONTH"] <= int(m.group(2)))]
            cols = ['SRamo', 'Pais', 'TipoRea', 'OfiRepPt', 'MonedaOri', 'CorrTom', 'CiaTom', 'CtoTom', 'Susc', 'Período',
                    'CALMONTH', 'IniVig', 'FinVig', 'PrimaTomadaOri', 'PmaTom_sEROri', 'PrimaCedidaOri', 'PrimaTomadaNal',
                    'PmaTom_sERNal', 'PrimaCedidaNal']
            return sel[cols].reset_index(drop=True)
        raise ValueError(f"SQL no simulado: {sql[:80]}")
    pd.read_sql = read_sql


def correr(script: Path, otra: Path, textos: list, marcos: list) -> dict:
    """Corre el script desde otra carpeta y sin argumentos (como el boton Run de VS Code); guarda lo que imprime."""
    import os

    def _print(*a, **k):
        for x in a:
            if isinstance(x, pd.DataFrame):
                marcos.append(x.copy())
        if not any(isinstance(x, pd.DataFrame) for x in a):
            textos.append(" ".join(str(x) for x in a))
            ORIG["print"](*a, **k)
    cwd = Path.cwd()
    otra.mkdir(parents=True, exist_ok=True)
    try:
        os.chdir(otra)
        sys.argv = [str(script)]
        sys.modules.pop("insumos_bd", None)
        return runpy.run_path(str(script), init_globals={"print": _print}, run_name="__main__")
    except BaseException:  # noqa: BLE001
        traceback.print_exc()
        return {}
    finally:
        os.chdir(cwd)


def _igual(a, b) -> bool:
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return False
    return (math.isnan(a) and math.isnan(b)) or abs(a - b) <= 1e-9 * (1 + abs(b))


def verificar_rrc(ins: InsumosBD, g: dict, capturas: dict, anio: int) -> list[str]:
    res = []
    if not g:
        return ["MAL RRC: el script no termino"]
    area_cat = {71: AREA["CAT71"], 73: AREA["CAT73"]}
    malos, n, fa, area_usada, mr_bd, mr_formula = 0, 0, 0, 0, 0, 0
    for m in range(1, 13):
        df = capturas.get(f"ConsultaPPTO_RRC_{m}_tradicional.xlsx")
        if df is None or not len(df):
            res.append(f"MAL RRC mes {m}: no hay ConsultaR")
            continue
        p = anio * 100 + m
        bc = pd.to_numeric(df["DESVIACION2026"]).sum() + g.get("BC_SONR_2026", 0)
        for _, r in df.iterrows():
            ramo = int(r["Ramo"])
            n += 1
            if ramo in (71, 73):
                v, _f = ins._is_crudo("RRC", ramo, int(r["CALMONTH"]))
                esp_media = v if not math.isnan(v) else area_cat[ramo]
            else:
                v, fuente = ins._is_crudo("RRC", ramo, p)
                esp_media = v if not math.isnan(v) else AREA["IS"]
                fa += fuente == "IS (FA)"
                area_usada += math.isnan(v)
            v99, _f = ins._is_crudo("RRC", ramo, p, "99.5")
            g_ = ins._factor_crudo("GTO", "RRC", ramo, p)
            fm = ins._factor_crudo("MR", "RRC", ramo, p)
            esp = {"BELMEDIA": esp_media, "BEL99": v99 if not math.isnan(v99) else AREA["IS99"],
                   "BELGASTO": g_ if not math.isnan(g_) else AREA["GTO"], "FACTORMR": fm}
            if not math.isnan(fm) and fm > 0:
                esp["MR2026_TCVal"] = r["MONTO_PI"] * r["PORC_ND"] * fm * r["TC_Valuación"]
                mr_bd += 1
            else:
                dur = r["DURMXN"] if r["MonedaOri"] == 1 else r["DUROTR"]
                esp["MR2026_TCVal"] = -1 * r["DESVIACION2026"] * g["RCS"] * g["COC"] * dur * (1 / bc)
                mr_formula += 1
            malos += sum(not _igual(r[c], v) for c, v in esp.items())
        if m >= 9:
            for ramo in (40, 50, 80, 90):
                sub = df[df["Ramo"] == ramo]
                if len(sub) and not all(_igual(x, ins.is_fa[("RRC", ramo, p)]["media"]) for x in sub["BELMEDIA"]):
                    malos += 1
                    res.append(f"MAL RRC {p} ramo {ramo}: no usa el IS de FA")
    res.append(f"{'OK ' if not malos else 'MAL'} RRC 12 meses: {n} contratos, {malos} indices distintos a los esperados; "
               f"IS de FA en {fa} contratos (40, 50, 80 y 90 proyectados), IS del area en {area_usada}; MR con FACTOR MR "
               f"de la BD en {mr_bd} y con la formula de capital en {mr_formula} (71 y 73)")
    return res


def verificar_sonr(ins: InsumosBD, g: dict, marcos: list, texto_script: str, anio: int) -> list[str]:
    if not g:
        return ["MAL SONR: el script no termino"]
    cuerpo = texto_script[texto_script.index("def Metodo_propio():"):texto_script.index("def Metodo_propio_reforecast")]
    bc = float(re.search(r"^\s*BC = (-?[\d.]+)", cuerpo, re.M).group(1))
    bc2 = float(re.search(r"^\s*BC2 = (-?[\d.]+)", cuerpo, re.M).group(1))
    meses = [f for f in marcos if "NoLAG" in f.columns and "Factor_MR" in f.columns]
    res = []
    if len(meses) != 12:
        res.append(f"MAL SONR: {len(meses)} meses de metodo propio (se esperaban 12)")
    malos, n, fa, area, mr_bd, mr_formula = 0, 0, 0, 0, 0, 0
    for df in meses:
        p = int(df["AñoMes"].iloc[0])
        for _, r in df.iterrows():
            ramo, k = int(r["Ramo"]), int(r["NoLAG"])
            n += 1
            v, fuente = ins._is_crudo("SONR", ramo, p)
            v99, _f = ins._is_crudo("SONR", ramo, p, "99.5")
            lag = ins._hp_crudo(f"LAG {k}", ramo, p)
            fr = ins._factor_crudo("RET", "SONR", ramo, p)
            fm = ins._factor_crudo("MR", "SONR", ramo, p)
            fa += fuente == "IS (FA)"
            area += sum(math.isnan(x) for x in (v, v99, lag, fr))
            esp = {"Ind Sin SONR Media": v if not math.isnan(v) else AREA["IS_SONR"],
                   "Ind Sin SONR 99.5%": v99 if not math.isnan(v99) else AREA["IS_SONR99"],
                   f"LAG {k}": lag if not math.isnan(lag) else AREA["LAG"],
                   "Factor_Ret": fr if not math.isnan(fr) else AREA["RET"], "Factor_MR": fm}
            if not math.isnan(fm) and fm > 0:
                esp["MR"] = r["Prima Dev"] * r["LAG"] * fm
                mr_bd += 1
            else:
                esp["MR"] = (r["Desviacion"] / -bc) * bc2
                mr_formula += 1
            malos += sum(not _igual(r[c], x) for c, x in esp.items())
    res.append(f"{'OK ' if not malos else 'MAL'} SONR 12 meses: {n} renglones (ramo x LAG x mes), {malos} parametros "
               f"distintos a los esperados; IS de FA en {fa}, {area} celdas del area; MR con Factor_MR de la BD en "
               f"{mr_bd} y con la formula del script en {mr_formula}")
    return res


def main():
    if len(sys.argv) < 2:
        raise SystemExit(__doc__)
    carpeta = Path(sys.argv[1]).resolve()
    rrc, sonr = next(carpeta.glob("*RRC*_BD.py")), next(carpeta.glob("*SONR*_BD.py"))
    bd = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else buscar_bd(None, [SCRIPTS])
    trabajo = (Path(sys.argv[3]) if len(sys.argv) > 3 else Path(tempfile.mkdtemp(prefix="sim_aod_"))).resolve()
    trabajo.mkdir(parents=True, exist_ok=True)
    for f in (rrc, sonr, SCRIPTS / "insumos_bd.py"):
        shutil.copy(f, trabajo / f.name)
    shutil.copy(bd, trabajo / NOMBRE_BD)
    if bd.with_name("Diagnostico_Proyeccion.xlsx").exists():
        shutil.copy(bd.with_name("Diagnostico_Proyeccion.xlsx"), trabajo / "Diagnostico_Proyeccion.xlsx")
    ins = InsumosBD(trabajo / NOMBRE_BD, verbose=False)
    anio = int(re.search(r"^zAño\s*=\s*(\d{4})", rrc.read_text(encoding="utf-8"), re.M).group(1))
    tablas_area(trabajo / "area", anio)
    lineas, ok = [], True
    for etiqueta, tc_hasta in (("TC en la base todo el ano", anio * 100 + 12), ("TC solo hasta el ultimo mes real", ins.ultimo_real)):
        print(f"\n######## {etiqueta}")
        capturas, textos, marcos = {}, [], []
        simular_access(ins, tc_hasta)
        redirigir(trabajo / "area", trabajo / f"salida_{tc_hasta}", capturas)
        try:
            g_rrc = correr(trabajo / rrc.name, trabajo / "cwd_ajena", textos, marcos)
            marcos_rrc = len(marcos)
            g_sonr = correr(trabajo / sonr.name, trabajo / "cwd_ajena", textos, marcos)
        finally:
            restaurar()
        res = [f"{etiqueta}:"]
        res += verificar_rrc(ins, g_rrc, capturas, anio)
        res += verificar_sonr(ins, g_sonr, marcos[marcos_rrc:], (trabajo / sonr.name).read_text(encoding="utf-8"), anio)
        for nombre in ("RRC_esc.xlsx", "SONR_esc.xlsx", "Parametros_usados_RRC.xlsx", "Parametros_usados_SONR.xlsx"):
            res.append(f"{'OK ' if (trabajo / f'salida_{tc_hasta}' / nombre).exists() else 'MAL'} {nombre}")
        sin_tc = sorted({int(t.split("cambio de ")[1].split(":")[0]) for t in textos if "no trae tipo de cambio" in t})
        esperado = [p for p in rango_meses(anio * 100 + 1, anio * 100 + 12) if p > tc_hasta]
        res.append(f"{'OK ' if sorted(set(sin_tc)) == esperado else 'MAL'} aviso de meses sin tipo de cambio en la base: "
                   f"{sin_tc or 'ninguno'}")
        lineas += res
        ok = ok and not any(x.startswith("MAL") for x in res)
    print("\n===== VERIFICACION")
    for x in lineas:
        print("  ", x)
    print(f"\nCarpeta de trabajo: {trabajo}")
    print("RESULTADO: todo OK" if ok else "RESULTADO: con fallas")
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
