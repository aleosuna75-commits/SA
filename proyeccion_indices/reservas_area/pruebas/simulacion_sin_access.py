# -*- coding: utf-8 -*-
"""Corre los dos scripts adaptados sin la base Access ni los archivos del area, con datos sinteticos, para probar el
flujo completo y que los indices que usan sean los de la BD proyectada.

Uso: python simulacion_sin_access.py "<ruta de la BD proyectada>" [carpeta de trabajo]

Simula pyodbc (tipo de cambio y contratos de la base de valuacion) y arma los catalogos y CSV auxiliares con pocos
renglones inventados; la BD proyectada es la real. Al final compara, renglon por renglon, los indices que tomaron los
scripts contra los de la BD (insumos_bd) y escribe un resumen. Los montos que salen NO son reservas: solo prueban que
el codigo corre y toma los insumos correctos.
"""
from __future__ import annotations

import re
import runpy
import shutil
import sys
import tempfile
import traceback
import types
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

AQUI = Path(__file__).resolve().parent
SCRIPTS = AQUI.parent
sys.path.insert(0, str(SCRIPTS))
from insumos_bd import InsumosBD, RAMOS_SCRIPT, RAMOS_SONR, mes_mas, rango_meses   # noqa: E402

RAMO_NOMBRE = {10: "Vida", 31: "Acc Per.", 35: "GMM", 39: "Salud", 40: "Resp. Civil", 50: "MyT", 60: "Incendio",
               71: "Terremoto", 73: "HyORH", 80: "Agropecuario", 90: "Autos", 100: "Crédito", 110: "Diversos"}
SUBRAMOS = {10: [10], 31: [30, 31, 32, 33], 35: [34, 35, 36], 39: [37, 38, 39], 40: [40, 46], 50: [50], 60: [60],
            71: [71], 73: [73], 80: [80], 90: [90], 100: [100], 110: [110]}
CEBE = {r: f"A{r:03d}" for r in RAMOS_SCRIPT}                      # centro de beneficio sintetico por ramo


def preparar(bd: Path, trabajo: Path) -> Path:
    trabajo.mkdir(parents=True, exist_ok=True)
    (trabajo / "csv").mkdir(exist_ok=True)
    shutil.copy(bd, trabajo / "BD_ BEL - IRR - MR_Proyeccion.xlsx")
    diag = bd.with_name("Diagnostico_Proyeccion.xlsx")
    if diag.exists():
        shutil.copy(diag, trabajo / "Diagnostico_Proyeccion.xlsx")
    for f in ("insumos_bd.py", "reforecastRRC_v11_insumosBD.py", "ReforecastSONR_v4_insumosBD.py"):
        shutil.copy(SCRIPTS / f, trabajo / f)
    (trabajo / "config_local.py").write_text(f'''# config sintetica de la simulacion
from pathlib import Path
CARPETA = Path(r"{trabajo}")
RUTA_BD = CARPETA / "BD_ BEL - IRR - MR_Proyeccion.xlsx"
RUTA_DIAGNOSTICO = CARPETA / "Diagnostico_Proyeccion.xlsx"
CARPETA_SALIDA = CARPETA / "salidas_area"
ANIO = None
MES = None
USAR_IS_FA = True
RAMOS_IS_FA = {{"RRC": (40, 50, 80, 90), "SONR": (40, 50, 80, 90)}}
MESES_FALTANTES_ESC3 = "BD"
ESCENARIO_BASE_CSV = {{}}
GUARDAR_INTERMEDIOS = True
ACCESS_DBQ = r"simulada.accdb"
CATALOGOS = str(CARPETA / "Catalogos.xlsx")
CSV_AUXILIARES_RRC = str(CARPETA / "csv")
CSV_AUXILIARES_SONR = str(CARPETA / "csv")
CSV_DURACION_RRC = str(CARPETA / "csv" / "ParametrosMens.csv")
CSV_IS_CAT = None
PPTO_TECNICO_RRC = str(CARPETA / "csv" / "PptoTecnico.csv")
PPTO_TECNICO_SONR = PPTO_TECNICO_RRC
FINVIG_AJUSTE_MANUAL = 45930
MR_RRC = {{"REAL": {{"RCS": 1.0e8, "COC": 0.1}}, "REAL_USD": {{"RCS": 5.0e6, "COC": 0.1}},
          "PPTO": {{"RCS": 1.0e8, "BC": -1.0e9, "COC": 0.1}}, "BC_SONR": 0.0}}
MR_SONR = {{"BC": -1.0e9, "BC2": 2.0e7}}
''', encoding="utf-8")
    # catalogos: hoja Valores con xRamo en J:M y xPais en O:T (encabezado en el renglon 2)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Valores"
    ws["J2"], ws["K2"], ws["L2"], ws["M2"] = "SR", "Ramo", "Desc", "Nota"
    ws["O2"], ws["P2"], ws["Q2"], ws["R2"], ws["S2"], ws["T2"] = "País", "TerrSAP", "Region", "x", "y", "z"
    fila = 3
    for clave, subs in SUBRAMOS.items():
        for sr in subs:
            ws.cell(fila, 10, sr)
            ws.cell(fila, 11, RAMO_NOMBRE[clave])
            fila += 1
    for i, (pais, terr) in enumerate([(1, "R01"), (2, "R03"), (3, "R05"), (4, "R04")], start=3):
        ws.cell(i, 15, pais)
        ws.cell(i, 16, terr)
    wb.save(trabajo / "Catalogos.xlsx")
    csv = trabajo / "csv"
    pd.DataFrame({"Llave": ["000000-0-0-0"]}).to_csv(csv / "LlavesPol.csv", index=False)
    cols_rrc = ['SRamo', 'Pais', 'TipoRea', 'OfiRepPt', 'MonedaOri', 'CorrTom', 'CiaTom', 'CtoTom', 'Susc', 'Período',
                'CALMONTH', 'IniVig', 'FinVig', 'PrimaTomadaOri', 'PmaTom_sEROri', 'PrimaCedidaOri', 'PrimaTomadaNal',
                'PmaTom_sERNal', 'PrimaCedidaNal', 'REGION', 'Ramo', 'LLAVE', 'LN2', 'FRECUENCIA', 'MONTO_PI', 'CESION',
                'BELMEDIA', 'BELGASTO', 'BEL99', 'DURMXN', 'DUROTR', 'RETMXN', 'RETOTR', 'PORC_ND', 'CEDIDA',
                'TC_Valuación', 'TC_CierreAnterior']
    pd.DataFrame(columns=cols_rrc).to_csv(csv / "AjManuales.csv", index=False)
    pd.DataFrame(columns=['CALMONTH', 'Ramo_filt', 'Pais', 'TipoRea', 'CorrTom', 'CiaTom', 'CtoTom', 'Susc', 'MonedaOri',
                          'IniVig', 'FinVig', 'Periodo', 'PmaTomOri', 'PmaTomNal']).to_csv(csv / "AjManuales_SONR.csv",
                                                                                           index=False)
    pd.DataFrame({"CeBe": [CEBE[r] for r in RAMOS_SCRIPT], "Ramo": RAMOS_SCRIPT,
                  "Ramo2": RAMOS_SCRIPT}).to_csv(csv / "Subramo.csv", index=False)
    anios = [str(a) for a in range(2020, 2030)]
    pd.DataFrame({"Ramo": RAMOS_SCRIPT, **{a: 0.05 for a in anios}}).to_csv(csv / "CesionPI.csv", index=False)
    pd.DataFrame({"AFUN": ["LN04001", "LN04005", "LN04008"],
                  "Linea de Negocio": ["Norte", "Sur", "Facultativos"]}).to_csv(csv / "AFUN.csv", index=False)
    pd.DataFrame({"Llave": ["1-1-1-1"], "Periodo": [1]}).to_csv(csv / "zFrecuencias.csv", index=False)
    llaves3 = [f"{r}-{reg}-{t}" for r in RAMOS_SCRIPT for reg in ("R01", "R03", "R05", "R04") for t in (1, 2, 3)]
    pd.DataFrame({"Llave": llaves3, **{a: 0.3 for a in anios}}).to_csv(csv / "TablaCesion_Esc1.csv", index=False)
    pd.DataFrame({"Llave": ["1-1-1-2024-1"], "Porcentaje Cedido": [0.2]}).to_csv(csv / "Cesion ID Esp.csv", index=False)
    pd.DataFrame({"Ramo": RAMOS_SCRIPT, "Pesos_dur": 1.1, "Resto Monedas_dur": 1.2, "Pesos_ret": 0.8,
                  "Resto Monedas_ret": 0.9}).to_csv(csv / "ParametrosMens.csv", index=False)
    pd.DataFrame([{"Ramo": r, "NoLAG": k} for r in RAMOS_SONR for k in range(1, 11)]).to_csv(
        csv / "TablaBase_MetodoPropio.csv", index=False)
    return trabajo


def presupuesto(trabajo: Path, anio: int):
    """PptoTecnico sintetico: prima mensual por ramo, con los nombres de columnas de BW (prueba _normalizar_ppto)."""
    filas = []
    rng = np.random.RandomState(1)
    for m in range(1, 13):
        for r in RAMOS_SCRIPT:
            for tipo in (1, 2):
                filas.append({"/ERP/GL_ACCT": 6101010000, "0CALMONTH": anio * 100 + m, "/ERP/PROFTCTR": CEBE[r],
                              "ZREGIONRP": "R01", "ZTIPOREAS": tipo, "ZOFICN_RP": 1, "/ERP/FUNCAREA": "LN04001",
                              "ZMGA": 1, "/ERP/TWAERS": "USD", "ZTIPOCES": 1, "/ERP/PRODUCT": 0.1, "ZCORREDOR": 1,
                              "ZCEDENTE": 1, "ZCONTRATO": 1, "ZSUSCYEAR": anio, "/ERP/AMOUNT": -float(rng.randint(50, 500) * 1000)})
    pd.DataFrame(filas).to_csv(trabajo / "csv" / "PptoTecnico.csv", index=False)


def contratos(anio: int, meses: list[int]):
    """Contratos sinteticos: por ramo y mes, uno proporcional (TipoRea 1) y uno facultativo (TipoRea 2)."""
    rng = np.random.RandomState(7)
    filas = []
    for m in meses:
        for clave, subs in SUBRAMOS.items():
            for tipo in (1, 2):
                prima = float(rng.randint(100, 900) * 1000)
                ini = pd.Timestamp(year=m // 100, month=m % 100, day=1)
                fin = ini + pd.DateOffset(months=12)
                filas.append({"SRamo": subs[0], "Pais": 1 if tipo == 1 else 2, "TipoRea": tipo, "OfiRepPt": 1,
                              "MonedaOri": 31 if clave % 20 else 1, "CorrTom": 1, "CiaTom": clave, "CtoTom": m % 7 + 1,
                              "Susc": m // 100, "Período": 1 if tipo == 1 else 0, "CALMONTH": m,
                              "IniVig": ini, "FinVig": fin, "PmaTomOri": prima, "PmaTomNal": prima * 18.0,
                              "PrimaTomadaOri": -prima, "PmaTom_sEROri": -prima, "PrimaCedidaOri": prima * 0.3,
                              "PrimaTomadaNal": -prima * 18.0, "PmaTom_sERNal": -prima * 18.0, "PrimaCedidaNal": prima * 5.4,
                              "Ramo_filt": clave, "Periodo": 1 if tipo == 1 else 0})
    return pd.DataFrame(filas)


def simular_access(ins: InsumosBD):
    """Inyecta un pyodbc falso y un pd.read_sql que responde con el TC y los contratos sinteticos."""
    tc_filas = []
    for p in rango_meses(mes_mas(ins.primer_periodo_hp, -12), ins.ultimo_real):
        usd = ins.tc_de(p)
        for mon, v in ((1, 1.0), (31, usd), (2, usd * 1.08)):
            tc_filas.append({"cTCAD_FecAMD": p, "cMON_Id": mon, "cTCAD_Mnt": v, "Llave": f"{p}-{mon}"})
    tc = pd.DataFrame(tc_filas)
    meses = rango_meses(mes_mas(ins.ultimo_real, -11 * 12), ins.ultimo_real)
    gonz = contratos(ins.anio, meses)

    class _Conn:
        def cursor(self):
            return self

        def close(self):
            pass

    falso = types.ModuleType("pyodbc")
    falso.connect = lambda conn_str: _Conn()
    sys.modules["pyodbc"] = falso

    def read_sql(sql, conn, *a, **k):
        if "aMOT_MovTipCambio" in sql:
            if "cMON_Id = 31" in sql:
                return tc[tc["cMON_Id"] == 31][["cTCAD_FecAMD", "cTCAD_Mnt"]].reset_index(drop=True)
            return tc.copy()
        if "aMOG_MovGonzalo" in sql:
            if "Ramo_filt" in sql:                   # consulta del SONR: meses entre zInicio y zFin (el WHERE del SQL)
                m = re.search(r"aPog_MesProc\) >= (\d+) And Val\(aPog_MesProc\) <= (\d+)", sql)
                sel = gonz[(gonz["CALMONTH"] >= int(m.group(1))) & (gonz["CALMONTH"] <= int(m.group(2)))]
                cols = ['CALMONTH', 'Ramo_filt', 'Pais', 'TipoRea', 'CorrTom', 'CiaTom', 'CtoTom', 'Susc', 'MonedaOri',
                        'IniVig', 'FinVig', 'Periodo', 'PmaTomOri', 'PmaTomNal']
                return sel[cols].reset_index(drop=True)
            # consulta del RRC: ultimos 12 meses o contratos en vigor a la fecha de valuacion
            m = re.search(r"aPog_MesProc\) > (\d+) and Val\(aPog_MesProc\) <= (\d+)\)  or \( FinVig > (\S+) and IniVig < ", sql)
            ini, fin = int(m.group(1)), int(m.group(2))
            fecha = pd.to_datetime(m.group(3)[:10], format="%d/%m/%Y")          # zFechaValuacion = 'dd/mm/aaaa'
            # (en la base los renglones en vigor caen tambien en la ventana de 12 meses: el script indexa xPND por
            # CALMONTH y solo trae esas 12 llaves)
            sel = gonz[(gonz["CALMONTH"] > ini) & (gonz["CALMONTH"] <= fin) & (gonz["FinVig"] > fecha)]
            cols = ['SRamo', 'Pais', 'TipoRea', 'OfiRepPt', 'MonedaOri', 'CorrTom', 'CiaTom', 'CtoTom', 'Susc', 'Período',
                    'CALMONTH', 'IniVig', 'FinVig', 'PrimaTomadaOri', 'PmaTom_sEROri', 'PrimaCedidaOri', 'PrimaTomadaNal',
                    'PmaTom_sERNal', 'PrimaCedidaNal']
            return sel[cols].reset_index(drop=True)
        raise ValueError(f"SQL no simulado: {sql[:80]}")
    pd.read_sql = read_sql


def correr(trabajo: Path, nombre: str) -> dict:
    print(f"\n===== {nombre}")
    cwd = Path.cwd()
    try:
        import os
        os.chdir(trabajo)
        sys.path.insert(0, str(trabajo))
        for m in ("config_local", "insumos_bd"):
            sys.modules.pop(m, None)
        return runpy.run_path(str(trabajo / nombre), run_name="__main__")
    except Exception:  # noqa: BLE001
        traceback.print_exc()
        return {}
    finally:
        os.chdir(cwd)


def verificar(ins: InsumosBD, g_rrc: dict, g_sonr: dict, trabajo: Path) -> list[str]:
    res = []
    salida = trabajo / "salidas_area"
    for nombre in ("RRC_esc.xlsx", "SONR_esc.xlsx", "Parametros_usados_RRC.xlsx", "Parametros_usados_SONR.xlsx"):
        res.append(f"{'OK ' if (salida / nombre).exists() else 'FALTA'} {nombre}")
    ult = ins.ultimo_real
    # RRC: ConsultaR del ultimo mes real (escenario 2) y la del escenario 4 (indices de diciembre)
    for etiqueta, llave, periodo, con_tc in (("RRC escenario 2 mes real", "df_Real_IS_Real", ult, True),
                                             ("RRC escenario 4 (real, indices de diciembre)", "xReforecast_Real", ins.anio * 100 + 12, False)):
        df = g_rrc.get(llave)
        if df is None or not len(df):
            res.append(f"SIN DATOS {etiqueta}")
            continue
        df = df[pd.to_numeric(df["Ramo"], errors="coerce").notna()]
        malos = 0
        for _, r in df.iterrows():
            ramo = int(r["Ramo"])
            esp = {"BELMEDIA": ins.indice("RRC", ramo, periodo) if ramo not in (71, 73) else ins.indice("RRC", ramo, int(r["CALMONTH"])),
                   "BEL99": ins.indice("RRC", ramo, periodo, "99.5"), "BELGASTO": ins.factor_gasto(ramo, periodo)}
            for c, v in esp.items():
                x = float(r[c]) if r[c] not in ("", None) else float("nan")
                if not (abs(x - v) <= 1e-9 or (np.isnan(x) and np.isnan(v))):
                    malos += 1
        res.append(f"{'OK ' if not malos else 'MAL'} {etiqueta}: {len(df)} contratos, {malos} indices distintos a la BD")
        col_mr = next((c for c in df.columns if str(c).startswith("MR") and str(c).endswith("_TCVal")), None)
        if col_mr is not None:
            tc = pd.to_numeric(df["TC_Valuación"]) if con_tc else 1.0     # (el escenario 4 ya trabaja en USD: sin TC)
            fmr = pd.Series([ins.factor_mr("RRC", int(x), periodo) for x in df["Ramo"]], index=df.index)
            esp_mr = pd.to_numeric(df["MONTO_PI"]) * pd.to_numeric(df["PORC_ND"]) * fmr * tc
            dif = int((abs(pd.to_numeric(df[col_mr]) - esp_mr) > 1e-6 * (1 + abs(esp_mr))).sum())
            res.append(f"{'OK ' if not dif else 'MAL'} {etiqueta}: MR = PND x FACTOR MR de la BD en {len(df) - dif} de {len(df)} contratos")
    # SONR: metodo propio del ultimo mes real
    df = g_sonr.get("df_Real_IS_Real")
    if df is None or not len(df):
        res.append("SIN DATOS SONR metodo propio")
    else:
        malos = 0
        for _, r in df.iterrows():
            ramo, k = int(r["Ramo"]), int(r["NoLAG"])
            esp_is = ins.indice("SONR", ramo, ult)
            esp_lag = 1 - ins.parametro(f"LAG {k}", ramo, ult)
            esp_fr = ins.factor_ret(ramo, ult)
            bel = r["Prima Dev"] * r["LAG"] * r["Ind Sin SONR Media"]
            esp_mr = r["Prima Dev"] * r["LAG"] * ins.factor_mr("SONR", ramo, ult)
            if abs(r["Ind Sin SONR Media"] - esp_is) > 1e-9 or abs(r["LAG"] - esp_lag) > 1e-9 \
                    or abs(r["Factor_Ret"] - esp_fr) > 1e-9 or abs(bel - r["BEL_RIESGO"]) > 1e-6 \
                    or abs(esp_mr - r["MR"]) > 1e-6 * (1 + abs(esp_mr)):
                malos += 1
        res.append(f"{'OK ' if not malos else 'MAL'} SONR metodo propio {ult}: {len(df)} renglones (ramo x LAG), {malos} con "
                   "IS, LAG, Factor_Ret o MR distintos a la BD")
    # escenarios y meses en las salidas
    for nombre, reserva in (("RRC_esc.xlsx", "RRC"), ("SONR_esc.xlsx", "SONR")):
        ruta = salida / nombre
        if not ruta.exists():
            continue
        out = pd.read_excel(ruta)
        esc = sorted(out["Escenario"].dropna().unique().tolist())
        per = out[out["Escenario"] == 3]["Periodo"]
        res.append(f"{reserva}: escenarios {esc}; renglones {len(out)}; escenario 3 de {per.min()} a {per.max()}; "
                   f"ramos {sorted(pd.to_numeric(out['Ramo'], errors='coerce').dropna().astype(int).unique().tolist())}")
    return res


def main():
    bd = Path(sys.argv[1]).resolve()
    trabajo = (Path(sys.argv[2]) if len(sys.argv) > 2 else Path(tempfile.mkdtemp(prefix="sim_area_"))).resolve()
    preparar(bd, trabajo)
    ins = InsumosBD(trabajo / "BD_ BEL - IRR - MR_Proyeccion.xlsx", verbose=False)
    presupuesto(trabajo, ins.anio)
    simular_access(ins)
    g_rrc = correr(trabajo, "reforecastRRC_v11_insumosBD.py")
    g_sonr = correr(trabajo, "ReforecastSONR_v4_insumosBD.py")
    print("\n===== VERIFICACION")
    lineas = verificar(ins, g_rrc, g_sonr, trabajo)
    for linea in lineas:
        print("  ", linea)
    print(f"\nCarpeta de trabajo: {trabajo}")
    if any(l.startswith(("MAL", "FALTA", "SIN DATOS")) for l in lineas):
        print("RESULTADO: con fallas")
        sys.exit(1)
    print("RESULTADO: todo OK")


if __name__ == "__main__":
    main()
