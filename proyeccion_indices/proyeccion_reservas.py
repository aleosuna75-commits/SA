# -*- coding: utf-8 -*-
"""
PROYECCION DE INDICES Y MONTOS DE RESERVAS  (202609 - 202712)
==============================================================

Proyecta, con exactamente el mismo formato de los archivos actuales:

  1. "BD_ BEL - IRR - MR.xlsx"
       * HParametros_2026  -> indices (Ind Sin RRC, Ind sin RRC 99.5%, Ind Sin SONR
                              Media, Ind Sin SONR 99.5%) y LAG 1..LAG 10, solo a partir
                              de los renglones "Real" (el mismo filtro de la hoja).
       * BD_Montos_RRC_SONR -> RRC (NETO, BRUTO, BEL, GTO, IRR, MR) y SONR (NETO,
                              BRUTO, BEL, IRR, MR) por ramo.
  2. "BD_ RFV.xlsx" (Fianzas, generado con llenar_bd_rfv.py)
       * BD_Montos_RRC_SONR -> RFV NETO, RFV BRUTO, RFV IRR y RCONT por ramo.

METODOLOGIA (resumen; el detalle queda en salidas/Diagnostico_Proyeccion.xlsx)
------------------------------------------------------------------------------
a) Coherencia contable. No se proyectan todas las lineas por separado: se proyectan
   los "drivers" y el resto se deriva, de modo que las identidades que cumple la
   historia se cumplen tambien en la proyeccion:
       RRC :  BEL (nivel),  g = GTO/BEL,  m = MR/BEL,  c = IRR/BRUTO (razones)
              GTO = g*BEL ; MR = m*BEL ; BRUTO = BEL+GTO+MR ; IRR = c*BRUTO ; NETO = BRUTO-IRR
       SONR:  BEL (nivel),  m = MR/BEL,  c = IRR/BRUTO
              MR = m*BEL ; BRUTO = BEL+MR ; IRR = c*BRUTO ; NETO = BRUTO-IRR
       RFV :  BRUTO (nivel), c = IRR/BRUTO ; IRR = c*BRUTO ; NETO = BRUTO-IRR
              RCONT: se proyecta el total y se reparte por ramo con la mezcla de los ultimos 8 meses.
b) Modelos: Ingenuo (caminata aleatoria), Ingenuo estacional, Media 12m, Suavizamiento
   exponencial simple (SES), Holt con tendencia amortiguada, Holt-Winters amortiguado
   (estacionalidad 12), Metodo Theta, ARIMA (orden por AICc) y Regresion log-lineal
   (tendencia + estacionalidad mensual). Montos e indices se modelan en logaritmos
   (efectos multiplicativos, garantiza positividad); razones y LAGs en escala original.
c) Seleccion con validacion fuera de muestra del METODO COMPLETO: se simula haber
   proyectado desde 8 cortes historicos a 1-16 meses con cada procedimiento candidato
   (modelos solos y combinaciones de pesos iguales). Montos: menor WAPE (pondera por USD),
   desempate por sesgo agregado. Indices, razones y LAGs: menor AvgRelMAE (error relativo
   al ingenuo) y, si su IC bootstrap al 90% incluye 1, el mejor procedimiento sin tendencia.
   Con la historia a 202608 resultaron: montos -> Theta + Holt amortiguado; indices ->
   Ingenuo + Media 12m; razones -> Ingenuo + SES; LAGs -> Ingenuo. Elegir el modelo serie
   por serie ("torneo") fue menos preciso y queda como modo alternativo. Los montos se
   modelan en USD: modelar en MXN y convertir con el TC real fue menos preciso (Danos y
   Fianzas); se puede cambiar en MODELAR_EN_MXN.
d) Backtest rolling-origin por serie (todos los modelos, mismos cortes, sin fuga): da el
   MASE de cada modelo por serie (hoja Series_Modelos), los intervalos al 80% y una red de
   seguridad: si en la propia serie el procedimiento validado es > 1.5x peor que el
   ingenuo (p.ej. cambio de regimen), se usa el ensamble propio de esa serie.
e) Reglas actuariales / de calidad de datos (todas reportadas en la hoja "Alertas"):
       - series en cero en los ultimos 6 meses -> se proyectan en cero;
       - parametros "en escalon" (se actualizan esporadicamente, >=50% de meses sin
         cambio) -> se mantiene el ultimo valor;
       - series cortas (<18 obs) -> SES;
       - razones y LAGs se acotan al rango historico observado y al dominio actuarial
         (cesion en [0, 1], %GTO y %MR >= 0);
       - RCONT con factores por mes del trimestre (acumula meses 1-2, libera en el 3);
       - alerta de saltos atipicos en el ultimo mes;
       - 99.5% >= media cuando asi ha sido siempre en la historia del ramo;
       - textos con espacios / celdas vacias se limpian; huecos <= 6 meses se interpolan
         y con huecos mayores solo se usa la historia posterior;
       - LAG en 0 despues de que el patron acumulado ya supero 50% se trata como faltante.

USO: abrir en VSCode y ejecutar (F5 o "Run Python File"). Los paquetes que falten se
instalan automaticamente en el interprete activo. Parametros en la seccion CONFIGURACION.
"""
from __future__ import annotations

import importlib
import os
import subprocess
import sys

# Un hilo de BLAS por proceso: el paralelismo se hace por series (evita sobresuscripcion de CPU).
# Debe definirse antes de importar numpy / statsmodels.
for _var in ("OMP_NUM_THREADS", "OPENBLAS_NUM_THREADS", "MKL_NUM_THREADS", "NUMEXPR_NUM_THREADS",
             "VECLIB_MAXIMUM_THREADS"):
    os.environ.setdefault(_var, "1")


def _asegurar_paquetes(paquetes: dict[str, str]) -> None:
    """Instala con el mismo interprete que ejecuta el script los paquetes que falten."""
    faltantes = []
    for modulo, nombre_pip in paquetes.items():
        try:
            importlib.import_module(modulo)
        except ImportError:
            faltantes.append(nombre_pip)
    if faltantes:
        print(f"Instalando paquetes faltantes: {', '.join(faltantes)} ...", flush=True)
        try:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", *faltantes])
        except subprocess.CalledProcessError as e:
            raise SystemExit(f"No se pudieron instalar {faltantes} (sin internet, proxy o permisos). Instalalos "
                             f"manualmente con: {sys.executable} -m pip install {' '.join(faltantes)}") from e
        # si pip instalo en la carpeta de usuario (Python 'para todos los usuarios' en Windows), agregarla
        import site
        usuario = site.getusersitepackages()
        if os.path.isdir(usuario) and usuario not in sys.path:
            site.addsitedir(usuario)
        importlib.invalidate_caches()


_asegurar_paquetes({
    "openpyxl": "openpyxl>=3.1",
    "numpy": "numpy",
    "pandas": "pandas",
    "scipy": "scipy",
    "statsmodels": "statsmodels>=0.14",
    "matplotlib": "matplotlib",
})

import math  # noqa: E402

import re  # noqa: E402
import time  # noqa: E402
import warnings  # noqa: E402
from concurrent.futures import ProcessPoolExecutor  # noqa: E402
from copy import copy  # noqa: E402
from dataclasses import dataclass, field  # noqa: E402
from datetime import datetime  # noqa: E402
from pathlib import Path  # noqa: E402

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.formula.translate import Translator  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from scipy.stats import norm as _normal  # noqa: E402
from statsmodels.tsa.arima.model import ARIMA  # noqa: E402
from statsmodels.tsa.forecasting.theta import ThetaModel  # noqa: E402
from statsmodels.tsa.holtwinters import ExponentialSmoothing, SimpleExpSmoothing  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from excel_fiel import guardar_libro, verificar_escritura  # noqa: E402
from tipo_cambio import TC_FCST  # noqa: E402

# =============================================================================
# CONFIGURACION
# =============================================================================
CARPETA = Path(__file__).resolve().parent
ENTRADAS = CARPETA / "entradas"
SALIDAS = CARPETA / "salidas"

ARCHIVO_BD_DANOS = ENTRADAS / "BD_ BEL - IRR - MR.xlsx"
ARCHIVO_BD_RFV = SALIDAS / "BD_ RFV.xlsx"        # lo genera llenar_bd_rfv.py (se corre solo si falta)

SALIDA_BD_DANOS = SALIDAS / "BD_ BEL - IRR - MR_Proyeccion.xlsx"
SALIDA_BD_RFV = SALIDAS / "BD_ RFV_Proyeccion.xlsx"
SALIDA_DIAGNOSTICO = SALIDAS / "Diagnostico_Proyeccion.xlsx"
SALIDA_GRAFICAS = SALIDAS / "Graficas_Proyeccion.pdf"

PERIODO_INICIO = 202609          # primer mes proyectado
PERIODO_FIN = 202712             # ultimo mes proyectado

HOJA_PARAMETROS = "HParametros_2026"
HOJA_MONTOS = "BD_Montos_RRC_SONR"
TIPO_INDICE_BASE = "Real"        # solo estos renglones se usan como historia
ETIQUETA_PROYECCION = "Proyección"  # "Tipo de Indice" de los renglones nuevos en HParametros
                                    # (se agrega al filtro de la hoja para que se vean junto a "Real")

INDICES = ["Ind Sin RRC", "Ind sin RRC 99.5%", "Ind Sin SONR Media", "Ind Sin SONR 99.5%"]
LAGS = [f"LAG {i}" for i in range(1, 11)]
PARES_MEDIA_995 = [("Ind Sin RRC", "Ind sin RRC 99.5%"), ("Ind Sin SONR Media", "Ind Sin SONR 99.5%")]

# Tipo de cambio que se escribe en la columna TC: supuesto de Inversiones (tipo_cambio.py, tomado de
# TC_Real_Esti.xlsx hoja TC: FCST para 2026 y FCST 2027 para 2027). Se aplica a los renglones nuevos y,
# si ACTUALIZAR_TC_EXISTENTES, tambien a los renglones que ya existen de esos meses (p.ej. 202606-202612,
# que en la BD traian una interpolacion). Meses sin dato toman el ultimo TC de la BD.
# Los montos se modelan en USD, asi que el TC no cambia las cifras proyectadas (salvo con MODELAR_EN_MXN).
TC_PROYECCION: dict[int, float] = dict(TC_FCST)
ACTUALIZAR_TC_EXISTENTES = True

RESALTAR_PROYECCION = False      # True = relleno azul claro en celdas proyectadas
SOBRESCRIBIR_PERIODOS_CON_DATOS = False  # False = se detiene si algun mes a proyectar ya trae cifras
REGENERAR_BD_RFV = True          # True = vuelve a llenar BD_ RFV en cada corrida (tarda ~20 s); False = solo si
                                 # no existe o alguna entrada tiene fecha mas reciente
COLOR_RESALTADO = "DDEBF7"
GENERAR_GRAFICAS = True
N_PROCESOS = None                # None = automatico (nucleos-1); 1 = sin paralelismo

# Parametros del torneo de modelos
CV_ORIGENES = 8                  # fechas de corte del backtest
CV_PASO = 2                      # meses entre fechas de corte
MIN_ENTRENAMIENTO = 12           # obs minimas para ajustar en cada corte
MIN_OBS_TORNEO = 18              # con menos obs no hay torneo (SES / ultimo valor)
MIN_OBS_ESTACIONAL = 24          # obs minimas para modelos estacionales (2 ciclos)
H_CV = 16                        # horizonte maximo evaluado en el backtest
MIN_CORTES_CV = 4                # cortes minimos del backtest
MAX_HUECO_INTERPOLABLE = 6       # huecos mas largos cortan la historia
# Seleccion del metodo de proyeccion:
#   "validado" = por tipo de serie se usa el procedimiento con menor error en la validacion fuera
#                de muestra del metodo completo (se recalcula en cada corrida)  <- recomendado
#   "torneo"   = cada serie elige su propio ensamble con su backtest (mas flexible, pero en la
#                validacion con estos datos resulto menos preciso que "validado")
MODO_SELECCION = "validado"
MIN_SERIES_VALIDACION = 5        # series minimas de un tipo para validar; si no, PROCEDIMIENTO_RESPALDO
MAX_INTERPOLADOS_VALIDACION = 3  # la validacion solo usa series con a lo mas 3 meses interpolados
TOLERANCIA_WAPE = 0.02           # montos: procedimientos a <= 2% del mejor WAPE se desempatan por sesgo
N_BOOTSTRAP = 2000               # remuestreos (por serie) para el IC90 del AvgRelMAE
CORTES_PISO_INTERVALO = 6        # con menos cortes de backtest, piso de caminata aleatoria a los intervalos
FACTOR_RESPALDO_SERIE = 1.5      # red de seguridad: si en el backtest de la propia serie el procedimiento
                                 # validado es > 1.5x peor que el ingenuo, se usa el ensamble propio de la serie
MAX_MODELOS_ENSAMBLE = 3
TOLERANCIA_ENSAMBLE = 1.5        # entran al ensamble los modelos con error <= 1.5 x el mejor
UMBRAL_CERO_MONTOS = 1.0         # |monto| < 1 USD se considera 0
MESES_CERO_EXTINTA = 6           # ultimos N meses en cero -> serie extinta
UMBRAL_ESCALON = 0.5             # >= 50% de meses sin cambio -> parametro en escalon
MESES_SIN_DATO_MAX = 12          # sin dato en los ultimos N meses -> no se proyecta
NIVEL_INTERVALO = 0.80

# Estructura de conceptos (drivers y derivados)
ESTRUCTURA = {
    "DANOS": {
        "RRC": {"nivel": "BEL", "razones_bel": ["GTO", "MR"], "cesion": "IRR"},
        "SONR": {"nivel": "BEL", "razones_bel": ["MR"], "cesion": "IRR"},
    },
    "FIANZAS": {
        "RFV": {"nivel": "BRUTO", "razones_bel": [], "cesion": "IRR", "totales_con_mezcla": ["RCONT"]},
    },
}
MESES_MEZCLA = 8                 # meses recientes para repartir por ramo los totales (RCONT)
# Moneda en que se modelan los montos. En USD por evidencia: en backtest, modelar en MXN (historia USD x
# TC real) y convertir con el TC real futuro fue MENOS preciso que modelar directo en USD, tanto en Danos
# (WAPE 22.5% vs 19.8%, 21 series) como en Fianzas (RFV BRUTO 7.4% vs 6.0%): las reservas se comportan como
# montos en dolares. True = modelar en MXN y convertir con el TC de cada mes proyectado de la BD.
MODELAR_EN_MXN = {"DANOS": False, "FIANZAS": False}
ESTACIONALIDAD_TRIMESTRAL = ("RCONT",)   # acumula en meses 1-2 del trimestre y libera en el 3
DOMINIO_CESION = (0.0, 1.0)      # IRR/BRUTO entre 0 y 100%
DOMINIO_RAZON_BEL = (0.0, None)  # GTO/BEL y MR/BEL no negativos

warnings.filterwarnings("ignore")


# =============================================================================
# UTILERIAS
# =============================================================================
def periodo_a_indice(p: int) -> int:
    return (p // 100) * 12 + (p % 100 - 1)


def indice_a_periodo(i: int) -> int:
    return (i // 12) * 100 + (i % 12) + 1


def rango_periodos(p_ini: int, p_fin: int) -> list[int]:
    return [indice_a_periodo(i) for i in range(periodo_a_indice(p_ini), periodo_a_indice(p_fin) + 1)]


def a_numero(v):
    """Convierte a float: numeros, textos con espacios (incluye \\xa0) o comas decimales.
    Devuelve (valor|nan, fue_texto)."""
    if v is None:
        return math.nan, False
    if isinstance(v, bool):
        return math.nan, False
    if isinstance(v, (int, float)):
        return float(v), False
    if isinstance(v, str):
        t = v.replace("\xa0", " ").strip().replace(",", "")
        if t == "":
            return math.nan, True
        try:
            return float(t), True
        except ValueError:
            return math.nan, True
    return math.nan, False


def norm(t) -> str:
    return re.sub(r"\s+", " ", str(t or "").replace("\xa0", " ")).strip().upper()


# =============================================================================
# MODELOS
# =============================================================================
def _m_ingenuo(z, h):
    return np.repeat(z[-1], h)


def _m_ingenuo_estacional(z, h):
    n = len(z)
    return np.array([z[n + i - 12 * (i // 12 + 1)] for i in range(h)])


def _m_media12(z, h):
    return np.repeat(np.mean(z[-12:]), h)


def _m_ses(z, h):
    return SimpleExpSmoothing(z, initialization_method="estimated").fit().forecast(h)


def _m_holt_amortiguado(z, h):
    return ExponentialSmoothing(z, trend="add", damped_trend=True,
                                initialization_method="estimated").fit().forecast(h)


def _m_hw_amortiguado(z, h):
    return ExponentialSmoothing(z, trend="add", damped_trend=True, seasonal="add", seasonal_periods=12,
                                initialization_method="estimated").fit().forecast(h)


def _m_theta(z, h):
    estacional = len(z) >= 24
    return np.asarray(ThetaModel(z, period=12, deseasonalize=estacional, method="additive")
                      .fit().forecast(h))


ORDENES_ARIMA = [((0, 1, 1), "n"), ((1, 1, 0), "n"), ((1, 1, 1), "n"), ((0, 1, 1), "t"),
                 ((1, 0, 0), "c"), ((2, 0, 0), "c")]


def _aicc(res, n):
    k = len(res.params)
    return res.aic + (2 * k * (k + 1)) / max(n - k - 1, 1)


def seleccionar_arima(z):
    mejor, mejor_aicc = None, math.inf
    for orden, tendencia in ORDENES_ARIMA:
        try:
            res = ARIMA(z, order=orden, trend=tendencia).fit()
            a = _aicc(res, len(z))
            if np.isfinite(a) and a < mejor_aicc:
                mejor, mejor_aicc = (orden, tendencia), a
        except Exception:  # noqa: BLE001
            continue
    return mejor


def _m_arima(z, h, orden=None):
    orden = orden or seleccionar_arima(z)
    if orden is None:
        raise ValueError("sin ARIMA valido")
    return ARIMA(z, order=orden[0], trend=orden[1]).fit().forecast(h)


def _m_regresion(z, h, mes0=0):
    """Regresion (en la escala transformada) con tendencia lineal y dummies de mes
    (si hay >=36 obs) sobre las ultimas 36 observaciones. mes0 = mes calendario (0-11)
    de la primera observacion de z."""
    n = len(z)
    k = min(n, 36)
    t = np.arange(n - k, n + h, dtype=float)
    meses = (mes0 + np.arange(n - k, n + h)) % 12
    X = [np.ones_like(t), t]
    if k >= 36:
        for m in range(1, 12):
            X.append((meses == m).astype(float))
    X = np.column_stack(X)
    beta, *_ = np.linalg.lstsq(X[:k], z[n - k:], rcond=None)
    return X[k:] @ beta


CATALOGO = {
    "Ingenuo": (_m_ingenuo, 1, False),
    "Ingenuo estacional": (_m_ingenuo_estacional, MIN_OBS_ESTACIONAL, True),
    "Media 12m": (_m_media12, 12, False),
    "SES": (_m_ses, 8, False),
    "Holt amortiguado": (_m_holt_amortiguado, 10, False),
    "Holt-Winters amortiguado": (_m_hw_amortiguado, MIN_OBS_ESTACIONAL, True),
    "Theta": (_m_theta, 10, False),
    "ARIMA": (_m_arima, 12, False),
    "Regresion log-lineal": (_m_regresion, 12, False),
}

MODELOS_POR_TIPO = {            # candidatos del backtest por serie (diagnostico / modo "torneo")
    "nivel": ["Ingenuo", "Ingenuo estacional", "SES", "Holt amortiguado", "Holt-Winters amortiguado",
              "Theta", "ARIMA", "Regresion log-lineal"],
    "indice": ["Ingenuo", "Ingenuo estacional", "Media 12m", "SES", "Holt amortiguado",
               "Holt-Winters amortiguado", "Theta", "ARIMA", "Regresion log-lineal"],
    "razon": ["Ingenuo", "Media 12m", "SES", "Holt amortiguado", "Holt-Winters amortiguado", "Theta", "ARIMA"],
    "lag": ["Ingenuo", "Media 12m", "SES", "Holt amortiguado", "Theta"],
}

# Procedimientos candidatos (combinaciones de pesos iguales) que se comparan en la validacion
# fuera de muestra del metodo completo; se elige, por tipo de serie, el de menor error.
MODELOS_BASE_VALIDACION = ["Ingenuo", "SES", "Holt amortiguado", "Theta", "Media 12m"]
PROCEDIMIENTOS_CANDIDATOS = {
    "Ingenuo": ["Ingenuo"],
    "SES": ["SES"],
    "Holt amortiguado": ["Holt amortiguado"],
    "Theta": ["Theta"],
    "Media 12m": ["Media 12m"],
    "Ingenuo + SES": ["Ingenuo", "SES"],
    "Ingenuo + Theta": ["Ingenuo", "Theta"],
    "Ingenuo + Media 12m": ["Ingenuo", "Media 12m"],
    "Theta + Holt amortiguado": ["Theta", "Holt amortiguado"],
    "Ingenuo + SES + Media 12m": ["Ingenuo", "SES", "Media 12m"],
    "SES + Holt amortiguado + Theta": ["SES", "Holt amortiguado", "Theta"],
}
PROCEDIMIENTOS_SIN_TENDENCIA = {"Ingenuo", "SES", "Media 12m", "Ingenuo + SES", "Ingenuo + Media 12m",
                                "Ingenuo + SES + Media 12m"}
# Respaldo si un tipo no tiene suficientes series para validar (resultado de la validacion con la
# historia a 202608)
PROCEDIMIENTO_RESPALDO = {
    "nivel": "Theta + Holt amortiguado",
    "indice": "Ingenuo + Media 12m",
    "razon": "Ingenuo + SES",
    "lag": "Ingenuo",
}


@dataclass
class Serie:
    clave: tuple                 # (libro, grupo, concepto, ramo)
    tipo: str                    # nivel | razon | indice | lag
    periodos: list[int]          # historia mensual continua
    valores: list[float]         # nan = sin dato
    ultimo_periodo: int          # ultimo mes real global (la proyeccion arranca al mes siguiente)
    h: int                       # meses a proyectar
    procedimiento: tuple = ()    # modelos a combinar (pesos iguales); vacio = torneo por serie
    trimestral: bool = False     # estacionalidad por mes del trimestre (p.ej. RCONT)
    dominio: tuple = (None, None)  # cotas actuariales (min, max) de la serie proyectada
    moneda: str = ""             # moneda en que se modela (montos)


@dataclass
class Resultado:
    clave: tuple
    tipo: str
    pronostico: list[float]
    li: list[float]
    ls: list[float]
    regla: str
    transformacion: str = ""
    n_obs: int = 0
    modelos: dict = field(default_factory=dict)     # nombre -> MASE backtest de la serie
    pesos: dict = field(default_factory=dict)        # nombre -> peso en la combinacion final
    mase_ensamble: float = math.nan                  # MASE backtest de la combinacion final
    alertas: list = field(default_factory=list)
    n_cortes: int = 0
    moneda: str = ""
    historia_periodos: list = field(default_factory=list)
    historia_valores: list = field(default_factory=list)


def _ajustar(nombre, z, h, mes0, orden_arima):
    f = CATALOGO[nombre][0]
    if nombre == "ARIMA":
        return np.asarray(f(z, h, orden_arima), dtype=float)
    if nombre == "Regresion log-lineal":
        return np.asarray(f(z, h, mes0), dtype=float)
    return np.asarray(f(z, h), dtype=float)


def _origenes_cv(n: int, minimo: int) -> list[int]:
    """Cortes del backtest (longitud del entrenamiento). Se combinan cortes con ventana de prueba
    completa de H_CV meses (miden el error a 16 meses con varias repeticiones) y cortes recientes
    (horizontes cortos con la informacion mas nueva). Todos los modelos se evaluan en los mismos."""
    fin_completo = n - H_CV
    completos = range(fin_completo - CV_PASO * (CV_ORIGENES - 1), fin_completo + 1, CV_PASO)
    recientes = range(fin_completo + CV_PASO, n, CV_PASO)
    return [o for o in sorted(set(completos) | set(recientes)) if o >= minimo]


def preparar(serie: Serie):
    """Reglas de calidad de datos. Devuelve (res, prep): si la serie queda resuelta por una regla
    (ceros, constante, escalon, sin datos) prep es None y res trae la proyeccion; si no, prep trae
    la historia limpia {y, per, brecha}."""
    tipo, h = serie.tipo, serie.h
    per = list(serie.periodos)
    y = np.array(serie.valores, dtype=float)
    res = Resultado(clave=serie.clave, tipo=tipo, pronostico=[math.nan] * h, li=[math.nan] * h,
                    ls=[math.nan] * h, regla="")

    def _constante(valor, regla):
        res.pronostico, res.li, res.ls = [float(valor)] * h, [float(valor)] * h, [float(valor)] * h
        res.regla = regla
        return res, None

    if tipo == "nivel":
        y[np.abs(y) < UMBRAL_CERO_MONTOS] = 0.0
    validos = np.where(~np.isnan(y))[0]
    if len(validos) == 0:
        res.regla = "Sin datos (no se proyecta)"
        return res, None
    # meses sin dato al final respecto al ultimo mes real global
    brecha = periodo_a_indice(serie.ultimo_periodo) - periodo_a_indice(per[validos[-1]])
    if brecha >= MESES_SIN_DATO_MAX:
        res.regla = f"Sin datos en los ultimos {MESES_SIN_DATO_MAX} meses (no se proyecta)"
        return res, None

    # 1) ceros / vacios iniciales (antes de que exista el negocio o el dato)
    if tipo in ("nivel", "razon"):
        nz = np.where((~np.isnan(y)) & (y != 0))[0]
        if len(nz) == 0:
            return _constante(0.0, "Serie en cero (se proyecta en cero)")
        inicio = nz[0]
    else:
        inicio = validos[0]
    fin = validos[-1]
    y, per = y[inicio:fin + 1], per[inicio:fin + 1]

    # 2) huecos largos: solo se usa el tramo posterior al ultimo hueco > MAX_HUECO_INTERPOLABLE
    corte, racha = 0, 0
    for i, es_nulo in enumerate(np.isnan(y)):
        racha = racha + 1 if es_nulo else 0
        if racha > MAX_HUECO_INTERPOLABLE:
            corte = i + 1
    if corte:
        res.alertas.append(f"Hueco de mas de {MAX_HUECO_INTERPOLABLE} meses sin dato: se usa solo la historia "
                           f"desde {per[corte]}")
        y, per = y[corte:], per[corte:]
        primero = np.where(~np.isnan(y))[0][0]
        y, per = y[primero:], per[primero:]

    # 3) parametro "en escalon" (se evalua con lo observado, sin interpolar)
    obs = y[~np.isnan(y)][-25:]
    escalon = len(obs) >= 7 and np.mean(np.abs(np.diff(obs)) < 1e-12) >= UMBRAL_ESCALON

    # 4) huecos cortos: interpolacion lineal
    n_interp = int(np.isnan(y).sum())
    if n_interp:
        idx = np.arange(len(y))
        ok = ~np.isnan(y)
        res.alertas.append(f"{n_interp} mes(es) sin dato interpolado(s) linealmente")
        y = np.interp(idx, idx[ok], y[ok])
    res.historia_periodos, res.historia_valores = per, [float(v) for v in y]
    n = len(y)
    res.n_obs = n
    if brecha > 0:
        res.alertas.append(f"Ultimo dato en {per[-1]}; se proyecta desde ahi ({brecha} mes(es) de rezago)")

    # 5) series extintas / constantes / en escalon
    if tipo == "nivel" and n >= MESES_CERO_EXTINTA and np.all(y[-MESES_CERO_EXTINTA:] == 0):
        return _constante(0.0, f"Serie en cero los ultimos {MESES_CERO_EXTINTA} meses (se proyecta en cero)")
    if n == 1 or np.all(np.abs(np.diff(y[-25:])) < 1e-12):
        return _constante(y[-1], "Serie constante (se mantiene el ultimo valor)")
    if escalon:
        return _constante(y[-1], "Parametro en escalon (se mantiene el ultimo valor)")
    return res, {"y": y, "per": per, "brecha": brecha, "n_interp": n_interp}


def pronosticar(serie: Serie) -> Resultado:
    """Proyecta una serie mensual aplicando reglas de calidad de datos, backtest por serie y el
    procedimiento validado para su tipo (o el torneo por serie si no se indica procedimiento)."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return _pronosticar(serie)


def _pronosticar(serie: Serie) -> Resultado:
    res, prep = preparar(serie)
    if prep is None:
        return _aplicar_dominio(res, serie.dominio)
    tipo, h = serie.tipo, serie.h
    y, per, brecha = prep["y"], prep["per"], prep["brecha"]
    n = len(y)
    hh = h + brecha                       # si el ultimo dato es anterior al ultimo mes real

    def _salida(pron, li=None, ls=None):
        pron = np.asarray(pron, dtype=float)[brecha:]
        li = pron if li is None else np.asarray(li, dtype=float)[brecha:]
        ls = pron if ls is None else np.asarray(ls, dtype=float)[brecha:]
        return pron, li, ls

    usar_log = tipo in ("nivel", "indice") and np.all(y > 0)
    res.transformacion = "log" if usar_log else "ninguna"
    res.moneda = serie.moneda
    z = np.log(y) if usar_log else y.copy()
    mes0 = per[0] % 100 - 1
    zq = _cuantil_normal()
    trimestral = serie.trimestral and n >= 12
    if trimestral:
        res.transformacion += " + factores por mes del trimestre"

    def inv(v):
        return np.exp(v) if usar_log else v

    def pred(m, z_tr, hz, orden):
        """Pronostico en escala original desde el entrenamiento z_tr (con estacionalidad trimestral
        estimada solo con z_tr, para no usar informacion futura)."""
        if trimestral:
            f = _factores_trimestrales(z_tr, mes0)
            q = (mes0 + np.arange(len(z_tr) + hz)) % 12 % 3
            s = f[q]
            return inv(_ajustar(m, z_tr - s[:len(z_tr)], hz, mes0, orden) + s[len(z_tr):])
        return inv(_ajustar(m, z_tr, hz, mes0, orden))

    # alerta de atipico en el ultimo mes (salto > 3 desviaciones de los cambios mensuales)
    if tipo in ("nivel", "indice") and n >= 13:
        d = np.diff(z)
        sd_d = np.std(d[:-1], ddof=1)
        if sd_d > 0 and abs(d[-1] - np.mean(d[:-1])) > 3 * sd_d:
            res.alertas.append(f"Salto atipico en el ultimo mes ({(y[-1] / y[-2] - 1):+.1%}); revisar si es un "
                               "evento puntual que se liberara")

    # series cortas: SES
    if n < MIN_OBS_TORNEO:
        try:
            pron = pred("SES", z, hh, None) if n >= 6 else inv(_m_ingenuo(z, hh))
            res.regla = "Serie corta: SES" if n >= 6 else "Serie corta: ultimo valor"
        except Exception:  # noqa: BLE001
            pron = inv(_m_ingenuo(z, hh))
            res.regla = "Serie corta: ultimo valor"
        sd = (np.std(np.diff(z), ddof=1) if n > 2 else 0.0) * np.sqrt(np.arange(1, hh + 1))
        li, ls = (pron * np.exp(-zq * sd), pron * np.exp(zq * sd)) if usar_log else (pron - zq * sd,
                                                                                     pron + zq * sd)
        p, li, ls = _salida(pron, li, ls)
        res.pronostico, res.li, res.ls = list(p), list(li), list(ls)
        res.alertas.append(f"Solo {n} observaciones: proyeccion de baja confiabilidad")
        return _post_proceso(res, y, tipo, serie.dominio)

    # ---------------- backtest rolling-origin por serie (mismos cortes para todos los modelos)
    procedimiento = [m for m in serie.procedimiento if n >= CATALOGO[m][1] + 2]
    candidatos = [m for m in MODELOS_POR_TIPO[tipo] if n >= CATALOGO[m][1] + 2]
    candidatos += [m for m in procedimiento if m not in candidatos]
    while True:
        minimo = max([MIN_ENTRENAMIENTO] + [CATALOGO[m][1] for m in candidatos])
        origenes = _origenes_cv(n, minimo)
        exigentes = [m for m in candidatos if CATALOGO[m][1] == minimo and CATALOGO[m][1] > MIN_ENTRENAMIENTO
                     and m not in procedimiento]
        if len(origenes) >= MIN_CORTES_CV or not exigentes:
            break
        candidatos = [m for m in candidatos if m not in exigentes]
    # orden ARIMA: en el backtest se elige solo con la historia hasta el primer corte (sin fuga de
    # informacion); para la proyeccion final, con toda la historia
    orden_cv = orden_final = None
    if "ARIMA" in candidatos and origenes:
        orden_cv = seleccionar_arima(z[:min(origenes)])
        orden_final = seleccionar_arima(z)
        if orden_cv is None or orden_final is None:
            candidatos.remove("ARIMA")
    escala = np.mean(np.abs(np.diff(y)))
    if not np.isfinite(escala) or escala <= 0:
        escala = max(np.mean(np.abs(y)), 1e-12)

    h_cv = min(hh, H_CV)
    errores = {m: [[] for _ in range(h_cv)] for m in candidatos}
    pred_cv = {m: {} for m in candidatos}
    for o in origenes:
        hz = min(h_cv, n - o)
        for m in candidatos:
            try:
                pr = pred(m, z[:o], hz, orden_cv)
                if not np.all(np.isfinite(pr)):
                    continue
                pred_cv[m][o] = pr
                for k in range(hz):
                    errores[m][k].append(abs(pr[k] - y[o + k]))
            except Exception:  # noqa: BLE001
                continue
    puntajes = {}
    for m in candidatos:
        por_h = [np.mean(e) for e in errores[m] if len(e) > 0]
        if len(pred_cv[m]) >= 0.8 * len(origenes) and por_h:      # debe ajustar en >= 80% de los cortes
            puntajes[m] = float(np.mean(por_h) / escala)
    res.n_cortes = len(origenes)
    res.modelos = {m: round(v, 4) for m, v in sorted(puntajes.items(), key=lambda kv: kv[1])}

    # ---------------- modelos de la proyeccion final
    mase_proc = _mase_combinacion(procedimiento, pred_cv, origenes, y, h_cv, escala) if procedimiento else math.nan
    respaldo = (procedimiento and "Ingenuo" in puntajes and np.isfinite(mase_proc)
                and len(origenes) >= MIN_CORTES_CV and mase_proc > FACTOR_RESPALDO_SERIE * puntajes["Ingenuo"])
    if respaldo:
        res.alertas.append(
            f"Red de seguridad: en el backtest de esta serie el procedimiento validado ({' + '.join(procedimiento)}, "
            f"MASE {mase_proc:.2f}) fue > {FACTOR_RESPALDO_SERIE}x peor que el ultimo valor (MASE "
            f"{puntajes['Ingenuo']:.2f}); se usa el ensamble propio de la serie")
    if procedimiento and not respaldo:
        pesos = {m: 1.0 / len(procedimiento) for m in procedimiento}
        res.regla = "Procedimiento validado: " + " + ".join(procedimiento)
    elif puntajes:
        orden = sorted(puntajes, key=puntajes.get)
        mejor = puntajes[orden[0]]
        elegidos = [m for m in orden if puntajes[m] <= TOLERANCIA_ENSAMBLE * max(mejor, 1e-12)]
        elegidos = elegidos[:MAX_MODELOS_ENSAMBLE]
        inv_err = {m: 1.0 / max(puntajes[m], 1e-12) for m in elegidos}
        pesos = {m: inv_err[m] / sum(inv_err.values()) for m in elegidos}
        res.regla = "Red de seguridad: ensamble propio de la serie" if respaldo else "Torneo por serie + ensamble"
    else:
        pesos = {"Ingenuo": 1.0}
        res.regla = "Sin backtest valido: ultimo valor"

    finales = {}
    for m in pesos:
        try:
            pr = pred(m, z, hh, orden_final)
            if np.all(np.isfinite(pr)):
                finales[m] = pr
        except Exception:  # noqa: BLE001
            res.alertas.append(f"No se pudo ajustar {m}; se omite de la combinacion")
    if not finales:
        finales = {"Ingenuo": inv(_m_ingenuo(z, hh))}
        res.alertas.append("Ningun modelo ajusto; se usa el ultimo valor")
    total = sum(pesos.get(m, 0) for m in finales) or 1.0
    pesos = {m: pesos.get(m, 0) / total for m in finales} if any(pesos.get(m) for m in finales) \
        else {m: 1.0 / len(finales) for m in finales}
    pron = sum(pesos[m] * finales[m] for m in finales)
    res.pesos = {m: round(w, 4) for m, w in pesos.items()}

    # ---------------- errores de la combinacion final en el backtest -> MASE e intervalos
    resid = [[] for _ in range(h_cv)]
    abs_ens = [[] for _ in range(h_cv)]
    for o in origenes:
        if not all(o in pred_cv.get(m, {}) for m in pesos):
            continue
        ens = sum(pesos[m] * pred_cv[m][o] for m in pesos)
        for k in range(len(ens)):
            real, est = y[o + k], ens[k]
            abs_ens[k].append(abs(est - real))
            if usar_log:
                if est > 0 and real > 0:
                    resid[k].append(math.log(real) - math.log(est))
            else:
                resid[k].append(real - est)
    por_h = [np.mean(e) for e in abs_ens if e]
    res.mase_ensamble = float(np.mean(por_h) / escala) if por_h else math.nan
    sd = _sd_por_horizonte(resid, hh)
    s_hist = s_proy = None
    if trimestral:
        f_q = _factores_trimestrales(z, mes0)
        s_todo = f_q[(mes0 + np.arange(n + hh)) % 12 % 3]
        s_hist, s_proy = s_todo[:n], s_todo[n:]
    if len(origenes) < CORTES_PISO_INTERVALO:
        # pocos cortes de backtest (series cortas): piso de caminata aleatoria con la volatilidad observada
        z_ds = z - s_hist if trimestral else z
        sd = np.maximum(sd, np.std(np.diff(z_ds), ddof=1) * np.sqrt(np.arange(1, hh + 1)))
    if usar_log:
        li, ls = pron * np.exp(-zq * sd), pron * np.exp(zq * sd)
    else:
        li, ls = pron - zq * sd, pron + zq * sd
    p, li, ls = _salida(pron, li, ls)
    res.pronostico, res.li, res.ls = list(p), list(li), list(ls)
    if "Ingenuo" in puntajes and np.isfinite(res.mase_ensamble) and res.mase_ensamble > puntajes["Ingenuo"] * 1.10:
        res.alertas.append("En el backtest de esta serie la proyeccion no supera al ultimo valor (ingenuo)")
    ajuste = None
    if trimestral and usar_log:
        ajuste = (np.exp(s_hist), np.exp(s_proy[brecha:]))
    return _post_proceso(res, y, tipo, serie.dominio, ajuste)


def _mase_combinacion(modelos, pred_cv, origenes, y, h_cv, escala):
    """MASE del backtest de una combinacion de pesos iguales (misma definicion que los puntajes)."""
    errores = [[] for _ in range(h_cv)]
    for o in origenes:
        if not all(o in pred_cv.get(m, {}) for m in modelos):
            continue
        ens = np.mean([pred_cv[m][o] for m in modelos], axis=0)
        for k in range(len(ens)):
            errores[k].append(abs(ens[k] - y[o + k]))
    por_h = [np.mean(e) for e in errores if e]
    return float(np.mean(por_h) / escala) if por_h else math.nan


# =============================================================================
# VALIDACION FUERA DE MUESTRA DEL METODO (seleccion del procedimiento por tipo)
# =============================================================================
def _trabajo_validacion(args):
    """Pronosticos de los modelos base desde un corte historico (para la validacion del metodo)."""
    clave, tipo, y, o, h = args
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        usar_log = tipo in ("nivel", "indice") and np.all(y > 0)
        z = np.log(y[:o]) if usar_log else np.asarray(y[:o], dtype=float)
        salida = {}
        for m in MODELOS_BASE_VALIDACION:
            try:
                pz = _ajustar(m, z, h, 0, None)
                if np.all(np.isfinite(pz)):
                    salida[m] = np.exp(pz) if usar_log else pz
            except Exception:  # noqa: BLE001
                continue
        return clave, tipo, o, salida


def validar_procedimientos(series: list[Serie]):
    """Backtest del metodo completo: para cada serie con historia suficiente se simula haber proyectado
    desde varios cortes pasados (ventanas completas de H_CV meses) con cada procedimiento candidato, y se
    compara contra lo real. Criterio de seleccion por tipo de serie:
      * montos (nivel): WAPE = suma de errores absolutos / suma de montos reales (pondera por USD); entre
        los procedimientos a menos de TOLERANCIA_WAPE del mejor, el de menor sesgo agregado a 13-16 meses.
      * indices, razones y LAGs (adimensionales): AvgRelMAE = media geometrica entre series del error
        relativo al ingenuo. Si su IC bootstrap al 90% incluye 1 (no mejora significativamente al ingenuo),
        se elige el mejor procedimiento SIN tendencia (principio de parsimonia)."""
    trabajos, reales = [], {}
    for s in series:
        res, prep = preparar(s)
        if prep is None or prep["brecha"] > 0 or prep["n_interp"] > MAX_INTERPOLADOS_VALIDACION:
            continue            # la validacion solo usa series con datos genuinos
        y = prep["y"]
        n = len(y)
        if s.tipo in ("nivel", "indice") and np.min(y) <= 0:
            continue
        origenes = [o for o in range(n - H_CV - CV_PASO * (CV_ORIGENES - 1), n - H_CV + 1, CV_PASO)
                    if o >= MIN_OBS_ESTACIONAL]
        if len(origenes) < 3:
            continue
        reales[s.clave] = (y, s.tipo, prep["per"])
        trabajos += [(s.clave, s.tipo, y, o, H_CV) for o in origenes]
    print(f"   Validacion del metodo: {len(reales)} series, {len(trabajos)} pronosticos fuera de muestra",
          flush=True)
    salidas = _mapear(_trabajo_validacion, trabajos)

    # errores por procedimiento, solo en cortes donde el procedimiento Y el ingenuo tienen pronostico
    filas = []                     # (tipo, procedimiento, clave, periodo_corte, h, pronostico, real)
    for clave, tipo, o, sal in salidas:
        y, _, per = reales[clave]
        if "Ingenuo" not in sal:
            continue
        for nombre, comp in PROCEDIMIENTOS_CANDIDATOS.items():
            if not all(m in sal for m in comp):
                continue
            f = np.mean([sal[m] for m in comp], axis=0)
            for k in range(H_CV):
                filas.append((tipo, nombre, clave, per[o - 1], k + 1, float(f[k]), float(y[o + k])))
    rng = np.random.default_rng(12345)
    tabla, seleccion = [], {}
    for tipo in ("nivel", "indice", "razon", "lag"):
        ft = [f for f in filas if f[0] == tipo]
        if not ft:
            continue
        cortes_ok = {}             # (clave, corte) disponibles por procedimiento
        for _, nombre, clave, corte, _, _, _ in ft:
            cortes_ok.setdefault(nombre, set()).add((clave, corte))
        resumen = []
        for nombre in PROCEDIMIENTOS_CANDIDATOS:
            fp = [f for f in ft if f[1] == nombre]
            if not fp:
                continue
            comunes = cortes_ok[nombre] & cortes_ok.get("Ingenuo", set())
            fp = [f for f in fp if (f[2], f[3]) in comunes]
            fb = [f for f in ft if f[1] == "Ingenuo" and (f[2], f[3]) in comunes]
            e_proc, e_base = {}, {}
            for f in fp:
                e_proc[f[2]] = e_proc.get(f[2], 0.0) + abs(f[5] - f[6])
            for f in fb:
                e_base[f[2]] = e_base.get(f[2], 0.0) + abs(f[5] - f[6])
            claves = [c for c in e_proc if e_base.get(c, 0) > 0 and e_proc[c] > 0]
            if not claves:
                continue
            logrel = np.log(np.array([e_proc[c] / e_base[c] for c in claves]))
            boot = [np.exp(np.mean(rng.choice(logrel, size=len(logrel)))) for _ in range(N_BOOTSTRAP)]
            # WAPE y sesgo agregado por libro (cada libro en su propia moneda) y promedio entre libros
            wape_libro, sesgo_libro = [], []
            for libro in sorted({f[2][0] for f in fp}):
                fl = [f for f in fp if f[2][0] == libro]
                suma_real = sum(abs(f[6]) for f in fl)
                if suma_real:
                    wape_libro.append(sum(abs(f[5] - f[6]) for f in fl) / suma_real * 100)
                agregado = {}
                for f in fl:
                    if f[4] >= 13:
                        a = agregado.setdefault((f[3], f[4]), [0.0, 0.0])
                        a[0] += f[5]
                        a[1] += f[6]
                if agregado:
                    sesgo_libro.append(float(np.mean([(F - A) / abs(A) for F, A in agregado.values() if A])) * 100)
            sesgo_agr = float(np.mean(sesgo_libro)) if sesgo_libro else math.nan
            ape = lambda h1, h2: float(np.nanmean(  # noqa: E731
                [abs(f[5] - f[6]) / abs(f[6]) for f in fp if h1 <= f[4] <= h2 and f[6]]) * 100)
            resumen.append({
                "Tipo": tipo, "Procedimiento": nombre,
                "WAPE %": float(np.mean(wape_libro)) if wape_libro else math.nan,
                "AvgRelMAE": float(np.exp(np.mean(logrel))),
                "IC90 inf": float(np.quantile(boot, 0.05)), "IC90 sup": float(np.quantile(boot, 0.95)),
                "% series mejor que ingenuo": float(np.mean(logrel < 0)), "Series": len(claves),
                "MAPE 1-6": ape(1, 6), "MAPE 7-12": ape(7, 12), "MAPE 13-16": ape(13, 16),
                "Sesgo agregado % 13-16": sesgo_agr,
                "Tendencia": "no" if nombre in PROCEDIMIENTOS_SIN_TENDENCIA else "si",
            })
        if not resumen:
            continue
        elegido, motivo = None, ""
        if tipo == "nivel":
            mejor = min(d["WAPE %"] for d in resumen)
            cerca = [d for d in resumen if d["WAPE %"] <= mejor * (1 + TOLERANCIA_WAPE)]
            elegido = min(cerca, key=lambda d: abs(d["Sesgo agregado % 13-16"]))
            motivo = (f"menor WAPE (error ponderado por monto); entre los {len(cerca)} a menos de "
                      f"{TOLERANCIA_WAPE:.0%} del mejor, el de menor sesgo agregado")
        else:
            orden = sorted(resumen, key=lambda d: d["AvgRelMAE"])
            elegido = orden[0]
            motivo = "menor AvgRelMAE"
            if elegido["IC90 sup"] >= 1 and elegido["Tendencia"] == "si":
                sin_tend = [d for d in orden if d["Tendencia"] == "no"]
                if sin_tend:
                    elegido = sin_tend[0]
                    motivo = ("la mejora del mejor procedimiento no es significativa (IC90 incluye 1): se usa el "
                              "mejor procedimiento sin tendencia")
        if elegido["Series"] >= MIN_SERIES_VALIDACION:
            seleccion[tipo] = elegido["Procedimiento"]
        orden_tabla = sorted(resumen, key=lambda d: d["WAPE %"] if tipo == "nivel" else d["AvgRelMAE"])
        for d in orden_tabla:
            d["Seleccionado"] = ("SI: " + motivo) if seleccion.get(tipo) == d["Procedimiento"] else ""
        tabla += orden_tabla
    return tabla, seleccion


def _mapear(funcion, trabajos):
    """map en paralelo por procesos (con respaldo secuencial)."""
    n_proc = N_PROCESOS or max(1, (os.cpu_count() or 2) - 1)
    if n_proc > 1 and len(trabajos) > 1:
        try:
            with ProcessPoolExecutor(max_workers=n_proc) as ex:
                return list(ex.map(funcion, trabajos, chunksize=4))
        except Exception as e:  # noqa: BLE001
            print(f"   Paralelismo no disponible ({e!r}); se continua en un solo proceso", flush=True)
    return [funcion(t) for t in trabajos]


def _cuantil_normal():
    return float(_normal.ppf(0.5 + NIVEL_INTERVALO / 2))


def _sd_por_horizonte(resid, hh):
    """Desviacion estandar del error por horizonte (RMSE del backtest); los horizontes sin
    observaciones se extrapolan con raiz(h) y se fuerza que no decrezca."""
    sd = np.full(hh, np.nan)
    for k, r in enumerate(resid[:hh]):
        if len(r) >= 2:
            sd[k] = math.sqrt(np.mean(np.square(r)))
    ult = np.where(~np.isnan(sd))[0]
    if len(ult) == 0:
        return np.zeros(hh)
    k0 = ult[-1]
    for k in range(hh):
        if np.isnan(sd[k]):
            ref = ult[ult <= k][-1] if np.any(ult <= k) else ult[0]
            sd[k] = sd[ref] * math.sqrt((k + 1) / (ref + 1))
    return np.maximum.accumulate(sd) if k0 >= 0 else sd


def _factores_trimestrales(z, mes0):
    """Factores (escala log) por mes del trimestre: promedio de la desviacion contra la media movil
    centrada de 3 meses, normalizados a suma cero."""
    q = (mes0 + np.arange(len(z))) % 12 % 3
    dev = [[] for _ in range(3)]
    for t in range(1, len(z) - 1):
        dev[q[t]].append(z[t] - np.mean(z[t - 1:t + 2]))
    f = np.array([np.mean(d) if d else 0.0 for d in dev])
    return f - f.mean()


def _recortar_dominio(res: Resultado, p, li, ls, dominio):
    lo_d, hi_d = dominio
    if lo_d is None and hi_d is None:
        return p, li, ls
    lo_d = -np.inf if lo_d is None else lo_d
    hi_d = np.inf if hi_d is None else hi_d
    if np.any(p < lo_d - 1e-12) or np.any(p > hi_d + 1e-12):
        res.alertas.append(f"Proyeccion fuera del dominio actuarial [{lo_d}, {hi_d}]: se acota "
                           "(la historia reciente ya estaba fuera; revisar el dato)")
    return np.clip(p, lo_d, hi_d), np.clip(li, lo_d, hi_d), np.clip(ls, lo_d, hi_d)


def _aplicar_dominio(res: Resultado, dominio) -> Resultado:
    """Dominio actuarial para las series resueltas por una regla (constante, escalon, etc.)."""
    p = np.array(res.pronostico, dtype=float)
    if np.all(np.isnan(p)):
        return res
    p, li, ls = _recortar_dominio(res, p, np.array(res.li, dtype=float), np.array(res.ls, dtype=float), dominio)
    res.pronostico, res.li, res.ls = [float(v) for v in p], [float(v) for v in li], [float(v) for v in ls]
    return res


def _post_proceso(res: Resultado, y, tipo, dominio=(None, None), ajuste=None):
    p = np.array(res.pronostico, dtype=float)
    li = np.array(res.li, dtype=float)
    ls = np.array(res.ls, dtype=float)

    if tipo in ("razon", "lag"):
        lo, hi = float(np.nanmin(y)), float(np.nanmax(y))
        if np.any(p < lo - 1e-12) or np.any(p > hi + 1e-12):
            res.alertas.append(f"Proyeccion acotada al rango historico [{lo:.4f}, {hi:.4f}]")
        p, li, ls = np.clip(p, lo, hi), np.clip(li, lo, hi), np.clip(ls, lo, hi)
    # dominio actuarial (despues del rango historico, para que siempre prevalezca)
    p, li, ls = _recortar_dominio(res, p, li, ls, dominio)
    if tipo == "nivel":
        p, li, ls = np.maximum(p, 0), np.maximum(li, 0), np.maximum(ls, 0)
    if tipo in ("nivel", "indice") and len(y) > len(p) and np.all(y > 0):
        k = len(p)
        yy, pp = (y, p) if ajuste is None else (y / ajuste[0], p / ajuste[1])   # sin estacionalidad trimestral
        cambios = np.abs(np.log(yy[k:] / yy[:-k]))
        if len(cambios) and pp[-1] > 0:
            cambio = abs(math.log(pp[-1] / yy[-1]))
            if cambio > np.max(cambios) * 1.0 + 1e-12:
                res.alertas.append(
                    f"Cambio proyectado a {k} meses ({(pp[-1] / yy[-1] - 1):+.1%}) mayor al maximo historico "
                    f"a {k} meses ({(math.exp(np.max(cambios)) - 1):.1%})")
    res.pronostico, res.li, res.ls = [float(v) for v in p], [float(v) for v in li], [float(v) for v in ls]
    return res


def correr_series(series: list[Serie]) -> dict:
    n_proc = N_PROCESOS or max(1, (os.cpu_count() or 2) - 1)
    resultados = {}
    t0 = time.time()
    if n_proc > 1 and len(series) > 1:
        try:
            with ProcessPoolExecutor(max_workers=n_proc) as ex:
                for i, r in enumerate(ex.map(pronosticar, series, chunksize=2), start=1):
                    resultados[r.clave] = r
                    if i % 25 == 0 or i == len(series):
                        print(f"   {i}/{len(series)} series ({time.time() - t0:,.0f} s)", flush=True)
            return resultados
        except Exception as e:  # noqa: BLE001
            print(f"   Paralelismo no disponible ({e!r}); se continua en un solo proceso", flush=True)
            resultados = {}
    for i, s in enumerate(series, start=1):
        r = pronosticar(s)
        resultados[r.clave] = r
        if i % 25 == 0 or i == len(series):
            print(f"   {i}/{len(series)} series ({time.time() - t0:,.0f} s)", flush=True)
    return resultados


# =============================================================================
# LECTURA DE LAS BD DE MONTOS
# =============================================================================
@dataclass
class BDMontos:
    ruta: Path
    wb: object
    ws: object
    col_concepto: int
    col_periodo: int
    col_tc: int
    cols_ramo: dict            # ramo(str) -> columna
    filas: dict                # (concepto, periodo) -> fila
    valores: dict              # (concepto, periodo, ramo) -> float
    tc: dict                   # periodo -> tc
    conceptos: list            # en el orden en que aparecen


def leer_bd_montos(ruta: Path) -> BDMontos:
    wb = openpyxl.load_workbook(ruta)
    ws = wb[HOJA_MONTOS]
    enc = {norm(ws.cell(3, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(3, c).value}
    cols_ramo = {h.split("_", 1)[1]: c for h, c in enc.items() if h.startswith("RAM_")}
    filas, valores, tc, conceptos = {}, {}, {}, []
    for r in range(4, ws.max_row + 1):
        concepto = ws.cell(r, enc["CONCEPTO"]).value
        periodo = ws.cell(r, enc["PERIODO"]).value
        if not concepto or not isinstance(periodo, (int, float)):
            continue
        concepto, periodo = norm(concepto), int(periodo)
        if concepto not in conceptos:
            conceptos.append(concepto)
        filas[(concepto, periodo)] = r
        for ramo, c in cols_ramo.items():
            v, _ = a_numero(ws.cell(r, c).value)
            valores[(concepto, periodo, ramo)] = 0.0 if math.isnan(v) else v
        t, _ = a_numero(ws.cell(r, enc["TC"]).value)
        if not math.isnan(t):
            tc[periodo] = t
    return BDMontos(ruta, wb, ws, enc["CONCEPTO"], enc["PERIODO"], enc["TC"], cols_ramo, filas, valores, tc,
                    conceptos)


def tc_para_periodo(bd: BDMontos, p: int) -> float:
    """TC con que se escribe (y se convierte) cada mes: TC_PROYECCION (supuesto de Inversiones) si lo trae
    (y, para renglones existentes, solo si ACTUALIZAR_TC_EXISTENTES); si no, el de la BD o su ultimo TC."""
    if p in TC_PROYECCION and (ACTUALIZAR_TC_EXISTENTES or p not in bd.tc):
        return TC_PROYECCION[p]
    if p in bd.tc:
        return bd.tc[p]
    return bd.tc[max(bd.tc)]


def leer_tc_real(bd: BDMontos, ultimo: int) -> dict:
    """TC real con que la fuente SAP convirtio cada mes a USD (fila 7 de BacktestingRRC, columnas SAP, de
    los Res_Rvas en entradas/). Se usa para regresar la historia a MXN; donde no hay dato se usa la BD."""
    tc = {p: v for p, v in bd.tc.items() if p <= ultimo}
    for ruta in sorted(ENTRADAS.glob("Res_Rvas_*.xlsx")):
        try:
            wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
            if "BacktestingRRC" not in wb.sheetnames:
                continue
            filas = list(wb["BacktestingRRC"].iter_rows(min_row=1, max_row=8, values_only=True))
            wb.close()
        except Exception:  # noqa: BLE001
            continue
        for c, escenario in enumerate(filas[0]):
            per, v = filas[7][c], filas[6][c]
            if norm(escenario) == "SAP" and isinstance(per, (int, float)) and isinstance(v, (int, float)) \
                    and int(per) <= ultimo:
                tc[int(per)] = float(v)
    return tc


def series_montos(bd: BDMontos, libro: str, reservas: dict, ultimo: int, h: int,
                  tc_hist: dict | None = None) -> list[Serie]:
    """reservas = {"RRC": {"nivel": "BEL", "razones_bel": [...], "cesion": "IRR"}, ...}. Construye las
    series driver (nivel y razones). Si tc_hist se indica, los montos se modelan en MXN (USD x TC real)."""
    periodos = sorted({p for (_, p) in bd.filas if p <= ultimo})
    periodos = rango_periodos(periodos[0], ultimo)
    moneda = "MXN" if tc_hist else "USD"

    def a_moneda(v, p):
        return v * tc_hist[p] if tc_hist and not math.isnan(v) else v

    series = []
    for pref, est in reservas.items():
        for ramo in bd.cols_ramo:
            def val(conc, p):
                return bd.valores.get((f"{pref} {conc}".strip(), p, ramo), math.nan)
            nivel = [a_moneda(val(est["nivel"], p), p) for p in periodos]
            series.append(Serie((libro, pref, est["nivel"], ramo), "nivel", periodos, nivel, ultimo, h,
                                moneda=moneda))
            for conc in est.get("razones_bel", []):
                raz = []
                for p in periodos:
                    b, x = val(est["nivel"], p), val(conc, p)
                    raz.append(x / b if (b and not math.isnan(b) and abs(b) >= UMBRAL_CERO_MONTOS
                                         and not math.isnan(x)) else math.nan)
                series.append(Serie((libro, pref, f"{conc}/{est['nivel']}", ramo), "razon", periodos, raz,
                                    ultimo, h, dominio=DOMINIO_RAZON_BEL))
            if est.get("cesion"):
                raz = []
                for p in periodos:
                    br = _bruto(bd, pref, est, ramo, p)
                    x = val(est["cesion"], p)
                    raz.append(x / br if (br and not math.isnan(br) and abs(br) >= UMBRAL_CERO_MONTOS
                                          and not math.isnan(x)) else math.nan)
                series.append(Serie((libro, pref, f"{est['cesion']}/BRUTO", ramo), "razon", periodos, raz,
                                    ultimo, h, dominio=DOMINIO_CESION))
        # conceptos que se proyectan como total y se reparten por ramo con la mezcla reciente (p.ej. RCONT)
        for conc in est.get("totales_con_mezcla", []):
            serie = [a_moneda(sum(bd.valores.get((norm(conc), p, r), 0.0) for r in bd.cols_ramo), p)
                     for p in periodos]
            series.append(Serie((libro, pref, f"{conc} TOTAL", "TOTAL"), "nivel", periodos, serie, ultimo, h,
                                trimestral=norm(conc) in {norm(c) for c in ESTACIONALIDAD_TRIMESTRAL},
                                moneda=moneda))
    return series


def mezcla_reciente(bd: BDMontos, concepto: str, ultimo: int, meses: int) -> dict:
    """Participacion de cada ramo en el concepto durante los ultimos `meses` meses con dato."""
    periodos = [p for p in sorted({p for (c, p) in bd.filas if c == norm(concepto) and p <= ultimo})
                if any(bd.valores.get((norm(concepto), p, r), 0.0) for r in bd.cols_ramo)][-meses:]
    totales = {r: sum(bd.valores.get((norm(concepto), p, r), 0.0) for p in periodos) for r in bd.cols_ramo}
    suma = sum(totales.values())
    return {r: (v / suma if suma else 0.0) for r, v in totales.items()}


def _bruto(bd, pref, est, ramo, p):
    """BRUTO historico tal como esta en la BD (en RRC/SONR la historia cumple BRUTO = BEL + GTO + MR)."""
    return bd.valores.get((f"{pref} BRUTO", p, ramo), math.nan)


def periodo_ultimo_real(periodos_proy: list[int]) -> int:
    return indice_a_periodo(periodo_a_indice(periodos_proy[0]) - 1)


def derivar_montos(bd: BDMontos, libro: str, reservas: dict, resultados: dict, periodos_proy: list[int],
                   en_mxn: bool = False):
    """Aplica las identidades contables y regresa {(concepto, periodo, ramo): valor en USD}. Si los montos
    se modelaron en MXN se convierten con el TC de cada mes proyectado (el mismo que se escribe en la BD)."""
    tc_proy = np.array([tc_para_periodo(bd, p) for p in periodos_proy]) if en_mxn else 1.0
    salida = {}
    for pref, est in reservas.items():
        for ramo in bd.cols_ramo:
            nivel = np.array(resultados[(libro, pref, est["nivel"], ramo)].pronostico, dtype=float)
            nivel = np.nan_to_num(nivel, nan=0.0) / tc_proy
            comps = {}
            for conc in est.get("razones_bel", []):
                r = np.nan_to_num(np.array(resultados[(libro, pref, f"{conc}/{est['nivel']}", ramo)].pronostico),
                                  nan=0.0)
                comps[conc] = r * nivel
            if est["nivel"] == "BRUTO":
                bruto = nivel
            else:
                bruto = nivel + sum(comps.values()) if comps else nivel.copy()
            c = np.nan_to_num(np.array(resultados[(libro, pref, f"{est['cesion']}/BRUTO", ramo)].pronostico),
                              nan=0.0)
            irr = c * bruto
            neto = bruto - irr
            valores = {f"{pref} BRUTO": bruto, f"{pref} {est['cesion']}": irr, f"{pref} NETO": neto}
            if est["nivel"] != "BRUTO":
                valores[f"{pref} {est['nivel']}"] = nivel
            for conc, v in comps.items():
                valores[f"{pref} {conc}"] = v
            for conc in est.get("totales_con_mezcla", []):
                total = np.nan_to_num(np.array(resultados[(libro, pref, f"{conc} TOTAL", "TOTAL")].pronostico),
                                      nan=0.0) / tc_proy
                mezcla = mezcla_reciente(bd, conc, periodo_ultimo_real(periodos_proy), MESES_MEZCLA)
                valores[conc] = total * mezcla[ramo]
            for conc, arr in valores.items():
                for i, p in enumerate(periodos_proy):
                    salida[(norm(conc), p, ramo)] = float(arr[i])
    return salida


# =============================================================================
# ESCRITURA DE LAS BD DE MONTOS (mismo formato)
# =============================================================================
def _copiar_estilo(origen, destino):
    if origen.has_style:
        destino.font = copy(origen.font)
        destino.border = copy(origen.border)
        destino.fill = copy(origen.fill)
        destino.number_format = origen.number_format
        destino.protection = copy(origen.protection)
        destino.alignment = copy(origen.alignment)


def escribir_bd_montos(bd: BDMontos, proy: dict, periodos_proy: list[int], ruta_salida: Path) -> dict:
    ws = bd.ws
    relleno = PatternFill("solid", fgColor=COLOR_RESALTADO) if RESALTAR_PROYECCION else None
    ultima_col = ws.max_column

    faltan = {c for (c, _, _) in proy} - set(bd.conceptos)
    if faltan:
        raise ValueError(f"Conceptos proyectados que no existen en la BD: {faltan}")
    conceptos_bd = [c for c in bd.conceptos if any((c, p, r) in proy for p in periodos_proy for r in bd.cols_ramo)]
    nuevos, actualizados = 0, 0
    # orden de los conceptos en los bloques nuevos: el del bloque del anio mas reciente
    anio_ref = max(p for (_, p) in bd.filas if p < PERIODO_INICIO) // 100
    primera_fila = {c: min((r for (cc, p), r in bd.filas.items() if cc == c and p // 100 == anio_ref),
                           default=10 ** 9) for c in conceptos_bd}
    orden = sorted(conceptos_bd, key=lambda c: primera_fila[c])
    fila_nueva = ws.max_row + 1
    # 0) TC de los renglones existentes segun el supuesto de Inversiones
    tc_cambiados = 0
    if ACTUALIZAR_TC_EXISTENTES:
        for (c, p), r in bd.filas.items():
            if p in TC_PROYECCION:
                nuevo = TC_PROYECCION[p]
                if ws.cell(r, bd.col_tc).value != nuevo:
                    ws.cell(r, bd.col_tc).value = nuevo
                    tc_cambiados += 1
    # 1) periodos que ya existen (p.ej. 202609-202612 en ceros)
    for (c, p), r in list(bd.filas.items()):
        if p in periodos_proy and c in conceptos_bd:
            for ramo, col in bd.cols_ramo.items():
                celda = ws.cell(r, col)
                celda.value = proy[(c, p, ramo)]
                if relleno:
                    celda.fill = relleno
            actualizados += 1
    # 2) periodos nuevos: se agregan al final, bloque por concepto (mismo orden que la BD)
    anios_nuevos = sorted({p // 100 for p in periodos_proy if all((c, p) not in bd.filas for c in conceptos_bd)})
    for anio in anios_nuevos:
        for c in orden:
            for p in [q for q in periodos_proy if q // 100 == anio]:
                if (c, p) in bd.filas:
                    continue
                plantilla = _fila_plantilla(bd, c, p)
                r = fila_nueva
                for col in range(1, ultima_col + 1):
                    _copiar_estilo(ws.cell(plantilla, col), ws.cell(r, col))
                    v = ws.cell(plantilla, col).value
                    if isinstance(v, str) and v.startswith("="):
                        ws.cell(r, col).value = Translator(v, origin=ws.cell(plantilla, col).coordinate) \
                            .translate_formula(ws.cell(r, col).coordinate)
                if ws.row_dimensions[plantilla].height:
                    ws.row_dimensions[r].height = ws.row_dimensions[plantilla].height
                ws.cell(r, bd.col_concepto).value = ws.cell(plantilla, bd.col_concepto).value
                ws.cell(r, bd.col_periodo).value = p
                ws.cell(r, bd.col_tc).value = tc_para_periodo(bd, p)
                for ramo, col in bd.cols_ramo.items():
                    ws.cell(r, col).value = proy[(c, p, ramo)]
                    if relleno:
                        ws.cell(r, col).fill = relleno
                bd.filas[(c, p)] = r
                fila_nueva += 1
                nuevos += 1
    if ws.auto_filter and ws.auto_filter.ref:
        ini, fin = ws.auto_filter.ref.split(":")
        col_fin = re.match(r"[A-Z]+", fin).group()
        ws.auto_filter.ref = f"{ini}:{col_fin}{ws.max_row}"
    return {"actualizados": actualizados, "nuevos": nuevos, "tc_cambiados": tc_cambiados}


def _fila_plantilla(bd: BDMontos, concepto: str, periodo: int) -> int:
    """Renglon del mismo concepto y mismo mes del anio mas reciente (para copiar formato)."""
    candidatos = [(p, r) for (c, p), r in bd.filas.items() if c == concepto and p % 100 == periodo % 100]
    if not candidatos:
        candidatos = [(p, r) for (c, p), r in bd.filas.items() if c == concepto]
    return max(candidatos)[1]


# =============================================================================
# HPARAMETROS
# =============================================================================
@dataclass
class HParam:
    ws: object
    cols: dict                 # nombre encabezado -> columna (primera aparicion)
    historia: dict             # (ramo_str, columna) -> {periodo: valor}
    ramos_activos: list        # (ramo_original, fila_plantilla) en el orden del ultimo bloque
    ultimo: int
    textos_limpiados: int
    lags_cero: list


def leer_hparametros(wb) -> HParam:
    ws = wb[HOJA_PARAMETROS]
    cols = {}
    for c in range(1, 60):
        v = ws.cell(3, c).value
        if v is not None and str(v).strip() not in cols:
            cols[str(v).strip()] = c
    necesarias = ["Tipo de Indice", "Llave", "Fecha", "Ramo"] + INDICES + LAGS
    faltan = [n for n in necesarias if n not in cols]
    if faltan:
        raise ValueError(f"{HOJA_PARAMETROS}: faltan columnas {faltan}")
    historia, textos, ultimos = {}, 0, {}
    lags_cero = []
    for r in range(4, ws.max_row + 1):
        if norm(ws.cell(r, cols["Tipo de Indice"]).value) != norm(TIPO_INDICE_BASE):
            continue
        fecha = ws.cell(r, cols["Fecha"]).value
        ramo = ws.cell(r, cols["Ramo"]).value
        if not isinstance(fecha, (int, float)) or ramo is None:
            continue
        fecha = int(fecha)
        clave_ramo = str(ramo).strip()
        ultimos.setdefault(fecha, []).append((ramo, r))
        max_previo = 0.0
        for nombre in INDICES + LAGS:
            v, fue_texto = a_numero(ws.cell(r, cols[nombre]).value)
            if fue_texto and not math.isnan(v):
                textos += 1
            if nombre in LAGS:
                # patron acumulado: un 0 despues de haber superado 50% es marcador de faltante
                if not math.isnan(v) and v == 0 and max_previo > 0.5:
                    lags_cero.append((fecha, clave_ramo, nombre))
                    v = math.nan
                if not math.isnan(v):
                    max_previo = max(max_previo, v)
            historia.setdefault((clave_ramo, nombre), {})[fecha] = v
    ultimo = max(ultimos)
    return HParam(ws, cols, historia, ultimos[ultimo], ultimo, textos, lags_cero)


def series_hparametros(hp: HParam, h: int, ultimo: int) -> list[Serie]:
    """ultimo = ultimo mes real global (PERIODO_INICIO - 1): si el ultimo 'Real' de la hoja es anterior, el
    mecanismo de rezago proyecta desde el ultimo dato y entrega exactamente los meses a proyectar."""
    series = []
    for ramo, _ in hp.ramos_activos:
        clave_ramo = str(ramo).strip()
        for nombre in INDICES + LAGS:
            hist = hp.historia.get((clave_ramo, nombre), {})
            fechas = [f for f, v in hist.items() if not math.isnan(v)]
            if not fechas:
                periodos, valores = [hp.ultimo], [math.nan]
            else:
                periodos = rango_periodos(min(fechas), hp.ultimo)
                valores = [hist.get(p, math.nan) for p in periodos]
            tipo = "indice" if nombre in INDICES else "lag"
            series.append(Serie(("HPARAM", HOJA_PARAMETROS, nombre, clave_ramo), tipo, periodos, valores,
                                ultimo, h))
    return series


def ajustar_orden_indices(hp: HParam, resultados: dict, alertas: list):
    """Si en la historia del ramo el 99.5% siempre ha sido >= la media, se respeta en la proyeccion."""
    for ramo, _ in hp.ramos_activos:
        clave_ramo = str(ramo).strip()
        for media, alto in PARES_MEDIA_995:
            hm = hp.historia.get((clave_ramo, media), {})
            ha = hp.historia.get((clave_ramo, alto), {})
            comunes = [f for f in hm if f in ha and not math.isnan(hm[f]) and not math.isnan(ha[f])]
            if not comunes or any(ha[f] < hm[f] for f in comunes):
                continue
            rm = resultados[("HPARAM", HOJA_PARAMETROS, media, clave_ramo)]
            ra = resultados[("HPARAM", HOJA_PARAMETROS, alto, clave_ramo)]
            cambios = 0
            for i, (vm, va) in enumerate(zip(rm.pronostico, ra.pronostico)):
                if not math.isnan(vm) and not math.isnan(va) and va < vm:
                    ra.pronostico[i] = vm
                    cambios += 1
            if cambios:
                alertas.append(("HPARAM", f"{alto} ramo {clave_ramo}",
                                f"{cambios} mes(es) ajustados para que {alto} >= {media}"))


def escribir_hparametros(hp: HParam, resultados: dict, periodos_proy: list[int]) -> int:
    ws = hp.ws
    relleno = PatternFill("solid", fgColor=COLOR_RESALTADO) if RESALTAR_PROYECCION else None
    ult_col = max(hp.cols.values())
    if ws.auto_filter and ws.auto_filter.ref:
        ult_col = max(ult_col, openpyxl.utils.column_index_from_string(
            re.match(r"[A-Z]+", ws.auto_filter.ref.split(":")[1]).group()))
    fila = ws.max_row + 1
    nuevas = 0
    for i, p in enumerate(periodos_proy):
        for ramo, plantilla in hp.ramos_activos:
            clave_ramo = str(ramo).strip()
            for col in range(1, ult_col + 1):
                _copiar_estilo(ws.cell(plantilla, col), ws.cell(fila, col))
            ws.cell(fila, hp.cols["Tipo de Indice"]).value = ETIQUETA_PROYECCION
            ws.cell(fila, hp.cols["Llave"]).value = f"{p}-{ramo}"
            ws.cell(fila, hp.cols["Fecha"]).value = p
            ws.cell(fila, hp.cols["Ramo"]).value = ramo
            for nombre in INDICES + LAGS:
                v = resultados[("HPARAM", HOJA_PARAMETROS, nombre, clave_ramo)].pronostico[i]
                celda = ws.cell(fila, hp.cols[nombre])
                celda.value = None if (v is None or math.isnan(v)) else float(v)
                if relleno and celda.value is not None:
                    celda.fill = relleno
            fila += 1
            nuevas += 1
    if ws.auto_filter and ws.auto_filter.ref:
        ini, fin = ws.auto_filter.ref.split(":")
        col_fin = re.match(r"[A-Z]+", fin).group()
        ws.auto_filter.ref = f"{ini}:{col_fin}{fila - 1}"
        for fc in ws.auto_filter.filterColumn:
            if fc.filters is not None and fc.filters.filter and TIPO_INDICE_BASE in fc.filters.filter \
                    and ETIQUETA_PROYECCION not in fc.filters.filter:
                fc.filters.filter.append(ETIQUETA_PROYECCION)
    return nuevas


# =============================================================================
# DIAGNOSTICO Y GRAFICAS
# =============================================================================
def escribir_diagnostico(resultados: dict, periodos_proy: list[int], alertas_generales: list, derivados: dict,
                         resumen: dict, tabla_validacion: list | None = None):
    wb = openpyxl.Workbook()
    negrita = Font(bold=True)
    encab = PatternFill("solid", fgColor="D9E1F2")

    ws = wb.active
    ws.title = "Resumen"
    lineas = [
        ("Proyeccion de indices y montos de reservas", ""),
        ("Fecha de ejecucion", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Periodos proyectados", f"{periodos_proy[0]} - {periodos_proy[-1]} ({len(periodos_proy)} meses)"),
        ("Ultimo mes real usado", str(resumen.get("ultimo"))),
        ("Series modeladas", str(len(resultados))),
        ("Tiempo de ejecucion (s)", f"{resumen.get('segundos', 0):,.0f}"),
        ("", ""),
        ("METODOLOGIA", ""),
        ("1. Coherencia contable", "Se proyectan drivers (BEL / BRUTO y razones) y se derivan las demas lineas: "
                                   "RRC: BRUTO = BEL+GTO+MR, IRR = %ces*BRUTO, NETO = BRUTO-IRR; SONR: BRUTO = BEL+MR; "
                                   "RFV: NETO = BRUTO-IRR; RCONT: total con factores por mes del trimestre, repartido "
                                   "por ramo con la mezcla de los ultimos 8 meses."),
        ("2. Moneda", "Montos modelados en USD: en backtest, modelar en MXN y convertir con el TC real fue menos "
                      "preciso (Danos WAPE 22.5% vs 19.8%; Fianzas 7.4% vs 6.0%). Libros en MXN (MODELAR_EN_MXN): "
                      + (", ".join(k for k, v in MODELAR_EN_MXN.items() if v) or "ninguno") + "."),
        ("3. Modelos evaluados", "Ingenuo, Ingenuo estacional, Media 12m, SES, Holt amortiguado, Holt-Winters "
                                 "amortiguado, Theta, ARIMA (AICc), Regresion log-lineal con estacionalidad."),
        ("4. Seleccion", "Validacion fuera de muestra del metodo completo (hoja Validacion_Metodo): se simula haber "
                         f"proyectado desde {CV_ORIGENES} cortes historicos a 1-16 meses. Montos: menor WAPE (error "
                         "ponderado por USD) con desempate por sesgo agregado. Indices, razones y LAGs: menor "
                         "AvgRelMAE; si la mejora vs el ingenuo no es significativa (IC90), procedimiento sin tendencia. "
                         "En cada serie hay ademas un backtest con todos los modelos (MASE en Series_Modelos) que da "
                         "los intervalos y la red de seguridad (si el procedimiento es > 1.5x peor que el ingenuo en la "
                         "propia serie, se usa el ensamble propio)."
                         + ("" if MODO_SELECCION == "validado" else " MODO TORNEO: cada serie usa su propio ensamble.")),
        ("5. Transformaciones", "Montos e indices en logaritmos (la proyeccion es la mediana); razones y LAGs en "
                                "escala original, acotadas al rango historico y al dominio actuarial (cesion en [0,1], "
                                "%GTO y %MR >= 0)."),
        ("6. Intervalos", f"{NIVEL_INTERVALO:.0%} a partir del error de la proyeccion en el backtest por horizonte "
                          "(con piso de caminata aleatoria si hay pocos cortes)."),
        ("7. Reglas", "Series en cero -> 0; parametros en escalon -> ultimo valor; series cortas -> SES; "
                      "99.5% >= media si siempre lo fue; limpieza de textos y LAG=0 marcadores; alerta de saltos "
                      "atipicos en el ultimo mes."),
        ("8. Tipo de cambio", "Columna TC = supuesto de Inversiones (TC_Real_Esti.xlsx, hoja TC: FCST 2026 y "
                              "FCST 2027), en tipo_cambio.py; 202601-202608 = TC real de SAP. Los montos se "
                              "modelan en USD, por lo que el TC no altera las cifras proyectadas."),
        ("9. Sesgo conocido", "En la validacion (2023-2026, periodo de fuerte crecimiento) las proyecciones de montos a "
                              "13-16 meses quedaron en promedio por debajo de lo real (ver 'Sesgo agregado' del "
                              "procedimiento elegido). Los modelos amortiguan la tendencia: si el plan de negocio "
                              "prevé un crecimiento sostenido, conviene contrastarlo."),
        ("Versiones", resumen.get("versiones", "")),
        ("", ""),
        ("Interpretacion MASE", "Error absoluto medio del backtest promediado en horizontes de 1 a 16 meses, "
                                "dividido entre la variacion mensual tipica de la serie. Al ser multi-horizonte, "
                                "valores > 1 son normales; lo relevante es compararlo contra 'MASE Ingenuo' "
                                "(caminata aleatoria): menor = mejor."),
    ]
    for i, (a, b) in enumerate(lineas, start=1):
        ws.cell(i, 1, a).font = negrita if a and b == "" or a in ("METODOLOGIA",) else Font()
        ws.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 120

    if resumen.get("seleccion"):
        fila = ws.max_row + 2
        ws.cell(fila, 1, "PROCEDIMIENTO SELECCIONADO POR TIPO").font = negrita
        for i, (tipo, nombre) in enumerate(sorted(resumen["seleccion"].items()), start=1):
            ws.cell(fila + i, 1, tipo)
            ws.cell(fila + i, 2, f"{nombre}  (combinacion de pesos iguales; ver hoja Validacion_Metodo)")

    # Validacion fuera de muestra del metodo
    if tabla_validacion:
        ws = wb.create_sheet("Validacion_Metodo")
        cols = ["Tipo", "Procedimiento", "Seleccionado", "WAPE %", "AvgRelMAE", "IC90 inf", "IC90 sup",
                "% series mejor que ingenuo", "Series", "MAPE 1-6", "MAPE 7-12", "MAPE 13-16",
                "Sesgo agregado % 13-16", "Tendencia"]
        ws.append(cols)
        for d in tabla_validacion:
            ws.append([d.get(c) for c in cols])
        _formato_tabla(ws, negrita, encab)
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                encabezado = ws.cell(1, c.column).value
                if encabezado in ("AvgRelMAE", "IC90 inf", "IC90 sup"):
                    c.number_format = "0.000"
                elif encabezado == "% series mejor que ingenuo":
                    c.number_format = "0%"
                elif encabezado in ("WAPE %", "MAPE 1-6", "MAPE 7-12", "MAPE 13-16", "Sesgo agregado % 13-16"):
                    c.number_format = "0.0"
        nota = ws.max_row + 2
        notas = [
            "Pronosticos fuera de muestra desde 8 cortes historicos, horizontes 1-16 meses (backtest del metodo "
            "completo). Solo series con datos genuinos (<= 3 meses interpolados).",
            "WAPE % = suma de errores absolutos / suma de valores reales (pondera por monto). Criterio para montos: "
            "menor WAPE; entre los que estan a <= 2% del mejor, el de menor |sesgo agregado|.",
            "AvgRelMAE = media geometrica entre series del error del procedimiento / error del ingenuo (< 1 = mejor "
            "que repetir el ultimo valor), con IC bootstrap al 90%. Criterio para indices, razones y LAGs: menor "
            "AvgRelMAE; si su IC incluye 1, el mejor procedimiento sin tendencia (parsimonia).",
            "Sesgo agregado % 13-16 = (suma proyectada - suma real) / suma real por corte y horizonte, promedio a "
            "13-16 meses. Negativo = la proyeccion quedo por debajo de lo real (en 2023-2026 hubo fuerte "
            "crecimiento).",
        ]
        for i, texto in enumerate(notas):
            ws.cell(nota + i, 1, texto)

    # Series y modelos
    ws = wb.create_sheet("Series_Modelos")
    todos_modelos = list(CATALOGO)
    cab = (["Libro", "Grupo", "Serie", "Ramo", "Tipo", "Moneda modelo", "Transformacion", "Obs", "Cortes backtest",
            "Desde", "Regla",
            "Modelos proyeccion (peso)", "MASE proyeccion"] + [f"MASE {m}" for m in todos_modelos]
           + ["Ultimo real", f"Proy {periodos_proy[0]}", f"Proy {periodos_proy[-1]}", "Var % vs ultimo real",
              "Alertas"])
    ws.append(cab)
    for k, r in sorted(resultados.items(), key=lambda kv: tuple(str(x) for x in kv[0])):
        ult = r.historia_valores[-1] if r.historia_valores else math.nan
        fin = r.pronostico[-1] if r.pronostico else math.nan
        var = (fin / ult - 1) if (ult and not math.isnan(ult) and not math.isnan(fin) and ult != 0) else None
        fila = [k[0], k[1], k[2], k[3], r.tipo, r.moneda, r.transformacion, r.n_obs, r.n_cortes or None,
                r.historia_periodos[0] if r.historia_periodos else None, r.regla,
                ", ".join(f"{m} ({w:.0%})" for m, w in r.pesos.items()),
                None if math.isnan(r.mase_ensamble) else r.mase_ensamble]
        fila += [r.modelos.get(m) for m in todos_modelos]
        fila += [None if math.isnan(ult) else ult,
                 None if not r.pronostico or math.isnan(r.pronostico[0]) else r.pronostico[0],
                 None if math.isnan(fin) else fin, var, " | ".join(r.alertas)]
        ws.append(fila)
    _formato_tabla(ws, negrita, encab)

    # Pronosticos (drivers con intervalos)
    ws = wb.create_sheet("Pronosticos_Drivers")
    ws.append(["Libro", "Grupo", "Serie", "Ramo", "Moneda", "Periodo", "Pronostico", f"LI {NIVEL_INTERVALO:.0%}",
               f"LS {NIVEL_INTERVALO:.0%}"])
    for k, r in sorted(resultados.items(), key=lambda kv: tuple(str(x) for x in kv[0])):
        for i, p in enumerate(periodos_proy):
            v = r.pronostico[i]
            if v is None or math.isnan(v):
                continue
            ws.append([k[0], k[1], k[2], k[3], r.moneda or "-", p, v, r.li[i], r.ls[i]])
    _formato_tabla(ws, negrita, encab)

    # Montos derivados (lo que se escribio en las BD)
    ws = wb.create_sheet("Montos_Proyectados")
    ws.append(["Libro", "Concepto", "Periodo", "Ramo", "Valor USD"])
    for libro, dic in derivados.items():
        for (c, p, ramo), v in sorted(dic.items()):
            ws.append([libro, c, p, ramo, v])
    _formato_tabla(ws, negrita, encab)

    ws = wb.create_sheet("Alertas")
    ws.append(["Libro", "Serie / Elemento", "Detalle"])
    for a in alertas_generales:
        ws.append(list(a))
    for k, r in sorted(resultados.items(), key=lambda kv: tuple(str(x) for x in kv[0])):
        for a in r.alertas:
            ws.append([k[0], f"{k[1]} | {k[2]} | ramo {k[3]}", a])
    _formato_tabla(ws, negrita, encab)
    wb.save(SALIDA_DIAGNOSTICO)


def _formato_tabla(ws, negrita, encab):
    for c in ws[1]:
        c.font = negrita
        c.fill = encab
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        letra = get_column_letter(col)
        largo = max((len(str(c.value)) for c in ws[letra][:200] if c.value is not None), default=8)
        ws.column_dimensions[letra].width = min(max(10, largo + 2), 60)
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            if isinstance(c.value, float):
                c.number_format = "0.0%" if ws.cell(1, c.column).value == "Var % vs ultimo real" else (
                    "#,##0.0000" if abs(c.value) < 100 else "#,##0")


def graficar(resultados: dict, derivados: dict, historia_montos: dict, periodos_proy: list[int]):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    def fechas(ps):
        return [datetime(p // 100, p % 100, 1) for p in ps]

    with PdfPages(SALIDA_GRAFICAS) as pdf:
        # 1) Montos: un concepto por pagina, un panel por ramo
        for libro, dic in derivados.items():
            conceptos = sorted({c for (c, _, _) in dic})
            ramos = sorted({r for (_, _, r) in dic}, key=lambda x: (len(x), x))
            for c in conceptos:
                fig, ejes = plt.subplots(int(math.ceil(len(ramos) / 4)), 4, figsize=(16, 10), squeeze=False)
                for ax, ramo in zip(ejes.flat, ramos):
                    hist = historia_montos[libro]
                    ph = sorted(p for (cc, p, rr) in hist if cc == c and rr == ramo)
                    ax.plot(fechas(ph), [hist[(c, p, ramo)] / 1e6 for p in ph], color="#1f4e79", lw=1.4)
                    ax.plot(fechas(periodos_proy), [dic[(c, p, ramo)] / 1e6 for p in periodos_proy],
                            color="#c55a11", lw=1.6, ls="--")
                    ax.set_title(f"RAM_{ramo}", fontsize=9)
                    ax.tick_params(labelsize=7)
                for ax in list(ejes.flat)[len(ramos):]:
                    ax.axis("off")
                fig.suptitle(f"{libro} - {c} (millones USD)   azul: real   naranja: proyeccion", fontsize=12)
                fig.tight_layout()
                pdf.savefig(fig)
                plt.close(fig)
        # 2) Drivers con intervalo (niveles e indices)
        grupos = {}
        for k, r in resultados.items():
            if r.tipo in ("nivel", "indice", "razon"):
                grupos.setdefault((k[0], k[1], k[2]), []).append(r)
        for (libro, grupo, serie), lista in sorted(grupos.items()):
            lista = sorted(lista, key=lambda r: (len(r.clave[3]), r.clave[3]))
            fig, ejes = plt.subplots(int(math.ceil(len(lista) / 4)), 4, figsize=(16, 10), squeeze=False)
            for ax, r in zip(ejes.flat, lista):
                if r.historia_periodos:
                    ax.plot(fechas(r.historia_periodos), r.historia_valores, color="#1f4e79", lw=1.3)
                pr = np.array(r.pronostico, dtype=float)
                if np.any(~np.isnan(pr)):
                    ax.plot(fechas(periodos_proy), pr, color="#c55a11", lw=1.5, ls="--")
                    ax.fill_between(fechas(periodos_proy), r.li, r.ls, color="#f4b183", alpha=0.35, lw=0)
                ax.set_title(f"{r.clave[3]}: {', '.join(r.pesos) or r.regla}"[:60], fontsize=7)
                ax.tick_params(labelsize=6)
            for ax in list(ejes.flat)[len(lista):]:
                ax.axis("off")
            fig.suptitle(f"{libro} | {grupo} | {serie}  (banda {NIVEL_INTERVALO:.0%})", fontsize=12)
            fig.tight_layout()
            pdf.savefig(fig)
            plt.close(fig)
        # 3) LAGs por ramo
        lags_ramo = {}
        for k, r in resultados.items():
            if r.tipo == "lag":
                lags_ramo.setdefault(k[3], []).append(r)
        for ramo, lista in sorted(lags_ramo.items(), key=lambda kv: (len(kv[0]), kv[0])):
            fig, ax = plt.subplots(figsize=(14, 7))
            for r in sorted(lista, key=lambda r: int(r.clave[2].split()[-1])):
                if r.historia_periodos:
                    linea, = ax.plot(fechas(r.historia_periodos), r.historia_valores, lw=1.1, label=r.clave[2])
                    ax.plot(fechas(periodos_proy), r.pronostico, lw=1.3, ls="--", color=linea.get_color())
            ax.set_title(f"HParametros - LAGs ramo {ramo} (continua: real, punteada: proyeccion)")
            ax.legend(fontsize=7, ncol=5)
            fig.tight_layout()
            pdf.savefig(fig)
            plt.close(fig)


# =============================================================================
# PROCESO PRINCIPAL
# =============================================================================
def _versiones() -> str:
    from importlib.metadata import PackageNotFoundError, version
    partes = [f"Python {sys.version.split()[0]}"]
    for paquete in ("openpyxl", "numpy", "pandas", "scipy", "statsmodels", "matplotlib"):
        try:
            partes.append(f"{paquete} {version(paquete)}")
        except PackageNotFoundError:
            continue
    return ", ".join(partes)


def _bd_rfv_desactualizada() -> bool:
    """La BD_ RFV llena se regenera siempre (REGENERAR_BD_RFV) o si no existe o alguna entrada es mas reciente."""
    if REGENERAR_BD_RFV or not ARCHIVO_BD_RFV.exists():
        return True
    entradas = list(ENTRADAS.glob("Res_Rvas_*.xlsx")) + [ENTRADAS / "BD_ RFV.xlsx"]
    t_salida = ARCHIVO_BD_RFV.stat().st_mtime
    return any(p.exists() and p.stat().st_mtime > t_salida for p in entradas)


def revisar_periodos(bd: BDMontos, nombre: str, ultimo: int, periodos_proy: list[int], alertas: list):
    """Evita proyectar sobre cifras reales y usar un mes historico vacio como si fuera cero real."""
    con_datos = sorted({p for (c, p, r), v in bd.valores.items() if p in periodos_proy and abs(v) > 0})
    if con_datos and not SOBRESCRIBIR_PERIODOS_CON_DATOS:
        raise SystemExit(f"{nombre}: los meses {con_datos} ya tienen cifras (reales?) y se sobrescribirian con la "
                         f"proyeccion. Mueve PERIODO_INICIO al mes siguiente al ultimo real o pon "
                         f"SOBRESCRIBIR_PERIODOS_CON_DATOS = True.")
    previo = indice_a_periodo(periodo_a_indice(ultimo) - 1)
    incompletos = []
    for concepto in bd.conceptos:
        tenia = any(abs(bd.valores.get((concepto, previo, r), 0.0)) > 0 for r in bd.cols_ramo)
        tiene = any(abs(bd.valores.get((concepto, ultimo, r), 0.0)) > 0 for r in bd.cols_ramo)
        if tenia and not tiene:
            incompletos.append(concepto)
    if incompletos:
        raise SystemExit(f"{nombre}: el mes {ultimo} (ultimo real, PERIODO_INICIO - 1) no tiene cifras en "
                         f"{incompletos}, aunque {previo} si. Carga ese mes completo o ajusta PERIODO_INICIO.")
    if not any(abs(v) > 0 for (c, p, r), v in bd.valores.items() if p <= ultimo):
        raise SystemExit(f"{nombre}: la historia esta vacia (todo en cero).")


def main():
    t0 = time.time()
    SALIDAS.mkdir(parents=True, exist_ok=True)
    periodos_proy = rango_periodos(PERIODO_INICIO, PERIODO_FIN)
    ultimo = indice_a_periodo(periodo_a_indice(PERIODO_INICIO) - 1)
    h = len(periodos_proy)
    print(f"Proyeccion {periodos_proy[0]} - {periodos_proy[-1]} ({h} meses). Historia hasta {ultimo}.", flush=True)
    salidas = [SALIDA_BD_DANOS, SALIDA_BD_RFV, SALIDA_DIAGNOSTICO] + ([SALIDA_GRAFICAS] if GENERAR_GRAFICAS else [])
    verificar_escritura(salidas + [ARCHIVO_BD_RFV])

    if _bd_rfv_desactualizada():
        print("Llenando BD_ RFV desde los Res_Rvas (llenar_bd_rfv.py) ...", flush=True)
        import llenar_bd_rfv  # noqa: WPS433
        llenar_bd_rfv.llenar()
    print(f"   BD_ RFV usada: {ARCHIVO_BD_RFV.name} "
          f"({datetime.fromtimestamp(ARCHIVO_BD_RFV.stat().st_mtime):%Y-%m-%d %H:%M})", flush=True)

    alertas = []
    print("Leyendo archivos ...", flush=True)
    bd_danos = leer_bd_montos(ARCHIVO_BD_DANOS)
    bd_rfv = leer_bd_montos(ARCHIVO_BD_RFV)
    revisar_periodos(bd_danos, "BD Daños", ultimo, periodos_proy, alertas)
    revisar_periodos(bd_rfv, "BD RFV", ultimo, periodos_proy, alertas)
    hp = leer_hparametros(bd_danos.wb)
    if hp.ultimo > ultimo:
        raise SystemExit(f"{HOJA_PARAMETROS}: ya hay renglones '{TIPO_INDICE_BASE}' de {hp.ultimo}, posteriores a "
                         f"{ultimo}. Mueve PERIODO_INICIO.")
    if hp.ultimo < ultimo:
        alertas.append(("HPARAM", "Ultimo mes Real", f"El ultimo '{TIPO_INDICE_BASE}' es {hp.ultimo}; los indices "
                                                     f"se proyectan desde ahi hasta cubrir {periodos_proy[0]}"))
    if hp.textos_limpiados:
        alertas.append(("HPARAM", "Limpieza", f"{hp.textos_limpiados} celdas con numeros guardados como texto "
                                              "(p.ej. espacios \\xa0) se convirtieron a numero"))
    for fecha, ramo, lag in hp.lags_cero:
        alertas.append(("HPARAM", f"{lag} ramo {ramo}", f"{fecha}: valor 0 tratado como faltante "
                                                        "(el patron acumulado ya superaba 50%)"))

    # Revision de identidades en la historia (deben cumplirse para que la derivacion sea valida)
    for libro, bd in (("DANOS", bd_danos), ("FIANZAS", bd_rfv)):
        for (c, p, ramo), v in bd.valores.items():
            if p > ultimo or not c.endswith("NETO"):
                continue
            pref = c.split()[0]
            bruto = bd.valores.get((f"{pref} BRUTO", p, ramo), 0.0)
            irr = bd.valores.get((f"{pref} IRR", p, ramo), 0.0)
            if abs(v - (bruto - irr)) > 1.0:
                alertas.append((libro, f"{c} {p} RAM_{ramo}", "La historia no cumple NETO = BRUTO - IRR"))

    tc_hist = {}
    for libro, bd in (("DANOS", bd_danos), ("FIANZAS", bd_rfv)):
        if MODELAR_EN_MXN.get(libro):
            tc_hist[libro] = leer_tc_real(bd, ultimo)
            alertas.append(("METODO", f"Moneda {libro}", "Montos modelados en MXN (historia USD x TC real de SAP) y "
                                                         "convertidos a USD con el TC de cada mes proyectado de la BD"))

    print("Construyendo series ...", flush=True)
    series = (series_montos(bd_danos, "DANOS", ESTRUCTURA["DANOS"], ultimo, h, tc_hist.get("DANOS"))
              + series_montos(bd_rfv, "FIANZAS", ESTRUCTURA["FIANZAS"], ultimo, h, tc_hist.get("FIANZAS"))
              + series_hparametros(hp, h, ultimo))
    tabla_validacion, seleccion = [], {}
    if MODO_SELECCION == "validado":
        print("Validando procedimientos fuera de muestra (backtest del metodo completo) ...", flush=True)
        tabla_validacion, seleccion = validar_procedimientos(series)
        for tipo, nombre in PROCEDIMIENTO_RESPALDO.items():
            if tipo not in seleccion:
                seleccion[tipo] = nombre
                alertas.append(("METODO", tipo, f"Sin series suficientes para validar; se usa {nombre}"))
        for s in series:
            s.procedimiento = tuple(PROCEDIMIENTOS_CANDIDATOS[seleccion[s.tipo]])
        for tipo, nombre in seleccion.items():
            print(f"   {tipo:<7} -> {nombre}", flush=True)
    print(f"   {len(series)} series a proyectar (con backtest por serie) ...", flush=True)
    resultados = correr_series(series)
    ajustar_orden_indices(hp, resultados, alertas)

    print("Aplicando identidades contables y escribiendo archivos ...", flush=True)
    proy_danos = derivar_montos(bd_danos, "DANOS", ESTRUCTURA["DANOS"], resultados, periodos_proy,
                                en_mxn="DANOS" in tc_hist)
    proy_rfv = derivar_montos(bd_rfv, "FIANZAS", ESTRUCTURA["FIANZAS"], resultados, periodos_proy,
                              en_mxn="FIANZAS" in tc_hist)
    validar(proy_danos, proy_rfv)

    info_danos = escribir_bd_montos(bd_danos, proy_danos, periodos_proy, SALIDA_BD_DANOS)
    n_hp = escribir_hparametros(hp, resultados, periodos_proy)
    guardar_libro(bd_danos.wb, SALIDA_BD_DANOS, original=ARCHIVO_BD_DANOS)
    info_rfv = escribir_bd_montos(bd_rfv, proy_rfv, periodos_proy, SALIDA_BD_RFV)
    guardar_libro(bd_rfv.wb, SALIDA_BD_RFV, original=ENTRADAS / "BD_ RFV.xlsx")

    historia = {
        "DANOS": {k: v for k, v in bd_danos.valores.items() if k[1] <= ultimo},
        "FIANZAS": {k: v for k, v in bd_rfv.valores.items() if k[1] <= ultimo},
    }
    segundos = time.time() - t0
    escribir_diagnostico(resultados, periodos_proy, alertas, {"DANOS": proy_danos, "FIANZAS": proy_rfv},
                         {"ultimo": ultimo, "segundos": segundos, "seleccion": seleccion,
                          "versiones": _versiones()},
                         tabla_validacion)
    graficas_ok = False
    if GENERAR_GRAFICAS:
        print("Generando graficas ...", flush=True)
        try:
            graficar(resultados, {"DANOS": proy_danos, "FIANZAS": proy_rfv}, historia, periodos_proy)
            graficas_ok = True
        except Exception as e:  # noqa: BLE001
            print(f"   No se pudieron generar las graficas: {e!r}")

    print("\nRESUMEN")
    print(f"   {SALIDA_BD_DANOS.name}: {info_danos['actualizados']} renglones actualizados, "
          f"{info_danos['nuevos']} renglones nuevos en {HOJA_MONTOS}; {n_hp} renglones nuevos en {HOJA_PARAMETROS}; "
          f"TC actualizado en {info_danos['tc_cambiados']} renglones")
    print(f"   {SALIDA_BD_RFV.name}: {info_rfv['actualizados']} renglones actualizados, "
          f"{info_rfv['nuevos']} renglones nuevos; TC actualizado en {info_rfv['tc_cambiados']} renglones")
    print(f"   Diagnostico: {SALIDA_DIAGNOSTICO.name}")
    if GENERAR_GRAFICAS:
        print(f"   Graficas: {SALIDA_GRAFICAS.name}" if graficas_ok
              else "   Graficas: NO se actualizaron (ver mensaje arriba)")
    n_alertas = len(alertas) + sum(len(r.alertas) for r in resultados.values())
    print(f"   Alertas a revisar: {n_alertas} (hoja 'Alertas' del diagnostico)")
    print(f"   Tiempo total: {time.time() - t0:,.0f} s")


def validar(proy_danos: dict, proy_rfv: dict):
    """Chequeos de consistencia de lo que se va a escribir."""
    for nombre, dic, reglas in (("DANOS", proy_danos, [("RRC", ["BEL", "GTO", "MR"]), ("SONR", ["BEL", "MR"])]),
                                ("FIANZAS", proy_rfv, [("RFV", None)])):
        for (c, p, ramo), v in dic.items():
            if not np.isfinite(v):
                raise AssertionError(f"{nombre}: valor no finito en {c} {p} {ramo}")
        for pref, comps in reglas:
            claves = {(p, r) for (c, p, r) in dic if c.startswith(pref + " ")}
            for p, r in claves:
                bruto = dic[(f"{pref} BRUTO", p, r)]
                neto = dic[(f"{pref} NETO", p, r)]
                irr = dic[(f"{pref} IRR", p, r)]
                assert abs(neto - (bruto - irr)) < 1e-6 * max(1.0, abs(bruto)), (nombre, pref, p, r)
                if comps:
                    suma = sum(dic[(f"{pref} {x}", p, r)] for x in comps)
                    assert abs(bruto - suma) < 1e-6 * max(1.0, abs(bruto)), (nombre, pref, p, r)
                assert bruto >= -1e-6, (nombre, pref, p, r, bruto)


if __name__ == "__main__":
    main()
