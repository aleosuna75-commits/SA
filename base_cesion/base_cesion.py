# -*- coding: utf-8 -*-
"""
================================================================================
  BASE DE % DE CESIÓN — Logouts del Presupuesto Técnico (hoja ParamPPTO)
================================================================================
Lee los logouts de una LN (o de todas) y arma una base con el % de cesión de
cada documento:

    LN · Nivel · TR · Cedente · Corredor · Contrato · MGA / Binder · % de cesión

USO: cambia la línea  LN = "..."  de abajo y dale "Run Python File" (▷) en VSCode.
     El script busca la carpeta con ese nombre (LN4001, LN4002, ...) en
     CARPETA_LNS, junto al script, o en Documentos / Escritorio / Descargas.
     Con LN = "TODAS" procesa todas las carpetas LN que encuentre y genera
     una sola base.

CARDINALIDAD: cada documento se clasifica según lo que traiga el logout
  * Binder   -> trae MGA (casilla MGA = Verdadero, o Binder general / segmentado)
  * Contrato -> sin MGA, pero con contrato
  * Cedente  -> sin contrato ni MGA
La hoja "Cardinalidad_LN" resume cuántos documentos de cada nivel tiene cada LN.

CÓMO SE LEE CADA LOGOUT (hoja ParamPPTO)
----------------------------------------
Los archivos se reconocen por su contenido (hoja ParamPPTO), no por su nombre,
y las filas se ubican por su ETIQUETA en la columna A.

  Línea de Negocio ........ B          -> LN
  Tipo Reas. .............. B          -> TR
  Compañía ................ B          -> Cedente
  Corredor (1a aparición) . B          -> Corredor del documento
  Contrato ................ B
  Tipo Venta .............. B  (y C = % Renov. cuando es "Comb. Nuevo y Renov.")
  MGA ..................... B = ¿es MGA?, C = Binder general (PRESUPUESTO!E18),
                             E = Binder segmentado, F = casilla GS
  Porcentaje .............. B..E       -> Tradicional, Retro Espec., Fronting, Retención
  Porcentaje de Cesión .... 5 años     -> 2027..2031
  % Comisiones del Cedido . 5 años     -> 2027..2031

AÑOS: el recuadro del logout es B..F (2027..2031), y así vienen, por ejemplo,
los de LN4006. Los de LN4003 (Fianzas) traen el 2027 en C, recorridos una
columna, y no traen 2031. Por eso la columna del primer año se detecta sola
para cada LN (según la LN del propio logout): 2027 = B si la mayoría de sus
logouts traen dato en B; 2027 = C si la mayoría empieza en C. Si hay mezcla,
se avisa. La base dice en "Col. 1er año" qué se usó en cada documento.

Los NOMBRES de LN, TR, cedente y corredor salen del catálogo del PptoTécnico que
Excel guarda en caché dentro de cada logout.

SALIDA: "Base_Cesion_<LN>_<fecha>.xlsx" junto a la carpeta de la LN, con las
hojas Base_Cesion, Cesion_Anual, Resumen, Cardinalidad_LN, Validaciones y Notas.

Si falta openpyxl o pandas (o están muy viejos), el script los instala solo con pip.
================================================================================
"""
from __future__ import annotations

import argparse
import importlib
import os
import re
import subprocess
import sys
import unicodedata
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree


# ==============================================================================
#   >>> CAMBIA SOLO ESTA LÍNEA <<<   LN a procesar: "LN4001", "LN4002", ... o "TODAS"
LN = "LN4006"
# ==============================================================================


# ------------------------------------------------------------------------------
# CONFIGURACIÓN (normalmente no hace falta tocar nada de aquí para abajo)
# ------------------------------------------------------------------------------
# Carpeta que contiene las carpetas LN4001, LN4002, ... Vacío = buscarlas junto al
# script y en Documentos / Escritorio / Descargas (hasta tres niveles abajo).
CARPETA_LNS = r""

# Ruta directa a una carpeta con logouts. Si se llena, se usa esta e ignora LN.
CARPETA_LOGOUTS = r""

# Dónde guardar el Excel de salida. Vacío = la carpeta que contiene a la(s) LN.
CARPETA_SALIDA = r""

HOJA_LOGOUT = "ParamPPTO"
PRIMER_ANIO = 2027
NUM_ANIOS = 5
# Columna donde cae el primer año en "Porcentaje de Cesión" y "% Comisiones":
# "AUTO" = se detecta por LN (lo que haga la mayoría de sus logouts: B o C); o fijar "B" / "C".
COLUMNA_PRIMER_ANIO = "AUTO"

# Si un documento trae la cesión solo en algunos años (p. ej. solo 2027):
# True = los años siguientes toman el último valor capturado; False = se dejan vacíos.
ARRASTRAR_ULTIMO_ANIO = False

# True = abrir el Excel generado al terminar (solo Windows / macOS).
ABRIR_AL_TERMINAR = True

# Catálogo de nombres: el logout guarda en caché la hoja "Valores" del PptoTécnico
# y trae los nombres definidos que apuntan a cada catálogo (p. ej.
# xCEDENTES = [1]Valores!$AR$2:$AV$999). Por llave de la base:
#   (nombre definido, columna de respaldo si no existe el nombre,
#    {columna en la base: desplazamiento desde la columna del código})
HOJA_CATALOGO = "Valores"
CATALOGOS = {
    "LN": ("xAFUN", "BH", {"Nombre LN": 1}),
    "TR": ("xTIPOREA", "V", {"TR desc.": 1}),
    "Cedente": ("xCEDENTES", "AR", {"Nombre Cedente": 1, "País Cedente": 3, "Grupo Cedente": 4}),
    "Corredor": ("xCORREDORES", "AX", {"Nombre Corredor": 1}),
}
TR_DESC_RESPALDO = {1: "Proporcional", 2: "No Proporcional", 3: "Facultativo"}

TOLERANCIA = 1e-6
MINIMOS = {"pandas": (1, 1), "openpyxl": (3, 0)}  # versión mínima de cada paquete


# ------------------------------------------------------------------------------
# PAQUETERÍA
# ------------------------------------------------------------------------------
def revisar_paquete(modulo):
    """'' si el paquete importa y cumple la versión mínima; si no, el problema."""
    try:
        paquete = importlib.import_module(modulo)
    except Exception as error:  # no instalado, o instalado pero roto (p. ej. choque con numpy)
        return f"{modulo}: {type(error).__name__}: {error}"
    texto_version = getattr(paquete, "__version__", None)
    if texto_version is None:  # p. ej. una carpeta "pandas" junto al script, no el paquete real
        return f"{modulo}: no está instalado (se encontró {getattr(paquete, '__path__', '?')})"
    version = tuple(int(x) for x in re.findall(r"\d+", str(texto_version))[:2])
    minimo = MINIMOS[modulo]
    if version < minimo:
        return f"{modulo} {texto_version} es muy viejo (se necesita {minimo[0]}.{minimo[1]} o más)"
    return ""


def asegurar_paquetes():
    """Instala con pip (en el mismo Python que corre el script) lo que falte y vuelve a arrancar."""
    problemas = {m: p for m in MINIMOS for p in [revisar_paquete(m)] if p}
    if not problemas:
        return
    requisitos = [f"{m}>={MINIMOS[m][0]}.{MINIMOS[m][1]}" for m in problemas]
    comando = f'"{sys.executable}" -m pip install --upgrade ' + " ".join(f'"{r}"' for r in requisitos)
    a_mano = ("\nInstálala a mano desde la terminal de VSCode con:\n"
              + (f"    & {comando}\n    (en PowerShell; en cmd, sin el '&' inicial)\n" if os.name == "nt"
                 else f"    {comando}\n")
              + "Si la red de la oficina usa proxy, agrega: --proxy http://usuario:clave@proxy:puerto")
    if os.environ.get("BASE_CESION_REINTENTO"):
        sys.exit("La paquetería sigue fallando después de instalarla:\n  "
                 + "\n  ".join(problemas.values()) + a_mano)

    print("Instalando / actualizando paquetería: " + ", ".join(requisitos) + " ...")
    print("(Si tarda más de un par de minutos, probablemente la red bloquea pypi.org.)")
    pip = [sys.executable, "-m", "pip"]
    silencio = {"stdout": subprocess.DEVNULL, "stderr": subprocess.DEVNULL}
    if subprocess.call(pip + ["--version"], **silencio) != 0:  # este Python no trae pip
        subprocess.call([sys.executable, "-m", "ensurepip", "--upgrade"], **silencio)
    base = pip + ["install", "--upgrade", "--disable-pip-version-check",
                  "--timeout", "20", "--retries", "1"] + requisitos
    intentos = [base]
    if sys.prefix == getattr(sys, "base_prefix", sys.prefix):  # --user no aplica dentro de un venv
        intentos.append(base[:4] + ["--user"] + base[4:])
    if not any(subprocess.call(cmd) == 0 for cmd in intentos):
        sys.exit("\nNo se pudo instalar la paquetería automáticamente." + a_mano)

    # Arrancar de nuevo en un proceso limpio para que tome lo recién instalado.
    print("Paquetería lista. Reiniciando el script...\n")
    entorno = dict(os.environ, BASE_CESION_REINTENTO="1")
    sys.exit(subprocess.call([sys.executable, str(Path(__file__).resolve())] + sys.argv[1:], env=entorno))


asegurar_paquetes()

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: E402


# ------------------------------------------------------------------------------
# UTILERÍAS
# ------------------------------------------------------------------------------
def normalizar(texto):
    if texto is None:
        return ""
    sin_acentos = unicodedata.normalize("NFKD", str(texto))
    sin_acentos = "".join(ch for ch in sin_acentos if not unicodedata.combining(ch))
    return " ".join(sin_acentos.lower().split())


def vacio(valor):
    return valor is None or (isinstance(valor, str) and valor.strip() == "")


def a_numero(valor):
    """Convierte a float: 0.8, '80%', '0,8', ' 80 % '. Devuelve (numero, es_invalido)."""
    if vacio(valor):
        return None, False
    if isinstance(valor, bool):  # VERDADERO/FALSO en una celda de porcentaje
        return None, True
    if isinstance(valor, (int, float)):
        return float(valor), False
    texto = str(valor).strip()
    es_pct = texto.endswith("%")
    texto = texto.rstrip("%").strip().replace("\u00a0", "").replace(",", ".")
    try:
        numero = float(texto)
    except ValueError:
        return None, True
    return (numero / 100 if es_pct else numero), False


def a_codigo(valor):
    """Códigos (TR, cedente, corredor...) como entero cuando se puede."""
    numero, _ = a_numero(valor)
    if numero is not None and float(numero).is_integer():
        return int(numero)
    return None if vacio(valor) else str(valor).strip()


def clave_codigo(valor):
    """Llave de texto para cruzar códigos con el catálogo: 39, '39', 39.0 -> '39'."""
    codigo = a_codigo(valor)
    return None if codigo is None else str(codigo).upper()


def a_si_no(valor):
    if vacio(valor):
        return None
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    texto = normalizar(valor)
    if texto in ("verdadero", "true", "si", "1"):
        return "Sí"
    if texto in ("falso", "false", "no", "0"):
        return "No"
    return str(valor)


# ------------------------------------------------------------------------------
# UBICAR LAS CARPETAS DE LAS LN
# ------------------------------------------------------------------------------
EXT_SOPORTADAS = (".xlsx", ".xlsm")
EXT_NO_SOPORTADAS = (".xls", ".xlsb")
MODO_COLUMNA = str(COLUMNA_PRIMER_ANIO).strip().upper()
if MODO_COLUMNA not in ("AUTO", "B", "C"):
    sys.exit(f'COLUMNA_PRIMER_ANIO debe ser "AUTO", "B" o "C" (dice "{COLUMNA_PRIMER_ANIO}").')


def carpetas_del_usuario():
    """Documentos, Escritorio y Descargas del usuario (incluye OneDrive), sin repetir."""
    casa = Path.home()
    candidatas = []
    if os.name == "nt":
        try:  # Ruta real de "Documentos" y "Escritorio", aunque estén redirigidas a OneDrive
            import ctypes
            import ctypes.wintypes
            for csidl in (5, 0):  # 5 = Documentos, 0 = Escritorio
                buf = ctypes.create_unicode_buffer(ctypes.wintypes.MAX_PATH)
                if ctypes.windll.shell32.SHGetFolderPathW(None, csidl, None, 0, buf) == 0:
                    candidatas.append(Path(buf.value))
        except Exception:
            pass
    nombres = ("Documents", "Documentos", "Desktop", "Escritorio", "Downloads", "Descargas")
    for var in ("OneDriveCommercial", "OneDrive", "OneDriveConsumer"):
        if os.environ.get(var):
            candidatas += [Path(os.environ[var]) / n for n in nombres]
    candidatas += [casa / n for n in nombres]
    for onedrive in sorted(casa.glob("OneDrive*")):
        candidatas += [onedrive / n for n in nombres]
    vistas, unicas = set(), []
    for c in candidatas:
        clave = str(c).lower()
        if clave not in vistas and c.is_dir():
            vistas.add(clave)
            unicas.append(c)
    return unicas


def es_excel_candidato(ruta):
    """Archivos que podrían ser logouts (no temporales de Office ni salidas de este script)."""
    nombre = ruta.name.lower()
    return not nombre.startswith("~$") and not nombre.startswith("base_cesion_")


def excels_en(carpeta, extensiones=EXT_SOPORTADAS):
    return sorted((p for p in carpeta.glob("**/*")
                   if p.suffix.lower() in extensiones and es_excel_candidato(p) and p.is_file()),
                  key=lambda p: str(p).lower())


def partes_ln(texto, solo_numero_ok=False):
    """'LN04006' / 'LN4006' / 'ln 4006' / 'LN4006 - copia' -> (4006, '');
    'LN4008-Agro' -> (4008, 'agro'); lo que no es LN -> None. Con solo_numero_ok, '4006' también."""
    texto = normalizar(texto)
    m = re.fullmatch(r"(?:ln)?[\s_-]*0*(\d+)(.*)", texto)
    if not m or (not texto.startswith("ln") and (not solo_numero_ok or m[2].strip())):
        return None
    sufijo = re.sub(r"[^a-z0-9]", "", m[2])
    sufijo = re.sub(r"^(copia|copy|respaldo|backup|old)?\d*$", "", sufijo)  # "- copia (2)", " (1)"
    return int(m[1]), sufijo


def etiqueta_ln(partes):
    return f"LN{partes[0]}" + (f"-{partes[1].title()}" if partes[1] else "")


def mas_reciente(carpeta):
    return max((p.stat().st_mtime for p in excels_en(carpeta)), default=0)


def buscar_carpetas_ln():
    """{(número, sufijo): [(prioridad, carpeta)]} con las carpetas LN#### que traen archivos Excel.
    Prioridad: 0 = CARPETA_LNS, 1 = junto al script, 2 = Documentos / Escritorio / Descargas."""
    raices = []
    if CARPETA_LNS:
        carpeta_lns = Path(CARPETA_LNS).expanduser()
        if not carpeta_lns.is_dir():
            sys.exit(f"CARPETA_LNS no existe: {carpeta_lns}")
        raices.append((carpeta_lns, 2, 0))
    junto_al_script = Path(__file__).resolve().parent
    raices += [(junto_al_script, 2, 1), (junto_al_script.parent, 1, 2)]
    raices += [(c, 3, 2) for c in carpetas_del_usuario()]  # ...\LN, ...\X\LN, ...\X\Y\LN

    halladas, vistas = [], set()
    for raiz, niveles, prioridad in raices:
        if not raiz.is_dir():
            continue
        candidatas = [raiz] + [c for n in range(1, niveles + 1) for c in raiz.glob("/".join(["*"] * n) + "/")]
        for c in candidatas:
            try:
                partes = partes_ln(c.name)
                real = c.resolve()
                if partes is None or str(real).lower() in vistas or not c.is_dir() \
                        or not excels_en(c, EXT_SOPORTADAS + EXT_NO_SOPORTADAS):
                    continue
                vistas.add(str(real).lower())
                halladas.append((partes, prioridad, real))
            except OSError:  # carpetas sin permiso de lectura
                continue
    # Una carpeta LN dentro de otra de la misma LN (zip descomprimido dos veces,
    # "LN4006 Líneas Especiales\LN4006") se lee junto con la de afuera.
    encontradas = {}
    for partes, prioridad, real in halladas:
        if not any(p[0] == partes[0] and r != real and r in real.parents for p, _, r in halladas):
            encontradas.setdefault(partes, []).append((prioridad, real))
    return encontradas


def elegir_carpeta_con_ventana(titulo):
    try:
        import tkinter as tk
        from tkinter import filedialog
        raiz = tk.Tk()
        raiz.withdraw()
        raiz.attributes("-topmost", True)
        ruta = filedialog.askdirectory(title=titulo)
        raiz.destroy()
        return Path(ruta) if ruta else None
    except Exception:
        return None


def pedir_carpeta(mensaje):
    print(mensaje)
    ruta = elegir_carpeta_con_ventana("Elige la carpeta de la LN (o la que contiene las carpetas LN)")
    if ruta is None:
        try:
            texto = input("Pega la ruta de la carpeta con los logouts: ")
        except EOFError:
            texto = ""
        # Quita comillas y el "& '...'" que agrega PowerShell al arrastrar una carpeta
        texto = re.sub(r"^&\s*", "", texto.strip()).strip().strip("'\"")
        ruta = Path(texto) if texto else None
    if ruta is None:
        sys.exit("No se indicó carpeta. Fin.")
    return ruta.expanduser().resolve()


def es_todas(ln_pedida):
    return normalizar(ln_pedida) in ("todas", "todos", "*")


def expandir(carpeta, ln_pedida, avisos, todas_si_no_esta=False):
    """Una carpeta dada a mano: si contiene carpetas LN####, se toman esas (todas o la pedida)."""
    if partes_ln(carpeta.name):
        return [(etiqueta_ln(partes_ln(carpeta.name)), carpeta)]
    subcarpetas = sorted((partes_ln(c.name), c.resolve()) for c in carpeta.iterdir()
                         if c.is_dir() and partes_ln(c.name) and excels_en(c, EXT_SOPORTADAS + EXT_NO_SOPORTADAS))
    if not subcarpetas:
        return [(carpeta.name, carpeta)]
    pedida = None if es_todas(ln_pedida) else partes_ln(ln_pedida, solo_numero_ok=True)
    grupos = [(etiqueta_ln(p), c) for p, c in subcarpetas if pedida is None or p[0] == pedida[0]]
    if not grupos and todas_si_no_esta:
        avisos.append(f"En {carpeta} no hay carpeta de {ln_pedida}; se leen todas las que trae.")
        grupos = [(etiqueta_ln(p), c) for p, c in subcarpetas]
    if not grupos:
        sys.exit(f"En {carpeta} no hay carpeta de {ln_pedida}; hay: " + ", ".join(c.name for _, c in subcarpetas))
    return grupos


def ubicar_grupos(ruta_argumento, ln_pedida):
    """Lista de (etiqueta, carpeta) a procesar y avisos sobre carpetas descartadas."""
    avisos = []
    if ruta_argumento or CARPETA_LOGOUTS:
        carpeta = Path(ruta_argumento or CARPETA_LOGOUTS).expanduser().resolve()
        if not carpeta.is_dir():
            sys.exit(f"La carpeta no existe: {carpeta}")
        return expandir(carpeta, ln_pedida, avisos, todas_si_no_esta=True), avisos

    encontradas = buscar_carpetas_ln()

    def elegir(partes, opciones):
        # Primero la fuente más explícita (CARPETA_LNS, junto al script); si empatan, la más reciente
        opciones = sorted(opciones, key=lambda o: (o[0], -mas_reciente(o[1])))
        for _, otra in opciones[1:]:
            avisos.append(f"{etiqueta_ln(partes)}: también encontré {otra}; se usa {opciones[0][1]}.")
        return opciones[0][1]

    if es_todas(ln_pedida):
        if not encontradas:
            carpeta = pedir_carpeta("No encontré carpetas LN#### (LN4001, LN4002, ...) junto al script "
                                    "ni en Documentos, Escritorio o Descargas.")
            return expandir(carpeta, ln_pedida, avisos), avisos
        return [(etiqueta_ln(p), elegir(p, o)) for p, o in sorted(encontradas.items())], avisos

    pedida = partes_ln(ln_pedida, solo_numero_ok=True)
    if pedida is None:
        sys.exit(f'LN = "{ln_pedida}" no parece una LN. Usa, por ejemplo, LN = "LN4006" o LN = "TODAS".')
    if pedida in encontradas:
        return [(etiqueta_ln(pedida), elegir(pedida, encontradas[pedida]))], avisos
    if not pedida[1]:
        # Sin carpeta exacta: "LN4006 Líneas Especiales" o sublíneas como LN4008-Agro y LN4008-Vida
        parecidas = {p: o for p, o in encontradas.items() if p[0] == pedida[0]}
        if parecidas:
            return [(etiqueta_ln(p), elegir(p, o)) for p, o in sorted(parecidas.items())], avisos
    carpeta = pedir_carpeta(f'No encontré la carpeta "{ln_pedida}" junto al script ni en Documentos, '
                            f'Escritorio o Descargas' + (f" ni en {CARPETA_LNS}" if CARPETA_LNS else "") + ".")
    return expandir(carpeta, ln_pedida, avisos), avisos


# ------------------------------------------------------------------------------
# LECTURA DE UN LOGOUT
# ------------------------------------------------------------------------------
PATRON_NOMBRE = re.compile(
    r"^Logout_(?P<ln>[^_]+)_r(?P<tr>\d+)_b(?P<corredor>\d+)_e(?P<cedente>\d+)"
    r"(?:_c(?P<contrato>\d+))?-(?P<venta>[^-]+)-v(?P<version>\d+)$",
    re.IGNORECASE,
)

# Etiqueta en columna A -> clave interna. Se toma la PRIMERA fila con esa etiqueta
# (la palabra "Corredor" aparece dos veces; la segunda es la distribución proyectada).
ETIQUETAS = {
    "linea de negocio": "ln",
    "tipo reas.": "tr",
    "of. rep.": "of_rep",
    "compania": "cedente",
    "corredor": "corredor",
    "contrato": "contrato",
    "tipo venta": "tipo_venta",
    "mga": "mga",
    "porcentaje": "retro",
    "porcentaje de cesion": "cesion",
    "% comisiones del cedido": "comision",
}
TIPOS_RETRO = ["Tradicional", "Retro Espec.", "Fronting", "Retención"]
ANIOS = list(range(PRIMER_ANIO, PRIMER_ANIO + NUM_ANIOS))
NIVELES = ["Cedente", "Contrato", "Binder"]  # de menos a más granular
MAX_FILAS_ENCABEZADO = 40
MAX_COLUMNAS = 26  # A..Z: se lee de más para detectar valores fuera de lugar


def leer_filas(ruta, formulas):
    """(hoja leída, primeras filas, ¿es logout?). formulas=True trae el texto de las fórmulas."""
    wb = load_workbook(ruta, read_only=True, data_only=not formulas, keep_links=False)
    try:
        ws = wb[HOJA_LOGOUT] if HOJA_LOGOUT in wb.sheetnames else wb.worksheets[0]
        filas = [list(f) + [None] * (MAX_COLUMNAS - len(f))
                 for f in ws.iter_rows(min_row=1, max_row=MAX_FILAS_ENCABEZADO,
                                       max_col=MAX_COLUMNAS, values_only=True)]
        es_logout = ws.title == HOJA_LOGOUT or all(k in ubicar_filas(filas) for k in ("ln", "cedente", "cesion"))
        return ws.title, filas, es_logout
    finally:
        wb.close()


def ubicar_filas(filas):
    """{clave interna: índice de fila} según las etiquetas de la columna A."""
    fila_de = {}
    for i, fila in enumerate(filas):
        clave = ETIQUETAS.get(normalizar(fila[0]))
        if clave and clave not in fila_de:
            fila_de[clave] = i
    return fila_de


def columna_primer_anio(lecturas):
    """(columna del primer año: 2 = B o 3 = C, logouts con dato en B, logouts con dato solo desde C)
    para los logouts de una LN. En AUTO decide la mayoría; si empatan (o no hay datos), B."""
    en_b = en_c = 0
    for _, filas, _ in lecturas:
        fila_de = ubicar_filas(filas)
        renglones = [filas[fila_de[k]] for k in ("cesion", "comision") if k in fila_de]
        if any(not vacio(r[1]) for r in renglones):
            en_b += 1
        elif any(not vacio(v) for r in renglones for v in r[2:2 + NUM_ANIOS]):
            en_c += 1
    if MODO_COLUMNA in ("B", "C"):
        return (2 if MODO_COLUMNA == "B" else 3), en_b, en_c
    return (3 if en_c > en_b else 2), en_b, en_c


def nombre_relativo(ruta, carpeta):
    """Identificador del archivo: su ruta dentro de la carpeta leída (distingue subcarpetas)."""
    try:
        return str(ruta.relative_to(carpeta))
    except ValueError:
        return ruta.name


def texto_limpio(valor):
    """Nombre como texto; vacío o 0 (sin nombre) -> None."""
    if vacio(valor) or str(valor).strip() in ("0", "0.0"):
        return None
    return str(valor).strip()


def leer_logout(ruta, archivo, lectura, col_ini):
    """Devuelve (registro, lista_de_validaciones) para un archivo ya leído."""
    avisos = []

    def avisar(nivel, tipo, detalle):
        avisos.append({"Archivo": archivo, "Nivel": nivel, "Tipo": tipo, "Detalle": detalle})

    registro = {"Archivo": archivo}
    hoja, filas, _ = lectura
    if hoja != HOJA_LOGOUT:
        avisar("Revisar", "Hoja no encontrada", f'No existe la hoja "{HOJA_LOGOUT}"; se leyó "{hoja}".')

    fila_de = ubicar_filas(filas)
    faltan = [k for k in ETIQUETAS.values() if k not in fila_de]
    if faltan:
        avisar("Error", "Etiquetas no encontradas",
               "No se encontraron en la columna A: " + ", ".join(faltan))

    # Fórmulas guardadas sin valor calculado (se leerían como celda vacía).
    _, crudas, _ = leer_filas(ruta, formulas=True)
    ultima_col = max(6, col_ini + NUM_ANIOS - 1)
    sin_valor = [f"{get_column_letter(j + 1)}{i + 1}"
                 for i in sorted(fila_de.values()) for j in range(1, ultima_col)
                 if isinstance(crudas[i][j], str) and crudas[i][j].startswith("=")
                 and filas[i][j] is None]
    if sin_valor:
        avisar("Error", "Fórmula sin valor calculado",
               "Celdas " + ", ".join(sin_valor) + ": abrir el logout en Excel, guardarlo y volver a correr")

    def celda(clave, col):  # col: 1 = A, 2 = B, ...
        return filas[fila_de[clave]][col - 1] if clave in fila_de else None

    def porcentaje(clave, col, nombre):
        valor = celda(clave, col)
        numero, invalido = a_numero(valor)
        if invalido:
            avisar("Error", "Valor no numérico", f"{nombre}: '{valor}'")
        elif numero is not None and numero < -TOLERANCIA:
            avisar("Revisar", "Porcentaje negativo", f"{nombre} = {numero:g}")
        elif numero is not None and numero > 1 + TOLERANCIA:
            pista = f" (¿capturado como {numero:g} en lugar de {numero / 100:g}?)" if numero <= 100 else ""
            avisar("Revisar", "Porcentaje mayor a 100%", f"{nombre} = {numero:g}{pista}")
        return numero

    def valores_fuera(clave, columnas_validas, nombre):
        fuera = [(get_column_letter(c), celda(clave, c)) for c in range(2, MAX_COLUMNAS + 1)
                 if c not in columnas_validas and not vacio(celda(clave, c))]
        if fuera:
            rango = f"{get_column_letter(min(columnas_validas))}..{get_column_letter(max(columnas_validas))}"
            avisar("Revisar", "Valor fuera de las columnas esperadas",
                   f"{nombre}: " + ", ".join(f"{c}={v}" for c, v in fuera) + f" (se esperan solo en {rango})")

    # --- Llaves del documento --------------------------------------------------
    registro["LN"] = a_codigo(celda("ln", 2))
    registro["TR"] = a_codigo(celda("tr", 2))
    registro["Of. Rep."] = a_codigo(celda("of_rep", 2))
    registro["Cedente"] = a_codigo(celda("cedente", 2))
    registro["Corredor"] = a_codigo(celda("corredor", 2))
    registro["Contrato"] = a_codigo(celda("contrato", 2))
    registro["Tipo Venta"] = celda("tipo_venta", 2)
    registro["% Renov."] = porcentaje("tipo_venta", 3, "% Renov.")
    registro["Es MGA"] = a_si_no(celda("mga", 2))
    registro["Binder general"] = texto_limpio(celda("mga", 3))     # PRESUPUESTO!E18 (yBind)
    registro["Binder segmentado"] = texto_limpio(celda("mga", 5))
    registro["GS"] = a_si_no(celda("mga", 6))
    vacias = [c for c in ("TR", "Cedente", "Corredor") if registro[c] is None]
    if vacias:
        avisar("Info", "Llave vacía", ", ".join(vacias) + " vacío en el logout (distinto de 0)")

    # --- Cardinalidad del documento -----------------------------------------------
    general, segmentado = registro["Binder general"], registro["Binder segmentado"]
    if registro["Es MGA"] == "Sí" or general or segmentado:
        registro["Nivel"] = "Binder"
    elif registro["Contrato"] not in (None, 0):
        registro["Nivel"] = "Contrato"
    else:
        registro["Nivel"] = "Cedente"
    if registro["Es MGA"] == "Sí" and not (general or segmentado):
        avisar("Revisar", "MGA sin nombre de binder",
               "La casilla MGA es Verdadero pero no trae Binder general ni Binder segmentado")
    elif registro["Nivel"] == "Binder" and not (general and segmentado):
        avisar("Info", "Binder con nombre incompleto",
               f"Binder general = {general or '(vacío)'}; Binder segmentado = {segmentado or '(vacío)'}")
    if registro["Es MGA"] == "No" and (general or segmentado):
        avisar("Revisar", "Binder con MGA = Falso",
               f"Trae binder ({general or segmentado}) pero la casilla MGA es Falso; se clasifica como Binder")

    # --- Tipo de retrocesión ----------------------------------------------------
    total = 0.0
    capturado = False
    for i, tipo in enumerate(TIPOS_RETRO):
        numero = porcentaje("retro", 2 + i, f"% {tipo}")
        registro[f"% {tipo}"] = numero
        if numero is not None:
            total += numero
            capturado = True
    registro["% Total Retrocesión"] = total if capturado else None
    valores_fuera("retro", range(2, 2 + len(TIPOS_RETRO)), "Porcentaje (Tipo Retrocesión)")

    # --- % de cesión y % comisiones del cedido por año --------------------------
    registro["Col. 1er año"] = get_column_letter(col_ini)
    cols_anios = range(col_ini, col_ini + NUM_ANIOS)
    for clave, prefijo in (("cesion", "% Cesión"), ("comision", "% Com. Cedido")):
        valores = [porcentaje(clave, c, f"{prefijo} {anio}") for c, anio in zip(cols_anios, ANIOS)]
        registro[f"_capturados_{clave}"] = [a for a, v in zip(ANIOS, valores) if v is not None]
        if ARRASTRAR_ULTIMO_ANIO:
            for i in range(1, NUM_ANIOS):
                if valores[i] is None:
                    valores[i] = valores[i - 1]
        for anio, v in zip(ANIOS, valores):
            registro[f"{prefijo} {anio}"] = v
        validas = list(cols_anios)
        if col_ini > 2:  # la columna B queda antes del primer año
            validas.append(2)
            if not vacio(celda(clave, 2)):
                avisar("Error", "Valor en la columna B de años",
                       f"{prefijo}: B={celda(clave, 2)}, pero en esta LN el {PRIMER_ANIO} cae en "
                       f"{get_column_letter(col_ini)}; revisar la alineación de años.")
        valores_fuera(clave, validas, prefijo)

    capturados = registro["_capturados_cesion"]
    cesion_capturada = bool(capturados)
    registro["Cesión capturada"] = "Sí" if cesion_capturada else "No"
    registro["Años con cesión"] = ", ".join(map(str, capturados)) or None

    # --- Consistencia con el nombre del archivo ---------------------------------
    # Sin sufijos de copia: "... -v2 (1)" (descarga repetida), "... -v2 - copia" (Explorador)
    nombre_limpio = re.sub(r"(\s*(\(\d+\)|-\s*(copia|copy)(\s*\(\d+\))?))+$", "", ruta.stem,
                           flags=re.IGNORECASE)
    version = re.findall(r"(?:^|[\s_-])v(\d+)\b", nombre_limpio, flags=re.IGNORECASE)
    registro["Versión archivo"] = int(version[-1]) if version else None
    registro["_es_copia"] = nombre_limpio != ruta.stem
    if nombre_limpio != ruta.stem:
        avisar("Info", "Copia de archivo", f"El nombre trae un sufijo de copia: '{ruta.stem[len(nombre_limpio):]}'")
    m = PATRON_NOMBRE.match(nombre_limpio)
    if m:  # solo los logouts que se llaman Logout_<LN>_r.._b.._e..: el nombre trae las llaves
        esperado = {"LN": m["ln"], "TR": int(m["tr"]), "Corredor": int(m["corredor"]),
                    "Cedente": int(m["cedente"]),
                    "Contrato": int(m["contrato"]) if m["contrato"] else None}
        for campo, valor in esperado.items():
            en_hoja = "" if registro[campo] is None else str(registro[campo]).upper()
            if en_hoja != ("" if valor is None else str(valor).upper()):
                avisar("Error", "Nombre vs contenido",
                       f"{campo}: archivo dice {valor}, la hoja dice {registro[campo]}")
        venta_archivo = normalizar(m["venta"])
        venta_hoja = normalizar(registro["Tipo Venta"])
        if (venta_archivo.startswith("renov") and not venta_hoja.startswith("renov")) or \
           (venta_archivo.startswith("nva") and not venta_hoja.startswith("nuevo")):
            avisar("Revisar", "Nombre vs contenido",
                   f"Tipo de venta: archivo dice {m['venta']}, la hoja dice {registro['Tipo Venta']}")

    # --- Reglas de negocio --------------------------------------------------------
    if not capturado:
        avisar("Revisar", "Tipo Retrocesión vacío", "No hay % en Tradicional/Retro Espec./Fronting/Retención")
    elif abs(total - 1) > TOLERANCIA:
        avisar("Revisar", "Tipo Retrocesión no suma 100%", f"Total = {total:.2%}")
    especifica = (registro["% Retro Espec."] or 0) + (registro["% Fronting"] or 0)
    if especifica > TOLERANCIA and not cesion_capturada:
        avisar("Revisar", "Falta % de cesión",
               f"Retro Espec. + Fronting = {especifica:.0%} pero no hay % de cesión capturado")
    if especifica > TOLERANCIA and cesion_capturada and \
            all(abs(registro[f"% Cesión {a}"] or 0) <= TOLERANCIA for a in ANIOS):
        avisar("Revisar", "Cesión 0% con Retro Espec./Fronting",
               f"Retro Espec. + Fronting = {especifica:.0%} pero el % de cesión capturado es 0%")
    if capturado and cesion_capturada and especifica <= TOLERANCIA:
        avisar("Revisar", "Cesión sin Retro Espec./Fronting",
               "Hay % de cesión capturado pero el tipo de retrocesión es 100% Tradicional/Retención")
    if registro["_capturados_comision"] and not cesion_capturada:
        avisar("Revisar", "Comisión sin % de cesión",
               "Hay % Comisiones del Cedido en " + ", ".join(map(str, registro["_capturados_comision"]))
               + " pero no hay % de cesión")
    if cesion_capturada and not registro["_capturados_comision"]:
        avisar("Info", "Cesión sin % comisión",
               "Hay % de cesión en " + ", ".join(map(str, capturados))
               + " pero la fila % Comisiones del Cedido está vacía")
    elif cesion_capturada and set(capturados) != set(registro["_capturados_comision"]):
        avisar("Info", "Años distintos en cesión y comisión",
               f"Cesión: {capturados}; Comisión: {registro['_capturados_comision']}")
    if cesion_capturada and capturados[0] > PRIMER_ANIO:
        avisar("Revisar", f"Cesión sin {PRIMER_ANIO}",
               f"La cesión empieza en {capturados[0]} (col. "
               f"{get_column_letter(col_ini + capturados[0] - PRIMER_ANIO)}); falta {PRIMER_ANIO}")
    elif cesion_capturada and len(capturados) < NUM_ANIOS:
        avisar("Info", "Cesión no viene en todos los años del logout",
               "Años con cesión: " + ", ".join(map(str, capturados)))
    es_combinada = "comb" in normalizar(registro["Tipo Venta"])
    if es_combinada and registro["% Renov."] is None:
        avisar("Revisar", "Falta % Renov.", "Tipo de venta combinado sin % de renovación")
    if not es_combinada and registro["% Renov."] is not None:
        avisar("Info", "% Renov. en venta no combinada",
               f"Tipo de venta {registro['Tipo Venta']} con % Renov. = {registro['% Renov.']:g}")

    registro["Ruta"] = str(ruta)
    return registro, avisos


# ------------------------------------------------------------------------------
# CATÁLOGO DE NOMBRES (caché de la hoja "Valores" del PptoTécnico en el logout)
# ------------------------------------------------------------------------------
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
NS_REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
NS_PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"
PATRON_RANGO_EXTERNO = re.compile(
    r"'?\[(?P<libro>\d+)\](?P<hoja>[^'!]+)'?!\$?(?P<col>[A-Z]+)\$?(?P<ini>\d+):\$?[A-Z]+\$?(?P<fin>\d+)")


def ubicar_catalogos(libro):
    """{clave: (parte externalLink, hoja, columna del código, fila inicial, fila final)}.

    Sigue los nombres definidos del logout (xCEDENTES = [1]Valores!$AR$2:$AV$999, ...)
    hasta la parte del zip que guarda el caché de ese libro externo. Si el logout no
    trae el nombre definido, usa la columna de respaldo en el primer vínculo que tenga
    la hoja HOJA_CATALOGO con datos."""
    raiz = ElementTree.fromstring(libro.read("xl/workbook.xml"))
    rels = ElementTree.fromstring(libro.read("xl/_rels/workbook.xml.rels"))
    destinos = {r.get("Id"): r.get("Target", "") for r in rels.iter(NS_PKG + "Relationship")}
    partes = []
    for ref in raiz.iter(NS + "externalReference"):
        destino = destinos.get(ref.get(NS_REL + "id"), "")
        partes.append(destino.lstrip("/") if destino.startswith("/") else "xl/" + destino)
    definidos = {d.get("name", "").lower(): (d.text or "").strip() for d in raiz.iter(NS + "definedName")}

    ubicaciones = {}
    for clave, (nombre_definido, col_respaldo, _) in CATALOGOS.items():
        m = PATRON_RANGO_EXTERNO.fullmatch(definidos.get(nombre_definido.lower(), ""))
        if m and 1 <= int(m["libro"]) <= len(partes):
            ubicaciones[clave] = (partes[int(m["libro"]) - 1], m["hoja"], m["col"], int(m["ini"]), int(m["fin"]))
    faltan = [c for c in CATALOGOS if c not in ubicaciones]
    if faltan:
        numero = lambda p: int(re.sub(r"\D", "", p) or 0)  # noqa: E731
        for parte in sorted((p for p in libro.namelist()
                             if re.fullmatch(r"xl/externalLinks/externalLink\d+\.xml", p)), key=numero):
            hoja = leer_hoja_en_cache(libro.read(parte), HOJA_CATALOGO)
            for clave in list(faltan):
                col = CATALOGOS[clave][1]
                if any(col in celdas for celdas in hoja.values()):
                    ubicaciones[clave] = (parte, HOJA_CATALOGO, col, 1, 10 ** 7)
                    faltan.remove(clave)
    return ubicaciones


def leer_hoja_en_cache(xml, hoja):
    """{fila: {columna: texto}} de una hoja guardada en caché en un externalLinkN.xml."""
    raiz = ElementTree.fromstring(xml)
    hojas = [h.get("val") for h in raiz.iter(NS + "sheetName")]
    if hoja not in hojas:
        return {}
    indice = str(hojas.index(hoja))
    tabla = {}
    for datos in raiz.iter(NS + "sheetData"):
        if datos.get("sheetId") != indice:
            continue
        for fila in datos.iter(NS + "row"):
            for c in fila.iter(NS + "cell"):
                v = c.find(NS + "v")
                m = re.fullmatch(r"([A-Z]+)(\d+)", c.get("r", ""))
                if m and v is not None and v.text and v.text.strip():
                    tabla.setdefault(int(m[2]), {})[m[1]] = v.text.strip()
    return tabla


def catalogo_del_logout(ruta, cache):
    """{'Cedente': {'39': {'Nombre Cedente': ..., ...}}, ...} con el caché del propio logout.

    `cache` guarda cada catálogo ya leído, identificado por el CRC de su parte en el
    zip: los logouts de un mismo PptoTécnico traen el mismo catálogo y se lee una vez."""
    catalogo = {clave: {} for clave in CATALOGOS}
    with zipfile.ZipFile(ruta) as libro:
        for clave, (parte, hoja, col, ini, fin) in ubicar_catalogos(libro).items():
            info = libro.getinfo(parte)
            llave_hoja = (info.CRC, info.file_size, hoja)
            if llave_hoja not in cache:
                cache[llave_hoja] = leer_hoja_en_cache(libro.read(parte), hoja)
            llave = llave_hoja + (clave, col, ini, fin)
            if llave not in cache:
                tabla, codigos = cache[llave_hoja], {}
                n_col = column_index_from_string(col)
                campos = CATALOGOS[clave][2]
                for fila in sorted(f for f in tabla if ini <= f <= fin):
                    codigo = clave_codigo(tabla[fila].get(col))
                    if codigo is not None:
                        codigos.setdefault(codigo, {campo: tabla[fila].get(get_column_letter(n_col + desp))
                                                    for campo, desp in campos.items()})
                cache[llave] = codigos
            catalogo[clave] = cache[llave]
    return catalogo


def agregar_nombres(registros):
    """Pone los nombres del catálogo del propio logout; si ahí falta un código, usa el de
    los demás logouts de la carpeta. Devuelve cuántos logouts traían catálogo."""
    cache, propios = {}, []
    combinado = {clave: {} for clave in CATALOGOS}
    for reg in registros:
        try:
            propio = catalogo_del_logout(Path(reg["Ruta"]), cache)
        except Exception:  # sin vínculos externos, zip raro, etc.
            propio = {clave: {} for clave in CATALOGOS}
        propios.append(propio)
        for clave, codigos in propio.items():
            for codigo, datos in codigos.items():
                combinado[clave].setdefault(codigo, datos)
    for reg, propio in zip(registros, propios):
        for clave, (_, _, campos) in CATALOGOS.items():
            codigo = clave_codigo(reg[clave])
            datos = propio[clave].get(codigo) or combinado[clave].get(codigo) or {}
            for campo in campos:
                reg[campo] = datos.get(campo)
        if reg["TR desc."] is None:
            reg["TR desc."] = TR_DESC_RESPALDO.get(reg["TR"])
    return sum(any(p.values()) for p in propios)


# ------------------------------------------------------------------------------
# ARMADO DE LA BASE
# ------------------------------------------------------------------------------
# Llave de un documento: dos archivos con la misma llave son versiones del mismo documento.
LLAVES = ["LN", "Nivel", "TR", "Cedente", "Corredor", "Contrato", "Binder general", "Binder segmentado",
          "Tipo Venta"]
COLUMNAS_BASE = (
    ["Archivo", "Carpeta LN", "LN", "Nombre LN", "Nivel", "TR", "TR desc.", "Cedente", "Nombre Cedente",
     "País Cedente", "Grupo Cedente", "Corredor", "Nombre Corredor", "Contrato", "Es MGA", "Binder general",
     "Binder segmentado", "Of. Rep.", "Tipo Venta", "% Renov."]
    + [f"% {t}" for t in TIPOS_RETRO] + ["% Total Retrocesión", "Cesión capturada", "Años con cesión"]
    + [f"% Cesión {a}" for a in ANIOS] + [f"% Com. Cedido {a}" for a in ANIOS]
    + ["GS", "Col. 1er año", "Versión archivo", "Versión vigente", "Observaciones", "Ruta"]
)
COLUMNAS_VALIDACIONES = ["Archivo", "Nivel", "Tipo", "Detalle"]


def numero_ln(ln, carpeta_ln):
    """Número de la LN de un logout (celda 'Línea de Negocio'); si no la trae, el de su carpeta."""
    partes = partes_ln(ln) or partes_ln(carpeta_ln)
    return partes[0] if partes else None


def leer_grupo(etiqueta, carpeta, raiz_relativa, validaciones):
    """Lee los logouts de la carpeta de una LN. Devuelve (registros, {número de LN: columna del 1er año})."""
    archivos = excels_en(carpeta)
    for p in excels_en(carpeta, EXT_NO_SOPORTADAS):
        validaciones.append({"Archivo": nombre_relativo(p, raiz_relativa), "Nivel": "Info",
                             "Tipo": "Formato no leído",
                             "Detalle": f"{p.suffix} no se puede leer; si es un logout, guárdalo como .xlsx"})
    print(f"{etiqueta}: {len(archivos)} archivos en {carpeta}")

    lecturas = []
    for n, ruta in enumerate(archivos, 1):
        print(f"  [{n}/{len(archivos)}] {ruta.name}")
        archivo = nombre_relativo(ruta, raiz_relativa)
        try:
            lectura = leer_filas(ruta, formulas=False)
        except Exception as error:  # archivo dañado, protegido, abierto, etc.
            validaciones.append({"Archivo": archivo, "Nivel": "Error", "Tipo": "No se pudo leer",
                                 "Detalle": f"{type(error).__name__}: {error}"})
            continue
        if not lectura[2]:
            validaciones.append({"Archivo": archivo, "Nivel": "Info", "Tipo": "No es logout",
                                 "Detalle": f'No tiene la hoja "{HOJA_LOGOUT}"; no se incluye en la base'})
            continue
        filas = lectura[1]
        fila_ln = ubicar_filas(filas).get("ln")
        lecturas.append((ruta, archivo, lectura,
                         numero_ln(filas[fila_ln][1] if fila_ln is not None else None, etiqueta)))

    # La columna del primer año se decide por LN del logout: un logout de otra LN que se
    # coló en la carpeta no cambia la de las demás.
    esperada = partes_ln(etiqueta)
    por_ln = {}
    for item in lecturas:
        por_ln.setdefault(item[3], []).append(item)
    registros, columnas = [], {}
    for num_ln, items in por_ln.items():
        col_ini, en_b, en_c = columna_primer_anio([lectura for _, _, lectura, _ in items])
        columnas[num_ln] = get_column_letter(col_ini)
        es_la_carpeta = len(por_ln) == 1 or (esperada is not None and num_ln == esperada[0])
        nombre = etiqueta if es_la_carpeta else f"{etiqueta} · LN{num_ln}"
        regs = []
        for ruta, archivo, lectura, _ in items:
            try:
                registro, avisos = leer_logout(ruta, archivo, lectura, col_ini)
            except Exception as error:
                validaciones.append({"Archivo": archivo, "Nivel": "Error", "Tipo": "No se pudo leer",
                                     "Detalle": f"{type(error).__name__}: {error}"})
                continue
            registro["Carpeta LN"] = etiqueta
            regs.append(registro)
            validaciones.extend(avisos)
            if esperada and num_ln is not None and num_ln != esperada[0]:
                validaciones.append({"Archivo": archivo, "Nivel": "Revisar", "Tipo": "Logout de otra LN",
                                     "Detalle": f"El logout es de {registro['LN']} y está en la carpeta {etiqueta}"})
        registros += regs

        # Avisos de la LN completa
        col = get_column_letter(col_ini)
        if en_b and en_c:
            validaciones.append({
                "Archivo": f"({nombre})", "Nivel": "Revisar", "Tipo": "Columna del primer año mixta",
                "Detalle": f"{en_b} logout(s) traen el % de cesión desde la col. B y {en_c} solo desde C; "
                           f"se usa {col} = {PRIMER_ANIO} para todos. Revisar los de la minoría "
                           f"(salen con 'Cesión sin {PRIMER_ANIO}' o 'Valor en la columna B de años')."})
        if MODO_COLUMNA == "AUTO" and col_ini == 3:
            validaciones.append({
                "Archivo": f"({nombre})", "Nivel": "Revisar", "Tipo": "Años recorridos una columna",
                "Detalle": f"La mayoría de los logouts de {nombre} traen el % de cesión desde la col. C: se "
                           f"toma C = {PRIMER_ANIO}. El recuadro del logout es B..F (así vienen otras LN), así "
                           f"que probablemente la herramienta de esta LN recorre los años y no exporta "
                           f"{ANIOS[-1]}."})
        ultimo = ANIOS[-1]
        multi = sum(len(r["_capturados_cesion"]) > 1 for r in regs)
        if multi and not any(ultimo in r["_capturados_cesion"] for r in regs):
            validaciones.append({
                "Archivo": f"({nombre})", "Nivel": "Revisar", "Tipo": f"Ningún logout trae {ultimo}",
                "Detalle": f"Ningún documento de {nombre} tiene % de cesión en la col. "
                           f"{get_column_letter(col_ini + NUM_ANIOS - 1)} ({ultimo}); {multi} documento(s) "
                           f"capturaron varios años y todos terminan antes."})
    return registros, columnas


def construir_base(grupos, raiz_relativa):
    registros, validaciones, columnas_por_ln = [], [], {}
    for etiqueta, carpeta in grupos:
        regs, columnas = leer_grupo(etiqueta, carpeta, raiz_relativa or carpeta.parent, validaciones)
        registros += regs
        columnas_por_ln.update({(etiqueta, num): col for num, col in columnas.items()})
        if not columnas:
            columnas_por_ln[(etiqueta, None)] = "-"

    if not registros:
        return (pd.DataFrame(columns=COLUMNAS_BASE),
                pd.DataFrame(validaciones, columns=COLUMNAS_VALIDACIONES), columnas_por_ln)

    if agregar_nombres(registros) == 0:
        validaciones.append({"Archivo": "(todos)", "Nivel": "Info", "Tipo": "Sin catálogo de nombres",
                             "Detalle": f'Los logouts no traen la hoja "{HOJA_CATALOGO}" en caché; '
                                        "las columnas de nombre quedan vacías."})
    else:
        for reg in registros:
            sin_nombre = [clave for clave in ("Cedente", "Corredor")
                          if reg[clave] is not None and reg[f"Nombre {clave}"] is None]
            if sin_nombre:
                validaciones.append({"Archivo": reg["Archivo"], "Nivel": "Info", "Tipo": "Código sin nombre",
                                     "Detalle": ", ".join(f"{c} {reg[c]}" for c in sin_nombre)
                                     + " no aparece en el catálogo de los logouts"})

    base = pd.DataFrame(registros)

    # Códigos como enteros (sin ".0") y porcentajes como número aunque vengan vacíos
    for c in ["TR", "Of. Rep.", "Cedente", "Corredor", "Contrato", "Versión archivo"]:
        numeros = pd.to_numeric(base[c], errors="coerce")
        if numeros.notna().sum() == base[c].notna().sum() and (numeros.dropna() % 1 == 0).all():
            base[c] = numeros.astype("Int64")
    for c in [c for c in COLUMNAS_BASE if c.startswith("%") and c in base.columns]:
        base[c] = pd.to_numeric(base[c], errors="coerce").astype("float64")

    # Versiones del mismo documento: la vigente es la de mayor versión; si empatan, la más
    # reciente, y si también empatan, el original antes que la copia. Los archivos sin llaves
    # (libros vacíos) y los binders sin ningún nombre no se agrupan con otros.
    claves = base[LLAVES].astype(object).where(base[LLAVES].notna(), "")
    binder_sin_nombre = base["Nivel"].eq("Binder") & base["Binder general"].isna() & base["Binder segmentado"].isna()
    sin_llave = base[[c for c in LLAVES if c != "Nivel"]].isna().all(axis=1) | binder_sin_nombre
    claves.loc[sin_llave, "Binder segmentado"] = "#" + base.loc[sin_llave, "Archivo"].astype(str)
    llave = claves.astype(str).agg("|".join, axis=1)
    fecha = base["Ruta"].map(lambda r: os.path.getmtime(r) if os.path.exists(r) else 0)
    orden = base.assign(_llave=llave, _fecha=fecha, _v=base["Versión archivo"].fillna(0),
                        _original=~base["_es_copia"].astype(bool))
    vigentes = set(orden.sort_values(["_v", "_fecha", "_original"], kind="mergesort")
                   .groupby("_llave").tail(1).index)
    base["Versión vigente"] = ["Sí" if i in vigentes else "No" for i in base.index]
    for _, grupo in base.groupby(llave):
        if len(grupo) < 2:
            continue
        for i, fila in grupo.iterrows():
            otros = grupo.drop(index=i)["Archivo"].astype(str)
            validaciones.append({"Archivo": fila["Archivo"], "Nivel": "Revisar", "Tipo": "Documento repetido",
                                 "Detalle": f"Versión vigente: {fila['Versión vigente']}. Misma LN/Nivel/TR/Cedente/"
                                            f"Corredor/Contrato/Binder/Venta que: " + ", ".join(otros)})

    val = pd.DataFrame(validaciones, columns=COLUMNAS_VALIDACIONES)
    resumen_obs = (val[val["Nivel"] != "Info"].groupby("Archivo")["Tipo"]
                   .apply(lambda s: "; ".join(dict.fromkeys(s))))
    base["Observaciones"] = base["Archivo"].map(resumen_obs)
    base = base.reindex(columns=COLUMNAS_BASE)
    texto = ("Carpeta LN", "LN", "Binder general", "Binder segmentado", "Archivo")
    base = base.sort_values(["Carpeta LN", "LN", "TR", "Cedente", "Corredor", "Contrato", "Binder general",
                             "Binder segmentado", "Archivo"],
                            key=lambda s: s.map(lambda v: "" if pd.isna(v) else str(v))
                            if s.name in texto else pd.to_numeric(s, errors="coerce"),
                            na_position="first")
    orden_nivel = {"Error": 0, "Revisar": 1, "Info": 2}
    val = val.sort_values(["Nivel", "Archivo", "Tipo"],
                          key=lambda s: s.map(orden_nivel) if s.name == "Nivel" else s)
    return base.reset_index(drop=True), val.reset_index(drop=True), columnas_por_ln


# ------------------------------------------------------------------------------
# HOJAS DE SALIDA
# ------------------------------------------------------------------------------
def construir_anual(base):
    ids = ["Archivo", "Carpeta LN", "LN", "Nivel", "TR", "TR desc.", "Cedente", "Nombre Cedente", "Corredor",
           "Nombre Corredor", "Contrato", "Binder general", "Binder segmentado", "Tipo Venta", "Cesión capturada",
           "Versión vigente"]
    filas = []
    for _, doc in base.iterrows():
        for anio in ANIOS:
            fila = {c: doc[c] for c in ids}
            fila["Año"] = anio
            fila["% Cesión"] = doc[f"% Cesión {anio}"]
            fila["% Com. Cedido"] = doc[f"% Com. Cedido {anio}"]
            filas.append(fila)
    anual = pd.DataFrame(filas, columns=ids + ["Año", "% Cesión", "% Com. Cedido"])
    for c in ["TR", "Cedente", "Corredor", "Contrato"]:
        if str(base[c].dtype) == "Int64":
            anual[c] = anual[c].astype("Int64")
    return anual


def construir_resumen(base):
    """Por LN, nivel y TR. Solo versiones vigentes; los % fuera de 0..100% no entran a mín./promedio/máx."""
    if base.empty:
        return pd.DataFrame()
    b = base[base["Versión vigente"].eq("Sí")]
    col = f"% Cesión {PRIMER_ANIO}"
    b = b.assign(_con=b["Cesión capturada"].eq("Sí"), _obs=b["Observaciones"].notna(),
                 _pct=b[col].where(b[col].between(0, 1)),
                 _desc=b["TR desc."].fillna(""), _nom=b["Nombre LN"].fillna(""), _ln=b["LN"].fillna(""),
                 _nivel=pd.Categorical(b["Nivel"], categories=NIVELES))
    nombres = {"Documentos": ("Archivo", "count"), "Con % cesión": ("_con", "sum"),
               "Con observaciones": ("_obs", "sum"), f"% Cesión {PRIMER_ANIO} mín.": ("_pct", "min"),
               f"% Cesión {PRIMER_ANIO} promedio simple": ("_pct", "mean"),
               f"% Cesión {PRIMER_ANIO} máx.": ("_pct", "max")}
    resumen = (b.groupby(["_ln", "_nom", "_nivel", "TR", "_desc"], dropna=False, observed=True)
               .agg(**nombres).reset_index()
               .rename(columns={"_ln": "LN", "_nom": "Nombre LN", "_nivel": "Nivel", "_desc": "TR desc."}))
    resumen["Nivel"] = resumen["Nivel"].astype(str)
    total = {"LN": "TOTAL", "Documentos": len(b), "Con % cesión": int(b["_con"].sum()),
             "Con observaciones": int(b["_obs"].sum()), f"% Cesión {PRIMER_ANIO} mín.": b["_pct"].min(),
             f"% Cesión {PRIMER_ANIO} promedio simple": b["_pct"].mean(),
             f"% Cesión {PRIMER_ANIO} máx.": b["_pct"].max()}
    return pd.concat([resumen, pd.DataFrame([total])], ignore_index=True)


def construir_cardinalidad(base, columnas_por_ln):
    """Una fila por carpeta y LN: cuántos documentos hay de cada nivel y cuál es el de la LN."""
    if not base.empty:
        numeros = [numero_ln(ln, carpeta) for ln, carpeta in zip(base["LN"], base["Carpeta LN"])]
    filas = []
    for (etiqueta, num_ln), col in columnas_por_ln.items():
        if base.empty:
            b = base
        else:
            elegidos = [c == etiqueta and n == num_ln for c, n in zip(base["Carpeta LN"], numeros)]
            b = base[pd.Series(elegidos, index=base.index) & base["Versión vigente"].eq("Sí")]
        conteo = {n: int(b["Nivel"].eq(n).sum()) for n in NIVELES} if not b.empty else dict.fromkeys(NIVELES, 0)
        presentes = [n for n in NIVELES if conteo[n]]
        if not presentes:
            cardinalidad = "Sin documentos"
        elif len(presentes) == 1:
            cardinalidad = presentes[0]
        else:
            cardinalidad = f"{presentes[-1]} (mixta: " + ", ".join(f"{n} {conteo[n]}" for n in presentes) + ")"
        filas.append({
            "Carpeta LN": etiqueta,
            "LN": ", ".join(sorted(b["LN"].dropna().astype(str).unique())) if not b.empty else "",
            "Nombre LN": ", ".join(sorted(b["Nombre LN"].dropna().astype(str).unique())) if not b.empty else "",
            "Documentos": len(b), **{f"Nivel {n}": conteo[n] for n in NIVELES},
            "Cardinalidad": cardinalidad,
            "Con % cesión": int(b["Cesión capturada"].eq("Sí").sum()) if not b.empty else 0,
            "Col. 1er año": col if len(b) else "-",
        })
    return pd.DataFrame(filas)


def construir_notas(grupos, ln_pedida, n_documentos, avisos_carpetas):
    filas = [
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("LN pedida", ln_pedida),
        *[(f"Carpeta leída ({etiqueta})", str(carpeta)) for etiqueta, carpeta in grupos],
        ("Avisos de carpetas", " | ".join(avisos_carpetas) or "Ninguno"),
        ("Documentos en la base", n_documentos),
        ("Hoja del logout", f"{HOJA_LOGOUT} (los archivos se reconocen por esta hoja, no por su nombre)"),
        ("Archivo", "Ruta del logout dentro de la carpeta leída"),
        ("Nivel", "Binder si la casilla MGA es Verdadero o trae Binder general / segmentado; si no, Contrato "
                  "si trae contrato; si no, Cedente"),
        ("Cardinalidad_LN", "Por carpeta y LN, cuántos documentos (versiones vigentes) hay de cada nivel; la "
                            "cardinalidad de la LN es el nivel más granular que aparece"),
        ("LN / TR / Cedente / Corredor / Contrato", "Línea de Negocio / Tipo Reas. / Compañía / Corredor "
                                                    "(primera aparición) / Contrato, col. B"),
        ("Es MGA / Binder general / Binder segmentado / GS",
         "Fila MGA: col. B (casilla MGA) / C (binder general del PRESUPUESTO, celda E18 = yBind) / "
         "E (binder segmentado) / F (casilla GS)"),
        ("Tipo Venta / % Renov.", "Tipo Venta, col. B / col. C"),
        ("Nombre LN, TR desc., Nombre / País / Grupo Cedente, Nombre Corredor",
         "Catálogo del PptoTécnico (nombres definidos " + ", ".join(c[0] for c in CATALOGOS.values())
         + ") que Excel guarda en caché dentro de cada logout; se usa el del propio logout"),
        ("% Tradicional ... % Retención", "Porcentaje, cols. B..E (tabla Tipo Retrocesión)"),
        ("% Cesión <año> / % Com. Cedido <año>",
         f"Porcentaje de Cesión / % Comisiones del Cedido, 5 columnas desde 'Col. 1er año' = "
         f"{ANIOS[0]}..{ANIOS[-1]}. Es el valor capturado por la LN; no está multiplicado por "
         "% Retro Espec. + % Fronting"),
        ("Col. 1er año", "Por LN, la mayoría de sus logouts: B si traen dato en la col. B (recuadro B..F de la "
                         "plantilla), C si lo traen solo desde C (herramienta que recorre los años, p. ej. "
                         "LN4003)" if MODO_COLUMNA == "AUTO" else f"Fijada en {MODO_COLUMNA}"),
        ("Años sin captura", "Se arrastró el último año capturado" if ARRASTRAR_ULTIMO_ANIO
         else "Se dejan vacíos, tal como vienen en el logout"),
        ("Versión archivo", "Sufijo -vN / vN del nombre del archivo, si lo trae"),
        ("Versión vigente", "Sí = la versión más alta de cada documento (misma LN/Nivel/TR/Cedente/Corredor/"
                            "Contrato/Binder/Venta); si empatan, la modificada más recientemente"),
        ("Resumen / Cardinalidad_LN", "Solo versiones vigentes; los % fuera de 0%..100% no entran a "
                                      "mín./promedio/máx."),
        ("Observaciones", "Resumen de la hoja Validaciones (niveles Error y Revisar)"),
    ]
    return pd.DataFrame(filas, columns=["Concepto", "Detalle"])


AZUL = PatternFill("solid", fgColor="1F4E78")
BLANCO_NEGRITA = Font(bold=True, color="FFFFFF")


def dar_formato(ws, df, anchos=None):
    for celda in ws[1]:
        celda.fill = AZUL
        celda.font = BLANCO_NEGRITA
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2" if "Archivo" in df.columns else "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    for i, nombre in enumerate(df.columns, 1):
        es_pct = str(nombre).startswith("%")
        if es_pct:
            for (celda,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                celda.number_format = "0.00%"
        muestra = [len(str(nombre)) * (0.6 if es_pct else 1)] + \
                  [len(str(v)) for v in df[nombre].head(200) if pd.notna(v)]
        ancho = (anchos or {}).get(nombre, min(max(muestra) + 2, 50))
        ws.column_dimensions[get_column_letter(i)].width = max(ancho, 9 if es_pct else 6)


def escribir_excel(salida, hojas):
    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        for nombre, df in hojas.items():
            df.to_excel(writer, sheet_name=nombre, index=False)
            dar_formato(writer.sheets[nombre], df,
                        anchos={"Archivo": 48, "Ruta": 80, "Detalle": 90, "Observaciones": 60})


# ------------------------------------------------------------------------------
# PRINCIPAL
# ------------------------------------------------------------------------------
def main():
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    parser = argparse.ArgumentParser(description="Base de % de cesión a partir de los logouts del PPTO.")
    parser.add_argument("carpeta", nargs="?", help="Carpeta con los logouts (opcional; si no, se usa LN)")
    parser.add_argument("--ln", help='LN a procesar, p. ej. LN4006 o TODAS (opcional; si no, se usa LN)')
    parser.add_argument("--salida", help="Carpeta donde guardar el Excel (opcional)")
    args = parser.parse_args()
    ln_pedida = args.ln or LN

    grupos, avisos_carpetas = ubicar_grupos(args.carpeta, ln_pedida)
    for aviso in avisos_carpetas:
        print(f"OJO: {aviso}")
    if not any(excels_en(c) for _, c in grupos):
        sys.exit("No hay archivos .xlsx en: " + ", ".join(str(c) for _, c in grupos)
                 + ("\n(Si los logouts están en .xlsb o .xls, ábrelos en Excel y guárdalos como .xlsx.)"
                    if any(excels_en(c, EXT_NO_SOPORTADAS) for _, c in grupos) else ""))

    # "Archivo" va relativo a la carpeta de la LN; con varias LN, relativo a la carpeta que las contiene
    padres = {c.parent for _, c in grupos}
    raiz_relativa = grupos[0][1] if len(grupos) == 1 else (padres.pop() if len(padres) == 1 else None)
    base, validaciones, columnas_por_ln = construir_base(grupos, raiz_relativa)

    if es_todas(ln_pedida):
        etiqueta = "TODAS"
    elif partes_ln(ln_pedida, solo_numero_ok=True) and not args.carpeta and not CARPETA_LOGOUTS:
        etiqueta = etiqueta_ln(partes_ln(ln_pedida, solo_numero_ok=True))
    elif len(grupos) == 1 and partes_ln(grupos[0][0]):
        etiqueta = grupos[0][0]
    else:
        lns = sorted(base["LN"].dropna().astype(str).unique()) if not base.empty else []
        etiqueta = lns[0] if len(lns) == 1 else ("varias_LN" if lns else "sin_LN")
    contenedora = grupos[0][1].parent if len(grupos) == 1 or raiz_relativa is None else raiz_relativa
    destino = Path(args.salida or CARPETA_SALIDA or contenedora).expanduser().resolve()
    salida = destino / f"Base_Cesion_{etiqueta}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    cardinalidad = construir_cardinalidad(base, columnas_por_ln)
    hojas = {
        "Base_Cesion": base,
        "Cesion_Anual": construir_anual(base),
        "Resumen": construir_resumen(base),
        "Cardinalidad_LN": cardinalidad,
        "Validaciones": validaciones,
        "Notas": construir_notas(grupos, ln_pedida, len(base), avisos_carpetas),
    }
    try:
        destino.mkdir(parents=True, exist_ok=True)
        escribir_excel(salida, hojas)
    except OSError as error:
        sys.exit(f"No se pudo guardar {salida}:\n  {error}\n"
                 f"Revisa que puedas escribir en {destino} o indica otra carpeta con --salida o CARPETA_SALIDA.")

    niveles = validaciones["Nivel"].value_counts() if not validaciones.empty else {}

    def nombre_fila(fila):  # "LN4003 · LN04006" si en la carpeta de una LN hay logouts de otra
        propia, del_logout = partes_ln(fila["Carpeta LN"]), partes_ln(fila["LN"])
        otra = propia and del_logout and propia[0] != del_logout[0]
        return f"{fila['Carpeta LN']} · {fila['LN']}" if otra else str(fila["Carpeta LN"])

    nombres = [nombre_fila(f) for _, f in cardinalidad.iterrows()]
    ancho = max([14] + [len(n) + 2 for n in nombres])
    print("\n" + "=" * 78)
    print(f"{'LN':<{ancho}}{'Docs':>6}{'Con cesión':>12}   {'Col. 1er año':<14}Cardinalidad")
    for nombre, (_, fila) in zip(nombres, cardinalidad.iterrows()):
        print(f"{nombre:<{ancho}}{fila['Documentos']:>6}{fila['Con % cesión']:>12}   "
              f"{fila['Col. 1er año']:<14}{fila['Cardinalidad']}")
    print("-" * 78)
    print("Validaciones .......... " + ", ".join(f"{k}: {niveles.get(k, 0)}" for k in ("Error", "Revisar", "Info")))
    if not validaciones.empty:
        for tipo in ("Columna del primer año mixta", "Años recorridos una columna", f"Ningún logout trae {ANIOS[-1]}"):
            for archivo in validaciones.loc[validaciones["Tipo"].eq(tipo), "Archivo"]:
                print(f"OJO {archivo}: {tipo.lower()}; ver hoja Validaciones.")
    print(f"Archivo generado ...... {salida}")
    print("=" * 78)

    if ABRIR_AL_TERMINAR:
        try:
            if os.name == "nt":
                os.startfile(salida)  # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                subprocess.Popen(["open", str(salida)])
        except Exception:
            pass


if __name__ == "__main__":
    main()
