# -*- coding: utf-8 -*-
"""Agrega al libro de Cesion la validacion de 05_FCST_2027.xlsx contra ER y sus vistas.

Seis hojas con el formato de las Val_* (mismos estilos):
  Val05_Resumen  renglon por renglon, a diciembre, cada hoja del 05 contra su referencia
  Val05_Mensual  PptoxMes_Red contra ER en MXN, mes a mes
  Val05_LN / Val05_Ramo / Val05_Region / Val05_TRea
                 cada bloque del 05 contra la vista evaluada para sus miembros (hojas Val_*),
                 en MXN; suma de bloques contra ER; 2025 contra la columna D de la vista;
                 Val05_LN trae ademas PptoxLN_Red (anual por LN)

Las cifras del 05 van capturadas (es un archivo externo); todo lo demas es formula:
la vista en MXN = importe del mes en Val_x × 'Parámetros'!F3:Q3 ÷ factor de la vista
(ER_x!E6:P6), acumulado. Con xMonEEFF en USD el factor es 1; en MXN es el mismo tipo de
cambio y la division da 1, asi que la comparacion vale en las dos monedas.

Cada formula se arma junto con su valor (se guardan ya calculadas, como las Val_*)."""
import pickle, struct, zipfile, re, os, shutil, subprocess, tempfile, math, sys
import openpyxl
import xlsbw as W
from biff import records, sheet_map, colrow
from reparar import records_raw
from biff import rd_xlwstr
def W_rd(pl): return rd_xlwstr(pl, 1)[0]

SRC, DST, F05 = 'cesion/FCST.xlsb', os.environ.get('DST_XLSB', 'FCST_2027_Cesion.xlsb'), '05_FCST_2027.xlsx'
CEL = pickle.load(open('cesion/celdas_v3.pkl', 'rb'))
DAT = pickle.load(open('cesion/datos05.pkl', 'rb'))
LL = pickle.load(open('llenado05.pkl', 'rb'))
W05 = openpyxl.load_workbook(F05, data_only=True)
NM = 12; MESES = list(range(202701, 202713))
NOMMES = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
ST = dict(TIT=1006, SUB=1007, NOTA=1008, BAR=1009, BAR_C=1010, HDR=1011, HDR_C=1012, TXT=1013, OBS=1014, NUM=1015,
          NUM_B=1016, NUM_G=1017, NUM_Y=1018, LBL_B=1019, LBL_G=1020, LBL_Y=1021, BOOL=1022, MES=1023, TXT_B=1024)
IX = dict(Param=4, ER_ln=6, ER_reg=7, ER_tre=8, ER_ram=28, ER1225_Real=32, Anexo2_1225=33, ER=56,
          Val_Ramo=63, Val_LN=64, Val_Region=65, Val_TRea=66)
NUEVAS = ['Val05_Resumen', 'Val05_Mensual', 'Val05_LN', 'Val05_Ramo', 'Val05_Region', 'Val05_TRea']
ITAB0 = 34                                                     # despues de Val_CAT (33)
for i, h in enumerate(NUEVAS[1:]): IX[h] = 69 + i              # ixti 69..73
XTI_EXTRA = [(0, ITAB0 + 1 + i, ITAB0 + 1 + i) for i in range(5)]
HOJA_IX = {'Parámetros': 'Param', 'ER1225_Real': 'ER1225_Real', 'Anexo2_1225': 'Anexo2_1225'}

C1, C2, C3, QC = 2, 15, 28, 40            # bloques de meses C..N, P..AA, AC..AN; AO = ¿Cuadra?
Y1, Y2, Y3, Y4 = 42, 43, 44, 45           # 2025: AQ del 05, AR referencia, AS diferencia, AT ¿Cuadra?
TOL = 0.005

# ============================================================ expresiones (rgce + valor)
class X:
    """rgce + valor + precedencia del operador raiz (9 atomo, 3 * /, 2 + -, 1 comparacion).
    Excel arma el texto de la formula desde los tokens y solo pone parentesis donde hay PtgParen,
    asi que se agregan donde la precedencia lo exige para que el texto diga lo mismo que se calcula."""
    __slots__ = ('g', 'v', 'p')
    def __init__(s, g, v, p=9): s.g = g; s.v = v; s.p = p
def lift(o): return o if isinstance(o, X) else X(W.num(o), float(o))
PREC = {W.ADD: 2, W.SUB: 2, W.MUL: 3, W.DIV: 3}
def _op(a, b, tok, fn):
    a, b = lift(a), lift(b); p = PREC[tok]
    ga = a.g + (W.PAREN if a.p < p else b'')
    gb = b.g + (W.PAREN if (b.p < p or (b.p == p and tok != W.ADD)) else b'')
    v = None if (a.v is None or b.v is None) else fn(a.v, b.v)
    return X(ga + gb + tok, v, p)
def add(a, b): return _op(a, b, W.ADD, lambda x, y: x + y)
def sub(a, b): return _op(a, b, W.SUB, lambda x, y: x - y)
def mul(a, b): return _op(a, b, W.MUL, lambda x, y: x * y)
def div(a, b): return _op(a, b, W.DIV, lambda x, y: x / y)
def sumar(xs):
    xs = list(xs); out = xs[0]
    for x in xs[1:]: out = add(out, x)
    return out
def cval(hoja, r, c):
    v = CEL[hoja].get((r, c), (None, None))[0]
    if isinstance(v, bool): return float(v)
    return float(v) if isinstance(v, (int, float)) else 0.0
def R3(hoja, r, c):
    return X(W.ref3d(IX[HOJA_IX.get(hoja, hoja)], r, c), cval(hoja, r, c))

class Hoja:
    def __init__(s, nombre): s.nombre = nombre; s.rows = {}; s.val = {}
    def _c(s, r, c, b, v): s.rows.setdefault(r, []).append(b); s.val[(r, c)] = v
    def t(s, r, c, txt, st): s._c(r, c, W.c_isst(c, P.add(txt), ST[st]), None)
    def n(s, r, c, v, st):
        if v is None: s._c(r, c, W.c_blank(c, ST[st]), None)
        else: s._c(r, c, W.c_num(c, v, ST[st]), float(v))
    def b(s, r, c, st): s._c(r, c, W.c_blank(c, ST[st]), None)
    def f(s, r, c, x, st):
        assert x.v is not None and math.isfinite(x.v), (s.nombre, r, c)
        v = 0.0 if x.v == 0 else x.v
        s._c(r, c, W.c_fmla_num(c, v, x.g, ST[st]), v)
    def fb(s, r, c, x, st):
        s._c(r, c, W.c_fmla_bool(c, bool(x.v), x.g, ST[st]), bool(x.v))
    def L(s, r, c):
        v = s.val.get((r, c)); v = 0.0 if v is None else v
        return X(W.ref(r, c, crel=True, rrel=True), v)
    def barra(s, r, txt, ultima, txt2=None):
        s.t(r, 0, txt, 'BAR')
        if txt2 is not None: s.t(r, 1, txt2, 'BAR')
        for c in range(1 if txt2 is None else 2, ultima + 1): s.b(r, c, 'BAR')
    def cuadra(s, r, c1, c2):
        vals = [s.val.get((r, c)) or 0.0 for c in range(c1, c2 + 1)]
        a = W.area(r, r, c1, c2)
        g = a + W.funcvar(1, 7) + W.func_fix(24) + a + W.funcvar(1, 6) + W.func_fix(24) + W.ADD + W.num(TOL) + W.LT
        return X(g, abs(max(vals)) + abs(min(vals)) < TOL, 1)
    def abs_lt(s, x): return X(x.g + W.func_fix(24) + W.num(TOL) + W.LT, abs(x.v) < TOL, 1)
    def suma_col(s, r1, r2, c):
        vals = [s.val.get((r, c)) or 0.0 for r in range(r1, r2 + 1)]
        return X(W.area(r1, r2, c, c) + W.funcvar(1, W.IF_SUM), sum(vals))

def AND_areas(areas_vals):
    g = b''.join(a for a, _ in areas_vals); vals = [v for _, vs in areas_vals for v in vs if v is not None]
    return X(g + W.funcvar(len(areas_vals), 36), all(vals))

def conv(ix_fact_hoja, j, fact_r, fact_c0):
    """'Parámetros'!{F+j}$3 / factor actual (vista E6:P6 o ER F3:Q3): tipo de cambio efectivo del mes."""
    return div(R3('Parámetros', 2, 5 + j), R3(ix_fact_hoja, fact_r, fact_c0 + j))

def num05(ws, r, c):
    v = ws.cell(r, c).value
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None

P = None

# ============================================================ Val05_Mensual
ER_FILA = LL['ER_FILA']
DIRECTAS = {60, 61, 64, 65, 66, 68}          # filas de ER que no pasan por el tipo de cambio
HOJA_IX['RIF'] = 'RIF'; IX['RIF'] = 10
RIF_ANUAL = {62: 232, 63: 233, 67: 236}      # ER!Q62/Q63/Q67 circulares -> anual de RIF!P233, P234, P237
def _m(h, f05, j): return h.L(MROW[f05], C2 + j)
def _max0(x): return X(x.g + W.num(0) + W.funcvar(2, 7), max(x.v, 0.0))
def _base(h, j): return _max0(sub(_m(h, 55, j), _m(h, 50, j)))
CALC_05 = {    # mismas formulas que ER (filas 70-81) sobre las cifras en MXN de esta hoja
    70: lambda h, j: X(W.area(MROW[44], MROW[52], C2 + j, C2 + j) + W.funcvar(1, W.IF_SUM),
                       sum(h.val.get((rr, C2 + j)) or 0.0 for rr in range(MROW[44], MROW[52] + 1))),
    71: lambda h, j: add(_m(h, 41, j), _m(h, 53, j)),
    73: lambda h, j: mul(_base(h, j), 0.3),
    74: lambda h, j: (lambda b, t: X(b.g + W.num(0.1) + W.MUL + t.g + W.funcvar(2, 6), min(b.v*0.1, t.v)))(_base(h, j), R3('Parámetros', 13, 16)),
    75: lambda h, j: mul(_m(h, 50, j), 0.28),
    76: lambda h, j: mul(_m(h, 50, j), 0.02),
    77: lambda h, j: X(W.area(MROW[58], MROW[61], C2 + j, C2 + j) + W.funcvar(1, W.IF_SUM),
                       sum(h.val.get((rr, C2 + j)) or 0.0 for rr in range(MROW[58], MROW[61] + 1))),
    78: lambda h, j: sub(_m(h, 55, j), _m(h, 62, j)),
    81: lambda h, j: add(_m(h, 64, j), _m(h, 66, j)),
}
DESC_SUP = {62: "ER fila 62 · anual de RIF!P233 ÷ 12 acumulado (ER!Q62 es circular)",
            63: "ER fila 63 · anual de RIF!P234 ÷ 12 acumulado (ER!Q63 es circular)",
            67: "ER fila 67 · anual de RIF!P237 ÷ 12 acumulado (ER!Q67 es circular)",
            70: "ER fila 70 · SUM de las filas 44 a 52 del 05 (filas de arriba en esta hoja; ER suma F60:F69, con la fila 65 Recargos = 0)", 71: "ER fila 71 · fila 41 + fila 53",
            73: "ER fila 73 · MAX(fila 55 − fila 50; 0) × 0.3", 74: "ER fila 74 · MIN(MAX(fila 55 − fila 50; 0) × 0.1; 'Parámetros'!Q14)",
            75: "ER fila 75 · fila 50 × 0.28", 76: "ER fila 76 · fila 50 × 0.02", 77: "ER fila 77 · SUM de las filas 58 a 61",
            78: "ER fila 78 · fila 55 − fila 62", 81: "ER fila 81 · fila 64 + fila 66 (ER fila 80 vacía)"}
MROW = {}

def hoja_mensual():
    h = Hoja('Val05_Mensual'); ws = W05['PptoxMes_Red']
    h.t(0, 0, "05_FCST_2027 · PptoxMes_Red CONTRA ER — PRESUPUESTO 2027 MENSUALIZADO (MXN, acumulado por mes)", 'TIT')
    h.t(1, 0, "Cada cifra del 05 (capturada tal como quedó en el archivo) contra el renglón equivalente de ER convertido a pesos: "
              "importe del mes en ER × 'Parámetros'!F3:Q3 ÷ ER!F3:Q3, acumulado. Es lo que da ER con xMonEEFF = \"MXN\".", 'SUB')
    h.t(2, 0, "Las filas financieras que en ER no pasan por el tipo de cambio (intereses sobre depósitos, dividendos, cambios, etc.) se toman tal cual. "
              "ER!Q62, Q63 y Q67 son circulares: esas filas usan el anual de RIF!P233/P234/P237 ÷ 12 y las de abajo (RIF, impuestos, utilidad) las fórmulas de ER sobre esta hoja.", 'SUB')
    h.t(3, 0, "Mes (CALMONTH)", 'TXT_B')
    for j in range(NM):
        for c0 in (C1, C2, C3): h.n(3, c0 + j, MESES[j], 'MES')
    h.t(3, Y1, "Dic 2025", 'HDR_C')
    r = 4
    h.t(r, 0, "Tipo de cambio MXN ('Parámetros'!F3:Q3)", 'LBL_G'); h.t(r, 1, "Lo que usa ER con xMonEEFF = MXN", 'OBS')
    for j in range(NM): h.f(r, C2 + j, R3('Parámetros', 2, 5 + j), 'NUM_G')
    r += 1
    h.t(r, 0, "Factor actual de ER (ER!F3:Q3)", 'LBL_G'); h.t(r, 1, "1.00 si xMonEEFF = USD", 'OBS')
    for j in range(NM): h.f(r, C2 + j, R3('ER', 2, 5 + j), 'NUM_G')
    r += 2
    h.barra(r, "PptoxMes_Red (bloque izquierdo, F:Q = 2027)  ↔  ER en MXN", Y4); r += 1
    h.t(r, 0, "Renglón del 05", 'HDR'); h.t(r, 1, "Renglón equivalente de ER", 'HDR')
    for j in range(NM):
        for c0 in (C1, C2, C3): h.t(r, c0 + j, NOMMES[j], 'HDR_C')
    h.t(r, QC, "¿Cuadra?", 'HDR_C'); h.t(r, Y1, "2025 en el 05 (W)", 'HDR_C'); h.t(r, Y2, "ER!E (Dic 2025)", 'HDR_C')
    h.t(r, Y3, "Diferencia", 'HDR_C'); h.t(r, Y4, "¿Cuadra?", 'HDR_C'); r += 1
    h.t(r, C1, "05_FCST_2027 (capturado)", 'TXT_B'); h.t(r, C2, "ER en MXN (fórmula)", 'TXT_B'); h.t(r, C3, "Diferencia (05 − ER)", 'TXT_B')
    r += 1
    malos = []
    for f, filas in ER_FILA.items():
        MROW[f] = r
        etq = ws.cell(f, 4).value or ws.cell(f, 3).value
        h.t(r, 0, f"fila {f} · {str(etq).strip()}", 'TXT')
        h.t(r, 1, DESC_SUP.get(filas[0]) or ' − '.join(f"ER fila {abs(x)}" for x in filas), 'OBS')
        v05 = [num05(ws, f, 6 + j) for j in range(NM)]
        definida = all(DAT['ER_MXN'].get(abs(x)) is not None for x in filas)
        for j in range(NM): h.n(r, C1 + j, v05[j], 'NUM')
        if definida:
            for j in range(NM):
                er_j = sumar([R3('ER', abs(x) - 1, 5 + j) if x > 0 else mul(R3('ER', abs(x) - 1, 5 + j), -1) for x in filas]) \
                    if len(filas) == 1 else sub(R3('ER', abs(filas[0]) - 1, 5 + j), R3('ER', abs(filas[1]) - 1, 5 + j))
                if filas[0] in RIF_ANUAL:                               # anual de RIF / 12, acumulado
                    paso = div(R3('RIF', RIF_ANUAL[filas[0]], 15), 12)
                    x = paso if j == 0 else add(h.L(r, C2 + j - 1), paso)
                elif filas[0] in CALC_05:
                    x = CALC_05[filas[0]](h, j)
                elif filas[0] in DIRECTAS:
                    x = er_j
                elif filas[0] == 79:                                   # = prima retenida en MXN × 'Parámetros'!F11:Q11
                    x = mul(h.L(MROW[10], C2 + j), R3('Parámetros', 10, 5 + j))
                else:
                    if j == 0:
                        x = mul(er_j, conv('ER', 0, 2, 5))
                    else:
                        ant = sumar([R3('ER', abs(filas[0]) - 1, 4 + j)]) if len(filas) == 1 else \
                            sub(R3('ER', abs(filas[0]) - 1, 4 + j), R3('ER', abs(filas[1]) - 1, 4 + j))
                        x = add(h.L(r, C2 + j - 1), mul(sub(er_j, ant), conv('ER', j, 2, 5)))
                h.f(r, C2 + j, x, 'NUM_G')
                h.f(r, C3 + j, sub(h.L(r, C1 + j), h.L(r, C2 + j)), 'NUM_Y')
                if v05[j] is None or abs(v05[j] - x.v) > TOL: malos.append((f, j, v05[j], x.v))
            h.fb(r, QC, h.cuadra(r, C3, C3 + NM - 1), 'BOOL')
        else:
            h.t(r, C2, "No se llenó: en ER esta fila depende de ER!Q62, Q63 o Q67, que tienen referencia circular (=$Q63/12+P63)", 'OBS')
        # 2025: bloque derecho del 05 (W) contra ER!E
        w25 = num05(ws, f, 23)
        h.n(r, Y1, w25, 'NUM')
        e25 = R3('ER', abs(filas[0]) - 1, 4) if len(filas) == 1 else sub(R3('ER', filas[0] - 1, 4), R3('ER', abs(filas[1]) - 1, 4))
        h.f(r, Y2, e25, 'NUM_G')
        h.f(r, Y3, sub(h.L(r, Y1), h.L(r, Y2)), 'NUM_Y')
        h.fb(r, Y4, h.abs_lt(h.L(r, Y3)), 'BOOL')
        r += 1
    assert not malos, malos[:5]
    r += 1
    h.t(r, 0, "El bloque derecho del 05 (X:AI) trae el presupuesto 2026 que ya tenía el archivo, recorrido una posición; el libro de Cesión no tiene esas cifras, por eso no se comparan aquí.", 'NOTA')
    cols = [(0, 0, 44), (1, 1, 34)] + [(C1 + j, C1 + j, 15) for j in range(NM)] + [(14, 14, 1.5)] + \
           [(C2 + j, C2 + j, 15) for j in range(NM)] + [(27, 27, 1.5)] + [(C3 + j, C3 + j, 12) for j in range(NM)] + \
           [(QC, QC, 10), (41, 41, 1.5), (Y1, Y2, 17), (Y3, Y3, 12), (Y4, Y4, 10)]
    return h, W.sheet_bin(h.rows, Y4 + 1, cols=cols, freeze=(2, 4))

# ============================================================ hojas por dimension
FILAS = [8, 9, 10, 11, 12, 15, 16, 17, 18, 19, 22, 23, 24, 25, 26, 27, 30, 31, 32, 34]
NOMBRE = {9: 'Prima Tomada', 10: 'Prima Retrocedida', 11: 'Prima Retenida', 12: 'Variación Rva. de Primas s/Prima Tomada',
          13: 'Variación Rva. de Primas s/Prima Retenida', 16: 'Prima Devengada', 17: 'Siniestros Ocurridos', 18: 'I.B.N.R.',
          19: 'Costos de Adquisición', 20: 'RESULTADO TÉCNICO TOTAL', 23: 'Prima Neta Devengada', 24: 'Siniestros Netos Ocurridos',
          25: 'I.B.N.R. (neto)', 26: 'Costos de Adquisición (neto)', 27: 'Costos Protecciones XL', 28: 'RESULTADO TÉCNICO A RETENCIÓN',
          31: 'Gastos Generales', 32: '(-) Contribución a los Gastos por Retrocesión', 33: 'TOTAL COSTOS DE OPERACIÓN', 35: 'RESULTADO DE OPERACIÓN'}
DIMINFO = {
 'LN':     dict(hoja05='PptoxLNxMes_Red', val='Val_LN',     vista='ER_ln',  nom='LÍNEA DE NEGOCIO', hoja='Val05_LN'),
 'RAMO':   dict(hoja05='PptoxRamo_Red',   val='Val_Ramo',   vista='ER_ram', nom='RAMO',             hoja='Val05_Ramo'),
 'REGION': dict(hoja05='PptoxRegión_Red', val='Val_Region', vista='ER_reg', nom='REGIÓN',           hoja='Val05_Region'),
 'TREA':   dict(hoja05='PptoxTR_Red',     val='Val_TRea',   vista='ER_tre', nom='TIPO DE REASEGURO', hoja='Val05_TRea'),
}
TIPO = {'LN': 'texto', 'RAMO': 'texto', 'REGION': 'texto', 'TREA': 'numero'}

def pos_val(hoja):
    """{(linea, miembro): fila0} de la hoja Val_x y columna de enero."""
    V = CEL[hoja]; out = {}; ln = None
    c0 = next(c for (r, c), (v, f) in V.items() if r == 3 and v == 202701.0)
    for r in sorted({r for r, c in V}):
        a = V.get((r, 0), (None, None))[0]
        if isinstance(a, str):
            m = re.match(r'^(\d+) · ', a)
            if m: ln = int(m.group(1)); continue
            if a in ('Miembro', 'Suma de los miembros', 'ER (global)', 'Diferencia (suma − ER)'): continue
        if ln is None or a is None: continue
        out[(ln, a)] = r
    return out, c0

# ---- traduccion de la formula de la columna D de la vista (real 2025) ----
def tokens(s):
    return re.findall(r"HLOOKUP\([^()]*\)|SUM\([A-Z]+\d+:[A-Z]+\d+\)|[A-Z]+\d+|\d+\.?\d*|[-+*/()]", s)

def formula_d(dim, texto, k, mbs, h, LAY, V25):
    """X de la columna D de la vista para el bloque k (suma de sus miembros) o None si no aplica."""
    toks = tokens(texto.lstrip('='))
    pos = [0]
    def peek(): return toks[pos[0]] if pos[0] < len(toks) else None
    def take(): t = toks[pos[0]]; pos[0] += 1; return t
    def local(n):                                 # D{n}: fila n de la vista -> fila n-1 del 05
        f = n - 1
        if (f, k) in LAY: return h_ref(LAY[(f, k)], Y2, V25.get((n, k)))
        return lift(0.0)
    def h_ref(r, c, v): return X(W.ref(r, c, crel=True, rrel=True), 0.0 if v is None else v)
    def hlook(t):
        m = re.match(r"HLOOKUP\(\$A\$6,'?([^!',]+)'?!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+),\$A(\d+),0\)", t)
        hoja, ca, ra, cb, rb, an = m.groups()
        c_a, c_b = ci(ca) - 1, ci(cb) - 1; r_a, r_b = int(ra) - 1, int(rb) - 1
        idx = CEL[DIMINFO[dim]['vista']].get((int(an) - 1, 0), (None,))[0]
        partes = []
        for mb in mbs:
            col = None
            for c in range(c_a, c_b + 1):
                kv = CEL[hoja].get((r_a, c), (None,))[0]
                if kv is None: continue
                kv = float(kv) if TIPO[dim] == 'numero' else str(kv)
                if kv == mb: col = c; break
            if col is None: continue                          # la vista daria #N/A para ese miembro
            v = cval(hoja, r_a + int(idx) - 1, col)
            look = W.num(mb) if isinstance(mb, float) else W.txt(mb)
            g = look + W.area3d(IX[HOJA_IX[hoja]], r_a, r_b, c_a, c_b) + W.num(idx) + W.num(0) + W.funcvar(4, 101)
            partes.append(X(g, v))
        return sumar(partes) if partes else None
    def factor():
        t = take()
        if t == '+': return factor()
        if t == '-':
            x = factor()
            return None if x is None else X(x.g + (W.PAREN if x.p < 9 else b'') + W.UMINUS, -x.v)
        if t == '(':
            x = expr(); take(); return None if x is None else X(x.g + W.PAREN, x.v)
        if t.startswith('HLOOKUP'): return hlook(t)
        if t.startswith('SUM('):
            m = re.match(r'SUM\([A-Z]+(\d+):[A-Z]+(\d+)\)', t); a, b = int(m.group(1)), int(m.group(2))
            xs = [local(n) for n in range(a, b + 1)]
            return X(sumar(xs).g + W.PAREN, sum(x.v for x in xs))
        if re.match(r'^[A-Z]+\d+$', t): return local(int(re.match(r'[A-Z]+(\d+)', t).group(1)))
        return lift(float(t))
    def term():
        x = factor()
        while peek() in ('*', '/'):
            op = take(); y = factor()
            if x is None or y is None: x = None; continue
            x = mul(x, y) if op == '*' else div(x, y)
        return x
    def expr():
        x = term()
        while peek() in ('+', '-'):
            op = take(); y = term()
            if x is None or y is None: x = None; continue
            x = add(x, y) if op == '+' else sub(x, y)
        return x
    return expr()

def ci(letras):
    n = 0
    for ch in letras: n = n * 26 + ord(ch) - 64
    return n

ANCLA = {}      # dim -> dict(SUMR, DIFR, first, last por concepto) para el resumen

def hoja_dim(dim):
    info = DIMINFO[dim]; h = Hoja(info['hoja']); ws = W05[info['hoja05']]
    bloques = LL['BLOQUES'][info['hoja05']]; vpos, c0v = pos_val(info['val']); vista = info['vista']
    ixv = IX[vista]; ixval = IX[info['val']]
    h.t(0, 0, f"05_FCST_2027 · {info['hoja05']} CONTRA LA VISTA {vista.upper()} — PRESUPUESTO 2027 POR {info['nom']} (MXN, acumulado por mes)", 'TIT')
    h.t(1, 0, f"Cada bloque del 05 (capturado tal como quedó en el archivo) contra la vista {vista} evaluada para los miembros de ese bloque "
              f"(hoja {info['val']}) y convertida a pesos: importe del mes × 'Parámetros'!F3:Q3 ÷ {vista}!E6:P6, acumulado. Es lo que da la vista con xMonEEFF = \"MXN\".", 'SUB')
    h.t(2, 0, "Abajo de cada renglón, la suma de los bloques del 05 contra ER en MXN (hoja Val05_Mensual). A la derecha, el 2025 que quedó en el 05 contra la columna D de la vista (real 2025).", 'SUB')
    h.t(3, 0, "Mes (CALMONTH)", 'TXT_B')
    for j in range(NM):
        for cc in (C1, C2, C3): h.n(3, cc + j, MESES[j], 'MES')
    h.t(3, Y1, "Dic 2025", 'HDR_C')
    # ---- layout previo
    LAY = {}; SUMR = {}; ERR = {}; DIFR = {}; BAR = {}; r = 5
    for f in FILAS:
        BAR[f] = r; r += 3
        for k in range(len(bloques)): LAY[(f, k)] = r; r += 1
        SUMR[f], ERR[f], DIFR[f] = r, r + 1, r + 2; r += 4
    # 2025 por bloque y renglon de la vista (suma de miembros con dato)
    V25 = {}
    for k, (_, _, _, mbs) in enumerate(bloques):
        for ln in list(NOMBRE) + [21, 22, 29, 30, 34]:
            vs = [DAT['VISTAS'][dim]['real25'][mb].get(ln) for mb in mbs]
            vs = [v for v in vs if v is not None]
            V25[(ln, k)] = sum(vs) if vs else None
    malos = []
    for f in FILAS:
        ln = f + 1; r = BAR[f]
        h.barra(r, f"{ln} · {NOMBRE[ln]}", Y4, f"05: fila {f} de cada bloque de {info['hoja05']}   ↔   vista {vista}, fila {ln}")
        r += 1
        h.t(r, 0, "Bloque del 05", 'HDR'); h.t(r, 1, "Miembros de la vista", 'HDR')
        for j in range(NM):
            for cc in (C1, C2, C3): h.t(r, cc + j, NOMMES[j], 'HDR_C')
        h.t(r, QC, "¿Cuadra?", 'HDR_C'); h.t(r, Y1, "2025 en el 05 (E)", 'HDR_C'); h.t(r, Y2, f"{vista} col. D", 'HDR_C')
        h.t(r, Y3, "Diferencia", 'HDR_C'); h.t(r, Y4, "¿Cuadra?", 'HDR_C')
        r += 1
        h.t(r, C1, "05_FCST_2027 (capturado)", 'TXT_B'); h.t(r, C2, f"Vista {vista} en MXN (fórmula)", 'TXT_B')
        h.t(r, C3, "Diferencia (05 − vista)", 'TXT_B')
        dtexto = CEL[vista].get((ln - 1, 3), (None, None))[1]
        for k, (r0, nombre, _, mbs) in enumerate(bloques):
            r = LAY[(f, k)]; f05 = f + (r0 - 6)
            h.t(r, 0, nombre, 'TXT')
            h.t(r, 1, ' + '.join(str(int(m)) if isinstance(m, float) else m for m in mbs) + f"   (05 fila {f05})", 'OBS')
            v05 = [num05(ws, f05, 7 + j) for j in range(NM)]
            for j in range(NM): h.n(r, C1 + j, v05[j], 'NUM')
            filas_val = [vpos[(ln, mb)] for mb in mbs if (ln, mb) in vpos]
            assert filas_val, (dim, ln, mbs)
            for j in range(NM):
                if ln == 32:
                    x = mul(h.L(LAY[(9, k)], C2 + j), R3('Parámetros', 9, 5 + j))
                else:
                    usd = sumar([X(W.ref3d(ixval, fv, c0v + j), cval(info['val'], fv, c0v + j)) for fv in filas_val])
                    cf = div(R3('Parámetros', 2, 5 + j), X(W.ref3d(ixv, 5, 4 + j), cval(vista, 5, 4 + j)))
                    if j == 0:
                        x = mul(usd, cf)
                    else:
                        ant = sumar([X(W.ref3d(ixval, fv, c0v + j - 1), cval(info['val'], fv, c0v + j - 1)) for fv in filas_val])
                        x = add(h.L(r, C2 + j - 1), mul(sub(usd, ant), cf))
                h.f(r, C2 + j, x, 'NUM_G')
                h.f(r, C3 + j, sub(h.L(r, C1 + j), h.L(r, C2 + j)), 'NUM_Y')
                if v05[j] is None or abs(v05[j] - x.v) > TOL: malos.append((dim, nombre, ln, j, v05[j], x.v))
            h.fb(r, QC, h.cuadra(r, C3, C3 + NM - 1), 'BOOL')
            # 2025
            e25 = num05(ws, f05, 5)
            xd = formula_d(dim, dtexto, k, mbs, h, LAY, V25) if dtexto else None
            if xd is None and e25 is None:
                for c in (Y1, Y2, Y3, Y4): h.b(r, c, 'NUM')
            else:
                h.n(r, Y1, e25, 'NUM')
                if xd is None: xd = lift(0.0)
                assert abs((xd.v or 0) - (V25.get((ln, k)) or 0)) < TOL, (dim, ln, nombre, xd.v, V25.get((ln, k)))
                h.f(r, Y2, xd, 'NUM_G'); h.f(r, Y3, sub(h.L(r, Y1), h.L(r, Y2)), 'NUM_Y')
                h.fb(r, Y4, h.abs_lt(h.L(r, Y3)), 'BOOL')
        # suma de bloques, ER, diferencia
        r1, r2 = LAY[(f, 0)], LAY[(f, len(bloques) - 1)]
        r = SUMR[f]
        h.t(r, 0, "Suma de los bloques", 'LBL_B'); h.t(r, 1, f"= SUM de los {len(bloques)} bloques de arriba", 'LBL_B')
        for j in range(NM):
            h.f(r, C1 + j, h.suma_col(r1, r2, C1 + j), 'NUM_B'); h.f(r, C2 + j, h.suma_col(r1, r2, C2 + j), 'NUM_B')
            h.f(r, C3 + j, sub(h.L(r, C1 + j), h.L(r, C2 + j)), 'NUM_Y')
        h.fb(r, QC, h.cuadra(r, C3, C3 + NM - 1), 'BOOL')
        h.f(r, Y1, h.suma_col(r1, r2, Y1), 'NUM_B'); h.f(r, Y2, h.suma_col(r1, r2, Y2), 'NUM_B')
        h.f(r, Y3, sub(h.L(r, Y1), h.L(r, Y2)), 'NUM_Y'); h.fb(r, Y4, h.abs_lt(h.L(r, Y3)), 'BOOL')
        r = ERR[f]
        h.t(r, 0, "ER (global) en MXN", 'LBL_G'); h.t(r, 1, f"= Val05_Mensual, fila del 05 {f}  ·  2025: ER!E (Dic 2025)", 'LBL_G')
        mr = MROW[f]
        for j in range(NM):
            h.f(r, C1 + j, X(W.ref3d(IX['Val05_Mensual'], mr, C2 + j), MENS.val[(mr, C2 + j)]), 'NUM_G')
        h.f(r, Y1, X(W.ref3d(IX['Val05_Mensual'], mr, Y2), MENS.val[(mr, Y2)]), 'NUM_G')
        r = DIFR[f]
        h.t(r, 0, "Diferencia (suma del 05 − ER)", 'LBL_Y')
        h.t(r, 1, "2027 (AC:AO): cero, los bloques del 05 suman el total de ER. 2025 (AS:AT): ver Val05_Resumen, sección 2025.", 'LBL_Y')
        for j in range(NM): h.f(r, C3 + j, sub(h.L(SUMR[f], C1 + j), h.L(ERR[f], C1 + j)), 'NUM_Y')
        h.fb(r, QC, h.cuadra(r, C3, C3 + NM - 1), 'BOOL')
        h.f(r, Y3, sub(h.L(SUMR[f], Y1), h.L(ERR[f], Y1)), 'NUM_Y'); h.fb(r, Y4, h.abs_lt(h.L(r, Y3)), 'BOOL')
    assert not malos, malos[:5]
    fin = DIFR[FILAS[-1]] + 2
    ANCLA[dim] = dict(SUMR=SUMR, ERR=ERR, DIFR=DIFR, LAY=LAY, nb=len(bloques), fin=fin)
    return h, fin

def seccion_anual_ln(h, r):
    """PptoxLN_Red: la columna 2027 de cada LN contra diciembre de la vista (Val05_LN) y el total contra ER."""
    ws = W05['PptoxLN_Red']; LAY = ANCLA['LN']['LAY']; bloques = LL['BLOQUES']['PptoxLNxMes_Red']
    grupos = LL['GRUPOS']                                     # [(col0, nombre, miembros)]
    kde = {tuple(mbs): k for k, (_, _, _, mbs) in enumerate(bloques)}
    ng = len(grupos); A0, B0, D0 = 2, 2 + ng + 1, 2 + 2 * (ng + 1)        # C.., L.., U..
    T0 = D0 + ng + 1                                                      # AD..
    h.barra(r, "PptoxLN_Red (ANUAL POR LN): la columna 2027 de cada LN contra diciembre de la vista; el total (columna E) contra ER", T0 + 3); r += 1
    h.t(r, 0, "Renglón", 'HDR'); h.t(r, 1, "Fila del 05", 'HDR')
    for i, (c0, nombre, _) in enumerate(grupos):
        h.t(r, A0 + i, str(nombre), 'HDR_C'); h.t(r, B0 + i, str(nombre), 'HDR_C'); h.t(r, D0 + i, str(nombre), 'HDR_C')
    h.t(r, T0, "Total 05 (E)", 'HDR_C'); h.t(r, T0 + 1, "ER dic (MXN)", 'HDR_C'); h.t(r, T0 + 2, "Diferencia", 'HDR_C'); h.t(r, T0 + 3, "¿Cuadra?", 'HDR_C')
    r += 1
    h.t(r, A0, "05: columna 2027 de cada LN (capturado)", 'TXT_B'); h.t(r, B0, "Vista, diciembre (Val05_LN)", 'TXT_B'); h.t(r, D0, "Diferencia", 'TXT_B')
    r += 1; r_ini = r; malos = []
    for f in FILAS:
        ln = f + 1
        h.t(r, 0, f"{ln} · {NOMBRE[ln]}", 'TXT'); h.t(r, 1, f"fila {f}", 'OBS')
        for i, (c0, nombre, mbs) in enumerate(grupos):
            k = kde[tuple(mbs)]
            v = num05(ws, f, c0 + 2); h.n(r, A0 + i, v, 'NUM')
            h.f(r, B0 + i, h.L(LAY[(f, k)], C2 + NM - 1), 'NUM_G')
            h.f(r, D0 + i, sub(h.L(r, A0 + i), h.L(r, B0 + i)), 'NUM_Y')
            if v is None or abs(v - h.val[(r, B0 + i)]) > TOL: malos.append((f, nombre, v, h.val[(r, B0 + i)]))
        h.n(r, T0, num05(ws, f, 5), 'NUM')
        mr = MROW[f]
        h.f(r, T0 + 1, X(W.ref3d(IX['Val05_Mensual'], mr, C2 + NM - 1), MENS.val[(mr, C2 + NM - 1)]), 'NUM_G')
        h.f(r, T0 + 2, sub(h.L(r, T0), h.L(r, T0 + 1)), 'NUM_Y')
        c1 = h.cuadra(r, D0, D0 + ng - 1); c2 = h.abs_lt(h.L(r, T0 + 2))
        h.fb(r, T0 + 3, X(c1.g + c2.g + W.funcvar(2, 36), bool(c1.v and c2.v)), 'BOOL')
        r += 1
    assert not malos, malos[:4]
    h.t(r + 1, 0, "La columna '2027 CA' se dejó vacía: en el 05 anterior '2026 CA' coincidía con 2026 en primas, siniestros, costos, XL y gastos (otro reparto entre LN) pero no en reservas, IBNR ni resultados, y el libro de Cesión no trae esa versión para 2027.", 'NOTA')
    ANCLA['LN']['anual'] = (r_ini, r - 1, T0 + 3)
    return r + 2, T0 + 3

# ============================================================ Val05_Resumen
def hoja_resumen():
    h = Hoja('Val05_Resumen')
    h.t(0, 0, "RESUMEN — ¿05_FCST_2027 CUADRA CONTRA ER Y SUS VISTAS?  Presupuesto 2027, diciembre acumulado (MXN)", 'TIT')
    h.t(1, 0, "El 05 se llenó recorriendo un año: 2027 sale de las vistas ER de este libro (en pesos), 2026 es el presupuesto 2026 que ya traía el archivo y 2025 el real que ya traía.", 'SUB')
    h.t(2, 0, "Cada cifra de esta hoja es fórmula sobre Val05_Mensual, Val05_LN, Val05_Ramo, Val05_Region y Val05_TRea, que comparan celda por celda los 12 meses.", 'SUB')
    grupos = [('MENSUAL', 'PptoxMes_Red vs ER'), ('LN', 'PptoxLNxMes_Red'), ('RAMO', 'PptoxRamo_Red'), ('REGION', 'PptoxRegión_Red'), ('TREA', 'PptoxTR_Red')]
    g0 = 2; r = 4
    h.t(r, 0, "Renglón", 'BAR'); h.t(r, 1, "Fila del 05", 'BAR')
    for i, (k, et) in enumerate(grupos):
        c = g0 + 4 * i
        h.t(r, c, et, 'BAR_C'); h.b(r, c + 1, 'BAR'); h.b(r, c + 2, 'BAR'); h.b(r, c + 3, 'BAR')
    r += 1
    h.t(r, 0, "", 'HDR'); h.t(r, 1, "", 'HDR')
    for i, (k, et) in enumerate(grupos):
        c = g0 + 4 * i
        h.t(r, c, "05 dic", 'HDR_C'); h.t(r, c + 1, "ER dic (MXN)" if k == 'MENSUAL' else "Vista dic (MXN)", 'HDR_C')
        h.t(r, c + 2, "Diferencia", 'HDR_C'); h.t(r, c + 3, "¿Cuadra (12 meses)?", 'HDR_C')
    r += 1; r1 = r
    for f in FILAS:
        ln = f + 1
        h.t(r, 0, f"{ln} · {NOMBRE[ln]}", 'TXT'); h.t(r, 1, f"fila {f}", 'OBS')
        mr = MROW[f]; c = g0
        h.f(r, c, X(W.ref3d(IX['Val05_Mensual'], mr, C1 + NM - 1), MENS.val[(mr, C1 + NM - 1)]), 'NUM')
        h.f(r, c + 1, X(W.ref3d(IX['Val05_Mensual'], mr, C2 + NM - 1), MENS.val[(mr, C2 + NM - 1)]), 'NUM_G')
        h.f(r, c + 2, X(W.ref3d(IX['Val05_Mensual'], mr, C3 + NM - 1), MENS.val[(mr, C3 + NM - 1)]), 'NUM_Y')
        h.fb(r, c + 3, X(W.ref3d(IX['Val05_Mensual'], mr, QC), MENS.val[(mr, QC)]), 'BOOL')
        for i, dim in enumerate(['LN', 'RAMO', 'REGION', 'TREA']):
            c = g0 + 4 * (i + 1); A = ANCLA[dim]; hd = DIMH[dim]; ix = IX[DIMINFO[dim]['hoja']]; sr = A['SUMR'][f]
            h.f(r, c, X(W.ref3d(ix, sr, C1 + NM - 1), hd.val[(sr, C1 + NM - 1)]), 'NUM')
            h.f(r, c + 1, X(W.ref3d(ix, sr, C2 + NM - 1), hd.val[(sr, C2 + NM - 1)]), 'NUM_G')
            h.f(r, c + 2, X(W.ref3d(ix, sr, C3 + NM - 1), hd.val[(sr, C3 + NM - 1)]), 'NUM_Y')
            a, b = A['LAY'][(f, 0)], A['DIFR'][f]
            vals = [hd.val.get((rr, QC)) for rr in range(a, b + 1)]
            area = struct.pack('<BHIIHH', 0x3B, ix, a, b, QC, QC)
            h.fb(r, c + 3, AND_areas([(area, vals)]), 'BOOL')
        r += 1
    r2 = r - 1
    r += 1
    # ---- totales por hoja
    h.barra(r, "¿CUADRAN TODAS LAS CELDAS?", 7); r += 1
    h.t(r, 0, "Hoja del 05", 'HDR'); h.t(r, 1, "Qué se compara", 'HDR'); h.t(r, 2, "Celdas del 05", 'HDR_C'); h.t(r, 3, "¿Todas cuadran?", 'HDR_C')
    h.t(r, 4, "Hoja de detalle", 'HDR'); r += 1
    conteo = {}
    def fila_total(etq, que, celdas, x, det):
        nonlocal r
        h.t(r, 0, etq, 'TXT'); h.t(r, 1, que, 'OBS'); h.n(r, 2, celdas, 'NUM'); h.fb(r, 3, x, 'BOOL'); h.t(r, 4, det, 'OBS')
        conteo[etq] = (celdas, x.v); r += 1
    defin = [f for f in ER_FILA if MENS.val.get((MROW[f], QC)) is not None]
    vals = [MENS.val.get((MROW[f], QC)) for f in ER_FILA]
    area = struct.pack('<BHIIHH', 0x3B, IX['Val05_Mensual'], min(MROW.values()), max(MROW.values()), QC, QC)
    fila_total("PptoxMes_Red · 2027", f"{len(defin)} renglones × 12 meses contra ER en MXN", len(defin) * NM,
               AND_areas([(area, vals)]), "Val05_Mensual, columna AO")
    for dim in ['LN', 'RAMO', 'REGION', 'TREA']:
        A = ANCLA[dim]; hd = DIMH[dim]; ix = IX[DIMINFO[dim]['hoja']]
        a, b = A['LAY'][(FILAS[0], 0)], A['DIFR'][FILAS[-1]]
        vals = [hd.val.get((rr, QC)) for rr in range(a, b + 1)]
        area = struct.pack('<BHIIHH', 0x3B, ix, a, b, QC, QC)
        fila_total(f"{DIMINFO[dim]['hoja05']} · 2027", f"{A['nb']} bloques × 20 renglones × 12 meses contra la vista; suma de bloques contra ER",
                   A['nb'] * 20 * NM, AND_areas([(area, vals)]), f"{DIMINFO[dim]['hoja']}, columna AO")
    a, b, cq = ANCLA['LN']['anual']
    vals = [DIMH['LN'].val.get((rr, cq)) for rr in range(a, b + 1)]
    area = struct.pack('<BHIIHH', 0x3B, IX['Val05_LN'], a, b, cq, cq)
    fila_total("PptoxLN_Red · 2027", "8 LN × 20 renglones contra diciembre de la vista; total contra ER", 8 * 20 + 20,
               AND_areas([(area, vals)]), "Val05_LN, sección PptoxLN_Red")
    r += 1
    # ---- 2025
    h.barra(r, "2025 (COLUMNA RECORRIDA): el real 2025 que quedó en el 05 contra la columna D de cada vista y contra ER!E", 7); r += 1
    h.t(r, 0, "Hoja del 05", 'HDR'); h.t(r, 1, "Qué se compara", 'HDR'); h.t(r, 2, "¿Cuadra con la vista?", 'HDR_C'); h.t(r, 3, "¿Suma = ER!E?", 'HDR_C')
    h.t(r, 4, "Nota", 'HDR'); r += 1
    vals = [MENS.val.get((MROW[f], Y4)) for f in ER_FILA]
    area = struct.pack('<BHIIHH', 0x3B, IX['Val05_Mensual'], min(MROW.values()), max(MROW.values()), Y4, Y4)
    x = AND_areas([(area, vals)])
    h.t(r, 0, "PptoxMes_Red (W)", 'TXT'); h.t(r, 1, "Columna W (2025) contra ER!E (Dic 2025)", 'OBS'); h.b(r, 2, 'BOOL'); h.fb(r, 3, x, 'BOOL')
    h.t(r, 4, "", 'OBS'); r += 1
    NOTA25 = {'RAMO': "ER_ram!D toma ER1225_Real, vínculo a ER_RPAT_Ramos_202512.xlsx cuyo valor guardado es el acumulado a febrero de 2025 (primas 3,188 M contra 20,691 M del año); "
                      "el 05 anterior ya venía así. Coincide con la vista pero no con ER!E. Si se actualizan los vínculos habría que volver a llenar esa columna.",
              'LN': 'Los bloques coinciden con la vista. La suma contra ER!E solo difiere en: IBNR y resultados técnicos por 1.85 y 1.29 pesos (redondeo entre Anexo2_1225 y ER!E), y Gastos Generales (703.49 M), Contribución (321.70 M), Total Costos de Operación (381.79 M) y Resultado de Operación, porque el real 2025 por dimensión (Anexo2_1225) no trae gastos. Por eso las filas por bloque de esos tres renglones no tienen ¿Cuadra? 2025 (AQ:AT en blanco) y el VERDADERO de la columna C no las cubre.', 'REGION': 'Los bloques coinciden con la vista. La suma contra ER!E solo difiere en: IBNR y resultados técnicos por 1.85 y 1.29 pesos (redondeo entre Anexo2_1225 y ER!E), y Gastos Generales (703.49 M), Contribución (321.70 M), Total Costos de Operación (381.79 M) y Resultado de Operación, porque el real 2025 por dimensión (Anexo2_1225) no trae gastos. Por eso las filas por bloque de esos tres renglones no tienen ¿Cuadra? 2025 (AQ:AT en blanco) y el VERDADERO de la columna C no las cubre.', 'TREA': 'Los bloques coinciden con la vista. La suma contra ER!E solo difiere en: IBNR y resultados técnicos por 1.85 y 1.29 pesos (redondeo entre Anexo2_1225 y ER!E), y Gastos Generales (703.49 M), Contribución (321.70 M), Total Costos de Operación (381.79 M) y Resultado de Operación, porque el real 2025 por dimensión (Anexo2_1225) no trae gastos. Por eso las filas por bloque de esos tres renglones no tienen ¿Cuadra? 2025 (AQ:AT en blanco) y el VERDADERO de la columna C no las cubre.'}
    for dim in ['LN', 'RAMO', 'REGION', 'TREA']:
        A = ANCLA[dim]; hd = DIMH[dim]; ix = IX[DIMINFO[dim]['hoja']]
        areas_b, areas_s = [], []
        for f in FILAS:
            a, b = A['LAY'][(f, 0)], A['LAY'][(f, A['nb'] - 1)]
            areas_b.append((struct.pack('<BHIIHH', 0x3B, ix, a, b, Y4, Y4), [hd.val.get((rr, Y4)) for rr in range(a, b + 1)]))
            d = A['DIFR'][f]
            areas_s.append((struct.pack('<BHIIHH', 0x3B, ix, d, d, Y4, Y4), [hd.val.get((d, Y4))]))
        h.t(r, 0, DIMINFO[dim]['hoja05'] + " (E)", 'TXT'); h.t(r, 1, f"Cada bloque contra {DIMINFO[dim]['vista']} columna D; suma de bloques contra ER!E", 'OBS')
        h.fb(r, 2, AND_areas(areas_b), 'BOOL'); h.fb(r, 3, AND_areas(areas_s), 'BOOL'); h.t(r, 4, NOTA25[dim], 'OBS'); r += 1
    r += 1
    # ---- notas
    h.barra(r, "LO QUE NO SE LLENÓ O CONVIENE REVISAR", 7); r += 1
    for txt in NOTAS:
        h.t(r, 0, txt, 'NOTA'); r += 1
    r += 1
    h.t(r, 0, "Cómo leer: '05 dic' es la cifra del archivo 05_FCST_2027 (capturada); 'ER dic' y 'Vista dic' son fórmulas sobre este libro en pesos. "
              "¿Cuadra (12 meses)? revisa los 12 meses de todos los bloques de ese renglón y la suma contra ER.", 'NOTA')
    cols = [(0, 0, 46), (1, 1, 60)] + [(g0 + 4 * i + j, g0 + 4 * i + j, 17 if j < 3 else 12) for i in range(5) for j in range(4)]
    return h, W.sheet_bin(h.rows, g0 + 20, cols=cols, freeze=(2, 6)), conteo

NOTAS = [
    "1) Filas financieras 2027 de PptoxMes_Red: ER!Q62, Q63 y Q67 tienen =$Q62/12+P62 (se refieren a sí mismas, el libro no calcula iterativo) y de enero a noviembre también dependen de ese Q. "
    "Se usó el anual del bloque 'ESTACIONALIDAD ACUMULADA MXN' de RIF, como ya hace la fila 60 (ER!Y60 = RIF!P231): RIF!P233 = 18.00 M (Productos de Inmuebles), P234 = 1,113.48 M (Intereses) y P237 = 411.55 M (Valuación), ÷ 12 acumulado. "
    "Con eso enero a noviembre coinciden con lo que ER tiene guardado; son los montos del presupuesto 2026. RIF, impuestos y utilidad salen con las fórmulas de ER (filas 70 a 81). "
    "Para corregir ER: Y62 = RIF!P233, Y63 = RIF!P234, Y67 = RIF!P237 y en F:Q usar $Y en lugar de $Q.",
    "2) Dividendos (ER fila 64): J64 = RIF!E22 × RIF!K23 (41.35 M) y se arrastra hasta O64; P64 = RIF!E22 (65.50 M) y Q64 = RIF!F22, que está en blanco (columna del año 2025 de RIF), así que diciembre da 0. "
    "El Resultado Integral de Financiamiento de diciembre (fila 53 del 05, 1,470.40 M) no los incluye, aunque la hoja RIF sí los trae (RIF!E22 y RIF!P235 = 65.50 M). "
    "El 05 trae lo que calcula ER; parece un corrimiento de una columna (Q64 debería ser RIF!E22; en el presupuesto 2026 diciembre traía 65.50 M).",
    "3) Costos Protecciones XL 2027 = 0 en todo el 05: ER (filas 29 y 43) y las vistas (fila 27) toman CtaMens!AA, que viene en 0 en los 201,231 renglones de 2027 (AB también); la hoja CostosXL solo llega a 202612 y CtaAnual también trae 0. "
    "En el presupuesto 2026 eran 1,561.79 M y en el real 2025 1,598.72 M. Por eso el Resultado Técnico a Retención 2027 (3,710.72 M) y el de Operación (3,223.88 M) no traen ese costo: hay que cargar el XL 2027.",
    "4) Prima 2027 que CtaMens no trae: no hay renglones de LN04009 (Daños Ultramar Londres), de LN04008 (Facultativos Sur sin Agro) ni del ramo GMM, y Salud solo trae prima en diciembre. Ninguna otra LN la recoge (Líneas Especiales baja de 7,147.7 M a 6,649.7 M). "
    "CtaAnual del mismo libro sí trae prima 2027 para ellas (en USD: LN04009 195.07 M, LN04008 31.85 M, GMM 3.00 M, Salud 31.69 M). Por eso esos bloques traen las reservas de 2_Reservas sin prima: "
    "Londres da prima 0 y resultado técnico a retención y de operación de −509.7 M, Fac. Sur y Agropecuario queda con 403.8 M de prima (solo Agro, contra 963.0 M en 2026) "
    "y GMM da resultado técnico a retención y de operación de −3.7 M.",
    "5) Gastos Generales y Gasto GAEF 2027 salen de la hoja Gastos, cuyos meses (Gastos!C5:N5) son 202601 a 202612: ER fila 46 = Gastos!C12:N12 = C6:N6 ÷ 'Parámetros'!F7:Q7 (N12 lleva +0.35) y ER fila 69 = −Gastos!C13:N13 = −C7:N7 ÷ F7:Q7. "
    "En pesos repiten el presupuesto 2026: Gastos Generales 697.16 M y GAEF −123.02 M.",
    "6) Daños Facultativos Sur y Agropecuario = LN04008 (reservas) + LN04008-Agro (CtaMens). Crédito = 'Crédito' (reservas) + 'Credito' (CtaMens). Así cada miembro de la vista cae en un bloque y los bloques suman ER.",
    "7) PptoxLN_Red: '2027 CA' se dejó vacía. En el 05 anterior '2026 CA' coincidía con 2026 en primas, siniestros, costos, XL y gastos (con otro reparto entre LN), pero no en reservas, IBNR ni resultados (columna A 'RESREVAS'); el libro de Cesión no trae esa segunda versión para 2027.",
    "8) Columnas T/U/V de las hojas por dimensión: T y U se recorrieron (eran U y V) y V se calculó para 2027 con la regla de cada renglón. En las filas de variación de reserva (11 y 12), "
    "V = variación ÷ (aumento de la prima tomada o retenida contra 2026), igual que el 05 anterior, y tienen cifra en todos los bloques. Quedan vacías las demás V cuyo denominador es la prima 2027 = 0 "
    "(Londres y GMM, filas 9, 10, 15, 18, 22, 25, 26, 30, 31 y 32 del bloque).",
    "9) Filas 37/38 (% combinado) de 2025 y 2026 se recalcularon con las cifras recorridas y la fórmula de la vista. En 32 celdas difieren de lo que imprimía el 05 anterior, que no cuadraba con sus propias cifras (Ramo 2025 de Resp. Civil a Fianzas y Región 2026). "
    "En 2027 la fila 38 queda vacía donde la prima es 0 (Londres, GMM y Salud de enero a noviembre).",
    "10) PptoxMes_Red, mayo 2026 (AB11): el 05 anterior traía '|' en lugar de la cifra; se puso Prima Tomada − Prima Devengada del mes, que es como sale en los demás meses.",
    "11) El 2026 del 05 es el presupuesto 2026 que ya traía el archivo (Integración2026_Dim_9), no el reforecast 9+3 de ER!D.",
    "12) El libro está en cálculo manual. Los valores guardados de ER_ln, ER_ram, ER_reg y ER_tre en las filas de CAT y gastos (17, 20, 24, 28, 31, 33 y 35) son anteriores al reparto de CAT: antes de comparar una vista contra el 05, recalcular con Ctrl+Alt+F9. Las hojas Val_* y Val05_* ya están calculadas. "
    "El libro trae además la marca de recálculo completo al abrir; todo recálculo completo muestra el aviso de referencia circular de ER!Q62, Q63 y Q67 (nota 1).",
    "13) En ER_ram el nombre xEvCat tiene una definición local que apunta a Inicio!C13 del libro externo Integración2025_2030 - Base_9+3_Esc1 (valor guardado 0), no al de este libro; ahí multiplica los eventos CAT (CtaMens!AC, 45 M USD en 2027) en las filas 17 y 24. "
    "Además la fila 24 de las cuatro vistas resta las recuperaciones AB × xEvCat con signo contrario a ER. Hoy no hay efecto (xEvCat = 0 y AB = 0); con los escenarios con eventos CAT, ER_ram dejaría fuera lo que sí toman ER y las otras vistas.",
    "14) ER!F70:Q81 guarda valores de un cálculo anterior (F70 = 1,379.76 M, igual a E70) y #REF! en G73:Q74, G77:Q78 y G81:Q81. Val05_Mensual no los usa: rehace RIF, impuestos y utilidad con las fórmulas de ER sobre su propia hoja. "
    "Se limpian al aplicar la corrección de la nota 1 y recalcular.",
]

# ============================================================ empaquetado
def empaquetar(src, dst, partes):
    shutil.copyfile(src, dst)
    with tempfile.TemporaryDirectory() as td:
        for nombre, data in partes.items():
            p = os.path.join(td, nombre); os.makedirs(os.path.dirname(p), exist_ok=True); open(p, 'wb').write(data)
        rr = subprocess.run(['zip', '-X', '-D', '-q', os.path.abspath(dst)] + list(partes), cwd=td, capture_output=True, text=True)
        if rr.returncode: raise RuntimeError(rr.stderr)

if __name__ == '__main__':
    import empaquetar2 as E2
    P = W.Cadenas(E2.leer_sst(SRC))
    _i = 0
    for _rid, _pl in records(zipfile.ZipFile(SRC).read('xl/sharedStrings.bin')):
        if _rid == 19:
            if _pl[0] == 0:                                   # solo cadenas simples (sin formato ni fonetica)
                _s = W_rd(_pl)
                P.idx.setdefault(_s, _i)
            _i += 1
    MENS, b_mens = hoja_mensual()
    DIMH = {}; BIN = {}
    for dim in ['LN', 'RAMO', 'REGION', 'TREA']:
        hd, fin = hoja_dim(dim)
        if dim == 'LN':
            fin, ultima = seccion_anual_ln(hd, fin)
        DIMH[dim] = hd
        cols = [(0, 0, 36), (1, 1, 44)] + [(C1 + j, C1 + j, 15) for j in range(NM)] + [(14, 14, 1.5)] + \
               [(C2 + j, C2 + j, 15) for j in range(NM)] + [(27, 27, 1.5)] + [(C3 + j, C3 + j, 12) for j in range(NM)] + \
               [(QC, QC, 10), (41, 41, 1.5), (Y1, Y2, 17), (Y3, Y3, 12), (Y4, Y4, 10)]
        BIN[dim] = W.sheet_bin(hd.rows, Y4 + 1, cols=cols, freeze=(2, 4))
    RES, b_res, conteo = hoja_resumen()
    hojas = [('Val05_Resumen', b_res), ('Val05_Mensual', b_mens), ('Val05_LN', BIN['LN']), ('Val05_Ramo', BIN['RAMO']),
             ('Val05_Region', BIN['REGION']), ('Val05_TRea', BIN['TREA'])]
    # ---- partes del paquete
    z = zipfile.ZipFile(SRC)
    wb = z.read('xl/workbook.bin'); rels = z.read('xl/_rels/workbook.bin.rels').decode('utf-8'); cts = z.read('[Content_Types].xml').decode('utf-8')
    usados = {int(m) for m in re.findall(r'worksheets/sheet(\d+)\.bin', ' '.join(z.namelist()))}
    n_rid = max(int(m) for m in re.findall(r'Id="rId(\d+)"', rels)) + 1
    max_tab = max(struct.unpack_from('<I', pl, 4)[0] for rid, pl in records(wb) if rid == 156)
    n_hojas = sum(1 for rid, pl in records(wb) if rid == 156); assert n_hojas == ITAB0, n_hojas
    nuevos = [dict(nombre=nm, parte=f'xl/worksheets/sheet{max(usados) + 1 + i}.bin', rid=f'rId{n_rid + i}', tab=max_tab + 1 + i, bin=b)
              for i, (nm, b) in enumerate(hojas)]
    partes = {}
    out = bytearray()
    for rid, pl, raw in records_raw(wb):
        if rid == 144:
            for n in nuevos: out += W.rec(156, struct.pack('<II', 0, n['tab']) + W.wstr(n['rid']) + W.wstr(n['nombre']))
        if rid == 362:
            (cn,) = struct.unpack_from('<I', pl, 0); assert cn == 69, cn
            cuerpo = pl[4:4 + 12 * cn] + b''.join(struct.pack('<Iii', *x) for x in XTI_EXTRA)
            out += W.rec(362, struct.pack('<I', cn + len(XTI_EXTRA)) + cuerpo); continue
        if rid == 157:
            (fl,) = struct.unpack_from('<H', pl, len(pl) - 2)
            out += W.rec(157, pl[:-2] + struct.pack('<H', fl | 0x0001)); continue
        out += raw
    partes['xl/workbook.bin'] = bytes(out)
    for n in nuevos:
        rels = rels.replace('</Relationships>', f'<Relationship Id="{n["rid"]}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="{n["parte"].replace("xl/", "")}"/></Relationships>')
        cts = cts.replace('</Types>', f'<Override PartName="/{n["parte"]}" ContentType="application/vnd.ms-excel.worksheet"/></Types>')
        partes[n['parte']] = n['bin']
    partes['xl/_rels/workbook.bin.rels'] = rels.encode('utf-8'); partes['[Content_Types].xml'] = cts.encode('utf-8')
    usos = sum(1 for n in nuevos for _rid, _pl in records(n['bin']) if _rid == 7)
    o = bytearray()
    for rid, pl, raw in records_raw(z.read('xl/sharedStrings.bin')):
        if rid == 159:
            tot, uni = struct.unpack_from('<II', pl, 0)
            o += W.rec(159, struct.pack('<II', tot + usos, uni + len(P.nuevas)) + pl[8:]); continue
        if rid == 160:
            for s in P.nuevas: o += W.rec(19, b'\x00' + W.wstr(s))
        o += raw
    partes['xl/sharedStrings.bin'] = bytes(o)
    app = z.read('docProps/app.xml').decode('utf-8')
    m = re.search(r'(<vt:lpstr>Hojas de c[^<]*</vt:lpstr></vt:variant><vt:variant><vt:i4>)(\d+)(</vt:i4>)', app)
    assert m and '<vt:lpstr>Val_CAT</vt:lpstr>' in app
    app = app[:m.start(2)] + str(int(m.group(2)) + len(nuevos)) + app[m.end(2):]
    app = re.sub(r'(<TitlesOfParts><vt:vector size=")(\d+)(")', lambda k: k.group(1) + str(int(k.group(2)) + len(nuevos)) + k.group(3), app)
    app = app.replace('<vt:lpstr>Val_CAT</vt:lpstr>', '<vt:lpstr>Val_CAT</vt:lpstr>' + ''.join(f'<vt:lpstr>{n["nombre"]}</vt:lpstr>' for n in nuevos))
    partes['docProps/app.xml'] = app.encode('utf-8')
    z.close()
    empaquetar(SRC, DST, partes)
    VALS = {'Val05_Resumen': RES.val, 'Val05_Mensual': MENS.val, **{DIMINFO[d]['hoja']: DIMH[d].val for d in DIMH}}
    pickle.dump(dict(VALS=VALS, MROW=MROW, ANCLA=ANCLA, conteo=conteo, nuevos=[(n['nombre'], n['parte']) for n in nuevos]),
                open('construccion05.pkl', 'wb'))
    nf = sum(1 for n in nuevos for _rid, _pl in records(n['bin']) if _rid in (8, 9, 10))
    print(f'{DST}: {len(nuevos)} hojas nuevas, {nf} fórmulas, {len(P.nuevas)} cadenas nuevas, {usos} usos de cadena')
    for k, (n, ok) in conteo.items(): print(f'  {k:<28} {n:>6} celdas  ¿todas cuadran? {ok}')
    booleanos = {h: sum(1 for v in VALS[h].values() if v is False) for h in VALS}
    print('  celdas FALSO por hoja:', booleanos)
