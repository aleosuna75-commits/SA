# -*- coding: utf-8 -*-
"""Llena 05_FCST_2027.xlsx recorriendo un año: 2026 -> 2027, 2025 -> 2026, 2024 -> 2025.

  2027  : vistas ER del libro de Cesion evaluadas por miembro, en MXN (datos05.pkl)
  2026  : el presupuesto 2026 que ya traia este mismo archivo, recorrido una posicion
  2025  : el real 2025 que ya traia este archivo (= columna D de las vistas)

Se edita el XML de cada hoja celda por celda (valor y valor guardado de las formulas);
estilos, anchos, formatos, nombres y demas partes del paquete se copian tal cual.

Entrada: 05_FCST_2027_orig.xlsx, cesion/datos05.pkl. Salida: 05_FCST_2027.xlsx y
llenado05.pkl (cada celda escrita con su origen, para la hoja de validacion)."""
import re, math, pickle, zipfile, shutil
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from lxml import etree

ORIG, OUT = '05_FCST_2027_orig.xlsx', '05_FCST_2027.xlsx'
D = pickle.load(open('cesion/datos05.pkl', 'rb'))
VIS, ER, TC = D['VISTAS'], D['ER_MXN'], D['TC']
NM = 12
NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
q = lambda t: f'{{{NS}}}{t}'

wv = openpyxl.load_workbook(ORIG, data_only=True)     # valores guardados
wf = openpyxl.load_workbook(ORIG)                     # formulas

def old(hoja, ref):
    v = wv[hoja][ref].value
    return v

def num(v): return isinstance(v, (int, float)) and not isinstance(v, bool)

# ------------------------------------------------------------------ plan de escritura
PLAN = {}        # hoja -> ref -> ('num', v) | ('blank',) | ('cached', v)
ORIGEN = {}      # (hoja, ref) -> texto del origen
def put(hoja, ref, v, origen):
    if v is None or (isinstance(v, float) and (math.isnan(v) or math.isinf(v))):
        PLAN.setdefault(hoja, {})[ref] = ('blank',)
    else:
        PLAN.setdefault(hoja, {})[ref] = ('num', float(v))
    ORIGEN[(hoja, ref)] = origen
def cached(hoja, ref, v, origen):
    PLAN.setdefault(hoja, {})[ref] = ('cached', v)
    ORIGEN[(hoja, ref)] = origen

def anio_mas_uno(v):
    if not num(v): return None
    v = int(v)
    if 2000 <= v <= 2100: return v + 1
    if 200001 <= v <= 210012: return v + 100
    return None

def recorrer_encabezado(hoja, fila, cols):
    for c in cols:
        ref = f'{c}{fila}'; nv = anio_mas_uno(old(hoja, ref))
        if nv is not None: put(hoja, ref, nv, 'encabezado recorrido un año')

# ------------------------------------------------------------------ hojas por dimension
FILAS = [8, 9, 10, 11, 12, 15, 16, 17, 18, 19, 22, 23, 24, 25, 26, 27, 30, 31, 32, 34]   # fila 05 = fila vista - 1
BASE = {9: 8, 10: 8, 15: 8, 18: 8, 30: 8, 31: 8, 32: 8, 16: 15, 17: 15, 19: 15,
        22: 10, 25: 10, 26: 10, 23: 22, 24: 22, 27: 22, 34: 22}                        # T/U/V = fila / base
MESES05 = [L(7+j) for j in range(NM)]                                                     # G..R

MAPA = {
 'PptoxLNxMes_Red': ('LN', {'DAÑOS CONTRATOS NORTE': ['LN04001'], 'DAÑOS FACULTATIVOS NORTE': ['LN04002'],
    'VIDA, ACCIDENTES Y ENFERMEDADES': ['LN04004'], 'DAÑOS CONTRATOS SUR': ['LN04005'],
    'DAÑOS FACULTATIVOS SUR Y AGROPECUARIO': ['LN04008', 'LN04008-Agro'], 'FIANZAS Y CRÉDITO': ['LN04003'],
    'DAÑOS LÍNEAS ESPECIALES': ['LN04006'], 'DAÑOS ULTRAMAR LONDRES': ['LN04009']}),
 'PptoxRamo_Red': ('RAMO', {'VIDA': ['Vida'], 'ACCIDENTES PERSONALES': ['Acc Per.'], 'GASTOS MÉDICOS': ['GMM'],
    'SALUD': ['Salud'], 'RESPONSABILIDAD CIVIL': ['Resp. Civil'], 'MARÍTIMO Y TRANSPORTES': ['MyT'],
    'INCENDIO': ['Incendio'], 'TERREMOTO': ['Terremoto'], 'RIESGOS HIDROMETEOROLÓGICOS': ['HyORH'],
    'AGROPECUARIO': ['Agropecuario'], 'AUTOMÓVILES': ['Autos'], 'CRÉDITO': ['Crédito', 'Credito'],
    'DIVERSOS': ['Diversos'], 'FIANZAS': ['Fianzas']}),
 'PptoxRegión_Red': ('REGION', {'MÉXICO': ['R01'], 'CENTROAMÉRICA': ['R02'], 'CONO SUR': ['R03'],
    'PACTO ANDINO Y VENEZUELA': ['R04'], 'ULTRAMAR': ['R05'], 'EL CARIBE': ['R06']}),
 'PptoxTR_Red': ('TREA', {'PROPORCIONAL': [1.0], 'NO PROPORCIONAL': [2.0], 'FACULTATIVO': [3.0]}),
}

def suma_miembros(dim, mbs, ln, j, clave='mxn'):
    vals = [VIS[dim][clave][mb][ln] for mb in mbs]
    vals = [v[j] for v in vals if v is not None]
    return sum(vals) if vals else None

def prev_miembros(dim, mbs, ln, clave):
    vals = [VIS[dim][clave][mb].get(ln) for mb in mbs]
    vals = [v for v in vals if v is not None]
    return sum(vals) if vals else None

BLOQUES = {}     # hoja -> [(fila_encabezado, nombre, dim, miembros)]
CHEQUEO_2025 = []
for hoja, (dim, mapa) in MAPA.items():
    ws = wv[hoja]; BLOQUES[hoja] = []
    usados = set()
    for k in range(40):
        r0 = 6 + 40*k; nombre = ws[f'C{r0}'].value
        if not nombre: break
        mbs = mapa[str(nombre).strip()]; usados.update(mbs)
        BLOQUES[hoja].append((r0, str(nombre).strip(), dim, mbs))
        o = r0 - 6
        recorrer_encabezado(hoja, r0, ['E', 'F'] + MESES05 + ['T', 'U', 'V'])
        nuevo = {}      # (fila05, col) -> valor nuevo, para ratios
        for f in FILAS:
            r = f + o; ln = f + 1
            e_new, f_new = old(hoja, f'F{r}'), old(hoja, f'R{r}')
            put(hoja, f'E{r}', e_new if num(e_new) else None, f'real 2025: {hoja}!F{r} del 05 anterior')
            put(hoja, f'F{r}', f_new if num(f_new) else None, f'Ppto 2026: {hoja}!R{r} del 05 anterior (dic 2026)')
            nuevo[(f, 'E')] = e_new if num(e_new) else None; nuevo[(f, 'F')] = f_new if num(f_new) else None
            for j, c in enumerate(MESES05):
                v = suma_miembros(dim, mbs, ln, j)
                put(hoja, f'{c}{r}', v, f'vista {dim} {"+".join(map(str, mbs))} fila {ln} mes {j+1} MXN')
                nuevo[(f, c)] = v
            # comprobacion: el 2025 recorrido es la columna D de la vista
            d25 = prev_miembros(dim, mbs, ln, 'real25'); d24 = prev_miembros(dim, mbs, ln, 'real24')
            CHEQUEO_2025.append((hoja, r, nombre, ln, e_new, d25, old(hoja, f'E{r}'), d24))
        # ratios combinados (filas 37 y 38): se recalculan con las cifras nuevas
        def g(f, c):
            v = nuevo.get((f, c)); return 0.0 if v is None else v
        for c in ['E', 'F'] + MESES05:
            r37, r38 = 37 + o, 38 + o
            if old(hoja, f'{c}{r37}') is None and c not in ('E', 'F') and old(hoja, f'G{r37}') is None:
                continue
            try: v37 = 1 - g(34, c)/g(22, c)
            except ZeroDivisionError: v37 = None
            try: v38 = (g(23, c)+g(24, c))/g(22, c) + (g(25, c)+g(26, c))/g(10, c) + g(32, c)/g(8, c)
            except ZeroDivisionError: v38 = None
            put(hoja, f'{c}{r37}', v37, '1 − Resultado de Operación / Prima Neta Devengada')
            put(hoja, f'{c}{r38}', v38, '(Sin+IBNR)/PND + (Costos+XL)/Prima Retenida + Gastos/Prima Tomada')
        # T/U/V: T <- U anterior, U <- V anterior, V = 2027 con la regla de cada fila
        for f in FILAS:
            r = f + o
            for dst, src in (('T', 'U'), ('U', 'V')):
                x = old(hoja, f'{src}{r}')
                put(hoja, f'{dst}{r}', x if num(x) else None, f'{hoja}!{src}{r} del 05 anterior (recorrido)')
            R = nuevo.get((f, 'R')); Fv = nuevo.get((f, 'F'))
            if f == 8:
                v = (R/Fv - 1) if (R is not None and Fv) else None; how = 'crecimiento 2027 vs 2026'
            elif f in BASE:
                b = nuevo.get((BASE[f], 'R'))
                v = ((R or 0.0)/b) if b else None; how = f'dic 2027 / fila {BASE[f]}'
            else:
                v = None; how = 'sin regla conocida (el 05 anterior traía la cifra capturada)'
            put(hoja, f'V{r}', v, how)
    sobran = [mb for mb in VIS[dim]['miembros'] if mb not in usados]
    assert not sobran, (hoja, sobran)

# ------------------------------------------------------------------ PptoxLN_Red (anual por LN)
H = 'PptoxLN_Red'; ws = wv[H]
LN_GRUPO = {'daños contratos norte': ['LN04001'], 'daños facultativos norte': ['LN04002'],
            'fianzas y crédito': ['LN04003'], 'vida, accidentes y enfermedades': ['LN04004'],
            'daños contratos sur': ['LN04005'], 'daños líneas especiales': ['LN04006'],
            'daños facultativos sur y agropecuario': ['LN04008', 'LN04008-Agro'],
            'daños ultramar londres': ['LN04009']}
FILAS_LN = [8, 9, 10, 11, 12, 15, 16, 17, 18, 19, 22, 23, 24, 25, 26, 27, 30, 31, 32, 34]
put(H, 'E6', 2027, 'encabezado recorrido un año')
GRUPOS = []
for k in range(8):
    c0 = 7 + 4*k; nombre = ws.cell(6, c0).value; mbs = LN_GRUPO[str(nombre).strip().lower()]
    GRUPOS.append((c0, nombre, mbs))
    for i in range(3):
        ref = f'{L(c0+i)}5'; nv = anio_mas_uno(old(H, ref))
        if nv is not None: put(H, ref, nv, 'encabezado recorrido un año')
    for f in FILAS_LN:
        a25, a26 = old(H, f'{L(c0+1)}{f}'), old(H, f'{L(c0+2)}{f}')
        put(H, f'{L(c0)}{f}', a25 if num(a25) else None, f'real 2025: {H}!{L(c0+1)}{f} del 05 anterior')
        put(H, f'{L(c0+1)}{f}', a26 if num(a26) else None, f'Ppto 2026: {H}!{L(c0+2)}{f} del 05 anterior')
        put(H, f'{L(c0+2)}{f}', suma_miembros('LN', mbs, f+1, NM-1),
            f'vista LN {"+".join(mbs)} fila {f+1} dic MXN')
        put(H, f'{L(c0+3)}{f}', None, '2027 CA: se deja vacía (no hay reparto con otra clasificación de LN)')
for f in FILAS_LN:
    tot = [suma_miembros('LN', mbs, f+1, NM-1) for _, _, mbs in GRUPOS]
    tot = [t for t in tot if t is not None]
    put(H, f'E{f}', sum(tot) if tot else None, 'suma de las LN 2027 (dic, MXN)')

# ------------------------------------------------------------------ PptoxMes_Red (total mensual)
H = 'PptoxMes_Red'
ER_FILA = {8: [5], 9: [6], 10: [7], 11: [8], 12: [9], 15: [11], 16: [13], 17: [15], 18: [17], 19: [20],
           22: [32], 23: [34, -37], 24: [36], 25: [39, -42], 26: [43], 27: [44], 30: [46], 31: [50], 32: [52],
           34: [53], 37: [55], 38: [56], 39: [57], 41: [58], 44: [60], 45: [61], 46: [62], 47: [63], 48: [64],
           49: [66], 50: [67], 51: [68], 52: [69], 53: [70], 55: [71], 58: [73], 59: [74], 60: [75], 61: [76],
           62: [77], 64: [78], 66: [79], 68: [81]}
IZQ = [L(6+j) for j in range(NM)]            # F..Q
DER = [L(24+j) for j in range(NM)]           # X..AI
def er_mxn(f, j):
    tot = 0.0
    for x in ER_FILA[f]:
        s = 1 if x > 0 else -1; serie = ER.get(abs(x))
        if serie is None: return None
        tot += s*serie[j]
    return tot

# comprobacion del mapeo: 2025 del 05 anterior (E) contra la columna E de ER (Dic 2025)
CEL = pickle.load(open('cesion/celdas_v3.pkl', 'rb'))
CHEQUEO_ER25 = []
for f, filas in ER_FILA.items():
    er25 = sum((1 if x > 0 else -1)*(CEL['ER'].get((abs(x)-1, 4), (0.0,))[0] or 0.0) for x in filas)
    CHEQUEO_ER25.append((f, old(H, f'E{f}'), er25))

recorrer_encabezado(H, 6, ['E'] + IZQ + ['W'] + DER)
FORM_DER = {}
for r in range(8, 91):
    for c in ['W'] + DER:
        fx = wf[H][f'{c}{r}'].value
        if isinstance(fx, str) and fx.startswith('='): FORM_DER[(c, r)] = fx[1:]
nuevo = {}
for r in range(8, 69):
    # bloque derecho: W <- E anterior (2025), X..AI <- F..Q anterior (Ppto 2026)
    for c_dst, c_src in [('W', 'E')] + list(zip(DER, IZQ)):
        if (c_dst, r) in FORM_DER: continue
        x = old(H, f'{c_src}{r}')
        if x is None and old(H, f'{c_dst}{r}') is None: continue
        if isinstance(x, str) and r == 11 and c_src == 'J':     # '|' capturado en may-26
            x = old(H, f'J8') - old(H, f'J15'); put(H, f'{c_dst}{r}', x, 'may-26: el 05 anterior traía "|"; = Prima Tomada − Prima Devengada')
        else:
            put(H, f'{c_dst}{r}', x if num(x) else None, f'{"real 2025" if c_src == "E" else "Ppto 2026"}: {H}!{c_src}{r} del 05 anterior')
        nuevo[(c_dst, r)] = x if num(x) else None
    # bloque izquierdo: E <- Q anterior (dic 2026), F..Q <- ER 2027 MXN
    if old(H, f'E{r}') is not None or old(H, f'Q{r}') is not None:
        x = old(H, f'Q{r}')
        put(H, f'E{r}', x if num(x) else None, f'Ppto 2026: {H}!Q{r} del 05 anterior (dic 2026)')
        nuevo[('E', r)] = x if num(x) else None
    if r in ER_FILA:
        for j, c in enumerate(IZQ):
            v = er_mxn(r, j)
            put(H, f'{c}{r}', v, f'ER fila(s) {ER_FILA[r]} mes {j+1} MXN' if v is not None
                else f'ER fila {ER_FILA[r]}: depende de la referencia circular de ER!Q62/Q63/Q67 → vacía')
            nuevo[(c, r)] = v

def g(c, r):
    v = nuevo.get((c, r)); return v
RAT = {74: ('div', 10, 8), 75: ('div', 23, 16), 76: ('div', 25, 18), 78: ('div', 16, 8), 79: ('div', 25, 8),
       81: ('div', 23, 22), 82: ('div', 24, 22), 83: ('sum', 81, 82), 84: ('div', 25, 10), 85: ('div', 26, 10),
       86: ('sum', 84, 85), 87: ('div', 32, 8), 89: ('sum3', 83, 86, 87), 90: ('uno_menos', 34, 22)}
def ratio(c, r):
    t = RAT[r]
    try:
        if t[0] == 'div':
            a, b = g(c, t[1]), g(c, t[2]); return None if a is None or not b else a/b
        if t[0] == 'sum':
            a, b = g(c, t[1]), g(c, t[2]); return None if a is None or b is None else a+b
        if t[0] == 'sum3':
            a, b, e = g(c, t[1]), g(c, t[2]), g(c, t[3]); return None if None in (a, b, e) else a+b+e
        a, b = g(c, t[1]), g(c, t[2]); return None if a is None or not b else 1 - a/b
    except ZeroDivisionError:
        return None
# formulas del bloque derecho (41, 55, 64, 68 y KPIs) y KPIs del izquierdo, en orden de dependencia
for r in (41, 55, 64, 68):
    for c in ['W'] + DER:
        if (c, r) not in FORM_DER: continue
        fx = FORM_DER[(c, r)]
        m = re.match(r'^\+?([A-Z]+)(\d+)([+-])([A-Z]+)(\d+)$', fx); assert m, fx
        a, b = g(m.group(1), int(m.group(2))), g(m.group(4), int(m.group(5)))
        v = None if a is None or b is None else (a + b if m.group(3) == '+' else a - b)
        cached(H, f'{c}{r}', v, f'fórmula ={fx}'); nuevo[(c, r)] = v
for r in sorted(RAT):
    for c in ['W'] + DER:
        if (c, r) in FORM_DER:
            v = ratio(c, r); cached(H, f'{c}{r}', v, f'fórmula ={FORM_DER[(c, r)]}'); nuevo[(c, r)] = v
    for c in IZQ:
        if old(H, f'{c}{r}') is None: continue
        v = ratio(c, r); put(H, f'{c}{r}', v, f'KPI fila {r} con las cifras 2027'); nuevo[(c, r)] = v

# ------------------------------------------------------------------ textos (tabla de cadenas)
TEXTOS = {'Fuente: Integración2026_Dim_9.xlsb': 'Fuente: FCST_2027_Cesion.xlsb (vistas ER, en MXN)',
          'ESTADO DE RESULTADOS PRESUPUESTO MENSUALIZADO 2026': 'ESTADO DE RESULTADOS PRESUPUESTO MENSUALIZADO 2027',
          'ESTADO DE RESULTADOS PRESUPUESTO 2026 POR LÍNEA DE NEGOCIO (NUEVAS)': 'ESTADO DE RESULTADOS PRESUPUESTO 2027 POR LÍNEA DE NEGOCIO (NUEVAS)',
          '2026 CA': '2027 CA',
          'ESTADO DE RESULTADOS PRESUPUESTO 2026 POR RAMO': 'ESTADO DE RESULTADOS PRESUPUESTO 2027 POR RAMO',
          'ESTADO DE RESULTADOS PRESUPUESTO 2026 POR REGIÓN (DE ORIGEN)': 'ESTADO DE RESULTADOS PRESUPUESTO 2027 POR REGIÓN (DE ORIGEN)',
          'ESTADO DE RESULTADOS PRESUPUESTO 2026 POR TIPO DE REASEGURO': 'ESTADO DE RESULTADOS PRESUPUESTO 2027 POR TIPO DE REASEGURO'}

# ------------------------------------------------------------------ escritura del XML
z = zipfile.ZipFile(ORIG)
rels = etree.fromstring(z.read('xl/_rels/workbook.xml.rels'))
wbx = etree.fromstring(z.read('xl/workbook.xml'))
rid2part = {r.get('Id'): 'xl/' + r.get('Target').lstrip('/').replace('xl/', '') for r in rels}
PARTE = {s.get('name'): rid2part[s.get(f'{{{"http://schemas.openxmlformats.org/officeDocument/2006/relationships"}}}id')]
         for s in wbx.find(q('sheets'))}

def ref_rc(ref):
    m = re.match(r'([A-Z]+)(\d+)', ref); return int(m.group(2)), CI(m.group(1))

def aplicar(xml, cambios):
    root = etree.fromstring(xml); sd = root.find(q('sheetData'))
    filas = {int(r.get('r')): r for r in sd.findall(q('row'))}
    for ref, acc in cambios.items():
        rr, cc = ref_rc(ref)
        row = filas.get(rr)
        if row is None:
            if acc[0] == 'blank': continue
            row = etree.Element(q('row')); row.set('r', str(rr))
            despues = [k for k in filas if k > rr]
            if despues: filas[min(despues)].addprevious(row)
            else: sd.append(row)
            filas[rr] = row
        cel = None
        for c in row.findall(q('c')):
            if c.get('r') == ref: cel = c; break
        if cel is None:
            if acc[0] == 'blank': continue
            cel = etree.Element(q('c')); cel.set('r', ref)
            sig = [c for c in row.findall(q('c')) if ref_rc(c.get('r'))[1] > cc]
            if sig: sig[0].addprevious(cel)
            else: row.append(cel)
            # estilo de la celda de la izquierda si la nueva no tiene
            prev = cel.getprevious()
            if prev is not None and prev.get('s'): cel.set('s', prev.get('s'))
        if acc[0] == 'cached':
            assert cel.find(q('f')) is not None, ref
            v = cel.find(q('v'))
            if acc[1] is None:
                if v is not None: cel.remove(v)
            else:
                if v is None: v = etree.SubElement(cel, q('v'))
                v.text = repr(float(acc[1]))
            if 't' in cel.attrib: del cel.attrib['t']
            continue
        assert cel.find(q('f')) is None, ref
        for ch in list(cel):
            cel.remove(ch)
        if 't' in cel.attrib: del cel.attrib['t']
        if acc[0] == 'num':
            v = etree.SubElement(cel, q('v')); v.text = repr(float(acc[1]))
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

def cadenas(xml):
    root = etree.fromstring(xml); hechos = set()
    for si in root.findall(q('si')):
        t = si.find(q('t'))
        if t is not None and t.text in TEXTOS:
            hechos.add(t.text); t.text = TEXTOS[t.text]
    assert hechos == set(TEXTOS), set(TEXTOS) - hechos
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

def libro(xml):
    root = etree.fromstring(xml); cp = root.find(q('calcPr'))
    cp.set('fullCalcOnLoad', '1')
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

nuevas = {}
for hoja, cambios in PLAN.items():
    nuevas[PARTE[hoja]] = aplicar(z.read(PARTE[hoja]), cambios)
usos = sum(len(re.findall(rb'<c [^>]*t="s"', x)) for x in
           [nuevas.get(PARTE[h], z.read(PARTE[h])) for h in PARTE])
sst = cadenas(z.read('xl/sharedStrings.xml'))
sst = re.sub(rb'count="\d+"', b'count="%d"' % usos, sst, count=1)
nuevas['xl/sharedStrings.xml'] = sst
nuevas['xl/workbook.xml'] = libro(z.read('xl/workbook.xml'))

with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as zo:
    for it in z.infolist():
        data = nuevas.get(it.filename, z.read(it.filename))
        zi = zipfile.ZipInfo(it.filename, date_time=it.date_time); zi.compress_type = zipfile.ZIP_DEFLATED
        zi.external_attr = it.external_attr
        zo.writestr(zi, data)

pickle.dump(dict(PLAN=PLAN, ORIGEN=ORIGEN, BLOQUES=BLOQUES, GRUPOS=GRUPOS, ER_FILA=ER_FILA,
                 CHEQUEO_2025=CHEQUEO_2025, CHEQUEO_ER25=CHEQUEO_ER25, MAPA=MAPA, LN_GRUPO=LN_GRUPO),
            open('llenado05.pkl', 'wb'))
n = sum(len(v) for v in PLAN.values())
print(f'{OUT}: {n} celdas escritas en {len(PLAN)} hojas')

# ------------------------------------------------------------------ comprobaciones del mapeo
peor25 = max((abs((a or 0) - (b or 0)), h, r, ln) for h, r, nm, ln, a, b, _, _ in CHEQUEO_2025)
peor24 = max((abs((c or 0) - (d or 0)), h, r, ln) for h, r, nm, ln, _, _, c, d in CHEQUEO_2025)
print('2025 recorrido vs columna D de la vista (peor):', peor25)
print('2024 del 05 anterior vs columna C de la vista (peor):', peor24)
malos = [(f, a, b) for f, a, b in CHEQUEO_ER25 if num(a) and abs(a - b) > 0.01]
print('PptoxMes_Red: 2025 del 05 anterior vs ER!E (filas con dif):', malos)
